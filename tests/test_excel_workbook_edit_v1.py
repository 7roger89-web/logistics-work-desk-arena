"""Generic XLSX upload and separate editable-workbook draft round trip."""
from __future__ import annotations

import base64
import hashlib
import json
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from core.gridfs_client import GridFSAttachmentStore
from views import logi_c1, logi_sandbox

CUSTOMER_ID = "CUST-WORKBOOK-EDIT-V1"
FILE_NAME = "freight-rate-workbook.xlsx"


def freight_workbook() -> bytes:
    wb = Workbook()
    rates = wb.active
    rates.title = "Ocean Freight"
    rates.merge_cells("B2:D2")
    rates["B2"] = "FREIGHT QUOTATION · NINGBO TO HAI PHONG"
    rates["B3"], rates["C3"], rates["D3"] = "Route", "Equipment", "Rate USD"
    rates["B4"], rates["C4"], rates["D4"] = "Ningbo → Hai Phong", "40HQ", 1750
    rates["B5"], rates["C5"], rates["D5"] = "Valid until", "2026-12-31", "=D4*1.05"
    rates["D4"].number_format = '"USD "#,##0.00'
    services = wb.create_sheet("Port charges")
    services["C4"] = "Terminal handling"
    services["D4"] = 210
    services["E4"] = "USD"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def mixed_customer_workbook() -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "客户说明"
    summary.merge_cells("B2:F2")
    summary["B2"] = "Customer workbook · mixed estimates"
    summary["B4"], summary["C4"] = "Route", "Ningbo → Hanoi"
    summary["B5"], summary["C5"] = "Terms", "Rates subject to confirmation"
    goods = wb.create_sheet("Customer goods tab")
    goods["C6"], goods["D6"], goods["E6"], goods["F6"], goods["G6"] = "SKU", "Product", "Qty", "Currency", "Price"
    goods["C7"], goods["D7"], goods["E7"], goods["F7"], goods["G7"] = "BX-01", "Cabinet", 12, "CNY", 870
    goods["G7"].comment = Comment("Confirm if this is customer sale price", "supplier")
    local = wb.create_sheet("Local charges")
    local["B4"], local["C4"], local["D4"] = "Charge", "Amount", "Currency"
    local["B5"], local["C5"], local["D5"] = "Destination service", 95, "USD"
    local["E5"] = "=C5*1.08"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def multi_sheet_mapping_workbook() -> bytes:
    wb = Workbook()
    ocean = wb.active
    ocean.title = "Ocean Freight"
    ocean.append(["Service", "Qty", "Unit", "Cost", "Currency", "Internal note"])
    ocean.append(["Sea freight 40HQ", 1, "container", 987654.32, "USD", "PRIVATE-COST-NOTE-OCEAN"])
    port = wb.create_sheet("Port charges")
    port.append(["Service", "Qty", "Unit", "Cost", "Currency", "Internal note"])
    port.append(["THC", 1, "container", 876543.21, "USD", "PRIVATE-COST-NOTE-THC"])
    port.append(["D/O", 1, "shipment", 765432.10, "USD", "PRIVATE-COST-NOTE-DO"])
    goods = wb.create_sheet("Customer goods tab")
    goods.append(["Description", "Currency", "Supplier cost"])
    goods.append(["Cabinet", "CNY", 654321.98])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _base36(value: int) -> str:
    alphabet = "0123456789abcdefghijklmnopqrstuvwxyz"
    if value == 0:
        return "0"
    digits = []
    while value:
        value, remainder = divmod(value, 36)
        digits.append(alphabet[remainder])
    return "".join(reversed(digits))


def _mapped_row_id(source: dict, sheet: str, row: int) -> str:
    identity = f"{source['sha256']}\0{sheet}\0{row}"
    first, second = 0x811C9DC5, 0x9E3779B9
    for index, char in enumerate(identity):
        code = ord(char)
        first = ((first ^ code) * 0x01000193) & 0xFFFFFFFF
        second = ((second ^ (code + index)) * 0x85EBCA6B) & 0xFFFFFFFF
    return f"map-{_base36(first)}-{_base36(second)}"


def _page_value(page: dict, row: int, column: int):
    cell = next((item for item in page.get("cells", []) if item["r"] == row and item["c"] == column), None)
    if not cell or (cell.get("kind") == "formula" and cell.get("unknown")):
        return None
    return cell.get("value")


def mapped_service_draft(source: dict, page: dict, rows: list[int]) -> dict:
    """Fixture equivalent of an explicit SERVICE/COST column map, not parser output."""
    lines = []
    for row in rows:
        name = _page_value(page, row, 1)
        if not name:
            continue
        lines.append({
            "id": _mapped_row_id(source, page["sheet"], row), "service_name": str(name), "code": "",
            "unit": str(_page_value(page, row, 3) or ""),
            "qty": str(_page_value(page, row, 2)) if _page_value(page, row, 2) is not None else None,
            "currency": str(_page_value(page, row, 5) or "").upper(),
            "cost_unit": str(_page_value(page, row, 4)) if _page_value(page, row, 4) is not None else None,
            "sell_unit": None, "sell_confirmed": False, "tax_pct": None,
            "internal_note": str(_page_value(page, row, 6) or ""), "public_note": "",
            "cost_source": f"INTERNAL-COST-SOURCE|{page['sheet']}|row {row}",
            "valid_until": "", "validity_status": "UNKNOWN", "source_warning": "",
            "condition": "UNCONFIRMED",
        })
    return excel_mapping_draft(source, service_lines=lines)


def mapped_goods_draft(source: dict, page: dict, row: int) -> dict:
    return excel_mapping_draft(source, goods=[{
        "id": _mapped_row_id(source, page["sheet"], row),
        "description": str(_page_value(page, row, 1) or ""), "model": "", "material": "",
        "length_mm": None, "width_mm": None, "height_mm": None, "qty": None, "unit_price": None,
        "currency": str(_page_value(page, row, 2) or "").upper(), "note": "",
        "source_sheet": page["sheet"], "source_row": row,
    }])


def excel_mapping_draft(source: dict, *, goods: list[dict] | None = None,
                        service_lines: list[dict] | None = None) -> dict:
    items = list(goods or [])
    lines = list(service_lines or [])
    return {
        "schema": "luc-giac/1", "goods_mode": "GOODS" if items else "SERVICE_ONLY", "goods": items,
        "supplier": "", "source": source, "source_warnings": [], "source_facts": {}, "source_costs": [],
        "disclosures": [], "additional_public_notes": "", "active_option": "A",
        "options": {"A": {"scope": "", "service_lines": lines}, "B": {"scope": "", "service_lines": []}},
        "shipment": {"cbm": None, "weight_kg": None, "pol": "", "pod": "", "pickup": "", "delivery": "", "load_model": "", "equipment": ""},
        "customs": {"confirmed": False, "basis_vnd": None, "basis_note": "", "duty_pct": None, "vat_pct": None, "hs_code": ""},
        "warning_notes": {},
    }


def _upsert_mapped_rows(existing: list[dict], incoming: list[dict]) -> list[dict]:
    rows = {}
    for index, item in enumerate(existing):
        rows[item.get("id") or f"existing:{index}"] = item
    for index, item in enumerate(incoming):
        rows[item.get("id") or f"incoming:{index}"] = item
    return list(rows.values())


def merge_mapping_for_dispatch(current: dict | None, mapped: dict) -> dict:
    """Mirror the UI row-upsert contract while exercising actual C1 writes below."""
    if not current or (current.get("source") or {}).get("sha256") != mapped["source"]["sha256"]:
        return mapped
    goods_in = mapped.get("goods") or []
    services_in = (mapped.get("options") or {}).get("A", {}).get("service_lines") or []
    incoming_ids = {item.get("id") for item in [*goods_in, *services_in] if item.get("id")}
    active = "B" if current.get("active_option") == "B" else "A"
    goods = _upsert_mapped_rows(current.get("goods") or [], goods_in)
    if services_in:
        goods = [item for item in goods if item.get("id") not in incoming_ids]
    options = {}
    for key in ("A", "B"):
        option = dict((current.get("options") or {}).get(key) or {"scope": "", "service_lines": []})
        lines = list(option.get("service_lines") or [])
        if goods_in:
            lines = [item for item in lines if item.get("id") not in incoming_ids]
        if services_in and key == active:
            lines = _upsert_mapped_rows(lines, services_in)
        options[key] = {**option, "service_lines": lines}
    return {**current, "schema": mapped["schema"], "source": mapped["source"], "goods": goods,
            "goods_mode": "GOODS" if goods else "SERVICE_ONLY", "active_option": active, "options": options}


def workbook_quote_payload(details: dict) -> dict:
    return {
        "customerId": CUSTOMER_ID, "salesOwner": "Sale test", "goodsDescription": "Cabinet",
        "origin": "Ningbo", "destination": "Hai Phong", "route": "Ningbo → Hai Phong",
        "incoterm": "FOB", "transportMode": "Sea", "fx": {}, "excelQuote": details,
    }


@pytest.fixture
def staging():
    import streamlit as st
    from core.database import get_db

    had_user = "current_user" in st.session_state
    previous_user = st.session_state.get("current_user")
    logi_sandbox.reset_backend()
    assert logi_sandbox.ensure_backend() == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    db = get_db()
    db.customers.insert_one({"id": CUSTOMER_ID, "name": "Freight customer"})
    try:
        yield db
    finally:
        logi_sandbox.reset_backend()
        if had_user:
            st.session_state["current_user"] = previous_user
        else:
            st.session_state.pop("current_user", None)


def dispatch(event: str, payload: dict):
    outcomes = logi_c1.consume_c1_triggers(
        {"seq": 0}, {event: {"seq": 1, "payload": {"requestId": "workbook-test", **payload}}}
    )
    assert len(outcomes) == 1
    return outcomes[0]


def import_generic(data: bytes):
    return dispatch("quote_import_excel", {
        "customerId": CUSTOMER_ID, "fileName": FILE_NAME,
        "base64": base64.b64encode(data).decode("ascii"),
    })


def test_non_template_workbook_without_quote_columns_uploads_views_edits_saves_and_reopens(staging):
    original = freight_workbook()
    original_sha = hashlib.sha256(original).hexdigest()
    source_book = load_workbook(BytesIO(original), data_only=False, read_only=True)
    source_cells = [str(value) for row in source_book["Ocean Freight"].iter_rows(values_only=True) for value in row if value is not None]
    source_book.close()
    assert not any(value.strip().casefold() in {"kích thước", "size", "số lượng", "quantity", "đơn giá rmb", "price rmb"}
                   for value in source_cells)
    imported = import_generic(original)
    assert imported["ok"], imported
    assert imported["readback"]["workbook_only"] is True
    assert imported["readback"]["workbook_message"] == (
        "Đã nhập Excel thành công. Chưa xác định được bảng để tính giá tự động; "
        "bạn có thể xem, sửa và chọn các cột cần dùng."
    )
    source = imported["readback"]["source"]
    assert source["sheet_names"] == ["Ocean Freight", "Port charges"]
    assert source["sha256"] == original_sha
    listed = dispatch("quote_workbook_sources_list", {"customerId": CUSTOMER_ID})
    assert listed["ok"], listed
    assert any(item["sha256"] == original_sha for item in listed["readback"]["sources"])

    store = GridFSAttachmentStore(staging)
    assert store.download_file(source["storage_file_id"]) == original
    page = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight",
    })
    assert page["ok"], page
    assert page["readback"]["overview"]["sheet_names"] == ["Ocean Freight", "Port charges"]
    cells = {(cell["r"], cell["c"]): cell for cell in page["readback"]["page"]["cells"]}
    assert cells[(2, 2)]["value"] == "FREIGHT QUOTATION · NINGBO TO HAI PHONG"
    assert cells[(4, 4)]["value"] == 1750
    assert cells[(5, 4)]["kind"] == "formula" and cells[(5, 4)]["f"] == "=D4*1.05"
    assert cells[(5, 4)]["unknown"] is True

    saved = dispatch("quote_workbook_draft_save", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight", "expectedVersion": 0,
        "edits": [
            {"sheet": "Ocean Freight", "r": 4, "c": 2, "type": "text", "value": "Ningbo → Hai Phong · checked"},
            {"sheet": "Ocean Freight", "r": 4, "c": 4, "type": "number", "value": "1900"},
        ],
    })
    assert saved["ok"], saved
    assert saved["readback"]["draft"]["saved"] is True
    assert saved["readback"]["draft"]["version"] == 1
    after_save = {(cell["r"], cell["c"]): cell for cell in saved["readback"]["page"]["cells"]}
    assert after_save[(4, 2)]["value"] == "Ningbo → Hai Phong · checked"
    assert after_save[(4, 4)]["value"] == 1900
    assert after_save[(5, 4)]["kind"] == "formula" and after_save[(5, 4)]["f"] == "=D4*1.05"
    assert store.download_file(source["storage_file_id"]) == original
    assert hashlib.sha256(store.download_file(source["storage_file_id"])).hexdigest() == original_sha

    reopened = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight",
    })
    assert reopened["ok"], reopened
    assert reopened["readback"]["draft"]["version"] == 1
    reopened_cells = {(cell["r"], cell["c"]): cell for cell in reopened["readback"]["page"]["cells"]}
    assert reopened_cells[(4, 2)]["value"] == "Ningbo → Hai Phong · checked"
    assert reopened_cells[(4, 4)]["value"] == 1900
    # openpyxl preserves the user formula source but deliberately does not evaluate it.
    drafts = list(store.bucket.find({"metadata.owner_type": "QUOTE_EXCEL_DRAFT", "metadata.source_file_id": source["storage_file_id"]}))
    assert len(drafts) == 1
    edited_bytes = store.download_file(str(drafts[0]._id))
    edited_wb = load_workbook(BytesIO(edited_bytes), data_only=False, keep_links=False)
    assert edited_wb["Ocean Freight"]["B4"].value == "Ningbo → Hai Phong · checked"
    assert edited_wb["Ocean Freight"]["D4"].value == 1900
    assert edited_wb["Ocean Freight"]["D5"].value == "=D4*1.05"
    cached_wb = load_workbook(BytesIO(edited_bytes), data_only=True, keep_links=False)
    assert cached_wb["Ocean Freight"]["D5"].value is None

    stale = dispatch("quote_workbook_draft_save", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight", "expectedVersion": 0,
        "edits": [{"sheet": "Ocean Freight", "r": 4, "c": 4, "type": "number", "value": "2000"}],
    })
    assert stale["ok"] is False and "đổi phiên bản" in stale["error"]

    restored = dispatch("quote_workbook_draft_save", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight", "expectedVersion": 1,
        "edits": [], "restoreOriginal": True,
    })
    assert restored["ok"], restored
    assert restored["readback"]["draft"]["version"] == 2
    restored_cells = {(cell["r"], cell["c"]): cell for cell in restored["readback"]["page"]["cells"]}
    assert restored_cells[(4, 4)]["value"] == 1750
    assert store.download_file(source["storage_file_id"]) == original


def test_workbook_cells_and_draft_are_grant_gated_and_customer_owned(staging):
    import streamlit as st

    original = freight_workbook()
    imported = import_generic(original)
    assert imported["ok"], imported
    source = imported["readback"]["source"]

    st.session_state["current_user"] = {"username": "viewer-workbook", "role": "viewer"}
    listed = dispatch("quote_workbook_sources_list", {"customerId": CUSTOMER_ID})
    assert listed["ok"] is False and "commercial_write" in listed["error"]
    locked = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight",
    })
    assert locked["ok"]
    assert locked["readback"]["page"]["locked"] is True
    assert locked["readback"]["page"]["cells"] == []
    refused = dispatch("quote_workbook_draft_save", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight", "expectedVersion": 0,
        "edits": [{"sheet": "Ocean Freight", "r": 4, "c": 4, "type": "number", "value": "1900"}],
    })
    assert refused["ok"] is False and "commercial_write" in refused["error"]

    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    wrong_customer = dispatch("quote_workbook_draft_save", {
        "customerId": "CUST-OTHER-WORKBOOK", "source": source, "sheet": "Ocean Freight", "expectedVersion": 0,
        "edits": [],
    })
    assert wrong_customer["ok"] is False and "khách hàng" in wrong_customer["error"]


def test_workbook_edit_rejects_formula_injection_and_merged_followers(staging):
    from core.quote_workbook_view import apply_cell_edits

    original = freight_workbook()
    with pytest.raises(ValueError, match="Công thức"):
        apply_cell_edits(original, [{"sheet": "Ocean Freight", "r": 3, "c": 2, "type": "text", "value": "=1+1"}])
    with pytest.raises(ValueError, match="gộp"):
        apply_cell_edits(original, [{"sheet": "Ocean Freight", "r": 2, "c": 3, "type": "text", "value": "not an anchor"}])
    with pytest.raises(ValueError, match="15 chữ số"):
        apply_cell_edits(original, [{"sheet": "Ocean Freight", "r": 4, "c": 4, "type": "number", "value": "1234567890123456"}])


def test_service_only_quote_draft_is_valid_and_customer_export_excludes_cost_and_source(staging):
    original = freight_workbook()
    imported = import_generic(original)
    assert imported["ok"], imported
    source = imported["readback"]["source"]
    details = {
        "schema": "luc-giac/1", "goods_mode": "SERVICE_ONLY", "goods": [], "supplier": "",
        "source": source, "source_warnings": [], "source_facts": {}, "source_costs": [], "disclosures": [],
        "additional_public_notes": "", "active_option": "A",
        "options": {"A": {"scope": "Ocean freight only", "service_lines": [{
            "id": "mapped-freight", "service_name": "Ocean freight 40HQ", "code": "", "unit": "container",
            "qty": 1, "currency": "USD", "cost_unit": 250, "sell_unit": 300, "sell_confirmed": True,
            "tax_pct": None, "internal_note": "PRIVATE-COST-NOTE", "public_note": "",
            "cost_source": "workbook mapping · Ocean Freight row 4", "valid_until": "", "validity_status": "UNKNOWN",
            "source_warning": "", "condition": "UNCONFIRMED",
        }]}, "B": {"scope": "", "service_lines": []}},
        "shipment": {"cbm": None, "weight_kg": None, "pol": "Ningbo", "pod": "Hai Phong", "pickup": "", "delivery": "", "load_model": "FCL", "equipment": "40HQ"},
        "customs": {"confirmed": False, "basis_vnd": None, "basis_note": "", "duty_pct": None, "vat_pct": None, "hs_code": ""},
        "warning_notes": {},
    }
    created = dispatch("quote_create", {
        "customerId": CUSTOMER_ID, "salesOwner": "Sale test", "goodsDescription": "",
        "origin": "Ningbo", "destination": "Hai Phong", "route": "Ningbo → Hai Phong",
        "incoterm": "FOB", "transportMode": "Sea", "fx": {}, "excelQuote": details,
    })
    assert created["ok"], created
    quote = created["readback"]
    assert quote["excelQuote"]["goods_mode"] == "SERVICE_ONLY"
    assert quote["excelQuote"]["goods"] == []
    assert quote["excelQuote"]["calculation"]["row_count"] == 0
    assert quote["excelQuote"]["calculation"]["quantity"] is None
    assert quote["excelQuote"]["calculation"]["estimate_vnd"] is None

    exported = dispatch("quote_export_xlsx", {"id": created["id"]})
    assert exported["ok"], exported
    workbook = load_workbook(BytesIO(base64.b64decode(exported["document"]["downloadBase64"])), data_only=False)
    all_text = "\n".join(str(cell.value or "") for sheet in workbook for row in sheet for cell in row)
    assert "Báo giá chỉ gồm phí dịch vụ; không có hàng hóa." in all_text
    assert "Ocean freight 40HQ" in all_text and "300" in all_text
    assert "250" not in all_text and "PRIVATE-COST-NOTE" not in all_text
    assert "FREIGHT QUOTATION" not in all_text and "Port charges" not in all_text
    assert "DANH MỤC HÀNG HÓA" not in all_text and "THUẾ NHẬP KHẨU DỰ KIẾN" not in all_text


def test_generic_workbook_with_table_far_down_and_right_can_be_opened_and_edited(staging):
    workbook = load_workbook(BytesIO(freight_workbook()))
    rates = workbook["Ocean Freight"]
    for index in range(7):
        workbook.create_sheet(f"Extra rate sheet {index + 1}")
    rates["AZ350"] = "Far table data"
    rates["AZ351"] = 987.5
    output = BytesIO()
    workbook.save(output)
    original = output.getvalue()

    imported = import_generic(original)
    assert imported["ok"], imported
    assert imported["readback"]["workbook_only"] is True
    source = imported["readback"]["source"]
    store = GridFSAttachmentStore(staging)
    assert store.download_file(source["storage_file_id"]) == original

    view = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight", "startRow": 340, "startCol": 48,
    })
    assert view["ok"], view
    page = view["readback"]["page"]
    assert len(view["readback"]["overview"]["sheet_names"]) == 9
    assert page["start_row"] == 340 and page["start_col"] == 48
    assert any(cell["r"] == 350 and cell["c"] == 52 and cell["value"] == "Far table data" for cell in page["cells"])

    saved = dispatch("quote_workbook_draft_save", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight",
        "startRow": 340, "startCol": 48, "expectedVersion": 0,
        "edits": [{"sheet": "Ocean Freight", "r": 351, "c": 52, "type": "number", "value": "1000"}],
    })
    assert saved["ok"], saved
    far = next(cell for cell in saved["readback"]["page"]["cells"] if cell["r"] == 351 and cell["c"] == 52)
    assert far["value"] == 1000
    assert store.download_file(source["storage_file_id"]) == original


def test_mixed_multisheet_cny_usd_formula_workbook_is_stored_and_kept_unmapped(staging):
    original = mixed_customer_workbook()
    original_sha = hashlib.sha256(original).hexdigest()
    imported = import_generic(original)
    assert imported["ok"], imported
    assert imported["readback"]["workbook_only"] is True
    source = imported["readback"]["source"]
    assert source["sheet_names"] == ["客户说明", "Customer goods tab", "Local charges"]
    assert source["sha256"] == original_sha
    assert not imported["readback"].get("goods")

    store = GridFSAttachmentStore(staging)
    assert store.download_file(source["storage_file_id"]) == original
    customer_page = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Customer goods tab",
    })
    assert customer_page["ok"], customer_page
    page = customer_page["readback"]["page"]
    cells = {(cell["r"], cell["c"]): cell for cell in page["cells"]}
    assert cells[(6, 3)]["value"] == "SKU"
    assert cells[(7, 6)]["value"] == "CNY"
    assert cells[(7, 7)]["value"] == 870
    assert customer_page["readback"]["overview"]["sheet_names"] == source["sheet_names"]

    charges = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Local charges",
    })
    assert charges["ok"], charges
    formula = next(cell for cell in charges["readback"]["page"]["cells"] if cell["r"] == 5 and cell["c"] == 5)
    assert formula["kind"] == "formula" and formula["f"] == "=C5*1.08" and formula["unknown"] is True
    assert charges["readback"]["overview"]["sheets"][1]["not_rendered"].get("comments") == 1
    assert store.download_file(source["storage_file_id"]) == original


def test_multisheet_mapping_upserts_through_c1_staging_save_reopen_preview_and_xlsx(staging):
    original = multi_sheet_mapping_workbook()
    imported = import_generic(original)
    assert imported["ok"], imported
    assert imported["readback"]["workbook_only"] is True
    source = imported["readback"]["source"]
    assert source["parser_version"] == "unmapped/1"
    assert source["sheet_names"] == ["Ocean Freight", "Port charges", "Customer goods tab"]
    store = GridFSAttachmentStore(staging)
    assert store.download_file(source["storage_file_id"]) == original

    ocean_view = dispatch("quote_source_workbook_view", {
        "customerId": CUSTOMER_ID, "source": source, "sheet": "Ocean Freight",
    })
    assert ocean_view["ok"], ocean_view
    ocean_page = ocean_view["readback"]["page"]
    ocean_mapping = mapped_service_draft(source, ocean_page, [2])
    ocean_line = ocean_mapping["options"]["A"]["service_lines"][0]
    assert ocean_line["service_name"] == "Sea freight 40HQ"
    assert ocean_line["qty"] == "1" and ocean_line["cost_unit"] == "987654.32"
    assert ocean_line["sell_unit"] is None and ocean_line["sell_confirmed"] is False
    assert ocean_line["tax_pct"] is None
    assert ocean_mapping["goods_mode"] == "SERVICE_ONLY" and ocean_mapping["goods"] == []

    # A real C1 preview confirms the mapper has not invented a customer price,
    # tax or FX rate before the operator explicitly confirms anything.
    preview = dispatch("quote_preview", workbook_quote_payload(ocean_mapping))
    assert preview["ok"], preview
    service_only_customer = preview["readback"]["customer"]
    assert service_only_customer["goods_mode"] == "SERVICE_ONLY"
    assert service_only_customer["service_lines"] == []
    assert service_only_customer["goods_calculation"]["quantity"] is None
    assert service_only_customer["fx_snapshot"]["rates_to_vnd"] == {"VND": 1.0}

    created = dispatch("quote_create", workbook_quote_payload(ocean_mapping))
    assert created["ok"], created
    quote_id = created["id"]
    saved = logi_c1.read_quote(quote_id)
    assert saved and saved["version"] == 1
    assert saved["excel_quote"]["goods_mode"] == "SERVICE_ONLY"
    assert len(saved["excel_quote"]["options"]["A"]["service_lines"]) == 1

    # Saved-quote source view is the same path the workspace uses for a valid
    # quote_edit; adding the overlapping Port charges range must retain Ocean.
    port_view = dispatch("quote_source_workbook_view", {"id": quote_id, "sheet": "Port charges"})
    assert port_view["ok"], port_view
    port_page = port_view["readback"]["page"]
    port_first = mapped_service_draft(source, port_page, [2])
    first_id = port_first["options"]["A"]["service_lines"][0]["id"]
    combined = merge_mapping_for_dispatch(saved["excel_quote"], port_first)
    edited = dispatch("quote_edit", {
        **workbook_quote_payload(combined), "id": quote_id, "expectedVersion": saved["version"],
    })
    assert edited["ok"], edited
    saved = logi_c1.read_quote(quote_id)
    assert saved and saved["version"] == 2
    assert saved["excel_quote"]["goods_mode"] == "SERVICE_ONLY"
    assert [line["service_name"] for line in saved["excel_quote"]["options"]["A"]["service_lines"]] == [
        "Sea freight 40HQ", "THC",
    ]

    # Extending the range to row 3 overlaps THC and adds D/O: row identity is
    # source SHA + sheet + source row, so THC is replaced rather than duplicated.
    port_full = mapped_service_draft(source, port_page, [2, 3])
    assert port_full["options"]["A"]["service_lines"][0]["id"] == first_id
    combined = merge_mapping_for_dispatch(saved["excel_quote"], port_full)
    edited = dispatch("quote_edit", {
        **workbook_quote_payload(combined), "id": quote_id, "expectedVersion": saved["version"],
    })
    assert edited["ok"], edited
    saved = logi_c1.read_quote(quote_id)
    names = [line["service_name"] for line in saved["excel_quote"]["options"]["A"]["service_lines"]]
    assert names == ["Sea freight 40HQ", "THC", "D/O"]
    assert len(names) == 3 and saved["excel_quote"]["goods_mode"] == "SERVICE_ONLY"

    reapplied = deepcopy(port_full)
    reapplied["options"]["A"]["service_lines"][0]["cost_unit"] = "876543.22"
    reapplied["options"]["A"]["service_lines"][0]["cost_source"] = "INTERNAL-COST-SOURCE|Port charges|updated row 2"
    combined = merge_mapping_for_dispatch(saved["excel_quote"], reapplied)
    edited = dispatch("quote_edit", {
        **workbook_quote_payload(combined), "id": quote_id, "expectedVersion": saved["version"],
    })
    assert edited["ok"], edited
    saved = logi_c1.read_quote(quote_id)
    service_lines = saved["excel_quote"]["options"]["A"]["service_lines"]
    assert len(service_lines) == 3
    assert [line["cost_unit"] for line in service_lines] == [987654.32, 876543.22, 765432.1]

    # Map one goods row from a third sheet with only name and currency selected.
    # Its source qty and supplier-cost columns are intentionally not mapped.
    goods_view = dispatch("quote_source_workbook_view", {"id": quote_id, "sheet": "Customer goods tab"})
    assert goods_view["ok"], goods_view
    goods_mapping = mapped_goods_draft(source, goods_view["readback"]["page"], 2)
    assert goods_mapping["goods"][0]["qty"] is None
    assert goods_mapping["goods"][0]["unit_price"] is None
    combined = merge_mapping_for_dispatch(saved["excel_quote"], goods_mapping)
    edited = dispatch("quote_edit", {
        **workbook_quote_payload(combined), "id": quote_id, "expectedVersion": saved["version"],
    })
    assert edited["ok"], edited
    saved = logi_c1.read_quote(quote_id)
    assert saved and saved["excel_quote"]["goods_mode"] == "GOODS"
    assert len(saved["excel_quote"]["goods"]) == 1
    assert saved["excel_quote"]["goods"][0]["description"] == "Cabinet"
    assert saved["excel_quote"]["goods"][0]["qty"] is None
    assert saved["excel_quote"]["goods"][0]["unit_price"] is None
    assert len(saved["excel_quote"]["options"]["A"]["service_lines"]) == 3

    # Explicitly emulate the operator confirming three selling prices; no mapper
    # value becomes a customer price until this separate action is supplied.
    confirmed = deepcopy(saved["excel_quote"])
    for line, sell in zip(confirmed["options"]["A"]["service_lines"], (3000, 350, 75)):
        line["sell_unit"] = sell
        line["sell_confirmed"] = True
        line["tax_pct"] = None
        line["internal_note"] = f"PRIVATE-COST-NOTE|{line['service_name']}"
    payload = workbook_quote_payload(confirmed)
    payload.update(id=quote_id, expectedVersion=saved["version"])
    edited = dispatch("quote_edit", payload)
    assert edited["ok"], edited
    saved = logi_c1.read_quote(quote_id)
    assert saved and saved["version"] == 6
    assert saved["excel_quote"]["goods_mode"] == "GOODS"

    public_preview = dispatch("quote_preview", workbook_quote_payload(saved["excel_quote"]))
    assert public_preview["ok"], public_preview
    customer = public_preview["readback"]["customer"]
    assert customer["goods_mode"] == "GOODS"
    assert [line["service_name"] for line in customer["service_lines"]] == [
        "Sea freight 40HQ", "THC", "D/O",
    ]
    assert [line["sell_unit"] for line in customer["service_lines"]] == [3000, 350, 75]
    assert all(line["tax_pct"] is None and line["tax_total"] is None for line in customer["service_lines"])
    assert customer["goods"][0]["qty"] is None
    assert customer["goods"][0]["unit_price"] is None
    assert customer["goods"][0]["line_total"] is None
    assert customer["goods_calculation"]["quantity"] is None
    assert customer["goods_calculation"]["estimate_vnd"] is None
    assert customer["fx_snapshot"]["rates_to_vnd"] == {"VND": 1.0}

    def keys(value):
        if isinstance(value, dict):
            for key, nested in value.items():
                yield key
                yield from keys(nested)
        elif isinstance(value, list):
            for nested in value:
                yield from keys(nested)

    public_json = json.dumps(customer, ensure_ascii=False)
    public_keys = set(keys(customer))
    assert not ({"cost_unit", "cost_source", "internal_note", "source_sheet", "source_row",
                 "source_facts", "source_costs", "storage_file_id", "sheet_names"} & public_keys)
    for internal in ("PRIVATE-COST-NOTE", "INTERNAL-COST-SOURCE", "multi-sheet-mapping.xlsx",
                     "Ocean Freight", "Port charges", "Customer goods tab"):
        assert internal.casefold() not in public_json.casefold()

    exported = dispatch("quote_export_xlsx", {"id": quote_id})
    assert exported["ok"], exported
    customer_xlsx = load_workbook(BytesIO(base64.b64decode(exported["document"]["downloadBase64"])), data_only=False)
    values = [cell.value for sheet in customer_xlsx for row in sheet.iter_rows() for cell in row if cell.value is not None]
    all_text = "\n".join(str(value) for value in values)
    assert all(name in all_text for name in ("Sea freight 40HQ", "THC", "D/O"))
    assert any(value == 3000 for value in values) and any(value == 350 for value in values) and any(value == 75 for value in values)
    assert all(secret not in all_text for secret in (
        "987654.32", "876543.22", "765432.1", "654321.98", "PRIVATE-COST-NOTE",
        "INTERNAL-COST-SOURCE", "multi-sheet-mapping.xlsx", "Ocean Freight", "Port charges",
        "Customer goods tab", "cost_unit", "cost_source",
    ))
    assert store.download_file(source["storage_file_id"]) == original
