"""LOGI_QUICK_QUOTE_LINE_INTEGRITY_AND_SERVICE_NAME_01

The quotation stage stores the operator's service wording. A pasted line may
have a blank code. A historical line keeps the code it already has. Nothing
in this path rewrites a name to "Dịch vụ khác" or a code to DVK.

Isolated staging transport only. Production MongoDB is not contacted.
"""

from __future__ import annotations

import base64
import io

import pytest
from openpyxl import load_workbook

from views import logi_c1, logi_sandbox

GENERAL_NOTE = "Nếu có phát sinh về xử lý hàng hóa, sẽ báo ngay tại thời điểm phát sinh."
EXPECTED = [
    "Cước biển",
    "Phụ phí hãng tàu",
    "Nâng cont, hạ vỏ",
    "Hạ tầng",
    "Thủ tục hải quan",
    "Vận tải về Sơn Tây, Hà Nội",
    "Kiểm hóa mở cont",
    "Kiểm hóa soi chiếu",
    "Công bố hợp quy hàng thuộc Bộ Công An",
]


@pytest.fixture()
def staging():
    import streamlit as st

    logi_sandbox.reset_backend()
    transport = logi_sandbox.ensure_backend()
    assert transport == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    yield transport
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    logi_sandbox.reset_backend()


@pytest.fixture()
def event_seq():
    state = {"n": 0}

    def _make(name: str, payload: dict) -> dict:
        state["n"] += 1
        return {name: {"seq": state["n"], "event": name, "payload": payload}}

    return _make


def _line(name: str, sell: int, currency: str = "VND", code: str = "") -> dict:
    return {
        "code": code,
        "serviceName": name,
        "unit": "cont 20'",
        "qty": 1,
        "currency": currency,
        "sellUnit": sell,
        "costUnit": "",
        "taxPct": "",
        "note": "",
    }


def _dispatch(seq_fn, name, payload):
    outcomes = logi_c1.consume_c1_triggers({"seq": 0}, seq_fn(name, payload))
    assert len(outcomes) == 1
    return outcomes[0]


def test_pasted_nine_lines_round_trip_without_service_code(staging, event_seq):
    lines = [
        _line(name, sell, currency)
        for name, sell, currency in [
            ("Cước biển", 950, "USD"),
            ("Phụ phí hãng tàu", 9_500_000, "VND"),
            ("Nâng cont, hạ vỏ", 2_600_000, "VND"),
            ("Hạ tầng", 250_000, "VND"),
            ("Thủ tục hải quan", 2_000_000, "VND"),
            ("Vận tải về Sơn Tây, Hà Nội", 6_000_000, "VND"),
            ("Kiểm hóa mở cont", 4_000_000, "VND"),
            ("Kiểm hóa soi chiếu", 2_000_000, "VND"),
            ("Công bố hợp quy hàng thuộc Bộ Công An", 10_000_000, "VND"),
        ]
    ]
    out = _dispatch(event_seq, "quote_create", {
        "customerId": "CUST-ANPHAT",
        "salesOwner": "Minh Tú",
        "origin": "NINGBO",
        "destination": "Hải Phòng",
        "route": "",
        "incoterm": "FOB",
        "transportMode": "Sea",
        "goodsDescription": "",
        "expectedDeliveryDate": "",
        "serviceLines": lines,
        "customerNotes": [GENERAL_NOTE],
    })
    assert out["ok"], out
    assert len(out["readback"]["serviceLines"]) == 9

    stored = logi_c1.read_quote(out["id"])
    assert stored is not None
    assert len(stored["service_lines"]) == 9
    assert [line["service_name"] for line in stored["service_lines"]] == EXPECTED
    assert [line.get("code") or "" for line in stored["service_lines"]] == [""] * 9
    assert all("Dịch vụ khác" not in line["service_name"] for line in stored["service_lines"])
    assert stored["customer_notes"] == [GENERAL_NOTE]
    assert not any("xử lý hàng hóa" in line["service_name"] for line in stored["service_lines"])

    projected = logi_c1.quote_projection(stored)
    assert [line["serviceName"] for line in projected["serviceLines"]] == EXPECTED
    assert [line["code"] for line in projected["serviceLines"]] == [""] * 9

    exported = _dispatch(event_seq, "quote_export_xlsx", {"id": out["id"], "requestId": "rq-9"})
    assert exported["ok"], exported
    book = load_workbook(io.BytesIO(base64.b64decode(exported["document"]["downloadBase64"])))
    flat = "\n".join(
        " ".join(str(cell) for cell in row if cell is not None)
        for row in book.active.iter_rows(values_only=True)
    )
    for name in EXPECTED:
        assert name in flat, f"XLSX Chi tiết lost {name}"
    assert "Dịch vụ khác" not in flat
    assert GENERAL_NOTE in flat


def test_historical_code_is_preserved_and_not_rewritten_to_dvk(staging, event_seq):
    created = _dispatch(event_seq, "quote_create", {
        "customerId": "CUST-ANPHAT",
        "salesOwner": "Minh Tú",
        "origin": "NINGBO",
        "destination": "Hải Phòng",
        "incoterm": "FOB",
        "serviceLines": [_line("Cước biển", 950, "USD", code="DV-CUOC-BIEN")],
    })
    assert created["ok"], created
    edited = _dispatch(event_seq, "quote_edit", {
        "id": created["id"],
        "serviceLines": [_line("Cước biển", 970, "USD", code="DV-CUOC-BIEN")],
        "customerNotes": [],
    })
    assert edited["ok"], edited
    stored = logi_c1.read_quote(created["id"])
    assert stored["service_lines"][0]["code"] == "DV-CUOC-BIEN"
    assert stored["service_lines"][0]["code"] != "DVK"
    assert stored["service_lines"][0]["service_name"] == "Cước biển"
    assert "Dịch vụ khác" not in stored["service_lines"][0]["service_name"]
