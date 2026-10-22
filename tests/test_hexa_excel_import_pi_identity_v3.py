"""HEXA-EXCEL-IMPORT-V0-PI-IDENTITY-FIX — the real PI code column must survive.

`REAL_WORKBOOK_TEST=NOT_RUN_IN_ARENA`: the Owner's workbook `PL1(1).xlsx` is not
in this repository. The fixture reproduces the STRUCTURE measured on it:

    sheet `PI (2)`, header row 7
      B `Mã`               -> 5 distinct product codes (rows 8-12)
      C `Model 型号`       -> the product NAMES
      E `Size尺寸(MM)` · F `Qt.数量` · G `Unit Price 单价 (CNY)` · H `Amount 金额 (CNY)`
    two of the five rows carry the SAME name (`ghế`) with DIFFERENT codes
    a second sheet `File tiền hàng` repeats the money table internally

Before the repair the parser read `Model 型号` as the model, never mapped `Mã`,
and therefore lost all five source codes while leaving `description` empty.
"""
from __future__ import annotations

import json
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
import pytest

from core.excel_utils import generate_quote_excel
from core.quote_excel_import import parse_luc_giac_workbook
from core.quotations import calculate_quote, customer_quote_snapshot
from scripts.make_hexa_layout_fixtures import (
    PI_IDENTITY_HEADER_ROW,
    PI_IDENTITY_HEADERS,
    PI_IDENTITY_MONEY_SHEET,
    PI_IDENTITY_ROWS,
    PI_IDENTITY_SHEET,
    SENTINEL,
    SHEET_CARTONS,
    SHEET_TOTAL,
    build_multilevel_packing_real_layout,
    build_packing_list_real_headers,
    build_proforma_invoice_pi_identity,
)

TODAY = date(2026, 10, 16)


@pytest.fixture(scope="module")
def pi_identity() -> bytes:
    return build_proforma_invoice_pi_identity()


def parse(data: bytes, file_name: str) -> dict:
    return parse_luc_giac_workbook(data, file_name, today=TODAY)


def warning_ids(proposal: dict) -> set[str]:
    return {w["id"] for w in proposal["source_warnings"]}


def unknown_codes(proposal: dict) -> set[str]:
    return {u["code"] for u in proposal["review"]["unknowns"]}


# ------------------------------------------------------------- fixture reality
def test_fixture_has_a_code_column_beside_a_name_column(pi_identity):
    wb = load_workbook(BytesIO(pi_identity))
    assert wb.sheetnames == [PI_IDENTITY_SHEET, PI_IDENTITY_MONEY_SHEET]
    sheet = wb[PI_IDENTITY_SHEET]
    assert sheet[f"B{PI_IDENTITY_HEADER_ROW}"].value == "Mã"
    assert sheet[f"C{PI_IDENTITY_HEADER_ROW}"].value == "Model 型号"
    codes = [sheet[f"B{row}"].value for row in range(PI_IDENTITY_HEADER_ROW + 1,
                                                     PI_IDENTITY_HEADER_ROW + 6)]
    names = [sheet[f"C{row}"].value for row in range(PI_IDENTITY_HEADER_ROW + 1,
                                                     PI_IDENTITY_HEADER_ROW + 6)]
    assert codes == [row[0] for row in PI_IDENTITY_ROWS]
    assert len(set(codes)) == 5
    # Two different products really do share one name.
    assert names.count("ghế") == 2
    assert len(set(names)) == 4


# ------------------------------------------------ identity preserved (the fix)
def test_all_five_codes_are_kept_in_the_identity_field(pi_identity):
    proposal = parse(pi_identity, "PL1-SYNTHETIC.xlsx")

    assert [item["model"] for item in proposal["goods"]] == [row[0] for row in PI_IDENTITY_ROWS]
    assert all(item["model"] for item in proposal["goods"])


def test_all_five_names_are_kept_in_the_description_field(pi_identity):
    proposal = parse(pi_identity, "PL1-SYNTHETIC.xlsx")

    assert [item["description"] for item in proposal["goods"]] == [row[1] for row in PI_IDENTITY_ROWS]


def test_two_products_with_the_same_name_stay_two_rows(pi_identity):
    proposal = parse(pi_identity, "PL1-SYNTHETIC.xlsx")

    assert len(proposal["goods"]) == 5
    same_name = [item for item in proposal["goods"] if item["description"] == "ghế"]
    assert len(same_name) == 2
    assert {item["model"] for item in same_name} == {"GH-001", "GH-002"}
    assert {item["source_row"] for item in same_name} == {8, 9}
    # Rows are never merged by name: quantities stay separate.
    assert sorted(item["qty"] for item in same_name) == [6.0, 10.0]


def test_source_sheet_row_and_column_evidence_are_kept(pi_identity):
    proposal = parse(pi_identity, "PL1-SYNTHETIC.xlsx")

    for item, row in zip(proposal["goods"], PI_IDENTITY_ROWS):
        assert item["source_sheet"] == PI_IDENTITY_SHEET
        assert item["source_row"] == PI_IDENTITY_HEADER_ROW + 1 + PI_IDENTITY_ROWS.index(row)
    evidence = {row["ref"]: row for row in proposal["review"]["rows"]
                if row["ref"].startswith(f"{PI_IDENTITY_SHEET}!")}
    assert evidence[f"{PI_IDENTITY_SHEET}!8"]["source_row"] == 8
    # The review evidence names the source columns the identity came from.
    assert evidence[f"{PI_IDENTITY_SHEET}!8"]["code_header"] == "Mã"
    assert evidence[f"{PI_IDENTITY_SHEET}!8"]["name_header"] == "Model 型号"
    assert evidence[f"{PI_IDENTITY_SHEET}!8"]["code"] == "GH-001"


def test_uncached_amount_formula_is_still_unknown_never_zero(pi_identity):
    proposal = parse(pi_identity, "PL1-SYNTHETIC.xlsx")

    assert [item["unit_price"] for item in proposal["goods"]] == [row[4] for row in PI_IDENTITY_ROWS]
    assert all(item["unit_price"] != 0 for item in proposal["goods"])
    assert "formula-gap" in warning_ids(proposal)
    assert all("line_total" not in item for item in proposal["goods"])


# ------------------------------------------------- internal sheet / privacy
def test_money_sheet_is_not_summed_twice_and_never_exported(pi_identity):
    proposal = parse(pi_identity, "PL1-SYNTHETIC.xlsx")

    assert proposal["source"]["goods_sheet"] == PI_IDENTITY_SHEET
    levels = {level["sheet"]: level for level in proposal["review"]["levels"]}
    assert levels[PI_IDENTITY_MONEY_SHEET]["role"] == "INTERNAL_CALC"
    assert "money-sheet-not-imported" in warning_ids(proposal)

    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["row_count"] == len(PI_IDENTITY_ROWS)
    assert calculation["quantity"] == sum(row[3] for row in PI_IDENTITY_ROWS)
    assert PI_IDENTITY_MONEY_SHEET not in json.dumps(proposal["goods"], ensure_ascii=False)

    quote = calculate_quote({"excel_quote": proposal, "status": "DRAFT",
                             "customer_name": "Công ty thử nghiệm (không có thật)"})
    snapshot = customer_quote_snapshot(quote)
    assert SENTINEL not in json.dumps(snapshot, ensure_ascii=False)
    exported = generate_quote_excel(quote).getvalue()
    assert SENTINEL.encode() not in exported
    assert PI_IDENTITY_MONEY_SHEET.encode() not in exported


# ------------------------------------------------------- ambiguous identity
def test_two_code_like_columns_are_flagged_not_guessed():
    """If both columns look like codes, keep both and ask for review."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "PI (2)"
    for column, title in {"B": "Mã", "C": "Model 型号", "F": "Qt.数量",
                          "G": "Unit Price 单价 (CNY)"}.items():
        ws[f"{column}7"] = title
    ws["B8"], ws["C8"], ws["F8"], ws["G8"] = "CODE-1", "SKU-9", 2, 50.0
    ws["B9"], ws["C9"], ws["F9"], ws["G9"] = "CODE-2", "SKU-8", 3, 60.0
    output = BytesIO()
    wb.save(output)

    proposal = parse_luc_giac_workbook(output.getvalue(), "ambiguous-pi.xlsx", today=TODAY)
    assert [item["model"] for item in proposal["goods"]] == ["CODE-1", "CODE-2"]
    # Both columns are kept for the operator, and the row is marked for review
    # instead of silently choosing one text as if it were a product name.
    assert all(item["description"] == "" for item in proposal["goods"])
    assert "pi-identity-ambiguous" in warning_ids(proposal)
    assert "identity_needs_review" in unknown_codes(proposal)
    evidence = {row["ref"] for row in proposal["review"]["rows"]}
    assert "PI (2)!8" in evidence
    kept = next(row for row in proposal["review"]["rows"] if row["ref"] == "PI (2)!8")
    assert kept["code"] == "CODE-1" and kept["name_column_text"] == "SKU-9"


# --------------------------------------------- no regression for file A and C
def test_file_a_packing_list_is_unchanged():
    proposal = parse(build_packing_list_real_headers(), "packing-list-REAL-HEADERS-SYNTHETIC.xlsx")

    assert proposal["source"]["layout"] == "PACKING_LIST"
    assert len(proposal["goods"]) == 3
    assert [item["model"] for item in proposal["goods"]] == ["BC-800", "BC-1000", "MC-600"]
    assert all(item["unit_price"] is None for item in proposal["goods"])


def test_file_c_keeps_ten_pallet_lines_and_thirty_six_carton_lines():
    proposal = parse(build_multilevel_packing_real_layout(), "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")

    assert proposal["source"]["layout"] == "MULTI_LEVEL_PACKING"
    levels = {level["sheet"]: level for level in proposal["review"]["levels"]}
    pallet_sheet = next(name for name in levels if levels[name]["role"] == "PALLET_DETAIL")
    carton_sheet = next(name for name in levels if levels[name]["role"] == "CARTON_DETAIL")
    assert levels[pallet_sheet]["row_count"] == 10
    assert levels[carton_sheet]["row_count"] == 36
    assert len(proposal["goods"]) == 10
    assert len(proposal["review"]["links"]) == 10
    assert SHEET_TOTAL in proposal["source"]["table_sheets"]
    assert SHEET_CARTONS == carton_sheet
