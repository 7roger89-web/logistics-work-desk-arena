"""Focused regression: synthetic Lục Giác source → existing C1 → writer → XLSX.

Storage tests use the repository's existing isolated mongomock/GridFS transport,
NOT a real production MongoDB. No network or customer workbook is required.
"""
from __future__ import annotations

import base64
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from core.quote_excel_import import parse_luc_giac_workbook
from core.quotations import calculate_quote, customer_quote_snapshot, new_quote_document, transition_quote
from scripts.make_luc_giac_quote_fixture import build_luc_giac_fixture
from views import logi_c1, logi_sandbox


@pytest.fixture
def source():
    return build_luc_giac_fixture()


@pytest.fixture
def proposal(source):
    return parse_luc_giac_workbook(source, "luc-giac-fixture.xlsx", today=date(2026, 10, 16))


@pytest.fixture
def staging():
    import streamlit as st
    from core.database import get_db

    had_user = "current_user" in st.session_state
    previous_user = st.session_state.get("current_user")
    logi_sandbox.reset_backend()
    assert logi_sandbox.ensure_backend() == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    db = get_db()
    # Synthetic acceptance identity, never a production insert.
    db.customers.insert_one({"id": "CUST-LUCGIAC-TEST", "name": "Công ty Lục Giác"})
    try:
        yield db
    finally:
        logi_sandbox.reset_backend()
        if had_user:
            st.session_state["current_user"] = previous_user
        else:
            st.session_state.pop("current_user", None)


def dispatch(event, payload):
    outcomes = logi_c1.consume_c1_triggers({"seq": 0}, {event: {"seq": 1, "payload": {"requestId": "test-request", **payload}}})
    assert len(outcomes) == 1
    assert outcomes[0]["requestId"] == payload.get("requestId", "test-request")
    return outcomes[0]


def imported(source):
    out = dispatch("quote_import_excel", {"customerId": "CUST-LUCGIAC-TEST", "fileName": "luc-giac-fixture.xlsx", "base64": base64.b64encode(source).decode()})
    assert out["ok"], out
    return out["readback"]


def payload(details, **extra):
    return {"customerId": "CUST-LUCGIAC-TEST", "salesOwner": "Sale test", "goodsDescription": "Bàn họp văn phòng",
            "incoterm": "EXW", "transportMode": "Sea", "excelQuote": details, "fx": {}, **extra}


def fee(**extra):
    return {"id": "fee-a", "service_name": "Vận chuyển", "qty": 1, "unit": "Lần", "currency": "VND", "cost_unit": 31800000,
            "sell_unit": None, "sell_confirmed": False, "tax_pct": None, "internal_note": "PRIVATE-SALE-NOTE", **extra}


def test_fixture_totals_sparse_identity_and_warnings(proposal):
    q = calculate_quote({"excel_quote": proposal})
    d = q["excel_quote"]
    assert d["calculation"]["row_count"] == 5
    assert d["calculation"]["quantity"] == 27
    assert d["calculation"]["goods_totals"] == [{"currency": "CNY", "amount": 607180}]
    assert d["goods"][0]["model"] == "MS-C0242S"
    assert all(not g["model"] and not g["material"] for g in d["goods"][1:])
    assert {w["id"] for w in d["source_warnings"]} >= {"equipment-length", "cargo-conflict", "dimension-conflict", "agent-rounding", "source-expiry", "vat-ambiguity", "formula-gap", "customs-basis"}
    assert "31/07/2026" in next(w["text"] for w in d["source_warnings"] if w["id"] == "source-expiry")
    assert d["calculation"]["estimate_vnd"] is None
    assert not q["service_lines"]
    assert len(d["disclosures"]) == 4 and all(x["selected"] for x in d["disclosures"])


def test_zero_goods_qty_preserved_and_missing_input_not_zero(proposal):
    proposal["goods"][0]["qty"] = 0
    q = calculate_quote({"excel_quote": proposal})
    assert q["excel_quote"]["goods"][0]["qty"] == 0
    assert q["excel_quote"]["goods"][0]["line_total"] == 0
    assert q["excel_quote"]["calculation"]["quantity"] == 11
    proposal["goods"][0]["qty"] = None
    q = calculate_quote({"excel_quote": proposal})
    assert q["excel_quote"]["goods"][0]["line_total"] is None
    assert q["excel_quote"]["calculation"]["goods_totals"][0]["amount"] is None


def test_formula_is_not_evaluated_or_zeroed(source):
    wb = load_workbook(BytesIO(source))
    wb["Khach gui"]["G7"] = "=10000+3055"
    data = BytesIO(); wb.save(data)
    p = parse_luc_giac_workbook(data.getvalue(), "formulas.xlsx")
    assert p["goods"][0]["unit_price"] is None
    assert calculate_quote({"excel_quote": p})["excel_quote"]["calculation"]["goods_totals"][0]["amount"] is None


@pytest.mark.parametrize("raw", ["31/02/2027", "2027-02-31", "2026-13-01", "2026-2-1", "banana"])
def test_invalid_calendar_date_rejected(proposal, raw):
    with pytest.raises(ValueError, match="Ngày"):
        calculate_quote({"excel_quote": proposal, "expected_delivery_date": raw})


def test_valid_leap_day(proposal):
    calculate_quote({"excel_quote": proposal, "expected_delivery_date": "2028-02-29"})


@pytest.mark.parametrize("field,value", [("qty", 0), ("qty", -1), ("cost_unit", "NaN"), ("sell_unit", "Infinity"), ("sell_unit", "abc"), ("tax_pct", 101)])
def test_bad_service_numbers_fail_closed(proposal, field, value):
    proposal["options"]["A"]["service_lines"] = [fee(**{field: value})]
    with pytest.raises(ValueError):
        calculate_quote({"excel_quote": proposal})


def test_draft_missing_sell_and_confirmation_are_unknown(proposal):
    proposal["options"]["A"]["service_lines"] = [fee()]
    q = new_quote_document({"excel_quote": proposal})
    assert q["service_lines"][0]["sell_unit"] is None
    assert q["service_lines"][0]["sell_total"] is None
    assert q["calculation"]["selling_price_before_tax_vnd"] is None
    assert q["calculation"]["totals_by_currency"][0]["sell_total"] is None
    with pytest.raises(ValueError, match="giá bán"):
        transition_quote(q, "READY_FOR_REVIEW")
    proposal["options"]["A"]["service_lines"][0]["sell_unit"] = 35000000
    q = calculate_quote({"excel_quote": proposal})
    assert q["service_lines"][0]["sell_unit"] == 35000000  # proposed input retained
    assert q["service_lines"][0]["sell_total"] is None  # not yet confirmed
    assert not customer_quote_snapshot(q)["service_lines"]


def test_options_not_copied_and_customs_not_derived_from_selling_price(proposal):
    proposal["options"]["A"].update({"scope": "Vận chuyển theo thỏa thuận", "service_lines": [fee(sell_unit=35000000, sell_confirmed=True, tax_pct=0)]})
    proposal["options"]["B"].update({"scope": "Ủy thác theo thỏa thuận", "service_lines": [fee(id="fee-b", sell_unit=40000000, sell_confirmed=True, tax_pct=0)]})
    proposal["customs"] = {"confirmed": True, "basis_vnd": 2400000000, "basis_note": "Giả định độc lập đã xác nhận cho fixture", "duty_pct": 5, "vat_pct": 8}
    q = calculate_quote({"excel_quote": proposal, "fx_snapshot": {"rates_to_vnd": {"CNY": 3908.55457227}}})
    assert q["excel_quote"]["customs"]["duty_vnd"] == 120000000
    assert q["excel_quote"]["customs"]["vat_vnd"] == 201600000
    taxes = deepcopy(q["excel_quote"]["customs"])
    proposal["options"]["A"]["service_lines"][0]["sell_unit"] = 99999999
    higher = calculate_quote({"excel_quote": proposal})
    assert higher["excel_quote"]["customs"] == taxes
    proposal["active_option"] = "B"
    b = calculate_quote({"excel_quote": proposal})
    assert [l["sell_unit"] for l in b["service_lines"]] == [40000000]
    assert b["excel_quote"]["options"]["A"]["service_lines"][0]["sell_unit"] == 99999999
    assert b["excel_quote"]["calculation"]["estimate_vnd"] is None  # missing goods FX
    proposal["customs"]["confirmed"] = False
    assert calculate_quote({"excel_quote": proposal})["excel_quote"]["customs"]["total_vnd"] is None


def test_mixed_currency_no_combined_total_without_fx(proposal):
    proposal["options"]["A"]["service_lines"] = [fee(currency="USD", sell_unit=100, sell_confirmed=True, tax_pct=0)]
    q = calculate_quote({"excel_quote": proposal})
    assert q["calculation"]["totals_by_currency"][0]["sell_total"] == 100
    assert q["excel_quote"]["calculation"]["goods_totals"][0]["amount"] == 607180
    assert q["excel_quote"]["calculation"]["before_tax_vnd"] is None
    assert q["calculation"]["selling_price_total_vnd"] is None


def test_real_c1_upload_gridfs_save_refresh_edit_export(staging, source):
    from core.database import load_quotes
    from core.gridfs_client import GridFSAttachmentStore

    d = imported(source)
    assert not staging.quotes.find_one({"customer_id": "CUST-LUCGIAC-TEST"})  # review, not auto-create
    assert GridFSAttachmentStore(staging).download_file(d["source"]["storage_file_id"]) == source
    d["options"]["A"]["service_lines"] = [fee()]
    d["shipment"].update({"pol": "NANSHA", "pod": "HAI PHONG", "delivery": "Thanh Trì, Hà Nội"})
    d["disclosures"][1]["selected"] = False
    d["additional_public_notes"] = "Dòng công bố một\nDòng công bố hai"
    created = dispatch("quote_create", payload(d))
    assert created["ok"], created
    qid = created["id"]
    fresh = load_quotes()[qid]  # independent fresh reader, no component-local truth
    assert fresh["status"] == "DRAFT" and fresh["service_lines"][0]["sell_unit"] is None
    assert fresh["excel_quote"]["disclosures"][1]["text"] == "Sản xuất dự kiến 30 ngày."
    assert not fresh["excel_quote"]["disclosures"][1]["selected"]
    assert fresh["excel_quote"]["additional_public_notes"].count("\n") == 1
    fresh_d = fresh["excel_quote"]
    fresh_d["goods"][0]["qty"] = 15
    edited = dispatch("quote_edit", {**payload(fresh_d), "id": qid, "expectedVersion": fresh["version"]})
    assert edited["ok"], edited
    assert edited["readback"]["excelQuote"]["calculation"]["quantity"] == 26
    assert edited["readback"]["excelQuote"]["calculation"]["goods_totals"][0]["amount"] == 594125
    stale = dispatch("quote_edit", {**payload(fresh_d), "id": qid, "expectedVersion": fresh["version"]})
    assert not stale["ok"] and "phiên bản" in stale["error"]
    exported = dispatch("quote_export_xlsx", {"id": qid})
    assert exported["ok"], exported
    wb = load_workbook(BytesIO(base64.b64decode(exported["document"]["downloadBase64"])))
    cells = [str(c.value or "") for ws in wb for row in ws for c in row]
    text = "\n".join(cells)
    assert "Sản xuất dự kiến 30 ngày." not in text
    assert "Thanh toán 80% trước khi giao hàng." in text
    assert "Dòng công bố một\nDòng công bố hai" in text
    assert "PRIVATE-SALE-NOTE" not in text
    assert "31800000" not in text
    assert "NANSHA → HAI PHONG" in text and "Thanh Trì, Hà Nội" in text
    assert not any(c.data_type == "f" for ws in wb for row in ws for c in row)
    downloaded = dispatch("quote_source_download", {"id": qid})
    assert base64.b64decode(downloaded["document"]["downloadBase64"]) == source


def test_reference_cost_source_is_rehydrated_not_selling_price(staging, source):
    staging.rate_cards.insert_one({"id": "RATE-EXPIRED", "record_type": "SUPPLIER_RATE_SNAPSHOT", "supplier": "PRIVATE-RATE-SUPPLIER", "status": "ACTIVE",
        "rows": [{"id": "r1", "kind": "CHARGE", "status": "AVAILABLE", "amount": 31800000, "currency": "VND", "unit": "Lần", "validUntil": "2026-07-31", "conditionalAmounts": ["+ phụ phí nếu phát sinh"], "reviewNote": "PRIVATE-RATE-NOTE"}]})
    d = imported(source)
    d["source_warnings"] = []  # forged attempt to erase workbook warnings
    d["options"]["A"]["service_lines"] = [fee(rate_snapshot_id="RATE-EXPIRED", rate_row_id="r1", cost_unit=0, validity_status="SOURCE_VALID")]
    out = dispatch("quote_preview", payload(d))
    assert out["ok"], out
    line = out["readback"]["quote"]["serviceLines"][0]
    assert line["costUnit"] == 31800000 and line["sellUnit"] is None
    calculated = out["readback"]["quote"]["excelQuote"]
    line = calculated["options"]["A"]["service_lines"][0]
    assert line["validity_status"] == "EXPIRED"
    assert line["condition"] == "IF_APPLICABLE"
    assert calculated["source_warnings"]
    assert "PRIVATE-RATE" not in json.dumps(out["readback"]["customer"], ensure_ascii=False)
    assert not calculated["options"]["B"]["service_lines"]


def test_sent_snapshot_immutable_revision_and_historical_xlsx(staging, source):
    d = imported(source)
    d["options"]["A"].update({"scope": "Vận chuyển đã thống nhất", "service_lines": [fee(sell_unit=35000000, sell_confirmed=True, tax_pct=0)]})
    d["additional_public_notes"] = "PUBLIC-V1\n=HYPERLINK(\"https://example.invalid\",\"literal text\")"
    out = dispatch("quote_create", payload(d, fx={"USD": 26500, "CNY_PER_USD": 6.78}))
    assert out["ok"], out
    qid = out["id"]
    assert out["readback"]["excelQuote"]["calculation"]["goods_total_vnd"] == 2373196165
    for status in ("READY_FOR_REVIEW", "SENT"):
        out = dispatch("quote_transition", {"id": qid, "nextStatus": status})
        assert out["ok"], out
    stored = logi_c1.read_quote(qid)
    sent = deepcopy(stored["sent_versions"])
    snapshot = sent[0]["customer_snapshot"]
    assert "PRIVATE-SALE-NOTE" not in json.dumps(snapshot)
    assert "estimated_margin" not in json.dumps(snapshot) and "cost_unit" not in json.dumps(snapshot)
    d["additional_public_notes"] = "PUBLIC-V2-DRAFT"
    revised = dispatch("quote_edit", {**payload(d), "id": qid, "expectedVersion": stored["version"]})
    assert revised["ok"], revised
    assert revised["readback"]["status"] == "DRAFT"
    assert logi_c1.read_quote(qid)["sent_versions"] == sent
    exported = dispatch("quote_export_xlsx", {"id": qid, "sentVersion": sent[0]["version"]})
    assert exported["ok"]
    wb = load_workbook(BytesIO(base64.b64decode(exported["document"]["downloadBase64"])))
    text = "\n".join(str(c.value or "") for ws in wb for row in ws for c in row)
    assert "PUBLIC-V1" in text and "PUBLIC-V2-DRAFT" not in text
    assert "PRIVATE-SALE-NOTE" not in text and "31800000" not in text
    assert not any(c.data_type == "f" for ws in wb for row in ws for c in row)


def test_source_ownership_and_rbac_fail_closed(staging, source):
    import streamlit as st
    d = imported(source)
    # Another existing customer cannot attach this customer's stored source.
    out = dispatch("quote_create", payload(d, customerId="CUST-ANPHAT"))
    assert not out["ok"] and "khách hàng" in out["error"]
    st.session_state["current_user"] = {"username": "viewer-test", "role": "viewer"}
    out = dispatch("quote_preview", payload(d))
    assert not out["ok"] and "RBAC" in out["error"]
    out = dispatch("quote_import_excel", {"customerId": "CUST-LUCGIAC-TEST", "fileName": "luc.xlsx", "base64": base64.b64encode(source).decode()})
    assert not out["ok"] and "RBAC" in out["error"]


@pytest.mark.parametrize("name,data", [("../bad.xlsx", b"abc"), ("bad.xlsm", b"abc"), ("bad.xlsx", b"not a zip"), ("big.xlsx", b"x" * (5 * 1024 * 1024 + 1))])
def test_unsafe_or_unsupported_workbooks_rejected(name, data):
    with pytest.raises(ValueError):
        parse_luc_giac_workbook(data, name)


def test_no_second_calculation_or_localstorage_authority_in_new_ui():
    root = Path(__file__).resolve().parents[1]
    source = (root / "component/src/components/ExcelQuoteWorkspace.tsx").read_text()
    assert 'dispatch("quote_preview"' in source
    assert "localStorage" not in source
    assert "* line.sell_unit" not in source and "unit_price *" not in source
    assert '"quote_import_excel"' in (root / "component/src/lib/dispatcher.ts").read_text()


def test_quote_create_uses_fresh_customer_not_render_cache(staging, source, monkeypatch):
    import core.database as database

    # Reproduce the actual Streamlit UI race: readback/picker already includes
    # a just-created customer while the render cache still contains old data.
    monkeypatch.setattr(database, "load_customers", lambda: {})
    created = dispatch("quote_create", payload(imported(source)))
    assert created["ok"], created
    # Conversely, a deleted record in a stale cache must never authorise writes.
    monkeypatch.setattr(database, "load_customers", lambda: {"CUST-LUCGIAC-TEST": {"id": "CUST-LUCGIAC-TEST", "name": "Stale"}})
    staging.customers.update_one({"id": "CUST-LUCGIAC-TEST"}, {"$set": {"is_deleted": True}})
    refused = dispatch("quote_create", payload(created["readback"]["excelQuote"]))
    assert not refused["ok"] and "Không tìm thấy khách hàng" in refused["error"]


def test_customer_tax_basis_is_public_but_internal_fx_and_assumptions_are_not(proposal):
    proposal["customs"].update({"confirmed": True, "basis_vnd": 2400000000, "basis_note": "PRIVATE-CUSTOMS-ASSUMPTION", "duty_pct": 5, "vat_pct": 8})
    q = calculate_quote({"excel_quote": proposal, "fx_snapshot": {"rates_to_vnd": {"CNY": 3900}, "source": "PRIVATE-FX-SOURCE"}})
    public = customer_quote_snapshot(q)
    assert public["customs"]["basis_vnd"] == 2400000000
    assert public["customs"]["total_vnd"] == 321600000
    assert "PRIVATE-" not in json.dumps(public)
    assert "source" not in public["fx_snapshot"]


def test_writer_success_without_fresh_readback_is_not_reported_as_saved(staging, source, monkeypatch):
    d = imported(source)
    monkeypatch.setattr(logi_c1, "read_quote", lambda _: None)
    outcome = dispatch("quote_create", payload(d))
    assert not outcome["ok"] and "chưa xác minh" in outcome["error"]


def test_direct_and_cross_fx_cannot_silently_overwrite_each_other(staging, source):
    out = dispatch("quote_preview", payload(imported(source), fx={"CNY": 4000, "USD": 26500, "CNY_PER_USD": 6.78}))
    assert not out["ok"] and "không tự ghi đè" in out["error"]


# Regression: the OWNER's actual workbook uses these header labels, not the
# earlier synthetic fixture's "Kích thước" / "Vật liệu" labels. The customer
# workbook itself is confidential and must not be checked into this repository.
def test_uploaded_original_luc_giac_column_headings_are_accepted():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "BÁO GIÁ"
    ws["B1"] = "BÁO GIÁ BÀN HỌP - NHẬP KHẨU TQ"
    for col, label in {
        "A": "STT", "B": "NSX", "C": "CHẤT LIỆU",
        "D": "KT SẢN PHẨM (mm)", "E": "SỐ LƯỢNG",
        "F": "ĐƠN VỊ", "G": "ĐƠN GIÁ RMB",
        "H": "THÀNH TIỀN RMB", "K": "GHI CHÚ",
    }.items():
        ws[f"{col}2"] = label
    for i, (dim, qty, unit_price) in enumerate((
        ("5200×1500×760", 16, 13055),
        ("7500×1600×760", 1, 18850),
        ("12000×1600×760", 3, 30150),
        ("16000×1600×760", 6, 40200),
        ("19000×1600×760", 1, 47800),
    ), start=3):
        ws[f"A{i}"] = i - 2
        ws[f"D{i}"] = dim
        ws[f"E{i}"] = qty
        ws[f"G{i}"] = unit_price
        ws[f"H{i}"] = f"=E{i}*G{i}"
    ws["C3"] = "Mã sp: MS-C0242S\nKhung: MDF\nBề mặt hoàn thiện: veneer"
    ws["K3"] = "kèm 4 ổ điện"

    stream = BytesIO()
    wb.save(stream)
    parsed = parse_luc_giac_workbook(
        stream.getvalue(), "BAO GIA BAN HOP.xlsx", today=date(2026, 10, 16)
    )
    calc = calculate_quote({"excel_quote": parsed})["excel_quote"]
    assert parsed["source"]["goods_sheet"] == "BÁO GIÁ"
    assert calc["calculation"]["row_count"] == 5
    assert calc["calculation"]["quantity"] == 27
    assert calc["calculation"]["goods_totals"] == [{"currency": "CNY", "amount": 607180}]
    assert "MS-C0242S" in parsed["goods"][0]["material"]
    assert parsed["goods"][0]["note"] == "kèm 4 ổ điện"
