"""HEXA-EXCEL-IMPORT-V0 — the three real PL/PI layout families, on SYNTHETIC files.

`REAL_WORKBOOK_TEST=NOT_RUN_IN_ARENA`: the Owner's three original workbooks live
in another conversation and are deliberately NOT in this repository. These tests
prove the importer against fixtures that reproduce the described STRUCTURE only:

  A. packing list, product photos, table at row 7, 257 declared columns, no price
  B. proforma invoice + a second money sheet that duplicates the goods table
  C. goods list → pallets → cartons, plus a hidden loading-instruction sheet

The negative test uses a FAKE internal marker (``TT_INTERNAL_SENTINEL``) to prove
the customer quotation never carries internal or hidden-sheet content.
"""
from __future__ import annotations

import hashlib
import json
from datetime import date
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from core.excel_utils import generate_quote_excel
from core.quote_excel_import import parse_luc_giac_workbook
from core.quote_workbook_view import overview, sheet_page
from core.quotations import calculate_quote, customer_quote_snapshot
from scripts.make_hexa_layout_fixtures import (
    CARTON_ROWS,
    PACKING_ROWS,
    PALLET_ROWS,
    PI_ROWS,
    SENTINEL,
    SHEET_CARTONS,
    SHEET_INSTRUCTION,
    SHEET_PACKING,
    SHEET_PALLETS,
    SHEET_PI,
    SHEET_PI_MONEY,
    SHEET_TOTAL,
    TOTAL_ROWS,
    build_multilevel_packing_list,
    build_oversized_width_workbook,
    build_packing_list_with_images,
    build_proforma_invoice,
)

TODAY = date(2026, 10, 16)


@pytest.fixture(scope="module")
def packing_list() -> bytes:
    return build_packing_list_with_images()


@pytest.fixture(scope="module")
def proforma() -> bytes:
    return build_proforma_invoice()


@pytest.fixture(scope="module")
def multilevel() -> bytes:
    return build_multilevel_packing_list()


def parse(data: bytes, file_name: str) -> dict:
    return parse_luc_giac_workbook(data, file_name, today=TODAY)


def warning_ids(proposal: dict) -> set[str]:
    return {w["id"] for w in proposal["source_warnings"]}


def unknown_codes(proposal: dict) -> set[str]:
    return {u["code"] for u in proposal["review"]["unknowns"]}


# ------------------------------------------------------- fixture reality checks
def test_fixtures_reproduce_the_three_shapes(packing_list, proforma, multilevel):
    wb = load_workbook(BytesIO(packing_list))
    assert wb.sheetnames == [SHEET_PACKING, "Sheet3"]
    assert wb[SHEET_PACKING].max_column == 257  # declared by formatting, data ends at T
    assert wb[SHEET_PACKING].max_row == 13
    assert wb[SHEET_PACKING]["B7"].value == "产品型号 Model"
    assert wb[SHEET_PACKING]._images

    wb = load_workbook(BytesIO(proforma))
    assert wb.sheetnames == [SHEET_PI, SHEET_PI_MONEY]
    assert wb[SHEET_PI]["B7"].value == "型号 Model"
    assert wb[SHEET_PI]._images
    assert wb[SHEET_PI_MONEY]["B12"].value == SENTINEL

    wb = load_workbook(BytesIO(multilevel))
    assert wb.sheetnames == [SHEET_TOTAL, SHEET_PALLETS, SHEET_CARTONS, SHEET_INSTRUCTION]
    assert wb[SHEET_INSTRUCTION].sheet_state == "hidden"
    assert wb[SHEET_INSTRUCTION]["A3"].value == SENTINEL


# ------------------------------------------------- A: packing list without price
def test_packing_list_without_price_is_importable_and_price_stays_unknown(packing_list):
    proposal = parse(packing_list, "packing-list-SYNTHETIC.xlsx")

    assert proposal["review"]["layout"] == "PACKING_LIST"
    assert proposal["source"]["goods_sheet"] == SHEET_PACKING
    assert len(proposal["goods"]) == len(PACKING_ROWS)
    first = proposal["goods"][0]
    assert (first["source_sheet"], first["source_row"]) == (SHEET_PACKING, 8)
    assert first["model"] == "BC-800" and first["qty"] == 6
    # A missing price is UNKNOWN, never 0.
    assert all(g["unit_price"] is None for g in proposal["goods"])
    assert "unit-price-missing" in warning_ids(proposal)
    assert "unit_price_missing" in unknown_codes(proposal)

    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["row_count"] == len(PACKING_ROWS)
    assert calculation["quantity"] == sum(row[4] for row in PACKING_ROWS)
    assert calculation["goods_totals"][0]["amount"] is None


def test_product_size_is_never_reported_as_package_or_carton_size(packing_list):
    proposal = parse(packing_list, "packing-list-SYNTHETIC.xlsx")
    # The goods table must not borrow the product dimension as shipment/package dims.
    assert all(g["length_mm"] is None and g["width_mm"] is None and g["height_mm"] is None
               for g in proposal["goods"])
    evidence = next(r for r in proposal["review"]["rows"] if r["ref"] == f"{SHEET_PACKING}!8")
    assert evidence["product_size"] == "800×480×850"
    assert evidence["package_size"] == "860×540×900"
    assert evidence["packing"] == "Carton + foam"
    assert evidence["cartons"] == 6
    assert evidence["gross_weight_kg"] == 108.0
    assert "dims-not-package" in warning_ids(proposal)
    # No CBM total is invented from product dimensions.
    assert proposal["shipment"]["cbm"] is None


def test_declared_257_columns_are_read_from_the_used_region_only(packing_list):
    proposal = parse(packing_list, "packing-list-SYNTHETIC.xlsx")
    bounds = next(b for b in proposal["source"]["sheet_bounds"] if b["sheet"] == SHEET_PACKING)
    assert bounds["declared_columns"] == 257
    assert bounds["used_columns"] <= 21  # real data ends at T
    assert bounds["declared_rows"] == 13
    assert bounds["images"] == 1


def test_absurdly_wide_workbook_is_still_refused():
    with pytest.raises(ValueError):
        parse(build_oversized_width_workbook(), "too-wide-SYNTHETIC.xlsx")


# ------------------------------------------- B: PI with a second money sheet
def test_pi_goods_come_from_one_sheet_and_are_never_counted_twice(proforma):
    proposal = parse(proforma, "proforma-invoice-SYNTHETIC.xlsx")

    assert proposal["review"]["layout"] == "PROFORMA_INVOICE"
    assert proposal["source"]["goods_sheet"] == SHEET_PI
    excluded = {t["sheet"]: t for t in proposal["source"]["excluded_tables"]}
    assert excluded[SHEET_PI_MONEY]["role"] == "INTERNAL_CALC"
    assert excluded[SHEET_PI_MONEY]["row_count"] == len(PI_ROWS)
    assert "money-sheet-not-imported" in warning_ids(proposal)

    assert len(proposal["goods"]) == len(PI_ROWS)
    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["row_count"] == len(PI_ROWS)
    expected = sum(qty * price for _model, _description, _size, qty, price in PI_ROWS)
    assert calculation["goods_totals"] == [{"currency": "CNY", "amount": expected}]
    # The duplicated money sheet would double this number; it must not be added.
    assert calculation["goods_totals"][0]["amount"] != expected * 2


def test_pi_keeps_source_row_for_every_goods_line(proforma):
    proposal = parse(proforma, "proforma-invoice-SYNTHETIC.xlsx")
    assert [(g["source_sheet"], g["source_row"]) for g in proposal["goods"]] == [
        (SHEET_PI, 8 + offset) for offset in range(len(PI_ROWS))
    ]
    assert proposal["goods"][0]["description"].startswith("Bathroom cabinet 800")


def test_uncached_price_formula_stays_unknown_never_zero(proforma):
    wb = load_workbook(BytesIO(proforma))
    wb[SHEET_PI]["G8"] = "=60+15"  # openpyxl stores the formula with no cached result
    stream = BytesIO()
    wb.save(stream)

    proposal = parse(stream.getvalue(), "proforma-invoice-formula-SYNTHETIC.xlsx")
    assert proposal["goods"][0]["unit_price"] is None
    assert "formula_no_cache" in unknown_codes(proposal)
    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["goods_totals"][0]["amount"] is None
    assert proposal["goods"][0]["unit_price"] != 0


# --------------------------------------------- C: goods list → pallets → cartons
def test_three_levels_are_kept_as_levels_not_added_as_three_goods_lists(multilevel):
    proposal = parse(multilevel, "multilevel-packing-SYNTHETIC.xlsx")

    assert proposal["review"]["layout"] == "MULTI_LEVEL_PACKING"
    assert proposal["source"]["goods_sheet"] == SHEET_TOTAL
    assert len(proposal["goods"]) == len(TOTAL_ROWS)
    roles = {level["sheet"]: level["role"] for level in proposal["review"]["levels"]}
    assert roles[SHEET_PALLETS] == "PALLET_DETAIL"
    assert roles[SHEET_CARTONS] == "CARTON_DETAIL"
    assert roles[SHEET_INSTRUCTION] == "HIDDEN_SHEET"
    counts = {level["sheet"]: level["row_count"] for level in proposal["review"]["levels"]}
    assert counts[SHEET_PALLETS] == len(PALLET_ROWS)
    assert counts[SHEET_CARTONS] == len(CARTON_ROWS)
    assert counts[SHEET_INSTRUCTION] == 0
    assert "level-not-summed" in warning_ids(proposal)

    # No CBM/weight total is produced by adding cartons + pallets + goods list.
    assert proposal["shipment"]["cbm"] is None
    assert proposal["shipment"]["weight_kg"] is None
    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["row_count"] == len(TOTAL_ROWS)
    assert calculation["quantity"] == sum(row[2] for row in TOTAL_ROWS)


def test_levels_keep_the_model_key_when_it_is_unambiguous(multilevel):
    proposal = parse(multilevel, "multilevel-packing-SYNTHETIC.xlsx")
    links = proposal["review"]["links"]
    assert links["BC-800"][SHEET_PALLETS] == ["PLT-01"]
    assert links["BC-800"][SHEET_CARTONS] == [3, 4, 5]
    assert links["BC-1000"][SHEET_PALLETS] == ["PLT-02"]
    # The row with no model code cannot be linked, so it waits for the operator.
    assert "level_link_missing" in unknown_codes(proposal)


def test_packing_evidence_is_review_data_with_its_own_source(multilevel):
    proposal = parse(multilevel, "multilevel-packing-SYNTHETIC.xlsx")
    assert all("cbm" not in g for g in proposal["goods"])
    refs = {row["ref"] for row in proposal["review"]["rows"]}
    assert f"{SHEET_TOTAL}!4" in refs
    assert f"{SHEET_CARTONS}!3" in refs
    assert proposal["review"]["rows"][0]["volume_cbm"] is not None


# ------------------------------------- internal / hidden content must not leak
@pytest.mark.parametrize("builder,file_name", [
    (build_proforma_invoice, "proforma-invoice-SYNTHETIC.xlsx"),
    (build_multilevel_packing_list, "multilevel-packing-SYNTHETIC.xlsx"),
])
def test_internal_sentinel_and_hidden_sheet_never_reach_the_customer_quote(builder, file_name):
    data = builder()
    proposal = parse(data, file_name)
    assert SENTINEL not in json.dumps(proposal["goods"], ensure_ascii=False)

    quote = calculate_quote({"excel_quote": proposal, "status": "DRAFT",
                             "customer_name": "Công ty thử nghiệm (không có thật)"})
    snapshot = customer_quote_snapshot(quote)
    assert SENTINEL not in json.dumps(snapshot, ensure_ascii=False)

    exported = generate_quote_excel(quote).getvalue()
    assert SENTINEL.encode() not in exported
    assert "Giá vốn nội bộ".encode() not in exported
    with ZipFile(BytesIO(exported)) as archive:
        package_text = "".join(archive.read(name).decode("utf-8", "ignore")
                               for name in archive.namelist() if name.endswith(".xml"))
    assert SENTINEL not in package_text
    assert SHEET_INSTRUCTION not in package_text


# ------------------------------------------------------- viewer on a wide source
def test_viewer_opens_the_wide_packing_list_read_only_and_grant_gated(packing_list):
    proposal = parse(packing_list, "packing-list-SYNTHETIC.xlsx")
    info = overview(packing_list, proposal)
    sheet = next(s for s in info["sheets"] if s["name"] == SHEET_PACKING)
    assert sheet["declared_column_count"] == 257
    assert sheet["column_count"] <= 21

    locked = sheet_page(packing_list, SHEET_PACKING, proposal=proposal)  # default deny
    assert locked["read_only"] is True and locked["locked"] is True and locked["cells"] == []

    granted = sheet_page(packing_list, SHEET_PACKING, proposal=proposal, can_view_internal=True)
    assert granted["read_only"] is True and granted["locked"] is False
    assert granted["limits"]["max_cols"] == 40
    assert max(cell["c"] for cell in granted["cells"]) <= 21


def test_hidden_loading_instruction_sheet_is_never_the_customer_sheet(multilevel):
    proposal = parse(multilevel, "multilevel-packing-SYNTHETIC.xlsx")
    info = overview(multilevel, proposal)
    by_name = {sheet["name"]: sheet for sheet in info["sheets"]}
    assert by_name[SHEET_INSTRUCTION]["state"] == "hidden"
    assert by_name[SHEET_INSTRUCTION]["classification"] != "CUSTOMER"
    assert info["customer_sheet"] == SHEET_TOTAL
    locked = sheet_page(multilevel, SHEET_INSTRUCTION, proposal=proposal)
    assert locked["locked"] is True and locked["cells"] == []


# ------------------------------------- end-to-end: import → store → quote → XLSX
@pytest.fixture
def staging():
    import streamlit as st

    from core.database import get_db
    from views import logi_sandbox

    had_user = "current_user" in st.session_state
    previous_user = st.session_state.get("current_user")
    logi_sandbox.reset_backend()
    assert logi_sandbox.ensure_backend() == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    db = get_db()
    db.customers.insert_one({"id": "CUST-HEXA-TEST", "name": "Công ty thử nghiệm (không có thật)"})
    try:
        yield db
    finally:
        logi_sandbox.reset_backend()
        if had_user:
            st.session_state["current_user"] = previous_user
        else:
            st.session_state.pop("current_user", None)


def dispatch(event, payload):
    from views import logi_c1

    outcomes = logi_c1.consume_c1_triggers(
        {"seq": 0}, {event: {"seq": 1, "payload": {"requestId": "hexa-test", **payload}}})
    assert len(outcomes) == 1
    return outcomes[0]


def test_packing_list_round_trip_keeps_source_hash_and_review_evidence(staging, packing_list):
    import base64

    imported = dispatch("quote_import_excel", {
        "customerId": "CUST-HEXA-TEST", "fileName": "packing-list-SYNTHETIC.xlsx",
        "base64": base64.b64encode(packing_list).decode(),
    })
    assert imported["ok"], imported
    details = imported["readback"]
    assert details["review"]["layout"] == "PACKING_LIST"
    assert details["source"]["sha256"] == hashlib.sha256(packing_list).hexdigest()
    # Re-validating the stored source (as quote_create does) restores the server's
    # own review evidence instead of trusting whatever the browser sent.
    forged = dict(details)
    forged["review"] = {"layout": "FORGED", "rows": [], "unknowns": [], "levels": [], "links": {}}
    created = dispatch("quote_create", {
        "customerId": "CUST-HEXA-TEST", "salesOwner": "Sale test",
        "goodsDescription": "Tủ lavabo", "incoterm": "EXW", "transportMode": "Sea",
        "excelQuote": forged, "fx": {},
    })
    assert created["ok"], created
    stored = created["readback"]["excelQuote"]
    assert stored["review"]["layout"] == "PACKING_LIST"
    assert stored["goods"][0]["unit_price"] is None
