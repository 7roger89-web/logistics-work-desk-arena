"""Phase G — khóa hành vi scripts/demo_bootstrap.py (clone tái lập được).

Bề mặt cần chứng minh:

* bootstrap tạo tariff.db rỗng qua initializer sẵn có, placeholder GNT/MA_NGAN_SACH/
  DOCX có gắn nhãn demo, và seed mongomock qua đúng transport sandbox của app;
* docgen giữ nguyên fail-closed: bootstrap KHÔNG đụng ``DOCGEN_TEMPLATE_DIR``;
* có ``MONGO_URI`` thật thì bootstrap bị chặn, không ghi gì (production untouched);
* chạy lại idempotent.
"""
from __future__ import annotations

import importlib.util
import sqlite3
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]

_spec = importlib.util.spec_from_file_location(
    "demo_bootstrap", ROOT / "scripts" / "demo_bootstrap.py"
)
demo_bootstrap = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(demo_bootstrap)  # type: ignore[union-attr]


@pytest.fixture
def _sandbox_cleanup():
    from core import config as core_config
    from core import database as core_database
    from views import logi_sandbox

    saved_config = {
        "MONGO_URI": core_config.MONGO_URI,
        "MONGO_DB_NAME": core_config.MONGO_DB_NAME,
        "ENV_IDENTITY": core_config.ENV_IDENTITY,
        "MONGO_ALLOWED_DB_NAMES": list(core_config.MONGO_ALLOWED_DB_NAMES),
    }
    saved_db = {
        "MONGO_URI": getattr(core_database, "MONGO_URI", None),
        "MONGO_DB_NAME": getattr(core_database, "MONGO_DB_NAME", None),
    }
    yield
    logi_sandbox.reset_backend()
    core_config.MONGO_URI = saved_config["MONGO_URI"]
    core_config.MONGO_DB_NAME = saved_config["MONGO_DB_NAME"]
    core_config.ENV_IDENTITY = saved_config["ENV_IDENTITY"]
    core_config.MONGO_ALLOWED_DB_NAMES = saved_config["MONGO_ALLOWED_DB_NAMES"]
    if saved_db["MONGO_URI"] is not None:
        core_database.MONGO_URI = saved_db["MONGO_URI"]
        core_database.MONGO_DB_NAME = saved_db["MONGO_DB_NAME"]


def test_bootstrap_creates_all_artifacts(tmp_path, _sandbox_cleanup):
    workdir = tmp_path / "demo"
    report = demo_bootstrap.bootstrap(workdir)

    # tariff.db rỗng, đúng schema initializer
    tariff_db = workdir / "data" / "tariff.db"
    assert tariff_db.is_file()
    with sqlite3.connect(tariff_db) as conn:
        tables = {row[0] for row in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        )}
        assert tables == {"tariff_table"}
        assert conn.execute("SELECT COUNT(*) FROM tariff_table").fetchone()[0] == 0

    templates = workdir / "templates"
    assert (templates / "GNT_MAU.placeholder.xlsx").is_file()
    assert (templates / "MA_NGAN_SACH.placeholder.xlsx").is_file()
    assert (templates / "docgen" / "HDNT-TT.placeholder.docx").is_file()
    assert (templates / "README.md").is_file()
    assert (templates / "manifest.json").is_file()

    # mongomock seed qua transport chính chủ
    assert report["mongomock"]["transport"] == "SANDBOX_STAGING"
    counts = report["mongomock"]["counts"]
    assert counts["customers"] >= 5 and counts["shipments"] >= 5

    # docgen vẫn fail-closed
    assert report["docgen_template_dir"] is None
    assert report["docgen_fail_closed_preserved"] is True


def test_placeholders_are_explicitly_marked_non_legal(tmp_path, _sandbox_cleanup):
    import json
    from openpyxl import load_workbook
    from docx import Document

    workdir = tmp_path / "demo"
    demo_bootstrap.bootstrap(workdir)
    templates = workdir / "templates"
    marker = demo_bootstrap.PLACEHOLDER_MARKER

    wb = load_workbook(templates / "GNT_MAU.placeholder.xlsx")
    assert marker in (wb.active["A1"].value or "")
    wb2 = load_workbook(templates / "MA_NGAN_SACH.placeholder.xlsx")
    assert marker in (wb2.active["A1"].value or "")

    doc = Document(templates / "docgen" / "HDNT-TT.placeholder.docx")
    assert any(marker in p.text for p in doc.paragraphs)

    manifest = json.loads((templates / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["legal_templates"] is False
    assert manifest["wired_to_docgen_template_dir"] is False
    assert manifest["docgen_owner_gate"] == "TEMPLATE_ASSET_STORAGE_OWNER_GATE"


def test_real_mongo_uri_blocks_everything(tmp_path, monkeypatch, _sandbox_cleanup):
    from core import config as core_config

    monkeypatch.setattr(core_config, "MONGO_URI", "mongodb://real-prod.invalid:27017")
    workdir = tmp_path / "demo"
    with pytest.raises(demo_bootstrap.BootstrapError):
        demo_bootstrap.bootstrap(workdir)
    # Không artifact nào được tạo sau khi chặn
    assert not workdir.exists()


def test_bootstrap_is_idempotent(tmp_path, _sandbox_cleanup):
    workdir = tmp_path / "demo"
    first = demo_bootstrap.bootstrap(workdir)
    second = demo_bootstrap.bootstrap(workdir)
    assert first["placeholders"] == second["placeholders"]
    assert first["mongomock"]["counts"] == second["mongomock"]["counts"]
    # Chạy lại không nhân bản document trong mongomock
    from core.database import get_db
    db = get_db()
    assert db.customers.count_documents({}) == first["mongomock"]["counts"]["customers"]
