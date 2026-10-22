"""HEXA-EXCEL-IMPORT-V0-REAL-LAYOUT-REPAIR — fixtures that mirror the REAL headers.

`REAL_WORKBOOK_TEST=NOT_RUN_IN_ARENA`: the Owner's three workbooks are still not
in this repository, so every PASS here is SYNTHETIC structure only.

What this file adds over `test_hexa_excel_import_layouts_v1.py` is the header
WORDING and row/column POSITIONS an independent review measured on the real
files, because the v1 fixtures used friendlier headers than the real files:

A. `Sheet1`, header row 7:
   `Item number` · `Product name` · `Size` · `Quantity` · `CTN` · `NW` · `GW` · `CBM`
   — no unit price at all.
B. `PI (2)`, header row 7:
   `Model 型号` · `Size尺寸(MM)` · `Qt.数量` · `Unit Price 单价 (CNY)` · `Amount 金额 (CNY)`
   plus the internal `File tiền hàng` money sheet.
C. `总清单`   header row 3: `Name product-Description` · `数量（件数）` · `净重` · `毛重` · `体积`,
   goods table rows 7-16.
   `托盘明细（2托）` header row 1 and REPEATED at row 9:
   `The pallet number` · `Pallet size` · `contract` · `House type` · `Room number` ·
   `Project name` · `PC` · `Total`  → PALLET level, never a goods list.
   `卫浴`      header row 2, repeated per group, with a `W` / `D` / `H` axis row
   directly underneath: `REF NO` · `Package NO` · `Desc.` · `Size(mm x mm x mm)` ·
   `NW(KG)` · `GW(KG)` · `VOL(CBM)` · `数量(箱)` · `货物中文名称` → CARTON level.
   `打托指令`  hidden → never read, never exported.
"""
from __future__ import annotations

import json
from datetime import date
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from core.excel_utils import generate_quote_excel
from core.quote_excel_import import parse_luc_giac_workbook
from core.quote_workbook_view import overview
from core.quotations import calculate_quote, customer_quote_snapshot
from scripts.make_hexa_layout_fixtures import (
    CARTON_GROUP_HEADER_ROWS_REAL,
    CARTON_GROUP_ONE_ROWS_REAL,
    CARTON_GROUP_THREE_ROWS_REAL,
    CARTON_GROUP_TWO_ROWS_REAL,
    CARTON_HEADER_ROW_REAL,
    CARTON_ROWS_REAL,
    CARTON_TOTAL_ROW_REAL,
    GOODS_CODES_REAL,
    PACKING_HEADER_ROW_REAL,
    PACKING_ROWS_REAL,
    PALLET_ALLOCATION_ROWS_REAL,
    PALLET_GROUP_ONE_ROWS_REAL,
    PALLET_GROUP_TWO_ROWS_REAL,
    PALLET_LINES_REAL,
    PALLET_REPEAT_HEADER_ROW_REAL,
    PI_HEADER_ROW_REAL,
    PI_ROWS_REAL,
    SENTINEL,
    SHEET_CARTONS,
    SHEET_INSTRUCTION,
    SHEET_PACKING_REAL,
    SHEET_PALLETS,
    SHEET_PI_MONEY_REAL,
    SHEET_PI_REAL,
    SHEET_TOTAL,
    TOTAL_FIRST_ROW_REAL,
    TOTAL_HEADER_ROW_REAL,
    TOTAL_ROWS_REAL,
    build_multilevel_packing_real_layout,
    build_packing_list_real_headers,
    build_proforma_invoice_real_headers,
)

TODAY = date(2026, 10, 16)



@pytest.fixture(scope="module")
def file_a() -> bytes:
    return build_packing_list_real_headers()


@pytest.fixture(scope="module")
def file_b() -> bytes:
    return build_proforma_invoice_real_headers()


@pytest.fixture(scope="module")
def file_c() -> bytes:
    return build_multilevel_packing_real_layout()


def parse(data: bytes, file_name: str) -> dict:
    return parse_luc_giac_workbook(data, file_name, today=TODAY)


def warning_ids(proposal: dict) -> set[str]:
    return {w["id"] for w in proposal["source_warnings"]}


def unknown_codes(proposal: dict) -> set[str]:
    return {u["code"] for u in proposal["review"]["unknowns"]}


def level_by_sheet(proposal: dict) -> dict[str, dict]:
    return {level["sheet"]: level for level in proposal["review"]["levels"]}


# ---------------------------------------------------------------- fixture reality
def test_fixtures_carry_the_real_header_wording(file_a, file_b, file_c):
    wb = load_workbook(BytesIO(file_a))
    assert wb.sheetnames == [SHEET_PACKING_REAL]
    assert wb[SHEET_PACKING_REAL][f"A{PACKING_HEADER_ROW_REAL}"].value == "Item number"
    assert wb[SHEET_PACKING_REAL][f"H{PACKING_HEADER_ROW_REAL}"].value == "CBM"

    wb = load_workbook(BytesIO(file_b))
    assert wb.sheetnames == [SHEET_PI_REAL, SHEET_PI_MONEY_REAL]
    assert wb[SHEET_PI_REAL][f"A{PI_HEADER_ROW_REAL}"].value == "Model 型号"
    assert wb[SHEET_PI_REAL][f"D{PI_HEADER_ROW_REAL}"].value == "Unit Price 单价 (CNY)"

    wb = load_workbook(BytesIO(file_c))
    assert wb.sheetnames == [SHEET_TOTAL, SHEET_PALLETS, SHEET_CARTONS, SHEET_INSTRUCTION]
    assert wb[SHEET_TOTAL][f"B{TOTAL_HEADER_ROW_REAL}"].value == "Name product-Description"
    assert wb[SHEET_TOTAL][f"C{TOTAL_HEADER_ROW_REAL}"].value == "数量（件数）"
    assert wb[SHEET_PALLETS]["A1"].value == "The pallet number"
    assert wb[SHEET_PALLETS][f"A{PALLET_REPEAT_HEADER_ROW_REAL}"].value == "The pallet number"
    assert wb[SHEET_CARTONS][f"D{CARTON_HEADER_ROW_REAL}"].value == "Size(mm x mm x mm)"
    assert wb[SHEET_CARTONS][f"D{CARTON_HEADER_ROW_REAL + 1}"].value == "W"
    assert wb[SHEET_CARTONS][f"F{CARTON_HEADER_ROW_REAL + 1}"].value == "H"
    assert wb[SHEET_CARTONS][f"A{CARTON_GROUP_HEADER_ROWS_REAL[1]}"].value == "REF NO"
    assert wb[SHEET_CARTONS][f"A{CARTON_GROUP_HEADER_ROWS_REAL[2]}"].value == "REF NO"
    assert wb[SHEET_INSTRUCTION].sheet_state == "hidden"


# ------------------------------------------- C: the three levels of the real file
def test_real_file_c_recognises_goods_pallet_and_carton_levels(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")

    assert proposal["source"]["layout"] == "MULTI_LEVEL_PACKING"
    assert proposal["review"]["layout"] == "MULTI_LEVEL_PACKING"
    assert proposal["source"]["goods_sheet"] == SHEET_TOTAL
    assert len(proposal["goods"]) == len(TOTAL_ROWS_REAL)

    levels = level_by_sheet(proposal)
    assert levels[SHEET_PALLETS]["role"] == "PALLET_DETAIL"
    assert levels[SHEET_CARTONS]["role"] == "CARTON_DETAIL"
    assert levels[SHEET_INSTRUCTION]["role"] == "HIDDEN_SHEET"
    assert levels[SHEET_INSTRUCTION]["hidden"] is True
    assert "level-not-summed" in warning_ids(proposal)


def test_real_file_c_pallet_level_is_counted_from_real_rows_only(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    pallet = level_by_sheet(proposal)[SHEET_PALLETS]

    assert pallet["row_count"] == len(PALLET_LINES_REAL) == len(PALLET_ALLOCATION_ROWS_REAL)
    # The header repeated at row 9 is structure, not a pallet line.
    assert pallet["source_rows"] == list(PALLET_ALLOCATION_ROWS_REAL)
    assert PALLET_REPEAT_HEADER_ROW_REAL not in pallet["source_rows"]

    refs = {row["ref"] for row in proposal["review"]["rows"]}
    assert f"{SHEET_PALLETS}!2" in refs
    assert f"{SHEET_PALLETS}!{PALLET_REPEAT_HEADER_ROW_REAL}" not in refs


def test_real_file_c_carton_level_skips_the_axis_row_and_repeated_headers(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    carton = level_by_sheet(proposal)[SHEET_CARTONS]

    # The premise of the whole test: this sheet IS the carton level. Before the
    # repair it was read as an extra goods list called UNCLASSIFIED_TABLE.
    assert carton["role"] == "CARTON_DETAIL"
    assert carton["row_count"] == len(CARTON_ROWS_REAL)
    assert carton["source_rows"] == list(
        CARTON_GROUP_ONE_ROWS_REAL + CARTON_GROUP_TWO_ROWS_REAL + CARTON_GROUP_THREE_ROWS_REAL)
    skipped = {CARTON_TOTAL_ROW_REAL}
    for header_row in CARTON_GROUP_HEADER_ROWS_REAL:
        skipped.update({header_row, header_row + 1})  # repeated header + W/D/H axis row
    assert set(carton["source_rows"]) & skipped == set()


def test_real_file_c_repeated_headers_never_become_goods_rows(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")

    assert [item["source_row"] for item in proposal["goods"]] == list(
        range(TOTAL_FIRST_ROW_REAL, TOTAL_FIRST_ROW_REAL + len(TOTAL_ROWS_REAL)))
    texts = " ".join(str(item["description"]) for item in proposal["goods"])
    for header_text in ("Name product-Description", "数量（件数）", "净重", "毛重", "体积"):
        assert header_text not in texts
    # The pallet/carton sheets never contribute a goods line of their own.
    assert all(item["source_sheet"] == SHEET_TOTAL for item in proposal["goods"])


def test_real_file_c_no_double_counting_of_qty_cbm_or_weight(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")

    # Weights and CBM stay review evidence; nothing is summed into the shipment.
    assert all("volume_cbm" not in item and "gross_weight_kg" not in item
               for item in proposal["goods"])
    assert proposal["shipment"]["cbm"] is None
    assert proposal["shipment"]["weight_kg"] is None

    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["row_count"] == len(TOTAL_ROWS_REAL)
    assert calculation["quantity"] == sum(row[1] for row in TOTAL_ROWS_REAL)

    carton_cbm_total = round(sum(row[8] for row in CARTON_ROWS_REAL), 3)
    goods_cbm_total = sum(row[4] for row in TOTAL_ROWS_REAL)
    assert carton_cbm_total != goods_cbm_total  # the levels really do differ
    serialized = json.dumps(proposal["shipment"])
    assert str(carton_cbm_total) not in serialized
    assert str(round(goods_cbm_total, 3)) not in serialized


def test_real_file_c_ten_pallet_allocation_lines_are_all_kept(file_c):
    """The real sheet names the pallet once per group; every line must survive."""
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    pallet_rows = [row for row in proposal["review"]["rows"]
                   if row["ref"].startswith(f"{SHEET_PALLETS}!")]

    assert len(pallet_rows) == 10
    assert [row["source_row"] for row in pallet_rows] == list(PALLET_ALLOCATION_ROWS_REAL)
    assert [row["contract"] for row in pallet_rows] == [line[0] for line in PALLET_LINES_REAL]

    # The pallet context is carried forward inside a group, and never across the
    # repeated group header at row 9.
    first = [row for row in pallet_rows if row["source_row"] in PALLET_GROUP_ONE_ROWS_REAL]
    second = [row for row in pallet_rows if row["source_row"] in PALLET_GROUP_TWO_ROWS_REAL]
    assert {row["pallet"] for row in first} == {"PLT-01"} and len(first) == 7
    assert {row["pallet"] for row in second} == {"PLT-02"} and len(second) == 3

    # PC / Total are the pallet level's own counts: kept apart, never added to qty.
    assert [row["pieces"] for row in pallet_rows] == [line[4] for line in PALLET_LINES_REAL]
    assert [row["total"] for row in pallet_rows] == [line[5] for line in PALLET_LINES_REAL]


def test_real_file_c_ten_contract_codes_link_goods_to_pallet(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    links = proposal["review"]["links"]

    assert sorted(links) == sorted(GOODS_CODES_REAL)
    expected_pallets = {"PLT-01": 7, "PLT-02": 3}
    seen: dict[str, int] = {}
    for code, per_sheet in links.items():
        assert list(per_sheet[SHEET_PALLETS]) , code
        pallet = per_sheet[SHEET_PALLETS][0]
        seen[pallet] = seen.get(pallet, 0) + 1
    assert seen == expected_pallets
    # The goods sheet is the source of the codes, not the level sheets.
    assert {item["model"] or item["description"] for item in proposal["goods"]} == set(GOODS_CODES_REAL)
    assert "level_link_missing" not in unknown_codes(proposal)


def test_real_file_c_thirty_six_carton_lines_link_by_ref_no(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    carton = level_by_sheet(proposal)[SHEET_CARTONS]

    assert carton["row_count"] == 36
    assert len(carton["codes"]) == len(GOODS_CODES_REAL)
    links = proposal["review"]["links"]
    carton_links = sum(len(per_sheet.get(SHEET_CARTONS, [])) for per_sheet in links.values())
    assert carton_links == 36
    assert {row["source_row"] for row in proposal["review"]["rows"]
            if row["ref"].startswith(f"{SHEET_CARTONS}!")} == set(
        CARTON_GROUP_ONE_ROWS_REAL + CARTON_GROUP_TWO_ROWS_REAL + CARTON_GROUP_THREE_ROWS_REAL)
    # Every REF NO in the source is one of the summary codes, and none is dropped.
    assert [row[0] for row in CARTON_ROWS_REAL] == [
        row["ref_no"] for row in proposal["review"]["rows"]
        if row["ref"].startswith(f"{SHEET_CARTONS}!")]


def test_real_file_c_ambiguous_or_unknown_code_is_never_guessed():
    """A duplicate / unknown reference code must fall back to NEEDS_REVIEW."""
    from io import BytesIO

    from openpyxl import Workbook

    from core.quote_excel_import import parse_luc_giac_workbook

    wb = Workbook()
    goods = wb.active
    goods.title = SHEET_TOTAL
    goods["A1"], goods["B1"], goods["C1"], goods["D1"] = (
        "序号", "Name product-Description", "数量（件数）", "体积")
    goods["A2"], goods["B2"], goods["C2"], goods["D2"] = 1, "CODE-A", 5, 1.0
    # The same code on two goods lines cannot be linked unambiguously.
    goods["A3"], goods["B3"], goods["C3"], goods["D3"] = 2, "CODE-A", 7, 1.5
    goods["A4"], goods["B4"], goods["C4"], goods["D4"] = 3, "CODE-B", 9, 2.0
    pallets = wb.create_sheet(SHEET_PALLETS)
    for column, title in {"A": "The pallet number", "B": "Pallet size", "C": "contract",
                          "G": "PC", "H": "Total"}.items():
        pallets[f"{column}1"] = title
    pallets["A2"], pallets["C2"], pallets["G2"], pallets["H2"] = "PLT-01", "CODE-A", 5, 5
    pallets["C3"], pallets["G3"], pallets["H3"] = "CODE-ZZ", 1, 1  # matches no goods line
    output = BytesIO()
    wb.save(output)

    proposal = parse_luc_giac_workbook(output.getvalue(), "ambiguous.xlsx", today=TODAY)
    links = proposal["review"]["links"]
    # The ambiguous code is not linked, and the unknown one creates no goods row.
    assert "CODE-A" not in links
    assert "CODE-ZZ" not in links
    assert len(proposal["goods"]) == 3
    assert "level-link-ambiguous" in warning_ids(proposal)
    assert "level-link-unmatched" in warning_ids(proposal)
    assert "level_link_missing" in unknown_codes(proposal)


def test_real_file_c_carton_axis_row_is_a_carton_size_not_a_product_size(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    carton_rows = {row["ref"]: row for row in proposal["review"]["rows"]
                   if row["ref"].startswith(f"{SHEET_CARTONS}!")}

    assert carton_rows[f"{SHEET_CARTONS}!4"]["carton_size_mm"] == [860.0, 540.0, 900.0]
    assert carton_rows[f"{SHEET_CARTONS}!4"]["volume_cbm"] == 0.42
    assert all("product_size" not in row for row in carton_rows.values())
    # A carton size is never promoted into the customer goods lines.
    assert all(item["length_mm"] is None and item["width_mm"] is None
               and item["height_mm"] is None for item in proposal["goods"])
    assert "dims-not-package" not in warning_ids(proposal)


def test_real_file_c_missing_price_stays_unknown_never_zero(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")

    assert all(item["unit_price"] is None for item in proposal["goods"])
    assert all(item["unit_price"] != 0 for item in proposal["goods"])
    assert "unit-price-missing" in warning_ids(proposal)
    assert "unit_price_missing" in unknown_codes(proposal)


def test_real_file_c_hidden_and_internal_data_never_reach_the_customer(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    assert SENTINEL not in json.dumps(proposal["goods"], ensure_ascii=False)
    assert SENTINEL not in json.dumps(proposal["review"], ensure_ascii=False)

    quote = calculate_quote({"excel_quote": proposal, "status": "DRAFT",
                             "customer_name": "Công ty thử nghiệm (không có thật)"})
    assert SENTINEL not in json.dumps(customer_quote_snapshot(quote), ensure_ascii=False)

    exported = generate_quote_excel(quote).getvalue()
    assert SENTINEL.encode() not in exported
    with ZipFile(BytesIO(exported)) as archive:
        package_text = "".join(archive.read(name).decode("utf-8", "ignore")
                               for name in archive.namelist() if name.endswith(".xml"))
    assert SENTINEL not in package_text
    assert SHEET_INSTRUCTION not in package_text
    assert SHEET_PALLETS not in package_text
    assert SHEET_CARTONS not in package_text


def test_real_file_c_viewer_treats_only_the_goods_sheet_as_customer(file_c):
    proposal = parse(file_c, "multilevel-REAL-LAYOUT-SYNTHETIC.xlsx")
    info = overview(file_c, proposal)

    assert info["customer_sheet"] == SHEET_TOTAL
    by_name = {sheet["name"]: sheet for sheet in info["sheets"]}
    assert by_name[SHEET_TOTAL]["classification"] == "CUSTOMER"
    for other in (SHEET_PALLETS, SHEET_CARTONS, SHEET_INSTRUCTION):
        assert by_name[other]["classification"] != "CUSTOMER"


# -------------------------------------------------------------- A: no regression
def test_real_file_a_packing_list_still_imports_with_price_unknown(file_a):
    proposal = parse(file_a, "packing-list-REAL-HEADERS-SYNTHETIC.xlsx")

    assert proposal["source"]["layout"] == "PACKING_LIST"
    assert proposal["source"]["goods_sheet"] == SHEET_PACKING_REAL
    assert [item["model"] for item in proposal["goods"]] == [row[0] for row in PACKING_ROWS_REAL]
    assert [item["qty"] for item in proposal["goods"]] == [row[3] for row in PACKING_ROWS_REAL]
    assert all(item["unit_price"] is None for item in proposal["goods"])
    assert "unit-price-missing" in warning_ids(proposal)

    # `Size` is the PRODUCT size: review evidence, never a carton dimension.
    review = {row["ref"]: row for row in proposal["review"]["rows"]}
    assert review[f"{SHEET_PACKING_REAL}!8"]["product_size"] == "800x480x850"
    assert review[f"{SHEET_PACKING_REAL}!8"]["cartons"] == 6.0
    assert review[f"{SHEET_PACKING_REAL}!8"]["volume_cbm"] == 2.51
    assert all(item["length_mm"] is None for item in proposal["goods"])
    assert "CARTON_DETAIL" not in {level["role"] for level in proposal["review"]["levels"]}
    assert "PALLET_DETAIL" not in {level["role"] for level in proposal["review"]["levels"]}


# -------------------------------------------------------------- B: no regression
def test_real_file_b_pi_reads_one_goods_sheet_and_never_sums_the_money_sheet(file_b):
    proposal = parse(file_b, "proforma-invoice-REAL-HEADERS-SYNTHETIC.xlsx")

    assert proposal["source"]["layout"] == "PROFORMA_INVOICE"
    assert proposal["source"]["goods_sheet"] == SHEET_PI_REAL
    assert len(proposal["goods"]) == len(PI_ROWS_REAL)
    assert [item["unit_price"] for item in proposal["goods"]] == [row[3] for row in PI_ROWS_REAL]
    assert [item["source_row"] for item in proposal["goods"]] == list(
        range(PI_HEADER_ROW_REAL + 1, PI_HEADER_ROW_REAL + 1 + len(PI_ROWS_REAL)))

    levels = level_by_sheet(proposal)
    assert levels[SHEET_PI_MONEY_REAL]["role"] == "INTERNAL_CALC"
    assert levels[SHEET_PI_MONEY_REAL]["hidden"] is False
    assert "money-sheet-not-imported" in warning_ids(proposal)

    calculation = calculate_quote({"excel_quote": proposal})["excel_quote"]["calculation"]
    assert calculation["row_count"] == len(PI_ROWS_REAL)
    assert calculation["quantity"] == sum(row[2] for row in PI_ROWS_REAL)


def test_real_file_b_internal_money_sheet_never_reaches_the_customer(file_b):
    proposal = parse(file_b, "proforma-invoice-REAL-HEADERS-SYNTHETIC.xlsx")
    quote = calculate_quote({"excel_quote": proposal, "status": "DRAFT",
                             "customer_name": "Công ty thử nghiệm (không có thật)"})
    assert SENTINEL not in json.dumps(customer_quote_snapshot(quote), ensure_ascii=False)
    exported = generate_quote_excel(quote).getvalue()
    assert SENTINEL.encode() not in exported


def test_real_file_b_uncached_amount_formula_is_not_invented(file_b):
    proposal = parse(file_b, "proforma-invoice-REAL-HEADERS-SYNTHETIC.xlsx")
    # The amount column is a formula with no cached value: no amount is summed
    # into the quotation from it, and the gap is reported instead of guessed.
    assert "formula-gap" in warning_ids(proposal)
    assert all("line_total" not in item for item in proposal["goods"])
