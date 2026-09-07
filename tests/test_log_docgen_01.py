"""Focused tests for LOG-DOCGEN-01 V1 Shipment Document Pack Generator.

Covers the locked owner rules, readiness (per-document identity + facts),
explicit PXK source, HDNT reuse, OOXML structural QA, RBAC, versioning,
linkage, ISSUED-writer immutability, finance safety, and template integrity.
Uses only synthetic data (no Production data) and never writes to Finance
Actual.
"""

from __future__ import annotations

import io
import os
from copy import deepcopy
from datetime import date
from decimal import Decimal
from unittest import TestCase
from unittest.mock import patch

from core.docgen import generator as gen
from core.docgen.rules import (
    BLOCKED,
    GENERATED,
    HDMB,
    HDNT,
    ISSUED,
    MISSING_INFO,
    NOT_APPLICABLE,
    PL,
    PXK,
    READY,
    VND_THRESHOLD,
    classify_valid_hdnt,
    compute_goods_totals,
    decide_applicable_documents,
    document_identity_missing,
    evaluate_document_readiness,
    normalize_goods_line,
    total_after_vat,
)

TPL_XLSX = "/tmp/log-docgen-01-assets/truong-thanh-agent.xlsx"
TPL_DOCX = "/tmp/log-docgen-01-assets/HDNT-TT-agent.docx"
_HAS_ASSETS = os.path.exists(TPL_XLSX) and os.path.exists(TPL_DOCX)


def _line(description: str, qty, price, unit="Cái", vat="0.08") -> dict:
    return {"description": description, "quantity": qty, "unit_price_vnd_excl_vat": price, "unit": unit, "vat_rate": Decimal(vat)}


def _agent() -> dict:
    return {
        "agent_id": "AGT-001",
        "legal_name": "CÔNG TY TNHH PHÚ ANH LGT",
        "legal_address": "Số 1 Đường ABC, Quận 1, TP. Hồ Chí Minh",
        "legal_code": "0123456789",
        "legal_representative": "Ông Nguyễn Văn An",
        "representative_title": "Giám đốc",
        "bank_account_holder": "PHÚ ANH LGT",
        "bank_account_number": "1234567890",
        "bank_name": "Vietcombank",
        "bank_branch": "CN Hà Nội",
        "phone": "0901234567",
        "email": "pa@example.com",
    }


def _company() -> dict:
    return {
        "company_id": "TT-001",
        "legal_name": "CÔNG TY TNHH GIAO NHẬN QUỐC TẾ TRƯỜNG THÀNH",
        "legal_address": "Số nhà 16 ngõ 41 đường Khương Đình, Phường Khương Đình, Hà Nội, Việt Nam",
        "tax_or_legal_code": "0105937262",
        "legal_representative": "Bà Lê Thị Thu Hương",
        "representative_title": "Giám đốc",
    }


def _hdnt_identity(no: str = "HDNT-001", d: str = "2026-09-01") -> dict:
    return {"hdnt_no": no, "hdnt_date": d, "valid_from": "2026-09-01", "valid_to": "2027-08-31", "currency": "VND"}


def _pl_identity(no: str = "PL-001", d: str = "2026-09-05") -> dict:
    return {"pl_no": no, "pl_date": d, "currency": "VND"}


def _hdmb_identity(no: str = "HDMB-001", d: str = "2026-09-05") -> dict:
    return {"hdmb_no": no, "hdmb_date": d, "currency": "VND"}


def _pxk_identity(no: str = "PXK-001", d: str = "2026-09-10") -> dict:
    return {"pxk_no": no, "pxk_date": d, "currency": "VND"}


def _pxk_source(doc_type=HDMB, line_count=1, line=None, source_id="SHP-001-HDMB-AGT-001", rev=1) -> dict:
    return {
        "source_document_type": doc_type,
        "source_contract_id": source_id,
        "source_revision": rev,
        "source_line_count": line_count,
        "source_goods_line": line or {"description": "Máy", "unit": "Cái", "quantity": "1", "unit_price_vnd_excl_vat": 60000000, "amount_vnd_excl_vat": 60000000},
    }


def _pxk_delivery(**overrides) -> dict:
    base = {"delivery_completed": True, "receiver_name": "Bà Hương", "receiver_title": "Giám đốc", "delivery_location": "Kho Trường Thành"}
    base.update(overrides)
    return base


def _parent(id_="HDNT-001", agent_id="AGT-001", company_id="TT-001") -> dict:
    return {
        "id": id_,
        "contract_no": id_,
        "hdnt_no": id_,
        "hdnt_date": "2026-09-01",
        "contract_date": "2026-09-01",
        "kind": HDNT,
        "agent_id": agent_id,
        "company_id": company_id,
        "valid_from": "2026-09-01",
        "valid_to": "2027-08-31",
    }


class ThresholdRulesTests(TestCase):
    def test_below_50m_applies_hdnt_pl(self) -> None:
        lines = [_line("Hàng A", 1, Decimal("4999999.99"))]
        decision = decide_applicable_documents(goods_lines=lines)
        self.assertEqual(decision[HDNT].state, READY)
        self.assertEqual(decision[PL].state, READY)
        self.assertEqual(decision[HDMB].state, NOT_APPLICABLE)

    def test_exactly_50m_applies_hdmb(self) -> None:
        lines = [_line("Hàng B", 1, Decimal("46296296"))]
        decision = decide_applicable_documents(goods_lines=lines)
        self.assertEqual(decision[HDMB].state, READY)
        self.assertIn(HDMB, [d for d in (HDMB, HDNT, PL) if decision[d].state == READY])

    def test_above_50m_applies_hdmb_only(self) -> None:
        lines = [_line("Máy móc", 1, Decimal("100000000"))]
        decision = decide_applicable_documents(goods_lines=lines)
        self.assertEqual(decision[HDMB].state, READY)
        self.assertEqual(decision[HDNT].state, NOT_APPLICABLE)
        self.assertEqual(decision[PL].state, NOT_APPLICABLE)

    def test_non_vnd_blocked(self) -> None:
        lines = [_line("Hàng A", 1, Decimal("100000"))]
        decision = decide_applicable_documents(goods_lines=lines, currency="USD")
        for doc_type in (HDNT, PL, HDMB):
            self.assertEqual(decision[doc_type].state, BLOCKED)

    def test_unknown_total_blocked(self) -> None:
        lines = [{"description": "", "quantity": -1, "unit_price_vnd_excl_vat": 100, "unit": "", "vat_rate": "0.08"}]
        decision = decide_applicable_documents(goods_lines=lines)
        for doc_type in (HDNT, PL, HDMB):
            self.assertEqual(decision[doc_type].state, BLOCKED)


class VatRulesTests(TestCase):
    def test_vat_8_percent_calc(self) -> None:
        lines = [_line("Hàng", 10, Decimal("100000"))]
        totals, _ = compute_goods_totals(lines)
        self.assertEqual(totals.subtotal, Decimal("1000000"))
        self.assertEqual(totals.vat_amount, Decimal("80000"))
        self.assertEqual(totals.total_after_vat, Decimal("1080000"))

    def test_mixed_vat_blocked(self) -> None:
        lines = [_line("A", 1, Decimal("100000")), _line("B", 1, Decimal("100000"), vat="0.10")]
        totals, normalized = compute_goods_totals(lines)
        self.assertTrue(any(not item["valid"] for item in normalized))

    def test_other_vat_blocked(self) -> None:
        lines = [_line("A", 1, Decimal("100000"), vat="0.10")]
        totals, normalized = compute_goods_totals(lines)
        self.assertFalse(normalized[0]["valid"])
        self.assertIn("Chỉ hỗ trợ thuế GTGT 8%", normalized[0]["validation_errors"])

    def test_unit_price_hr_rounding(self) -> None:
        lines = [_line("Hàng", 3, Decimal("33333.33"))]
        totals, _ = compute_goods_totals(lines)
        self.assertEqual(totals.subtotal, Decimal("100000"))
        self.assertEqual(totals.total_after_vat, Decimal("108000"))


class AuthorityRulesTests(TestCase):
    def test_missing_agent_master_blocked(self) -> None:
        readiness = evaluate_document_readiness(
            HDNT, goods_lines=[_line("A", 1, Decimal("100000"))], agent=None, company=_company()
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("Đại lý" in item for item in readiness.missing))

    def test_customer_cannot_silently_substitute_for_agent(self) -> None:
        customer = {"name": "Công ty Khách", "address": "abc", "tax_code": "999"}
        readiness = evaluate_document_readiness(
            HDNT, goods_lines=[_line("A", 1, Decimal("100000"))], agent=customer, company=_company()
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("Đại lý" in item for item in readiness.missing))

    def test_factory_cannot_silently_substitute_for_agent(self) -> None:
        factory = {"factory": "Nhà máy X", "contact": "abc"}
        readiness = evaluate_document_readiness(
            HDNT, goods_lines=[_line("A", 1, Decimal("100000"))], agent=factory, company=_company()
        )
        self.assertEqual(readiness.state, MISSING_INFO)


class LegalIdentityTests(TestCase):
    """Each legal document owns its own number/date; absent = MISSING_INFO."""

    def test_hdnt_missing_number_not_ready(self) -> None:
        identity = dict(_hdnt_identity(no="", d="2026-09-01"))
        readiness = evaluate_document_readiness(
            HDNT, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(), document_identity=identity
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("số" in m.lower() and "nguyên tắc" in m for m in readiness.missing))

    def test_hdnt_missing_validity_not_ready(self) -> None:
        identity = dict(_hdnt_identity())
        identity["valid_to"] = ""
        readiness = evaluate_document_readiness(
            HDNT, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(), document_identity=identity
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("hiệu lực" in m for m in readiness.missing))

    def test_pl_missing_number_not_ready(self) -> None:
        readiness = evaluate_document_readiness(
            PL, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(),
            parent_contract=_parent(), document_identity={"pl_no": "", "pl_date": "", "currency": "VND"},
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("Phụ lục" in m or "PL" in m or "Số" in m for m in readiness.missing))

    def test_hdmb_separate_number_not_shared_with_hdnt(self) -> None:
        # HDNT identity must not satisfy HDMB.
        hdnt_identity = dict(_hdnt_identity(no="HDNT-ABC"))
        hdmb_readiness = evaluate_document_readiness(
            HDMB, goods_lines=[_line("Máy", 1, Decimal("60000000"))], agent=_agent(), company=_company(), document_identity=hdnt_identity
        )
        self.assertEqual(hdmb_readiness.state, MISSING_INFO)
        self.assertTrue(any("mua bán" in m for m in hdmb_readiness.missing))

    def test_document_identity_missing_returns_hdnt_and_hdmb_own_fields(self) -> None:
        missing_hdnt = document_identity_missing(HDNT, {})
        missing_hdmb = document_identity_missing(HDMB, {})
        self.assertTrue(any("nguyên tắc" in m for m in missing_hdnt))
        self.assertTrue(any("mua bán" in m for m in missing_hdmb))


class UnknownNotZeroTests(TestCase):
    def test_blank_quantity_is_unknown_not_zero(self) -> None:
        line = normalize_goods_line({"description": "A", "unit": "Cái", "quantity": "", "unit_price_vnd_excl_vat": "100000", "vat_rate": "0.08"}, 1)
        self.assertFalse(line["valid"])
        self.assertIn("UNKNOWN", line["validation_errors"][0])

    def test_blank_price_is_unknown_not_zero(self) -> None:
        line = normalize_goods_line({"description": "A", "unit": "Cái", "quantity": "1", "unit_price_vnd_excl_vat": "", "vat_rate": "0.08"}, 1)
        self.assertFalse(line["valid"])
        self.assertEqual(line["unit_price_vnd_excl_vat"], None)

    def test_explicit_zero_price_is_valid(self) -> None:
        line = normalize_goods_line({"description": "A", "unit": "Cái", "quantity": "1", "unit_price_vnd_excl_vat": "0", "vat_rate": "0.08"}, 1)
        self.assertTrue(line["valid"])
        self.assertEqual(line["unit_price_vnd_excl_vat"], 0)

    def test_amount_words_zero_guard(self) -> None:
        self.assertEqual(gen.amount_in_words_vnd(0), "không đồng")

    def test_unknown_total_is_not_zero(self) -> None:
        lines = [{"description": "", "quantity": -1, "unit_price_vnd_excl_vat": 100, "unit": "", "vat_rate": "0.08"}]
        self.assertIsNone(total_after_vat(lines))


class GoodsCountRulesTests(TestCase):
    def test_hdmb_one_line_ready(self) -> None:
        decision = decide_applicable_documents(goods_lines=[_line("Máy", 1, Decimal("60000000"))])
        self.assertEqual(decision[HDMB].state, READY)

    def test_hdmb_more_than_one_blocked(self) -> None:
        lines = [_line("A", 1, Decimal("30000000")), _line("B", 1, Decimal("30000000"))]
        decision = decide_applicable_documents(goods_lines=lines)
        self.assertEqual(decision[HDMB].state, BLOCKED)

    def test_pl_18_lines_ready(self) -> None:
        lines = [_line(f"H{i}", 1, Decimal("100000")) for i in range(18)]
        decision = decide_applicable_documents(goods_lines=lines)
        self.assertEqual(decision[PL].state, READY)

    def test_pl_19_lines_blocked(self) -> None:
        lines = [_line(f"H{i}", 1, Decimal("100000")) for i in range(19)]
        decision = decide_applicable_documents(goods_lines=lines)
        self.assertEqual(decision[PL].state, BLOCKED)


class HdntPlRulesTests(TestCase):
    def test_pl_without_parent_blocked(self) -> None:
        readiness = evaluate_document_readiness(
            PL, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(), parent_contract=None
        )
        self.assertEqual(readiness.state, BLOCKED)

    def test_pl_parent_agent_mismatch_blocked(self) -> None:
        parent = _parent(agent_id="AGT-XYZ")
        readiness = evaluate_document_readiness(
            PL, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(), parent_contract=parent, document_identity=_pl_identity()
        )
        self.assertEqual(readiness.state, BLOCKED)


class HdntReuseTests(TestCase):
    def test_single_valid_hdnt_selects_reuse_single(self) -> None:
        candidates = {"HDNT-001": _parent("HDNT-001")}
        decision = classify_valid_hdnt(candidates, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 1)
        self.assertFalse(decision["needs_selection"])
        self.assertEqual(decision["reuse_target"]["hdnt_no"], "HDNT-001")

    def test_two_simultaneously_valid_hdnt_requires_explicit_selection(self) -> None:
        candidates = {
            "HDNT-0001": _parent("HDNT-0001"),
            "HDNT-0002": _parent("HDNT-0002"),
        }
        decision = classify_valid_hdnt(candidates, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 2)
        self.assertTrue(decision["needs_selection"])
        self.assertIsNone(decision["reuse_target"])
        self.assertEqual(len(decision["valid_hdnt"]), 2)

    def test_wrong_agent_hdnt_not_valid(self) -> None:
        candidates = {"HDNT-999": _parent("HDNT-999", agent_id="AGT-OTHER")}
        decision = classify_valid_hdnt(candidates, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)

    def test_expired_hdnt_not_valid(self) -> None:
        parent = _parent("HDNT-EXP")
        parent["valid_to"] = "2020-01-01"
        decision = classify_valid_hdnt({"HDNT-EXP": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)


class PxkRulesTests(TestCase):
    def test_not_delivered_blocked(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(), delivery_completed=False
        )
        self.assertEqual(readiness.state, BLOCKED)

    def test_missing_source_requires_explicit_selection(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(),
            delivery_completed=True, document_identity=_pxk_identity(), delivery=_pxk_delivery(), pxk_source=None,
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("nguồn" in m.lower() for m in readiness.missing))

    def test_missing_receiver_not_ready(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(),
            delivery_completed=True, document_identity=_pxk_identity(),
            delivery=_pxk_delivery(receiver_name=""), pxk_source=_pxk_source(),
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("người nhận" in m.lower() for m in readiness.missing))

    def test_missing_location_not_ready(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(),
            delivery_completed=True, document_identity=_pxk_identity(),
            delivery=_pxk_delivery(delivery_location=""), pxk_source=_pxk_source(),
        )
        self.assertEqual(readiness.state, MISSING_INFO)
        self.assertTrue(any("địa điểm" in m.lower() for m in readiness.missing))

    def test_confirmed_delivery_full_facts_ready(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("Máy", 1, Decimal("60000000"))], agent=_agent(), company=_company(),
            delivery_completed=True, document_identity=_pxk_identity(), delivery=_pxk_delivery(), pxk_source=_pxk_source(),
        )
        self.assertEqual(readiness.state, READY)

    def test_multiline_source_pxk_blocked(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(),
            delivery_completed=True, document_identity=_pxk_identity(), delivery=_pxk_delivery(),
            pxk_source=_pxk_source(line_count=2),
        )
        self.assertEqual(readiness.state, BLOCKED)

    def test_pl_source_with_multiple_lines_blocked(self) -> None:
        readiness = evaluate_document_readiness(
            PXK, goods_lines=[_line("A", 1, Decimal("100000"))], agent=_agent(), company=_company(),
            delivery_completed=True, document_identity=_pxk_identity(), delivery=_pxk_delivery(),
            pxk_source=_pxk_source(doc_type=PL, line_count=2),
        )
        self.assertEqual(readiness.state, BLOCKED)


class TemplateGenerationTests(TestCase):
    def setUp(self) -> None:
        if not _HAS_ASSETS:
            self.skipTest("Template assets not present")

    def test_source_hashes_unchanged_after_generation(self) -> None:
        before_xlsx = _sha256_file(TPL_XLSX)
        before_docx = _sha256_file(TPL_DOCX)
        gen.generate_xlsx_file(TPL_XLSX, PL, contract_no="PL-001", contract_date="2026-09-05", agent=_agent(), company=_company(), goods_lines=[_line("Hàng", 5, Decimal("100000"))], parent_contract={"contract_no": "HDNT-001", "contract_date": "2026-09-01", "id": "HDNT-001"})
        gen.generate_docx_file(TPL_DOCX, contract_no="HDNT-001", contract_date="2026-09-01", valid_from="2026-09-01", valid_to="2027-08-31", agent=_agent(), company=_company())
        self.assertEqual(_sha256_file(TPL_XLSX), before_xlsx)
        self.assertEqual(_sha256_file(TPL_DOCX), before_docx)

    def test_pl_output_opens_and_preserves_formulas(self) -> None:
        import openpyxl

        out = gen.generate_xlsx_file(TPL_XLSX, PL, contract_no="PL-001", contract_date="2026-09-05", agent=_agent(), company=_company(), goods_lines=[_line("Hàng", 5, Decimal("100000"))], parent_contract={"contract_no": "HDNT-001", "contract_date": "2026-09-01", "id": "HDNT-001"})
        wb = openpyxl.load_workbook(io.BytesIO(out), data_only=False)
        ws = wb["pl"]
        self.assertEqual(ws["F28"].value, "=SUM(F10:F27)")
        self.assertEqual(ws["F29"].value, "=F28*8%")
        self.assertEqual(ws["F30"].value, "=SUM(F28:F29)")
        self.assertEqual(ws["B10"].value, "Hàng")
        self.assertGreaterEqual(len(ws.merged_cells.ranges), 10)
        self.assertEqual(ws.page_setup.orientation, "portrait")

    def test_pxk_writes_explicit_values_not_formula(self) -> None:
        import openpyxl

        lines = [_line("Máy móc", 5, Decimal("20000000"), unit="Bộ")]
        out = gen.generate_xlsx_file(TPL_XLSX, PXK, contract_no="PXK-001", contract_date="2026-09-10", agent=_agent(), company=_company(), goods_lines=lines, delivery={"receiver_name": "Bà Hương", "receiver_title": "Giám đốc", "delivery_location": "Kho TT"})
        wb = openpyxl.load_workbook(io.BytesIO(out), data_only=False)
        ws = wb["pxk"]
        self.assertEqual(str(ws["B15"].value), "Máy móc")
        self.assertEqual(str(ws["C15"].value), "Bộ")
        self.assertEqual(int(ws["D15"].value), 5)
        self.assertFalse(str(ws["B15"].value).startswith("="))

    def test_pxk_explicit_source_drives_b15(self) -> None:
        import openpyxl

        payload = _pxk_source(doc_type=HDMB, source_id="SHP-001-HDMB-AGT-001", line={"description": "Máy chính hãng", "unit": "Bộ", "quantity": "2", "unit_price_vnd_excl_vat": 10000000})
        with _with_user({"role": "admin", "username": "admin", "user_id": "u1"}), patch("core.database.get_db", return_value=_Database()):
            result = gen.generate_document(
                PXK,
                shipment_id="SHP-001",
                agent=_agent(),
                company=_company(),
                goods_lines=[_line("Máy chính hãng", 2, Decimal("10000000"), unit="Bộ")],
                contract=_pxk_identity(),
                delivery=_pxk_delivery(),
                pxk_source=payload,
                template_path_xlsx=TPL_XLSX,
            )
        wb = openpyxl.load_workbook(io.BytesIO(result["bytes"]), data_only=False)
        ws = wb["pxk"]
        self.assertEqual(str(ws["B15"].value), "Máy chính hãng")
        self.assertEqual(str(ws["C15"].value), "Bộ")
        self.assertEqual(int(ws["D15"].value), 2)
        # source snapshot provenance contains the explicit source
        self.assertEqual(result["metadata"]["source_snapshot"]["source_document_type"], HDMB)
        self.assertEqual(result["metadata"]["source_snapshot"]["source_contract_id"], "SHP-001-HDMB-AGT-001")
        self.assertTrue(str(ws["B15"].value) != "" and not str(ws["B15"].value).startswith("="))

    def test_hdnt_docx_opens_with_expected_text(self) -> None:
        import docx

        out = gen.generate_docx_file(TPL_DOCX, contract_no="HDNT-001", contract_date="2026-09-01", valid_from="2026-09-01", valid_to="2027-08-31", agent=_agent(), company=_company())
        document = docx.Document(io.BytesIO(out))
        self.assertIn("Số: HDNT-001", document.paragraphs[2].text)
        self.assertIn("CÔNG TY TNHH PHÚ ANH LGT", document.paragraphs[5].text)
        self.assertEqual(len(document.sections), 1)
        self.assertIsNotNone(document.sections[0].top_margin)

    def test_sheet_parts_retained(self) -> None:
        import zipfile

        out = gen.generate_xlsx_file(TPL_XLSX, HDMB, contract_no="HDMB-001", contract_date="2026-09-05", agent=_agent(), company=_company(), goods_lines=[_line("Máy", 1, Decimal("100000000"))])
        names = zipfile.ZipFile(io.BytesIO(out)).namelist()
        self.assertIn("xl/printerSettings/printerSettings1.bin", names)
        self.assertIn("xl/calcChain.xml", names)
        self.assertIn("xl/sharedStrings.xml", names)


class VersionAndLinkageTests(TestCase):
    def test_revision_increments_across_regenerations(self) -> None:
        logical = gen._document_id(HDNT, "SHP-001", "AGT-001")
        existing = {
            f"{logical}::R1": {"document_id": logical, "generation_revision": 1},
            f"{logical}::R2": {"document_id": logical, "generation_revision": 2},
        }
        next_rev = gen._next_revision(HDNT, "SHP-001", "AGT-001", existing)
        self.assertEqual(next_rev, 3)

    def test_issued_immutable_prevents_second_issue(self) -> None:
        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "D::R1", "document_id": "D", "status": GENERATED, "generation_revision": 1}))
            result = gen.issue_document("D")
            self.assertEqual(result["status"], ISSUED)
            with self.assertRaises(gen.DocumentGenerationError):
                gen.issue_document("D")

    def test_direct_writer_cannot_mutate_issued(self) -> None:
        from core.database import save_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "E::R1", "document_id": "E", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("E")
            record = gen_load_latest("E")
            self.assertEqual(record["status"], ISSUED)
            # Attempt an arbitrary rewrite of content on the ISSUED record.
            hacked = deepcopy(record)
            hacked["source_snapshot"] = {"document_no": "HACKED"}
            self.assertFalse(save_docgen_document(hacked))
            # Attempt to rewrite the legal identity field.
            hacked2 = deepcopy(record)
            hacked2["output_hash"] = "deadbeef"
            self.assertFalse(save_docgen_document(hacked2))
            # The stored record is unchanged.
            self.assertEqual(gen_load_latest("E").get("output_hash"), record.get("output_hash"))

    def test_supersede_creates_new_logical_id_and_links_old(self) -> None:
        db = _Database()
        old_doc_id = gen._document_id(HDMB, "SHP-001", "AGT-001")  # deterministic old id
        new_doc_id = f"{old_doc_id}-V2"
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            record = {"id": f"{old_doc_id}::R1", "document_id": old_doc_id, "status": GENERATED, "generation_revision": 1, "shipment_id": "SHP-001"}
            self.assertTrue(_saved_docgen_document(record))
            gen.issue_document(old_doc_id)
            result = gen.supersede_document(
                old_doc_id,
                new_document_id=new_doc_id,
                document_type=HDMB,
                shipment_id="SHP-001",
                contract_no="HDMB-002",
                contract_date="2026-09-06",
                agent=_agent(),
                company=_company(),
                goods_lines=[_line("Máy", 1, Decimal("100000000"))],
                template_path_xlsx=TPL_XLSX,
            )
            # A NEW logical id is minted; it must differ from the old one.
            self.assertNotEqual(result["document_id"], old_doc_id)
            self.assertEqual(result["document_id"], new_doc_id)
            self.assertEqual(result["supersedes_document_id"], old_doc_id)
            new_rec = gen_load_latest(new_doc_id)
            self.assertEqual(new_rec.get("supersedes_document_id"), old_doc_id)
            old = gen_load_latest(old_doc_id)
            self.assertEqual(old.get("superseded_by_document_id"), new_doc_id)
            # Old record remains ISSUED (immutable) and points to the replacement.
            self.assertEqual(old.get("status"), ISSUED)


class ContractLinkageTests(TestCase):
    def test_nm_shipment_linkage_additive_and_preserves_scalar(self) -> None:
        from core.database import link_contract_to_shipment, load_contracts_for_shipment

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db), patch("core.auth.st", type("AuthStreamlit", (), {"session_state": {"current_user": {"role": "support", "username": "support", "user_id": "u1"}}})()):
            # Two distinct contract records (HDNT and PL) for the same shipment.
            self.assertTrue(_save_contract({"id": "HDNT-001", "contract_no": "HDNT-001", "kind": HDNT, "agent_id": "AGT-001", "company_id": "TT-001", "shipment_id": "SHP-001"}))
            self.assertTrue(_save_contract({"id": "PL-001", "contract_no": "PL-001", "kind": PL, "agent_id": "AGT-001", "company_id": "TT-001"}))
            # Additively link HDNT then PL to the same shipment.
            self.assertTrue(link_contract_to_shipment("HDNT-001", "SHP-001"))
            self.assertTrue(link_contract_to_shipment("PL-001", "SHP-001"))
            linked = load_contracts_for_shipment("SHP-001")
            self.assertIn("HDNT-001", linked)
            self.assertIn("PL-001", linked)
            # Both remain readable and the legacy scalar linkage is preserved.
            self.assertEqual(linked["HDNT-001"].get("shipment_id"), "SHP-001")
            self.assertEqual(linked["PL-001"].get("shipment_id"), "SHP-001")
            # Adding PL did NOT overwrite the HDNT linkage.
            self.assertIn("SHP-001", linked["HDNT-001"].get("shipment_ids", []))
            self.assertIn("SHP-001", linked["PL-001"].get("shipment_ids", []))


class RbacTests(TestCase):
    def _parent_contract(self) -> dict:
        return _parent()

    def _run_generate(self, user, doc_type=PL):
        with _with_user(user):
            try:
                contract = _pl_identity() if doc_type == PL else (_hdmb_identity() if doc_type == HDMB else _hdnt_identity())
                contract["id"] = contract.get(f"{doc_type.lower()}_no")
                return gen.generate_document(
                    doc_type,
                    shipment_id="SHP-001",
                    agent=_agent(),
                    company=_company(),
                    goods_lines=[_line("A", 1, Decimal("100000"))] if doc_type != HDMB else [_line("Máy", 1, Decimal("60000000"))],
                    contract=contract,
                    parent_contract=self._parent_contract() if doc_type == PL else None,
                    template_path_xlsx=TPL_XLSX,
                    template_path_docx=TPL_DOCX if doc_type == HDNT else None,
                )
            except Exception as exc:  # noqa: BLE001
                return exc

    def test_admin_generate_ok(self) -> None:
        if not _HAS_ASSETS:
            self.skipTest("no assets")
        result = self._run_generate({"role": "admin", "username": "admin", "user_id": "u1"})
        self.assertIsInstance(result, dict)
        self.assertEqual(result["status"], GENERATED)

    def test_support_generate_ok(self) -> None:
        if not _HAS_ASSETS:
            self.skipTest("no assets")
        result = self._run_generate({"role": "support", "username": "support", "user_id": "u2"})
        self.assertIsInstance(result, dict)

    def test_sale_can_generate(self) -> None:
        if not _HAS_ASSETS:
            self.skipTest("no assets")
        result = self._run_generate({"role": "sale", "username": "sale", "user_id": "u3"})
        self.assertIsInstance(result, dict)

    def test_sale_cannot_issue(self) -> None:
        from core.auth import PermissionDeniedError

        with _with_user({"role": "sale", "username": "sale", "user_id": "u3"}), patch("core.database.get_db", return_value=_Database()):
            with self.assertRaises(PermissionDeniedError):
                gen.issue_document("D")

    def test_admin_can_issue(self) -> None:
        db = _Database()
        with _with_user({"role": "admin", "username": "admin", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "D::R1", "document_id": "D", "status": GENERATED, "generation_revision": 1}))
            result = gen.issue_document("D")
            self.assertEqual(result["status"], ISSUED)

    def test_support_can_issue(self) -> None:
        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u2"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "D::R1", "document_id": "D", "status": GENERATED, "generation_revision": 1}))
            result = gen.issue_document("D")
            self.assertEqual(result["status"], ISSUED)

    def test_viewer_generate_denied(self) -> None:
        result = self._run_generate({"role": "viewer", "username": "viewer", "user_id": "u4"})
        self.assertIsInstance(result, PermissionError)


class FinanceSafetyTests(TestCase):
    def test_generator_does_not_import_finance_writer(self) -> None:
        import core.docgen.generator as gmod

        source = inspect_source(gmod)
        self.assertNotIn("financial_ledger", source)
        self.assertNotIn("finance_review", source)
        self.assertNotIn("save_transaction", source)


# ---------------------------------------------------------------------------
# LOG-DOCGEN-01-PR1 final correctness repair regression tests
# ---------------------------------------------------------------------------


class PxkSourceIndependentOfEditorTests(TestCase):
    """Repair 1: PXK readiness/generation must NOT depend on current goods editor."""

    def test_empty_editor_with_hdmb_source_pxk_ready(self) -> None:
        """current goods_lines=[] with valid HDMB source => PXK READY."""
        src = _pxk_source(doc_type=HDMB, line_count=1, line={
            "description": "Máy", "unit": "Cái", "quantity": "1",
            "unit_price_vnd_excl_vat": 60000000, "amount_vnd_excl_vat": 60000000,
        })
        readiness = evaluate_document_readiness(
            PXK,
            goods_lines=[],  # empty editor
            agent=_agent(),
            company=_company(),
            delivery_completed=True,
            document_identity=_pxk_identity(),
            delivery=_pxk_delivery(),
            pxk_source=src,
            selected_contract={"currency": "VND"},
        )
        self.assertEqual(readiness.state, READY, f"Expected READY but got {readiness.state}: {readiness.reason}")

    def test_empty_editor_with_pl_source_pxk_ready(self) -> None:
        """current goods_lines=[] with valid PL source => PXK READY."""
        src = _pxk_source(doc_type=PL, line_count=1, line={
            "description": "Máy", "unit": "Cái", "quantity": "1",
            "unit_price_vnd_excl_vat": 60000000, "amount_vnd_excl_vat": 60000000,
        })
        readiness = evaluate_document_readiness(
            PXK,
            goods_lines=[],  # empty editor
            agent=_agent(),
            company=_company(),
            delivery_completed=True,
            document_identity=_pxk_identity(),
            delivery=_pxk_delivery(),
            pxk_source=src,
            selected_contract={"currency": "VND"},
        )
        self.assertEqual(readiness.state, READY, f"Expected READY but got {readiness.state}: {readiness.reason}")


class HdntValidityFailClosedTests(TestCase):
    """Repair 2: HDNT validity must fail closed on all missing/invalid conditions."""

    def test_expired_hdnt_cannot_parent_pl(self) -> None:
        parent = _parent("HDNT-EXP")
        parent["valid_to"] = "2020-01-01"
        decision = classify_valid_hdnt({"HDNT-EXP": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)
        self.assertIsNone(decision["reuse_target"])

    def test_future_hdnt_cannot_parent_pl(self) -> None:
        parent = _parent("HDNT-FUT")
        parent["valid_from"] = "2027-01-01"
        parent["valid_to"] = "2028-01-01"
        decision = classify_valid_hdnt({"HDNT-FUT": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)

    def test_missing_validity_cannot_parent_pl(self) -> None:
        parent = _parent("HDNT-NOVAL")
        parent["valid_from"] = ""
        parent["valid_to"] = ""
        decision = classify_valid_hdnt({"HDNT-NOVAL": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)

    def test_cancelled_hdnt_cannot_parent_pl(self) -> None:
        parent = _parent("HDNT-CANC")
        parent["status"] = "cancelled"
        decision = classify_valid_hdnt({"HDNT-CANC": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)

    def test_inactive_hdnt_cannot_parent_pl(self) -> None:
        parent = _parent("HDNT-INACT")
        parent["status"] = "inactive"
        decision = classify_valid_hdnt({"HDNT-INACT": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)

    def test_deleted_hdnt_cannot_parent_pl(self) -> None:
        parent = _parent("HDNT-DEL")
        parent["is_deleted"] = True
        decision = classify_valid_hdnt({"HDNT-DEL": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 0)

    def test_one_valid_hdnt_reuse(self) -> None:
        parent = _parent("HDNT-OK")
        decision = classify_valid_hdnt({"HDNT-OK": parent}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 1)
        self.assertFalse(decision["needs_selection"])
        self.assertEqual(decision["reuse_target"]["hdnt_no"], "HDNT-OK")

    def test_two_valid_hdnt_explicit_selection(self) -> None:
        p1 = _parent("HDNT-A")
        p2 = _parent("HDNT-B")
        decision = classify_valid_hdnt({"HDNT-A": p1, "HDNT-B": p2}, agent_id="AGT-001", company_id="TT-001", today=date(2026, 9, 1))
        self.assertEqual(decision["valid_count"], 2)
        self.assertTrue(decision["needs_selection"])
        self.assertIsNone(decision["reuse_target"])


class IssuedFullImmutabilityTests(TestCase):
    """Repair 4: After ISSUED only superseded_by_document_id may change."""

    def test_cannot_change_issued_at_after_issued(self) -> None:
        from core.database import save_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "IM1::R1", "document_id": "IM1", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("IM1")
            record = gen_load_latest("IM1")
            self.assertEqual(record["status"], ISSUED)
            # Attempt to change issued_at
            hacked = deepcopy(record)
            hacked["issued_at"] = "2099-01-01T00:00:00"
            self.assertFalse(save_docgen_document(hacked))

    def test_cannot_change_issued_by_after_issued(self) -> None:
        from core.database import save_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "IM2::R1", "document_id": "IM2", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("IM2")
            record = gen_load_latest("IM2")
            self.assertEqual(record["status"], ISSUED)
            hacked = deepcopy(record)
            hacked["issued_by"] = "attacker"
            self.assertFalse(save_docgen_document(hacked))

    def test_cannot_change_generated_by_after_issued(self) -> None:
        from core.database import save_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "IM3::R1", "document_id": "IM3", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("IM3")
            record = gen_load_latest("IM3")
            self.assertEqual(record["status"], ISSUED)
            hacked = deepcopy(record)
            hacked["generated_by"] = "attacker"
            self.assertFalse(save_docgen_document(hacked))


class LegacyScalarPreservedTests(TestCase):
    """Repair 3: link_contract_to_shipment must not overwrite legacy scalar."""

    def test_legacy_scalar_preserved_after_second_linkage(self) -> None:
        from core.database import link_contract_to_shipment, load_contracts_for_shipment

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db), patch("core.auth.st", type("AuthStreamlit", (), {"session_state": {"current_user": {"role": "support", "username": "support", "user_id": "u1"}}})()):
            # Initial contract with shipment_id = SHP-001
            self.assertTrue(_save_contract({
                "id": "C-NM-001", "contract_no": "C-NM-001", "kind": HDNT,
                "agent_id": "AGT-001", "company_id": "TT-001",
                "shipment_id": "SHP-001",
            }))
            # Link SHP-002 — must NOT overwrite the legacy scalar
            self.assertTrue(link_contract_to_shipment("C-NM-001", "SHP-002"))
            linked = load_contracts_for_shipment("SHP-002")
            self.assertIn("C-NM-001", linked)
            # shipment_ids contains BOTH
            record = linked["C-NM-001"]
            self.assertIn("SHP-001", record.get("shipment_ids", []))
            self.assertIn("SHP-002", record.get("shipment_ids", []))
            # Legacy scalar remains SHP-001 (the original)
            self.assertEqual(record.get("shipment_id"), "SHP-001")


class ContractCacheImmediateReadbackTests(TestCase):
    """Repair 5: After contract save, immediate read must see new data."""

    def test_cache_invalidation_after_save(self) -> None:
        from core.database import _invalidate_contracts_cache, load_contracts

        # Outside Streamlit runtime, load_contracts goes directly to _load_contracts_impl.
        # We test that _invalidate_contracts_cache is callable and doesn't raise.
        _invalidate_contracts_cache()  # Must not raise


class PlParentValidityServiceGateTests(TestCase):
    """PL readiness must reject invalid HDNT parent at service layer."""

    def test_expired_parent_blocked(self) -> None:
        expired_parent = _parent("HDNT-EXP")
        expired_parent["valid_to"] = "2020-01-01"
        readiness = evaluate_document_readiness(
            PL,
            goods_lines=[_line("A", 1, Decimal("100000"))],
            agent=_agent(),
            company=_company(),
            parent_contract=expired_parent,
            delivery_completed=False,
            today=date(2026, 9, 1),
        )
        self.assertEqual(readiness.state, BLOCKED)
        self.assertIn("hiệu lực", readiness.reason.lower())

    def test_future_parent_blocked(self) -> None:
        future_parent = _parent("HDNT-FUT")
        future_parent["valid_from"] = "2027-01-01"
        future_parent["valid_to"] = "2028-01-01"
        readiness = evaluate_document_readiness(
            PL,
            goods_lines=[_line("A", 1, Decimal("100000"))],
            agent=_agent(),
            company=_company(),
            parent_contract=future_parent,
            delivery_completed=False,
            today=date(2026, 9, 1),
        )
        self.assertEqual(readiness.state, BLOCKED)
        self.assertIn("hiệu lực", readiness.reason.lower())

    def test_missing_validity_parent_blocked(self) -> None:
        no_validity_parent = _parent("HDNT-NOVAL")
        no_validity_parent["valid_from"] = ""
        no_validity_parent["valid_to"] = ""
        readiness = evaluate_document_readiness(
            PL,
            goods_lines=[_line("A", 1, Decimal("100000"))],
            agent=_agent(),
            company=_company(),
            parent_contract=no_validity_parent,
            delivery_completed=False,
            today=date(2026, 9, 1),
        )
        self.assertEqual(readiness.state, BLOCKED)

    def test_cancelled_parent_blocked(self) -> None:
        cancelled_parent = _parent("HDNT-CANC")
        cancelled_parent["status"] = "cancelled"
        readiness = evaluate_document_readiness(
            PL,
            goods_lines=[_line("A", 1, Decimal("100000"))],
            agent=_agent(),
            company=_company(),
            parent_contract=cancelled_parent,
            delivery_completed=False,
            today=date(2026, 9, 1),
        )
        self.assertEqual(readiness.state, BLOCKED)

    def test_inactive_parent_blocked(self) -> None:
        inactive_parent = _parent("HDNT-INACT")
        inactive_parent["status"] = "inactive"
        readiness = evaluate_document_readiness(
            PL,
            goods_lines=[_line("A", 1, Decimal("100000"))],
            agent=_agent(),
            company=_company(),
            parent_contract=inactive_parent,
            delivery_completed=False,
            today=date(2026, 9, 1),
        )
        self.assertEqual(readiness.state, BLOCKED)

    def test_valid_parent_passes(self) -> None:
        valid_parent = _parent("HDNT-OK")
        readiness = evaluate_document_readiness(
            PL,
            goods_lines=[_line("A", 1, Decimal("100000"))],
            agent=_agent(),
            company=_company(),
            parent_contract=valid_parent,
            delivery_completed=False,
            today=date(2026, 9, 1),
        )
        # Should not be BLOCKED due to parent validity
        self.assertNotEqual(readiness.state, BLOCKED)


class ControlledSupersedeLinkHelperTests(TestCase):
    """Issued supersede link must use controlled helper, not generic writer."""

    def test_generic_writer_cannot_change_superseded_by_on_issued(self) -> None:
        from core.database import save_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "SUP1::R1", "document_id": "SUP1", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("SUP1")
            record = gen_load_latest("SUP1")
            self.assertEqual(record["status"], ISSUED)
            # Attempt to change superseded_by_document_id via generic writer
            hacked = deepcopy(record)
            hacked["superseded_by_document_id"] = "NEW-DOC"
            self.assertFalse(save_docgen_document(hacked))

    def test_controlled_helper_admin_pass(self) -> None:
        from core.database import link_superseded_docgen_document

        db = _Database()
        with _with_user({"role": "admin", "username": "admin", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "SUP2::R1", "document_id": "SUP2", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("SUP2")
            record = gen_load_latest("SUP2")
            self.assertEqual(record["status"], ISSUED)
            # Controlled helper should succeed
            self.assertTrue(link_superseded_docgen_document("SUP2", "NEW-DOC-2"))
            updated = gen_load_latest("SUP2")
            self.assertEqual(updated.get("superseded_by_document_id"), "NEW-DOC-2")
            self.assertEqual(updated.get("status"), ISSUED)

    def test_controlled_helper_support_pass(self) -> None:
        from core.database import link_superseded_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "SUP3::R1", "document_id": "SUP3", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("SUP3")
            record = gen_load_latest("SUP3")
            self.assertEqual(record["status"], ISSUED)
            # Controlled helper should succeed for support role
            self.assertTrue(link_superseded_docgen_document("SUP3", "NEW-DOC-3"))

    def test_controlled_helper_sale_denied(self) -> None:
        from core.database import link_superseded_docgen_document

        db = _Database()
        # First issue the document with support role
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "SUP4::R1", "document_id": "SUP4", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("SUP4")
            record = gen_load_latest("SUP4")
            self.assertEqual(record["status"], ISSUED)
        
        # Now test that sale role is denied by the link helper
        with _with_user({"role": "sale", "username": "sale", "user_id": "u3"}), patch("core.database.get_db", return_value=db):
            # Controlled helper should deny sale role
            self.assertFalse(link_superseded_docgen_document("SUP4", "NEW-DOC-4"))

    def test_controlled_helper_rejects_second_different_replacement(self) -> None:
        from core.database import link_superseded_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "SUP5::R1", "document_id": "SUP5", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("SUP5")
            record = gen_load_latest("SUP5")
            self.assertEqual(record["status"], ISSUED)
            # First linkage should succeed
            self.assertTrue(link_superseded_docgen_document("SUP5", "NEW-DOC-5A"))
            # Second different linkage should be rejected
            self.assertFalse(link_superseded_docgen_document("SUP5", "NEW-DOC-5B"))
            # Record should still point to first replacement
            updated = gen_load_latest("SUP5")
            self.assertEqual(updated.get("superseded_by_document_id"), "NEW-DOC-5A")

    def test_controlled_helper_idempotent_same_id(self) -> None:
        from core.database import link_superseded_docgen_document

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db):
            self.assertTrue(_saved_docgen_document({"id": "SUP6::R1", "document_id": "SUP6", "status": GENERATED, "generation_revision": 1}))
            gen.issue_document("SUP6")
            record = gen_load_latest("SUP6")
            self.assertEqual(record["status"], ISSUED)
            # First linkage
            self.assertTrue(link_superseded_docgen_document("SUP6", "NEW-DOC-6"))
            # Same linkage again should be idempotent
            self.assertTrue(link_superseded_docgen_document("SUP6", "NEW-DOC-6"))


class CacheImmediateReadbackEvidenceTests(TestCase):
    """Strengthened cache evidence: save -> invalidate -> read sees new data."""

    def test_save_invalidate_read_sees_new_contract(self) -> None:
        from core.database import save_contract_to_cloud, load_contracts

        db = _Database()
        with _with_user({"role": "support", "username": "support", "user_id": "u1"}), patch("core.database.get_db", return_value=db), patch("core.auth.st", type("AuthStreamlit", (), {"session_state": {"current_user": {"role": "support", "username": "support", "user_id": "u1"}}})()):
            # Save a new HDNT
            self.assertTrue(_save_contract({
                "id": "HDNT-NEW-001",
                "contract_no": "HDNT-NEW-001",
                "kind": HDNT,
                "agent_id": "AGT-001",
                "company_id": "TT-001",
                "valid_from": "2026-09-01",
                "valid_to": "2027-08-31",
            }))
            # Immediate read should see it
            contracts = load_contracts()
            self.assertIn("HDNT-NEW-001", contracts)
            self.assertEqual(contracts["HDNT-NEW-001"].get("kind"), HDNT)


# ---------------------------------------------------------------------------
# In-memory DB + auth helpers
# ---------------------------------------------------------------------------


class _Result:
    matched_count = 1
    modified_count = 1
    upserted_id = None


class _Collection:
    def __init__(self) -> None:
        self.docs: dict[str, dict] = {}

    @staticmethod
    def _matches(document: dict, query: dict) -> bool:
        for key, expected in query.items():
            actual = document.get(key)
            if isinstance(expected, dict) and "$ne" in expected:
                if actual == expected["$ne"]:
                    return False
            elif actual != expected:
                return False
        return True

    def find_one(self, query: dict, projection=None, **kwargs):
        del projection, kwargs
        for document in self.docs.values():
            if self._matches(document, query):
                return deepcopy(document)
        return None

    def find(self, query: dict | None = None, projection=None, **kwargs):
        del projection, kwargs
        return [deepcopy(document) for document in self.docs.values() if self._matches(document, query or {})]

    def update_one(self, query: dict, update: dict, upsert=False, **kwargs):
        del kwargs
        current = self.find_one(query)
        if current is None and not upsert:
            return type("NoMatch", (), {"matched_count": 0, "modified_count": 0})()
        key = str(query.get("id") or query.get("username"))
        payload = current or {name: value for name, value in query.items() if not isinstance(value, dict)}
        payload.update(deepcopy(update.get("$set", {})))
        if current is None:
            payload.update(deepcopy(update.get("$setOnInsert", {})))
        self.docs[key] = payload
        return _Result()

    def insert_one(self, document: dict, **kwargs):
        del kwargs
        key = str(document.get("event_id") or document.get("id") or len(self.docs))
        self.docs[key] = deepcopy(document)
        return _Result()


class _Database:
    def __init__(self) -> None:
        self.contracts = _Collection()
        self.contract_transactions = _Collection()
        self.authority_companies = _Collection()
        self.authority_agents = _Collection()
        self.docgen_documents = _Collection()
        self.audit_logs = _Collection()

    def __getitem__(self, name: str) -> _Collection:
        return getattr(self, name)

    def __getattr__(self, name: str) -> _Collection:
        if name.startswith("_"):
            raise AttributeError(name)
        collection = _Collection()
        setattr(self, name, collection)
        return collection


def _with_user(user) -> object:
    return patch("core.auth.st", type("AuthStreamlit", (), {"session_state": {"current_user": user}})())


def _saved_docgen_document(document: dict) -> bool:
    from core.database import save_docgen_document

    return save_docgen_document(document)


def _save_contract(document: dict) -> bool:
    from core.database import save_contract_to_cloud

    return save_contract_to_cloud(document)


def gen_load_latest(document_id: str):
    from core.database import load_docgen_document_latest

    return load_docgen_document_latest(document_id)


def _sha256_file(path: str) -> str:
    import hashlib

    with open(path, "rb") as handle:
        return hashlib.sha256(handle.read()).hexdigest()


def inspect_source(module) -> str:
    import inspect

    return inspect.getsource(module)
