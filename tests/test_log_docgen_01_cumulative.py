"""One offline cumulative acceptance sweep; no private template assets required.

Real XLSX/DOCX packages are built from explicitly SYNTHETIC fixture code.
They cover every recovered repair through the production generators. Source
files are read-only and hash/stat checked, and persistence is confined to a
fail-closed in-memory document/audit store. Run alongside the original focused
DOCGEN, OOXML and RBAC suites. This is NOT a Native Office acceptance pass.
"""

from __future__ import annotations

import io
from copy import copy, deepcopy
from datetime import date
from decimal import Decimal
from hashlib import sha256
from types import SimpleNamespace
from unittest.mock import patch

import docx
import openpyxl
import pytest
from lxml import etree

from core import auth, database
from core.docgen import generator as gen, rules
from core.docgen.rules import GENERATED, HDMB, HDNT, ISSUED, PL, PXK
from core.docgen.template_engine import CachedFormulaValue, mutate_xlsx
from tests.fixtures.build_synthetic_docgen import (
    FORMULAS, NS, SHEET_PARTS, build_synthetic_docx, build_synthetic_xlsx, zip_parts,
)
from tests.test_log_docgen_01 import (
    _Collection, _agent, _company, _hdmb_identity, _hdnt_identity, _line,
    _parent, _pl_identity, _pxk_delivery, _pxk_identity, _pxk_source, _with_user,
)

DOCUMENT_TYPES = (HDNT, PL, HDMB, PXK)
ROLES = ("admin", "support", "sale", "viewer", "unknown", None)


class _AcceptanceDate(date):
    @classmethod
    def today(cls):
        return cls(2026, 9, 8)


@pytest.fixture(autouse=True)
def fixed_date_and_no_mongo(monkeypatch):
    monkeypatch.setattr(rules, "date", _AcceptanceDate)
    monkeypatch.setattr(auth, "AUTH_REQUIRED", True)
    with patch("pymongo.MongoClient", side_effect=AssertionError("No Mongo connections in acceptance tests")):
        yield


@pytest.fixture
def template_file(tmp_path):
    """Every synthetic input, including repeat-generation inputs, stays immutable."""
    originals = []

    def create(raw, suffix):
        path = tmp_path / f"synthetic-{len(originals)}{suffix}"
        path.write_bytes(raw)
        path.chmod(0o444)
        originals.append((path, sha256(raw).hexdigest(), path.stat()))
        return path

    yield create
    for path, digest, before in originals:
        assert sha256(path.read_bytes()).hexdigest() == digest
        after = path.stat()
        assert (after.st_mtime_ns, after.st_size, after.st_mode) == (
            before.st_mtime_ns, before.st_size, before.st_mode,
        )


@pytest.fixture
def templates(template_file):
    return SimpleNamespace(
        xlsx=template_file(build_synthetic_xlsx(), ".xlsx"),
        docx=template_file(build_synthetic_docx(), ".docx"),
    )


def _user(role):
    return {"role": role, "username": f"synthetic.{role}", "user_id": f"synthetic-{role}"} if role else {}


def _generation_args(document_type, templates):
    identity = {HDNT: _hdnt_identity, PL: _pl_identity, HDMB: _hdmb_identity, PXK: _pxk_identity}
    return dict(
        shipment_id="SHP-SYNTHETIC",
        agent={**_agent(), "legal_representative": "Ông Nguyễn Văn Anh"},
        company=_company(),
        goods_lines=[_line("SYNTHETIC Máy", 1, Decimal(125000000 if document_type == HDMB else 100000))],
        contract=identity[document_type](),
        parent_contract=_parent() if document_type == PL else None,
        delivery=_pxk_delivery() if document_type == PXK else None,
        pxk_source=_pxk_source() if document_type == PXK else None,
        template_path_xlsx=str(templates.xlsx),
        template_path_docx=str(templates.docx),
    )


def _cell(sheet, coord):
    cell = sheet.find(f".//x:c[@r='{coord}']", NS)
    assert cell is not None, coord
    return cell


def _assert_calc_chain(parts, present):
    assert ("xl/calcChain.xml" in parts) is present
    rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    types = etree.fromstring(parts["[Content_Types].xml"])
    calc_rels = [rel for rel in rels if "calcChain" in str(dict(rel.attrib))]
    calc_types = [item for item in types if "calcChain" in str(dict(item.attrib))]
    assert len(calc_rels) == len(calc_types) == int(present)
    if present:
        assert calc_rels[0].get("Target") == "calcChain.xml"
        assert calc_types[0].get("PartName") == "/xl/calcChain.xml"


def _assert_full_recalculation(parts):
    workbook = etree.fromstring(parts["xl/workbook.xml"])
    calc = workbook.find("x:calcPr", NS)
    assert calc is not None
    assert {key: calc.get(key) for key in ("calcMode", "fullCalcOnLoad", "forceFullCalc")} == {
        "calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1",
    }
    tags = [etree.QName(child).localname for child in workbook]
    assert tags.index("sheets") < tags.index("calcPr")
    if "definedNames" in tags:
        assert tags.index("sheets") < tags.index("definedNames") < tags.index("calcPr")
    return workbook


def _assert_unmapped_parts_unchanged(before, after, allowed):
    for name, raw in before.items():
        if name not in allowed:
            assert after[name] == raw, name
    assert not (set(after) - set(before))


@pytest.mark.parametrize("agent_rep,normalized", [
    ("Ông Nguyễn Văn Anh", "Nguyễn Văn Anh"),
    ("Bà Nguyễn Thị An", "Nguyễn Thị An"),
    ("Nguyễn Văn Anh", "Nguyễn Văn Anh"),
])
@pytest.mark.parametrize("company_rep", ["Bà Lê Thị Thu Hương", "Ông Lê Văn Hùng"])
def test_hdnt_exact_paragraphs_and_run_properties(templates, agent_rep, normalized, company_rep):
    agent = {**_agent(), "legal_representative": agent_rep}
    company = {**_company(), "legal_representative": company_rep}
    before_masters = deepcopy((agent, company))
    output = gen.generate_docx_file(
        str(templates.docx), contract_no="HDNT-SYNTHETIC", contract_date="2026-09-08",
        valid_from="2026-09-08", valid_to="2027-09-08", agent=agent, company=company,
    )
    source = docx.Document(str(templates.docx))
    document = docx.Document(io.BytesIO(output))
    agent_paragraph, company_paragraph = document.paragraphs[9], document.paragraphs[15]
    assert agent_paragraph.text == f"Đại diện: Ông/Bà {normalized}  Chức vụ: Giám đốc"
    assert company_paragraph.text == f"Đại diện : {company_rep}  Chức vụ: Giám đốc"
    assert [run.text for run in agent_paragraph.runs] == [
        "Đại diện: Ông", f"/Bà {normalized}", "  ", "Chức vụ: Giám đốc",
    ]
    assert [run.text for run in company_paragraph.runs[2:]] == [
        company_rep, "  ", "", "Chức vụ: Giám đốc",
    ]
    for index in (9, 15):
        before, after = source.paragraphs[index], document.paragraphs[index]
        assert len(before.runs) == len(after.runs)
        assert etree.tostring(before._p.pPr) == etree.tostring(after._p.pPr)
        for old_run, new_run in zip(before.runs, after.runs):
            assert etree.tostring(old_run._r.rPr) == etree.tostring(new_run._r.rPr)
        assert "Ông/Bà Ông" not in after.text
        assert "Ông/Bà Bà" not in after.text
        assert "   " not in after.text
        assert not any(char in after.text for char in "\t\r\n")
        assert after.text == after.text.rstrip()
    assert len(document.paragraphs) == len(source.paragraphs)
    assert len(document.sections) == len(source.sections) == 1
    assert etree.tostring(source.sections[0]._sectPr) == etree.tostring(document.sections[0]._sectPr)
    assert (agent, company) == before_masters


@pytest.mark.parametrize("document_type,sheet_name,expected_totals", [
    (HDMB, "HDMB", (125000000, 10000000, 135000000)),
    (PL, "pl", (1800000, 144000, 1944000)),
])
@pytest.mark.parametrize("shared_strings", [True, False], ids=["shared", "inline"])
@pytest.mark.parametrize("existing_names", [True, False], ids=["existing-names", "new-names"])
def test_totals_caches_recalc_layout_and_unmapped_parts(
    template_file, document_type, sheet_name, expected_totals, shared_strings, existing_names,
):
    raw = build_synthetic_xlsx(shared_strings=shared_strings, existing_defined_names=existing_names)
    path = template_file(raw, ".xlsx")
    lines = ([_line("SYNTHETIC Máy", 1, Decimal(125000000))] if document_type == HDMB else
             [_line(f"SYNTHETIC Hàng {i}", 3, Decimal("33333.33")) for i in range(18)])
    kwargs = dict(contract_no=f"{document_type}-SYNTHETIC", contract_date="2026-09-08",
                  agent=_agent(), company=_company(), goods_lines=lines, parent_contract=_parent())
    output = gen.generate_xlsx_file(str(path), document_type, **kwargs)
    before, after = zip_parts(raw), zip_parts(output)
    sheet = etree.fromstring(after[SHEET_PARTS[sheet_name]])
    formulas = openpyxl.load_workbook(io.BytesIO(output), data_only=False)
    values = openpyxl.load_workbook(io.BytesIO(output), data_only=True)
    for (coord, formula), expected in zip(FORMULAS[sheet_name].items(), expected_totals):
        cell = _cell(sheet, coord)
        assert cell.find("x:f", NS).text == formula
        assert Decimal(cell.find("x:v", NS).text) == expected
        assert formulas[sheet_name][coord].value == "=" + formula
        assert values[sheet_name][coord].value == expected
    _assert_calc_chain(after, present=True)
    assert after["xl/calcChain.xml"] == before["xl/calcChain.xml"]
    workbook_xml = _assert_full_recalculation(after)
    assert formulas[sheet_name].page_setup.orientation == "portrait"
    allowed = {"xl/workbook.xml", SHEET_PARTS[sheet_name], "xl/sharedStrings.xml"}

    if document_type == PL:
        row6 = sheet.find("x:sheetData/x:row[@r='6']", NS)
        assert row6.get("ht") == "75" and row6.get("customHeight") == "1"
        names = workbook_xml.findall("x:definedNames/x:definedName", NS)
        titles = [node for node in names if node.get("name") == "_xlnm.Print_Titles"]
        assert len(titles) == 1
        assert titles[0].get("localSheetId") == "1"
        assert titles[0].text == "pl!$9:$9"
        assert formulas["pl"].print_title_rows == "$9:$9"
        assert any(node.get("name") == "synthetic_keep" for node in names) is existing_names
        # Repeat mutation must preserve row 9 and the schema order, not add duplicates.
        repeat = gen.generate_xlsx_file(str(template_file(output, ".xlsx")), PL, **kwargs)
        repeat_wb = _assert_full_recalculation(zip_parts(repeat))
        assert len(repeat_wb.findall("x:definedNames/x:definedName[@name='_xlnm.Print_Titles']", NS)) == 1
    else:
        allowed.add("xl/styles.xml")
        old_sheet = etree.fromstring(before[SHEET_PARTS["HDMB"]])
        original_style = _cell(old_sheet, "E24").get("s")
        assert original_style == _cell(old_sheet, "F24").get("s")
        assert _cell(sheet, "F24").get("s") == original_style
        assert _cell(sheet, "E24").get("s") != original_style
        old_xfs = etree.fromstring(before["xl/styles.xml"]).find("x:cellXfs", NS)
        new_xfs = etree.fromstring(after["xl/styles.xml"]).find("x:cellXfs", NS)
        assert len(new_xfs) == len(old_xfs) + 1 == int(new_xfs.get("count"))
        for old_xf, retained_xf in zip(old_xfs, new_xfs):
            assert etree.tostring(old_xf) == etree.tostring(retained_xf)
        e24, f24 = formulas["HDMB"]["E24"], formulas["HDMB"]["F24"]
        assert e24.value == 125000000
        assert e24.alignment.shrinkToFit is True
        assert not e24.alignment.wrapText
        assert f24.alignment.wrapText is True
        assert not f24.alignment.shrinkToFit
        assert e24.number_format == f24.number_format == "#,##0"
        assert e24.alignment.horizontal == f24.alignment.horizontal == "right"
        assert e24.font == copy(f24.font)
        assert e24.fill == copy(f24.fill)
        assert e24.border == copy(f24.border)
        assert formulas["HDMB"]["A13"].value == "Đại diện bởi Ông/Bà: Nguyễn Văn An    Chức vụ: Giám đốc"
        assert formulas["HDMB"]["A20"].value == "Đại diện bởi Ông/Bà: Lê Thị Thu Hương    Chức vụ: Giám đốc"
    _assert_unmapped_parts_unchanged(before, after, allowed)
    formulas.close()
    values.close()


@pytest.mark.parametrize("source_type", [HDMB, PL])
@pytest.mark.parametrize("shared_strings", [True, False], ids=["shared", "inline"])
@pytest.mark.parametrize("formula_inputs", [True, False], ids=["formula-inputs", "literal-inputs"])
@pytest.mark.parametrize("empty_editor", [True, False], ids=["empty-editor", "conflicting-editor"])
def test_pxk_source_literals_conditional_chain_cleanup_and_long_layout(
    templates, template_file, source_type, shared_strings, formula_inputs, empty_editor,
):
    raw = build_synthetic_xlsx(shared_strings=shared_strings, pxk_formula_inputs=formula_inputs)
    path = template_file(raw, ".xlsx")
    args = _generation_args(PXK, templates)
    args["template_path_xlsx"] = str(path)
    args["goods_lines"] = [] if empty_editor else [_line("WRONG EDITOR", 999, 999)]
    args["pxk_source"] = _pxk_source(source_type, rev=7, line={
        "description": "SYNTHETIC Máy từ nguồn đã chọn", "unit": "Bộ", "quantity": 3,
        "unit_price_vnd_excl_vat": 1000000,
    })
    args["agent"]["legal_name"] = "SYNTHETIC ĐẠI LÝ VẬN TẢI QUỐC TẾ " * 5
    args["agent"]["legal_address"] = "SYNTHETIC Địa chỉ rất dài tại khu công nghiệp " * 6
    args["delivery"]["receiver_name"] = "SYNTHETIC Người đại diện tiếp nhận hàng hóa " * 5
    args["delivery"]["delivery_location"] = "SYNTHETIC Kho nhận hàng tại khu công nghiệp " * 5
    before_args = deepcopy(args)
    with _with_user(_user("support")):
        result = gen.generate_document(PXK, **args)
    before, after = zip_parts(raw), zip_parts(result["bytes"])
    sheet = etree.fromstring(after[SHEET_PARTS["pxk"]])
    workbook = openpyxl.load_workbook(io.BytesIO(result["bytes"]), data_only=False)
    pxk = workbook["pxk"]
    for coord, expected in (("B15", "SYNTHETIC Máy từ nguồn đã chọn"), ("C15", "Bộ"), ("D15", 3)):
        assert pxk[coord].value == expected
        assert pxk[coord].data_type != "f"
        assert _cell(sheet, coord).find("x:f", NS) is None
    _assert_calc_chain(after, present=not formula_inputs)
    _assert_full_recalculation(after)
    for row, height in ((6, 40), (7, 40), (10, 30), (11, 30)):
        assert pxk.row_dimensions[row].height == height
        assert pxk.row_dimensions[row].customHeight
        assert pxk[f"A{row}"].alignment.wrapText is True
    for coord, expected in (
        ("A6", args["agent"]["legal_name"].strip()), ("A7", args["agent"]["legal_address"].strip()),
        ("A10", args["delivery"]["receiver_name"].strip()), ("A11", args["delivery"]["delivery_location"].strip()),
    ):
        assert expected in pxk[coord].value
    assert {"A6:D6", "A7:D7", "A10:D10", "A11:D11"} <= {str(m) for m in pxk.merged_cells.ranges}
    assert pxk.page_setup.fitToWidth == 1 and pxk.page_setup.fitToHeight == 0
    assert pxk.sheet_properties.pageSetUpPr.fitToPage is True
    assert pxk.page_setup.orientation == "portrait"
    tags = [etree.QName(child).localname for child in sheet]
    assert tags.index("sheetData") < tags.index("mergeCells") < tags.index("pageMargins") < tags.index("pageSetup")
    snapshot = result["metadata"]["source_snapshot"]
    assert snapshot["source_document_type"] == source_type
    assert snapshot["source_revision"] == 7
    assert snapshot["source_goods_snapshot"] == args["pxk_source"]["source_goods_line"]
    assert snapshot["goods_lines"] == [args["pxk_source"]["source_goods_line"]]
    assert args == before_args
    allowed = {"xl/workbook.xml", SHEET_PARTS["pxk"], "xl/styles.xml", "xl/sharedStrings.xml"}
    if formula_inputs:
        allowed |= {"xl/calcChain.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
        for part in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
            retained = [dict(node.attrib) for node in etree.fromstring(before[part])
                        if "calcChain" not in str(dict(node.attrib))]
            assert [dict(node.attrib) for node in etree.fromstring(after[part])] == retained
    old_xfs = etree.fromstring(before["xl/styles.xml"]).find("x:cellXfs", NS)
    new_xfs = etree.fromstring(after["xl/styles.xml"]).find("x:cellXfs", NS)
    assert len(new_xfs) == len(old_xfs) + 4
    for old_xf, retained_xf in zip(old_xfs, new_xfs):
        assert etree.tostring(old_xf) == etree.tostring(retained_xf)
    _assert_unmapped_parts_unchanged(before, after, allowed)
    workbook.close()


@pytest.mark.parametrize("fill", [
    {}, {"missing-sheet": {"B15": "literal"}}, {"pxk": {}},
    {"pxk": {"Z999": "missing cell"}}, {"pxk": {"A6": "literal text only"}},
    {"pl": {"F28": CachedFormulaValue(123)}},
])
@pytest.mark.parametrize("shared_strings", [True, False], ids=["shared", "inline"])
def test_calc_chain_is_not_removed_without_formula_to_literal(template_file, fill, shared_strings):
    raw = build_synthetic_xlsx(shared_strings=shared_strings)
    output = mutate_xlsx(str(template_file(raw, ".xlsx")), fill)
    before, after = zip_parts(raw), zip_parts(output)
    _assert_calc_chain(after, present=True)
    for part in ("xl/calcChain.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
        assert before[part] == after[part]
    _assert_full_recalculation(after)


class _RecordingCollection(_Collection):
    def __init__(self, name, writes):
        super().__init__()
        self.name, self.writes = name, writes

    def update_one(self, query, update, **kwargs):
        self.writes.append((self.name, "update_one", deepcopy(update)))
        return super().update_one(query, update, **kwargs)

    def insert_one(self, document, **kwargs):
        self.writes.append((self.name, "insert_one", deepcopy(document)))
        return super().insert_one(document, **kwargs)


class _DocumentOnlyDatabase:
    """Unknown/Finance collections fail closed instead of being auto-created."""
    def __init__(self):
        self.writes = []
        self.denied_access = []
        self.docgen_documents = _RecordingCollection("docgen_documents", self.writes)
        self.audit_logs = _RecordingCollection("audit_logs", self.writes)

    def __getitem__(self, name):
        if name not in ("docgen_documents", "audit_logs"):
            self.denied_access.append(name)
            raise AssertionError(f"Forbidden non-document collection: {name}")
        return getattr(self, name)

    def __getattr__(self, name):
        self.denied_access.append(name)
        raise AssertionError(f"Forbidden non-document database attribute: {name}")


@pytest.fixture(autouse=True)
def document_db(monkeypatch):
    db = _DocumentOnlyDatabase()
    monkeypatch.setattr(database, "get_db", lambda: db)
    yield db
    assert not db.denied_access
    assert {name for name, _, _ in db.writes} <= {"docgen_documents", "audit_logs"}


@pytest.mark.parametrize("document_type", DOCUMENT_TYPES)
@pytest.mark.parametrize("role", ROLES)
def test_generate_download_rbac_metadata_and_zero_finance_writes(templates, document_db, document_type, role):
    args = _generation_args(document_type, templates)
    before_args = deepcopy(args)
    with _with_user(_user(role)):
        allowed = role in ("admin", "support", "sale")
        assert auth.can_generate_documents() is allowed
        if not allowed:
            with pytest.raises(PermissionError, match="sales_write"):
                gen.generate_document(document_type, **args)
        else:
            result = gen.generate_document(document_type, **args)
            assert result["status"] == GENERATED
            assert result["output_hash"] == sha256(result["bytes"]).hexdigest()
            source_path = templates.docx if document_type == HDNT else templates.xlsx
            assert result["template_hash"] == sha256(source_path.read_bytes()).hexdigest()
            assert result["metadata"]["generated_by"] == f"synthetic.{role}"
            assert result["metadata"]["source_snapshot"]["document_no"] == args["contract"][f"{document_type.lower()}_no"]
            assert result["bytes"] and "bytes" not in result["metadata"]
    assert args == before_args
    assert document_db.writes == []


def _generated_and_saved(templates, document_type=HDMB):
    result = gen.generate_document(document_type, **_generation_args(document_type, templates))
    assert database.save_docgen_document(result["metadata"])
    return result


def _supersede_args(templates, old_id):
    return dict(
        new_document_id=old_id + "-REPLACEMENT", document_type=HDMB,
        shipment_id="SHP-SYNTHETIC", contract_no="HDMB-SYNTHETIC-REPLACEMENT",
        contract_date="2026-09-08", agent=_agent(), company=_company(),
        goods_lines=[_line("SYNTHETIC replacement", 1, Decimal(125000000))],
        template_path_xlsx=str(templates.xlsx),
    )


@pytest.mark.parametrize("role", ROLES)
def test_issue_and_supersede_rbac_and_frozen_original(templates, document_db, role):
    with _with_user(_user("admin")):
        result = _generated_and_saved(templates)
    old_id = result["document_id"]
    before_writes = deepcopy(document_db.writes)
    with _with_user(_user(role)):
        allowed = role in ("admin", "support")
        assert auth.can_issue_documents() is allowed
        if not allowed:
            with pytest.raises(auth.PermissionDeniedError):
                gen.issue_document(old_id)
            with pytest.raises(auth.PermissionDeniedError):
                gen.supersede_document(old_id, **_supersede_args(templates, old_id))
            assert document_db.writes == before_writes
            assert database.load_docgen_document_latest(old_id)["status"] == GENERATED
            return
        assert gen.issue_document(old_id)["status"] == ISSUED
        issued = database.load_docgen_document_latest(old_id)
        with pytest.raises(gen.DocumentGenerationError):
            gen.issue_document(old_id)
        assert database.load_docgen_document_latest(old_id) == issued
        replacement = gen.supersede_document(old_id, **_supersede_args(templates, old_id))
        new_id = replacement["document_id"]
        assert new_id == old_id + "-REPLACEMENT"
        assert database.load_docgen_document_latest(old_id) == {**issued, "superseded_by_document_id": new_id}
        new_record = database.load_docgen_document_latest(new_id)
        assert new_record["status"] == GENERATED
        assert new_record["supersedes_document_id"] == old_id
        assert new_record["source_snapshot"]["hdmb_no"] == "HDMB-SYNTHETIC-REPLACEMENT"
        assert new_record["source_snapshot"] != issued["source_snapshot"]
        writes = deepcopy(document_db.writes)
        assert database.link_superseded_docgen_document(old_id, new_id)
        assert not database.link_superseded_docgen_document(old_id, new_id + "-OTHER")
        assert document_db.writes == writes
        assert any(event[2].get("action") == "supersede_link" for event in document_db.writes if event[0] == "audit_logs")


@pytest.mark.parametrize("field,value", [
    ("output_hash", "altered"), ("output_file_name", "altered.xlsx"),
    ("template_hash", "altered"), ("source_reference", "altered"),
    ("source_snapshot", {"hdmb_no": "ALTERED-LEGAL-NUMBER"}),
    ("generation_revision", 99), ("issued_at", "2000-01-01"),
    ("issued_by", "impostor"), ("generated_by", "impostor"),
    ("superseded_by_document_id", "UNCONTROLLED"), ("status", GENERATED),
    ("unexpected_metadata", "not allowed"),
])
def test_issued_writer_immutability_and_no_side_effects(templates, document_db, field, value):
    with _with_user(_user("support")):
        result = _generated_and_saved(templates)
        gen.issue_document(result["document_id"])
        issued = database.load_docgen_document_latest(result["document_id"])
        before_writes = deepcopy(document_db.writes)
        assert not database.save_docgen_document({**issued, field: value})
        assert database.load_docgen_document_latest(result["document_id"]) == issued
        assert document_db.writes == before_writes


@pytest.mark.parametrize("document_type", DOCUMENT_TYPES)
def test_regeneration_keeps_previous_revision_and_master_inputs(templates, document_db, document_type):
    args = _generation_args(document_type, templates)
    with _with_user(_user("support")):
        first = gen.generate_document(document_type, **args)
        assert database.save_docgen_document(first["metadata"])
        previous = database.load_docgen_document(first["metadata"]["id"])
        second = gen.generate_document(document_type, existing_documents={previous["id"]: previous}, **args)
        assert second["revision"] == 2
        assert first["document_id"] == second["document_id"]
        assert database.save_docgen_document(second["metadata"])
        assert database.load_docgen_document(previous["id"]) == previous
        assert database.load_docgen_document_latest(first["document_id"])["generation_revision"] == 2
