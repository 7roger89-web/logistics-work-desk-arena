"""LOGI_QUOTE_COST_WORKFLOW_UX_01_REPAIR_01 — canonical truth of the quick-quote path.

PR #95 made the Báo Giá desk able to turn a pasted customer message into a
quote. This file challenges what actually reaches (and comes back from) the
canonical store, because the desk looking right before Save proves nothing.

Four defects are pinned here:

A. The pasted general sentence ("Nếu có phát sinh về xử lý hàng hóa, sẽ báo
   ngay tại thời điểm phát sinh.") must survive create → fresh readback →
   edit → reopen, and must reach the customer's XLSX "Ghi chú" block. It is
   carried on a dedicated ``customer_notes`` field, NOT dumped into the quote's
   INTERNAL ``assumptions`` control list — those internal statements must never
   be printed on a customer document.

B. UNKNOWN != ZERO. A line whose Giá vốn / VAT was never stated must NOT be
   persisted as a canonical 0 — it persists as None. (REPAIR_02 reversed this
   file's original fail-closed enforcement: only Giá bán is required, because
   blocking the save merely pushed the operator into typing a fake 0.)

C. The XLSX must read the REAL authority field names (legal_name /
   legal_address / legal_representative / representative_title / phone /
   email) and enrich the "To" block from the existing customer master.

D. A quote's date authority is its own: explicit date → canonical created_at →
   today only as a genuine last resort. Re-exporting a June quote in September
   must not restamp it.

Everything runs on the isolated staging transport (mongomock) through the REAL
writers. The production database is never touched.
"""

from __future__ import annotations

import base64
import io

import pytest
from openpyxl import load_workbook

from views import logi_c1
from views import logi_sandbox

GENERAL_NOTE = "Nếu có phát sinh về xử lý hàng hóa, sẽ báo ngay tại thời điểm phát sinh."
SECOND_NOTE = "Giá chưa bao gồm lưu kho quá 5 ngày."

# The exact internal control statements _quote_payload() stamps on every quote.
INTERNAL_ASSUMPTIONS = (
    "Đây là estimate/quotation, chưa phải nghĩa vụ thanh toán.",
    "Giá agent/đối tác phải được đối chiếu với nguồn báo giá thực tế.",
    "Không tự động hóa khai báo hải quan sai lệch.",
)


@pytest.fixture()
def staging():
    import streamlit as st

    logi_sandbox.reset_backend()
    transport = logi_sandbox.ensure_backend()
    assert transport == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    yield transport
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    logi_sandbox.reset_backend()


@pytest.fixture()
def event_seq():
    state = {"n": 0}

    def _make(name: str, payload: dict) -> dict:
        state["n"] += 1
        return {name: {"seq": state["n"], "event": name, "payload": payload}}

    return _make


def _dispatch(seq_fn, name, payload):
    outcomes = logi_c1.consume_c1_triggers({"seq": 0}, seq_fn(name, payload))
    assert len(outcomes) == 1
    return outcomes[0]


def _line(**over):
    """A line as the quick-paste form emits it (sell stated, the rest typed in)."""
    base = {
        "code": "DV-CUOC-BIEN", "serviceName": "Cước biển", "unit": "Cont 20'",
        "qty": 1, "currency": "VND", "costUnit": 0, "sellUnit": 9_500_000,
        "taxPct": 0, "note": "",
    }
    base.update(over)
    return base


def _create_payload(lines, **over):
    payload = {
        "customerId": "CUST-ANPHAT", "salesOwner": "Minh Tú",
        "origin": "NINGBO", "destination": "Hải Phòng", "route": "Biển",
        "incoterm": "FOB", "transportMode": "Sea", "goodsDescription": "Linh kiện",
        "expectedDeliveryDate": "2026-10-05", "serviceLines": lines,
    }
    payload.update(over)
    return payload


# ---------------------------------------------------------------------------
# BLOCKER A — the customer-facing general note must survive persistence
# ---------------------------------------------------------------------------

def test_pasted_general_note_survives_create_and_fresh_readback(staging, event_seq):
    out = _dispatch(event_seq, "quote_create",
                    _create_payload([_line()], customerNotes=[GENERAL_NOTE]))
    assert out["ok"], out
    qid = out["id"]

    # The dispatcher's own fresh readback already carries it...
    assert out["readback"]["customerNotes"] == [GENERAL_NOTE]

    # ...and so does an INDEPENDENT fresh read of the canonical document.
    stored = logi_c1.read_quote(qid)
    assert stored is not None
    assert stored["customer_notes"] == [GENERAL_NOTE]
    assert logi_c1.quote_projection(stored)["customerNotes"] == [GENERAL_NOTE]


def test_general_note_is_not_dumped_into_the_internal_assumption_list(staging, event_seq):
    out = _dispatch(event_seq, "quote_create",
                    _create_payload([_line()], customerNotes=[GENERAL_NOTE]))
    stored = logi_c1.read_quote(out["id"])
    # Internal control statements stay internal and untouched...
    assert list(stored["assumptions"]) == list(INTERNAL_ASSUMPTIONS)
    # ...and the customer sentence is NOT smuggled into them.
    assert GENERAL_NOTE not in stored["assumptions"]


def test_edit_preserves_notes_it_does_not_touch_and_replaces_ones_it_does(staging, event_seq):
    qid = _dispatch(event_seq, "quote_create",
                    _create_payload([_line()], customerNotes=[GENERAL_NOTE]))["id"]

    # An edit that never mentions the notes must not erase them.
    out = _dispatch(event_seq, "quote_edit", {"id": qid, "serviceLines": [_line()]})
    assert out["ok"], out
    assert out["readback"]["customerNotes"] == [GENERAL_NOTE]

    # An edit that DOES carry notes replaces them wholesale (operator intent).
    out = _dispatch(event_seq, "quote_edit",
                    {"id": qid, "customerNotes": [GENERAL_NOTE, SECOND_NOTE]})
    assert out["ok"], out
    assert out["readback"]["customerNotes"] == [GENERAL_NOTE, SECOND_NOTE]
    assert logi_c1.read_quote(qid)["customer_notes"] == [GENERAL_NOTE, SECOND_NOTE]

    # Explicitly clearing them is allowed too — but only when asked.
    out = _dispatch(event_seq, "quote_edit", {"id": qid, "customerNotes": []})
    assert out["readback"]["customerNotes"] == []


def test_blank_and_duplicate_notes_are_normalised_not_stored_as_noise(staging, event_seq):
    out = _dispatch(event_seq, "quote_create", _create_payload(
        [_line()], customerNotes=[GENERAL_NOTE, "   ", GENERAL_NOTE, SECOND_NOTE]))
    assert out["readback"]["customerNotes"] == [GENERAL_NOTE, SECOND_NOTE]


# ---------------------------------------------------------------------------
# BLOCKER B — UNKNOWN must never become a canonical ZERO
#
# SUPERSEDED IN PART by LOGI_QUOTE_COST_WORKFLOW_UX_01_REPAIR_02: the invariant
# "UNKNOWN != ZERO" stands, but the way it was enforced here (refuse the save)
# was wrong for the workflow. Blocking the save made the Sale type a fake 0 to
# get past the form, which produces the very lie the rule exists to prevent.
# The canonical rule is now UNKNOWN != ZERO **and** UNKNOWN != SAVE BLOCKER:
# a blank cost/VAT persists as None. Full coverage lives in
# tests/test_logi_quote_cost_workflow_ux_01_repair_02.py.
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("field", ["costUnit", "taxPct"])
def test_blank_optional_line_number_persists_as_unknown_not_zero(
    staging, event_seq, field
):
    """This is the exact shape the quick-paste form produces: sell only."""
    out = _dispatch(event_seq, "quote_create",
                    _create_payload([_line(**{field: ""})]))
    assert out["ok"], out
    stored = logi_c1.read_quote(out["id"])["service_lines"][0]
    key = {"costUnit": "cost_unit", "taxPct": "tax_pct"}[field]
    assert stored[key] is None, f"a blank {field} must stay UNKNOWN, never become 0"


def test_blank_sell_price_is_still_refused(staging, event_seq):
    """Giá bán is the number quoted to the customer — it cannot be UNKNOWN."""
    out = _dispatch(event_seq, "quote_create", _create_payload([_line(sellUnit="")]))
    assert not out["ok"]
    assert "Dòng dịch vụ 1" in out["error"] and "Giá bán" in out["error"]
    assert out["id"] == ""


def test_missing_optional_key_entirely_is_also_unknown(staging, event_seq):
    line = _line()
    del line["costUnit"]
    out = _dispatch(event_seq, "quote_create", _create_payload([line]))
    assert out["ok"], out
    assert logi_c1.read_quote(out["id"])["service_lines"][0]["cost_unit"] is None


def test_an_explicit_zero_is_still_a_valid_answer(staging, event_seq):
    """UNKNOWN and ZERO are different answers; 0 must remain usable."""
    out = _dispatch(event_seq, "quote_create",
                    _create_payload([_line(costUnit=0, taxPct=0)]))
    assert out["ok"], out
    line = logi_c1.read_quote(out["id"])["service_lines"][0]
    assert float(line["cost_unit"]) == 0.0
    assert float(line["tax_pct"]) == 0.0


def test_quote_edit_also_keeps_a_blank_as_unknown(staging, event_seq):
    qid = _dispatch(event_seq, "quote_create", _create_payload([_line()]))["id"]
    out = _dispatch(event_seq, "quote_edit",
                    {"id": qid, "serviceLines": [_line(taxPct="")]})
    assert out["ok"], out
    assert logi_c1.read_quote(qid)["service_lines"][0]["tax_pct"] is None


# ---------------------------------------------------------------------------
# BLOCKERS C + D and the REAL export path (create → read → export handler)
# ---------------------------------------------------------------------------

def _export(seq_fn, qid):
    out = _dispatch(seq_fn, "quote_export_xlsx", {"id": qid, "requestId": f"rq-{qid}"})
    assert out["ok"], out
    doc = out["document"]
    assert doc["byteSize"] > 0
    return load_workbook(io.BytesIO(base64.b64decode(doc["downloadBase64"]))).active, doc


def _flat(ws):
    return " \n".join(
        " ".join(str(c) for c in row if c is not None)
        for row in ws.iter_rows(values_only=True)
    )


def test_real_export_path_carries_authority_customer_notes_and_date(staging, event_seq):
    """ONE end-to-end pass: canonical create → fresh read → export handler."""
    # Enrich the EXISTING customer master through its own canonical writer, so
    # the export has real contact data to find (staging only).
    enriched = _dispatch(event_seq, "customer_edit", {
        "id": "CUST-ANPHAT", "name": "An Phát JSC",
        # `customer_edit` takes the canonical snake_case keys (existing writer).
        "contact_person": "Chị Lan (Mua hàng)",
        "email": "muahang@anphat.vn", "phone": "0912345678",
    })
    assert enriched["ok"], enriched

    created = _dispatch(event_seq, "quote_create", _create_payload(
        [
            _line(note="cước tại thời điểm này"),
            _line(code="DV-LOCAL-CN", serviceName="Phí local đầu Trung Quốc",
                  currency="CNY", sellUnit=1200, costUnit=1100, taxPct=0,
                  note="Có thể phát sinh (nếu có)"),
        ],
        customerNotes=[GENERAL_NOTE],
        fx={"CNY": 3600},
    ))
    assert created["ok"], created
    qid = created["id"]

    ws, doc = _export(event_seq, qid)
    text = _flat(ws)

    # --- C: company authority read from the REAL field names -----------------
    assert ws["A1"].value == "CÔNG TY TNHH GIAO NHẬN QUỐC TẾ TRƯỜNG THÀNH"
    assert ws["A2"].value == "Số 12, Đường Láng Hạ, Ba Đình, Hà Nội"   # legal_address
    assert ws["A3"].value == "info@truongthanh.vn"
    assert "Nguyễn Thị Thu Hằng" in str(ws["B7"].value)                # legal_representative
    assert "Giám đốc" in str(ws["B7"].value)                            # representative_title
    assert ws["B9"].value == "0241234567"                               # phone

    # --- C: customer contact enrichment from the customer master ------------
    assert ws["G7"].value == "An Phát JSC"
    assert ws["B8"].value == "Chị Lan (Mua hàng)"
    assert ws["G8"].value == "muahang@anphat.vn"
    assert ws["G9"].value == "0912345678"

    # --- D: the quote's OWN date, not "today" -------------------------------
    stored = logi_c1.read_quote(qid)
    assert str(ws["G10"].value) == str(stored["created_at"]).split("T")[0]

    # --- A: the pasted sentence reaches the customer's Ghi chú --------------
    assert GENERAL_NOTE in text

    # --- A: internal control statements are NOT leaked to the customer ------
    for internal in INTERNAL_ASSUMPTIONS:
        assert internal not in text, f"internal assumption leaked to customer: {internal}"

    # --- per-line notes stay on their own lines -----------------------------
    assert "cước tại thời điểm này" in str(ws.cell(row=18, column=9).value)
    assert "nếu có" in str(ws.cell(row=19, column=9).value).lower()
    assert "cước tại thời điểm này" not in str(ws.cell(row=19, column=9).value)

    # --- currencies stay apart ----------------------------------------------
    assert ws.cell(row=18, column=8).value == "VND"
    assert ws.cell(row=19, column=8).value == "CNY"
    assert "Tổng VND" in text and "Tổng CNY" in text

    # --- workbook hygiene ----------------------------------------------------
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith("="), f"{cell.coordinate} is a formula"
                for broken in ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!"):
                    assert broken not in cell.value
    lowered = text.lower()
    assert "phú anh" not in lowered and "phu anh" not in lowered

    assert doc["fileName"].endswith(".xlsx")


def test_reexporting_an_older_quote_keeps_its_own_date(staging, event_seq):
    """An old quote re-exported today must not be restamped with today."""
    from datetime import date

    qid = _dispatch(event_seq, "quote_create", _create_payload([_line()]))["id"]

    # Age the canonical document the way a real June quote would look.
    from core.database import load_quotes, save_quotes_to_cloud

    stored = dict(load_quotes()[qid])
    stored["created_at"] = "2026-06-12T09:31:07+00:00"
    assert save_quotes_to_cloud({qid: stored})

    ws, _ = _export(event_seq, qid)
    assert str(ws["G10"].value) == "2026-06-12"
    assert str(ws["G10"].value) != str(date.today())


def test_an_explicit_quote_date_outranks_created_at(staging, event_seq):
    from core.database import load_quotes, save_quotes_to_cloud

    qid = _dispatch(event_seq, "quote_create", _create_payload([_line()]))["id"]
    stored = dict(load_quotes()[qid])
    stored["created_at"] = "2026-06-12T09:31:07+00:00"
    stored["date"] = "2026-06-15"
    assert save_quotes_to_cloud({qid: stored})

    ws, _ = _export(event_seq, qid)
    assert str(ws["G10"].value) == "2026-06-15"


def test_export_of_a_missing_quote_fails_closed(staging, event_seq):
    out = _dispatch(event_seq, "quote_export_xlsx", {"id": "Q-NOPE", "requestId": "rq-1"})
    assert not out["ok"]
    assert "Q-NOPE" in out["error"]
