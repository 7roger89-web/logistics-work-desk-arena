"""LOGI_QUOTE_COST_WORKFLOW_UX_01_REPAIR_02 — UNKNOWN is not ZERO, and not a blocker.

REPAIR_01 correctly refused to persist a blank Giá vốn / VAT as a canonical 0.
It enforced that by refusing the Save. In the Owner's real workflow that is the
wrong end of the trade: the Sale pastes a customer message that states SELL
prices only, so a hard block just teaches the operator to type a fake 0 (or a
fake FX rate) to get the DRAFT saved — manufacturing exactly the false data the
rule exists to prevent.

The canonical rule this file pins is therefore:

    UNKNOWN != ZERO   and   UNKNOWN != SAVE BLOCKER

A DRAFT is savable from what is actually known: service name, quantity, SELL
price, currency. ``cost_unit``, ``tax_pct`` and the quote-time FX rate may all
stay unknown; each one that is unknown makes its OWN derived figures unknown
(never 0), including the margin — the quote never claims a 100% margin because
nobody typed a cost.

Coverage A–G:
  A  Owner paste shape (9 lines, sell only, no cost/VAT/FX) saves.
  B  a canonical line keeps cost_unit / tax_pct as None, not 0.
  C  mixed USD (no FX) + VND: originals preserved, no fake aggregate/margin.
  D  an explicit 0 is still a real zero.
  E  a later edit supplying cost/VAT/FX resolves the totals, same quote identity.
  F  XLSX: per-currency subtotals, unknown VAT is not "0%", explicit 0 is "0%".
  G  quote → shipment draft carries UNKNOWN through, never a fabricated 0.

Everything runs on the isolated staging transport (mongomock) through the REAL
writers. The production database is never touched.
"""

from __future__ import annotations

import base64
import io

import pytest
from openpyxl import load_workbook

from views import logi_c1
from views import logi_sandbox

GENERAL_NOTE = "Nếu có phát sinh về xử lý hàng hóa, sẽ báo ngay tại thời điểm phát sinh."


@pytest.fixture()
def staging():
    import streamlit as st

    logi_sandbox.reset_backend()
    transport = logi_sandbox.ensure_backend()
    assert transport == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    yield transport
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    logi_sandbox.reset_backend()


@pytest.fixture()
def event_seq():
    state = {"n": 0}

    def _make(name: str, payload: dict) -> dict:
        state["n"] += 1
        return {name: {"seq": state["n"], "event": name, "payload": payload}}

    return _make


def _dispatch(seq_fn, name, payload):
    outcomes = logi_c1.consume_c1_triggers({"seq": 0}, seq_fn(name, payload))
    assert len(outcomes) == 1
    return outcomes[0]


def _create(seq_fn, lines, **over):
    payload = {
        "customerId": "CUST-ANPHAT", "salesOwner": "Minh Tú",
        "origin": "NINGBO", "destination": "Hải Phòng", "route": "Biển",
        "incoterm": "FOB", "transportMode": "Sea", "goodsDescription": "Linh kiện",
        "expectedDeliveryDate": "2026-10-05", "serviceLines": lines,
    }
    payload.update(over)
    return _dispatch(seq_fn, "quote_create", payload)


# The Owner's FOB NINGBO paste, as the parser hands it to the form: every line
# states a SELL price and a currency, and nothing else. No cost. No VAT. No FX.
OWNER_PASTE_LINES = [
    {"serviceName": n, "unit": u, "qty": 1, "currency": c,
     "sellUnit": s, "costUnit": "", "taxPct": "", "note": ""}
    for n, u, c, s in [
        ("Cước biển Ningbo - Hải Phòng", "Cont 20'", "USD", 950),
        ("THC đầu Ningbo", "Cont", "USD", 120),
        ("Phí chứng từ", "Bộ", "USD", 45),
        ("Phí seal", "Cái", "USD", 12),
        ("THC đầu Hải Phòng", "Cont", "VND", 3_100_000),
        ("Phí D/O", "Bộ", "VND", 850_000),
        ("Khai hải quan", "Tờ khai", "VND", 900_000),
        ("Vận chuyển nội địa", "Chuyến", "VND", 4_200_000),
        ("Phí nâng hạ", "Lần", "VND", 650_000),
    ]
]


# ---------------------------------------------------------------------------
# A — the Owner's paste saves as a DRAFT with no cost, no VAT and no FX
# ---------------------------------------------------------------------------

def test_a_owner_paste_saves_without_cost_vat_or_fx(staging, event_seq):
    out = _create(event_seq, OWNER_PASTE_LINES, customerNotes=[GENERAL_NOTE])
    assert out["ok"], out
    stored = logi_c1.read_quote(out["id"])
    assert stored["status"] == "DRAFT"
    assert len(stored["service_lines"]) == 9
    assert stored["customer_notes"] == [GENERAL_NOTE]
    # Not one line acquired a cost or a VAT nobody stated.
    assert all(line["cost_unit"] is None for line in stored["service_lines"])
    assert all(line["tax_pct"] is None for line in stored["service_lines"])
    # No FX rate was invented for the USD lines.
    assert "USD" not in dict(stored["fx_snapshot"]["rates_to_vnd"])


# ---------------------------------------------------------------------------
# B — the canonical line keeps UNKNOWN as UNKNOWN
# ---------------------------------------------------------------------------

def test_b_canonical_line_keeps_unknown_cost_and_vat(staging, event_seq):
    qid = _create(event_seq, OWNER_PASTE_LINES)["id"]
    line = logi_c1.read_quote(qid)["service_lines"][0]
    assert line["currency"] == "USD"
    assert float(line["sell_unit"]) == 950.0
    assert float(line["sell_total"]) == 950.0
    assert line["cost_unit"] is None and line["cost_total"] is None
    assert line["tax_pct"] is None and line["tax_total"] is None
    # …and the projection the desk renders says UNKNOWN too, not 0.
    projected = logi_c1.quote_projection(logi_c1.read_quote(qid))["serviceLines"][0]
    assert projected["costUnit"] is None and projected["taxPct"] is None
    assert projected["sellUnit"] == 950.0


# ---------------------------------------------------------------------------
# C — mixed currency with no FX: originals survive, aggregates stay honest
# ---------------------------------------------------------------------------

def test_c_missing_fx_preserves_originals_and_fakes_nothing(staging, event_seq):
    qid = _create(event_seq, OWNER_PASTE_LINES)["id"]
    stored = logi_c1.read_quote(qid)
    usd, vnd = stored["service_lines"][0], stored["service_lines"][4]

    # Original-currency truth is intact on both sides.
    assert usd["currency"] == "USD" and float(usd["sell_total"]) == 950.0
    assert vnd["currency"] == "VND" and float(vnd["sell_total"]) == 3_100_000.0
    # The USD line simply has no VND value; the VND line does.
    assert usd["sell_vnd"] is None
    assert float(vnd["sell_vnd"]) == 3_100_000.0

    calc = stored["calculation"]
    # No fake combined VND total, and above all no fake margin.
    assert calc["selling_price_before_tax_vnd"] is None
    assert calc["selling_price_total_vnd"] is None
    assert calc["estimated_cost_vnd"] is None
    assert calc["estimated_output_tax_vnd"] is None
    assert calc["estimated_margin_vnd"] is None
    assert calc["estimated_margin_pct"] is None
    assert calc["unknown_reasons"], "an UNKNOWN total must say why"

    # Per-currency totals remain fully usable without any rate.
    by_ccy = {row["currency"]: row for row in calc["totals_by_currency"]}
    assert set(by_ccy) == {"USD", "VND"}
    assert by_ccy["USD"]["sell_total"] == pytest.approx(950 + 120 + 45 + 12)
    assert by_ccy["VND"]["sell_total"] == pytest.approx(
        3_100_000 + 850_000 + 900_000 + 4_200_000 + 650_000)


def test_c2_margin_is_unknown_when_only_the_cost_is_unknown(staging, event_seq):
    """Even in pure VND with a full FX picture, unknown cost => unknown margin."""
    qid = _create(event_seq, [{
        "serviceName": "Khai hải quan", "unit": "Tờ khai", "qty": 1, "currency": "VND",
        "sellUnit": 900_000, "costUnit": "", "taxPct": 8, "note": "",
    }])["id"]
    calc = logi_c1.read_quote(qid)["calculation"]
    assert calc["selling_price_before_tax_vnd"] == pytest.approx(900_000)
    assert calc["estimated_output_tax_vnd"] == pytest.approx(72_000)
    assert calc["estimated_cost_vnd"] is None
    assert calc["estimated_margin_vnd"] is None, "margin must not assume cost = 0"
    assert calc["estimated_margin_pct"] is None


# ---------------------------------------------------------------------------
# D — an explicit zero is a real zero
# ---------------------------------------------------------------------------

def test_d_explicit_zero_cost_and_vat_persist_as_zero(staging, event_seq):
    qid = _create(event_seq, [{
        "serviceName": "Phí miễn phí", "unit": "Lần", "qty": 1, "currency": "VND",
        "sellUnit": 1_000_000, "costUnit": 0, "taxPct": 0, "note": "",
    }])["id"]
    stored = logi_c1.read_quote(qid)
    line = stored["service_lines"][0]
    assert float(line["cost_unit"]) == 0.0 and float(line["cost_total"]) == 0.0
    assert float(line["tax_pct"]) == 0.0 and float(line["tax_total"]) == 0.0
    calc = stored["calculation"]
    # Every figure is known here, including a real 100% margin.
    assert calc["estimated_cost_vnd"] == pytest.approx(0.0)
    assert calc["estimated_margin_vnd"] == pytest.approx(1_000_000)
    assert calc["estimated_margin_pct"] == pytest.approx(100.0)


# ---------------------------------------------------------------------------
# E — supplying the facts later resolves the totals, same quote
# ---------------------------------------------------------------------------

def test_e_later_edit_resolves_unknowns_without_a_new_quote(staging, event_seq):
    created = _create(event_seq, [{
        "serviceName": "Cước biển", "unit": "Cont 20'", "qty": 2, "currency": "USD",
        "sellUnit": 950, "costUnit": "", "taxPct": "", "note": "",
    }])
    qid, quote_no = created["id"], logi_c1.read_quote(created["id"])["quote_no"]
    assert logi_c1.read_quote(qid)["calculation"]["estimated_margin_vnd"] is None

    out = _dispatch(event_seq, "quote_edit", {
        "id": qid,
        "serviceLines": [{
            "serviceName": "Cước biển", "unit": "Cont 20'", "qty": 2, "currency": "USD",
            "sellUnit": 950, "costUnit": 800, "taxPct": 10, "note": "",
        }],
        "fx": {"USD": 25_500},
    })
    assert out["ok"], out

    stored = logi_c1.read_quote(qid)
    assert stored["id"] == qid and stored["quote_no"] == quote_no, "identity is stable"
    calc = stored["calculation"]
    assert calc["selling_price_before_tax_vnd"] == pytest.approx(2 * 950 * 25_500)
    assert calc["estimated_cost_vnd"] == pytest.approx(2 * 800 * 25_500)
    assert calc["estimated_output_tax_vnd"] == pytest.approx(2 * 950 * 0.10 * 25_500)
    assert calc["estimated_margin_vnd"] == pytest.approx(2 * 150 * 25_500)
    assert calc["unknown_reasons"] == []


# ---------------------------------------------------------------------------
# F — the customer XLSX tells the same truth
# ---------------------------------------------------------------------------

def _export(seq_fn, qid):
    out = _dispatch(seq_fn, "quote_export_xlsx", {"id": qid, "requestId": f"rq-{qid}"})
    assert out["ok"], out
    data = base64.b64decode(out["document"]["downloadBase64"])
    return load_workbook(io.BytesIO(data)).active


def _rows(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def _flat(ws):
    return " \n".join(" ".join(str(c) for c in row if c is not None) for row in _rows(ws))


def test_f_xlsx_unknown_vat_is_not_zero_percent_and_no_fake_combined_total(
    staging, event_seq
):
    qid = _create(event_seq, OWNER_PASTE_LINES, customerNotes=[GENERAL_NOTE])["id"]
    ws = _export(event_seq, qid)
    text = _flat(ws)

    # Unknown VAT renders as an explicit UNKNOWN mark, never as a 0% promise.
    vat_cells = [row[5] for row in _rows(ws)[17:26]]
    assert vat_cells and all(str(v) == "—" for v in vat_cells), vat_cells
    assert "0%" not in text, "an unfilled VAT must never read as 0%"

    # One subtotal per currency, and no mixed-currency grand total.
    assert "TỔNG CỘNG THEO TỪNG LOẠI TIỀN (không quy đổi gộp)" in text
    assert "Tổng USD" in text and "Tổng VND" in text
    # No FX rate was supplied, so the sheet says so instead of converting.
    assert "chưa có tỷ giá báo giá" in text
    # The customer's own note still reaches the document (REPAIR_01 stays green).
    assert GENERAL_NOTE in text


def test_f2_xlsx_shows_an_explicit_zero_vat_as_zero_percent(staging, event_seq):
    qid = _create(event_seq, [{
        "serviceName": "Phí miễn thuế", "unit": "Lần", "qty": 1, "currency": "VND",
        "sellUnit": 1_000_000, "costUnit": 0, "taxPct": 0, "note": "",
    }])["id"]
    ws = _export(event_seq, qid)
    line_row = _rows(ws)[17]
    assert str(line_row[5]) == "0%", "an explicit zero VAT is a real answer"
    assert line_row[6] == pytest.approx(1_000_000), "0% VAT => after-VAT equals pre-tax"


def test_f3_pre_tax_line_amount_is_shown_even_when_vat_is_unknown(staging, event_seq):
    qid = _create(event_seq, [{
        "serviceName": "Cước biển", "unit": "Cont 20'", "qty": 2, "currency": "USD",
        "sellUnit": 950, "costUnit": "", "taxPct": "", "note": "",
    }])["id"]
    ws = _export(event_seq, qid)
    line_row = _rows(ws)[17]
    assert line_row[2] == pytest.approx(950)        # Đơn giá
    assert line_row[3] == pytest.approx(2)          # Số lượng
    assert line_row[4] == pytest.approx(1900)       # Thành tiền (trước VAT)
    assert str(line_row[5]) == "—"                  # VAT UNKNOWN
    assert str(line_row[6]) == "—"                  # after-VAT depends on it


def test_f4_internal_assumptions_never_reach_the_customer_sheet(staging, event_seq):
    qid = _create(event_seq, OWNER_PASTE_LINES, customerNotes=[GENERAL_NOTE])["id"]
    text = _flat(_export(event_seq, qid))
    assert "Giá agent/đối tác" not in text
    assert "Không tự động hóa khai báo hải quan" not in text


# ---------------------------------------------------------------------------
# G — the handoff to a shipment must not resurrect zeros
# ---------------------------------------------------------------------------

def test_g_shipment_draft_carries_unknown_through(staging, event_seq):
    from core.quotations import build_shipment_draft

    qid = _create(event_seq, OWNER_PASTE_LINES)["id"]
    draft = build_shipment_draft(logi_c1.read_quote(qid), "SHP-TEST-001")
    details = draft["data"]["3.4"]["details"]
    assert len(details) == 9
    for detail in details:
        assert detail["cost"] is None, "an unknown quote cost must not become 0đ"
        assert detail["vat"] is None and detail["vat_line"] is None
        assert detail["quote_estimate"] is True
    # The USD lines have no VND price either — but their original amount travels.
    usd = details[0]
    assert usd["price"] is None
    assert usd["quote_currency"] == "USD"
    assert usd["quote_sell_total"] == pytest.approx(950)
    # A VND line keeps its real number.
    assert details[4]["price"] == pytest.approx(3_100_000)


def test_g2_shipment_draft_keeps_a_real_zero_as_zero(staging, event_seq):
    from core.quotations import build_shipment_draft

    qid = _create(event_seq, [{
        "serviceName": "Phí miễn phí", "unit": "Lần", "qty": 1, "currency": "VND",
        "sellUnit": 1_000_000, "costUnit": 0, "taxPct": 0, "note": "",
    }])["id"]
    detail = build_shipment_draft(logi_c1.read_quote(qid), "SHP-TEST-002")["data"]["3.4"]["details"][0]
    assert detail["cost"] == pytest.approx(0.0)
    assert detail["vat"] == pytest.approx(0.0)
