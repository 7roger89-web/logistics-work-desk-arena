"""LOGI_QUOTE_COST_WORKFLOW_UX_01 — the customer quotation XLSX export.

Outcome C: the Owner can hand a customer a real quotation file instead of a
chat message. The evidence here is about TRUTH in the workbook, not decoration:

* the Owner's template hierarchy is reproduced in the canonical anchor rows
  (header / From-To / subject / table header row 17 / notes-signature tail);
* every service line keeps its OWN currency and there is exactly ONE subtotal
  per currency — different currencies are never summed into one number;
* a VND conversion appears ONLY when the quote itself stores a quote-time FX
  rate, and that exact rate is disclosed on the sheet;
* conditional ("nếu có") and estimated ("khoảng") charges stay visibly
  conditional / estimated rather than reading as committed prices;
* NO formula is written at all, so the sample workbook's broken #REF! cells
  cannot be reproduced;
* the sample's Phú Anh branding is never carried over.

Run: python3 -m pytest tests/test_logi_quote_cost_workflow_ux_01.py -q
"""

from __future__ import annotations

from openpyxl import load_workbook

from core.excel_utils import generate_quote_excel

OWNER_QUOTE = {
    "quote_no": "BG-2026-0311",
    "customer_name": "CÔNG TY TNHH AN PHÁT",
    "customer_email": "muahang@anphat.vn",
    "origin": "NINGBO",
    "destination": "Hải Phòng",
    "incoterm": "FOB",
    "goods_description": "Linh kiện điện tử",
    "service_lines": [
        {
            "service_name": "Cước biển",
            "sell_unit": 950, "qty": 1, "currency": "USD", "tax_pct": 0,
            "note": "cước tại thời điểm này, nếu vượt 20 tấn cộng thêm 80 USD",
        },
        {
            "service_name": "Phí local đầu Trung Quốc",
            "sell_unit": 9_500_000, "qty": 1, "currency": "VND", "tax_pct": 8,
        },
        {
            "service_name": "Phí kiểm hóa",
            "sell_unit": 4_000_000, "qty": 1, "currency": "VND", "tax_pct": 8,
            "condition": "IF_APPLICABLE",
        },
        {
            "service_name": "Phí xử lý hàng",
            "sell_unit": 10_000_000, "qty": 1, "currency": "VND", "tax_pct": 0,
            "status": "ESTIMATE_REQUIRES_CONFIRMATION",
            "note": "KHOẢNG 10 triệu — chưa làm bao giờ nên chưa nắm rõ",
        },
    ],
    "assumptions": ["Nếu có phát sinh về xử lý hàng hóa tại cảng thì sẽ báo lại sau."],
}


def _sheet(quote):
    return load_workbook(generate_quote_excel(quote)).active


def _rows(ws):
    return [
        [c for c in row if c is not None]
        for row in ws.iter_rows(values_only=True)
    ]


def _flat(ws):
    return " \n".join(" ".join(str(c) for c in row) for row in _rows(ws) if row)


def test_template_hierarchy_anchor_rows():
    """The Owner's layout: header, From/To, subject, row-17 table, notes tail."""
    ws = _sheet(OWNER_QUOTE)
    assert ws["A5"].value == "BÁO GIÁ"
    assert ws["A7"].value == "From" and ws["F7"].value == "To"
    assert ws["G7"].value == "CÔNG TY TNHH AN PHÁT"
    assert ws["A8"].value == "Attn" and ws["F8"].value == "Email"
    assert ws["A9"].value == "Tel/Fax"
    assert ws["A10"].value == "Mobile" and ws["F10"].value == "Date"
    assert str(ws["A11"].value).startswith("Subject:")
    assert "NINGBO - Hải Phòng" in str(ws["A11"].value)
    assert "Mặt hàng" in str(ws["A13"].value)
    assert "FOB" in str(ws["A14"].value)

    headers = [ws.cell(row=17, column=i).value for i in range(1, 10)]
    assert headers == [
        "STT", "Chi tiết", "Đơn giá", "Số lượng", "Thành tiền",
        "VAT", "thành tiền", "Đơn vị tiền", "Phát sinh (nếu có)",
    ]
    # The line block starts immediately under the header row.
    assert ws.cell(row=18, column=1).value == 1
    assert ws.cell(row=18, column=2).value == "Cước biển"
    assert "Ghi chú:" in _flat(ws)


def test_every_line_keeps_its_own_currency_and_totals_never_merge():
    ws = _sheet(OWNER_QUOTE)
    assert ws.cell(row=18, column=8).value == "USD"
    assert ws.cell(row=19, column=8).value == "VND"

    text = _flat(ws)
    assert "TỔNG CỘNG THEO TỪNG LOẠI TIỀN" in text
    assert "Tổng USD" in text and "Tổng VND" in text

    rows = _rows(ws)
    usd_total = next(r for r in rows if r and r[0] == "Tổng USD")
    vnd_total = next(r for r in rows if r and r[0] == "Tổng VND")
    assert usd_total[1] == 950                      # 950 USD, VAT 0
    # 9,500,000*1.08 + 4,000,000*1.08 + 10,000,000 = 24,580,000
    assert vnd_total[1] == 9_500_000 * 1.08 + 4_000_000 * 1.08 + 10_000_000
    # The USD amount is NEVER folded into the VND subtotal.
    assert usd_total[1] != vnd_total[1]


def test_no_fx_stored_means_no_invented_conversion():
    ws = _sheet(OWNER_QUOTE)  # this quote stores no fx_snapshot
    text = _flat(ws)
    assert "chưa có tỷ giá báo giá" in text
    assert "Tỷ giá báo giá đã dùng" not in text


def test_a_stored_quote_time_fx_is_used_and_disclosed():
    quote = dict(OWNER_QUOTE)
    quote["fx_snapshot"] = {
        "rates_to_vnd": {"USD": 25_500},
        "source": "Tỷ giá chốt lúc báo giá 22/09/2026",
    }
    ws = _sheet(quote)
    rows = _rows(ws)
    disclosure = next(r for r in rows if r and str(r[0]).startswith("USD → VND"))
    assert "25,500" in str(disclosure[0])            # the exact rate is shown
    assert disclosure[1] == 950 * 25_500             # converted from the USD total only
    assert "Tỷ giá chốt lúc báo giá 22/09/2026" in _flat(ws)


def test_conditional_and_estimated_charges_stay_qualified():
    ws = _sheet(OWNER_QUOTE)
    conditional = ws.cell(row=20, column=9).value      # Phí kiểm hóa
    estimated = ws.cell(row=21, column=9).value        # Phí xử lý hàng
    assert "nếu có" in str(conditional).lower()
    assert "ước tính" in str(estimated).lower()
    assert "chưa phải chi phí chắc chắn" in str(estimated).lower()
    # The Owner's own uncertainty wording survives verbatim.
    assert "chưa nắm rõ" in str(estimated)


def test_no_formula_is_written_so_no_broken_ref_can_exist():
    """The sample workbook carried #REF! in L19/L20/L22/L23. We write values."""
    ws = _sheet(OWNER_QUOTE)
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str):
                assert not value.startswith("="), f"{cell.coordinate} is a formula"
                for broken in ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!"):
                    assert broken not in value, f"{cell.coordinate} carries {broken}"


def test_sample_branding_is_not_reproduced():
    text = _flat(_sheet(OWNER_QUOTE)).lower()
    assert "phú anh" not in text and "phu anh" not in text
    assert "trường thành" in text


def test_empty_quote_exports_without_inventing_anything():
    ws = _sheet({"customer_name": "Khách mới"})
    text = _flat(ws)
    assert "BÁO GIÁ" in text
    assert "Khách mới" in text
    assert "TỔNG CỘNG" in text
    # No line, no total, no invented currency row.
    assert ws.cell(row=18, column=1).value is None
