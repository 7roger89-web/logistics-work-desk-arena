"""Issue #168 — FORM_07 Biên bản giao nhận pilot acceptance.

Proves on pure unit seams AND on the real staging dispatcher:
  * the exact Owner source receipt is pinned (SHA-256 / blob / size) and the
    Git-resident sanitized derivative matches its own pin, carries no banned
    artifact and no personal metadata from the original;
  * a drifted/missing asset fails closed (catalog unavailable, generation
    refused);
  * FORM_07 enters through the server-owned catalog (no hard-coded checkbox);
  * selection persistence/readback and deselection history retention reuse the
    existing picker event;
  * server-owned canonical readiness fails closed on every missing fact;
  * RBAC refuses a viewer with zero storage writes;
  * an authorized generation renders a real, structurally valid, deterministic
    DOCX with repeating goods rows, stores it through the EXISTING durable
    attachment path, writes immutable revision metadata with bounded
    provenance, survives deselection, and re-downloads the exact bytes after a
    remount through the EXISTING attachment_download event.
"""

from __future__ import annotations

import base64
import hashlib
import io
import zipfile
from pathlib import Path

import pytest

from core import database
from core.docgen import form07 as form07_mod
from core.docgen.form07 import (
    FORM07_DOCUMENT_TYPE,
    build_form07_context,
    evaluate_form07_readiness,
    load_template_spec,
    render_form07_docx,
    scan_package_safety,
    template_asset_available,
)
from core.shipment_attachments import fetch_attachment_bytes
from core.shipment_document_catalog import SHIPMENT_DOCUMENT_CATALOG, project_catalog
from views import logi_real_data as rd
from views import logi_sandbox
from views import logi_view_model as vm

REPO_ROOT = Path(__file__).resolve().parents[1]

# Durable Owner source receipt (Issue #168 / owner-input/form07-source-168).
SOURCE_SHA256 = "cd2a1df5434a217bdae34e292844ad1bacbcc737b2282137b835af46663161b5"
SOURCE_SIZE_BYTES = 18997
SOURCE_GIT_BLOB_SHA = "e2c3dea169bff006f4073f7fae2b2c285c912f3f"

SHIPMENT_ID = "LO-0295"          # staging seed: delivered, contract HDMB-0295
UNDELIVERED_ID = "LO-0306"       # staging seed: not delivered
CONTRACT_ID = "HDMB-0295"

VIEWER_ACTOR = {
    "username": "staging.viewer",
    "full_name": "Viewer (Test)",
    "role": "VIEWER",
    "user_id": "STAGING-VIEWER",
}

GOODS_LINES = [
    {
        "description": "Kính nổi 5mm",
        "specification_or_code": "1200x600",
        "unit": "Tấm",
        "quantity": "100",
        "unit_price_vnd_excl_vat": "350000",
        "vat_rate": "0.08",
    },
    {
        "description": "Khung nhôm định hình",
        "unit": "Bộ",
        "quantity": "20.5",
        "unit_price_vnd_excl_vat": None,   # FORM_07 prints no money: UNKNOWN price must not block
    },
]

MANUAL = {"form07_no": "BBGN-2026-0295", "order_date": "25/09/2026"}


@pytest.fixture()
def staging():
    import streamlit as st

    logi_sandbox.reset_backend()
    assert logi_sandbox.ensure_backend() == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    st.session_state.pop(vm.LAST_WRITES_KEY, None)
    yield
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    logi_sandbox.reset_backend()


def _dispatch(name: str, payload: dict, seq: int = 1) -> dict:
    outcomes = vm.consume_phase_b_triggers(
        {"seq": seq - 1},
        {name: {"seq": seq, "event": name, "payload": payload}},
    )
    assert len(outcomes) == 1
    return outcomes[0]


def _prepare_source_contract() -> None:
    """Give the seeded HDMB-0295 the canonical parties + goods an operator books."""
    db = database.get_db()
    db["contracts"].update_one(
        {"id": CONTRACT_ID},
        {"$set": {"agent_id": "AGT-PHUANH", "company_id": "TT-COMPANY", "goods_lines": GOODS_LINES}},
    )


def _full_facts() -> dict:
    return {
        "delivery": {
            "is_delivered": True,
            "delivered_at": "2026-09-09",
            "delivery_location": "Kho Thiên Ân — KP.2, Thủ Đức, TP.HCM",
            "receiver_name": "A. Kiệt",
            "receiver_title": "Thủ kho",
        },
        "agent": {
            "agent_id": "AGT-PHUANH",
            "legal_name": "CÔNG TY TNHH PHÚ ANH LOGISTICS",
            "legal_address": "Số 45, Nguyễn Văn Linh, Quận 7, TP. Hồ Chí Minh",
            "phone": "0281234567",
            "legal_representative": "Phạm Văn Thành",
            "representative_title": "Giám đốc",
        },
        "company": {
            "company_id": "TT-COMPANY",
            "legal_name": "CÔNG TY TNHH GIAO NHẬN QUỐC TẾ TRƯỜNG THÀNH",
            "legal_address": "Số 12, Đường Láng Hạ, Ba Đình, Hà Nội",
            "phone": "0241234567",
        },
        "contract": {"id": CONTRACT_ID, "goods_lines": GOODS_LINES},
    }


# ---------------------------------------------------------------------------
# 1. Source receipt + sanitized asset pinning + safety
# ---------------------------------------------------------------------------

def test_source_receipt_and_sanitized_asset_are_pinned_and_safe():
    spec = load_template_spec()
    receipt = spec["source_receipt"]
    assert receipt["sha256"] == SOURCE_SHA256
    assert receipt["size_bytes"] == SOURCE_SIZE_BYTES
    assert receipt["git_blob_sha"] == SOURCE_GIT_BLOB_SHA
    assert receipt["transport_branch"] == "owner-input/form07-source-168"
    assert receipt["status"] == "VERIFIED_EXACT_BINARY"

    asset_path = REPO_ROOT / spec["asset_path"]
    raw = asset_path.read_bytes()
    assert hashlib.sha256(raw).hexdigest() == spec["asset_sha256"]
    assert len(raw) == spec["asset_size_bytes"]
    assert template_asset_available() is True

    scan = scan_package_safety(raw)
    assert scan["ok"], scan["issues"]

    archive = zipfile.ZipFile(io.BytesIO(raw))
    names = archive.namelist()
    for banned in ("word/vbaProject.bin", "word/comments.xml", "word/people.xml",
                   "word/_rels/document.xml.rels.external"):
        assert banned not in names
    assert not any(name.startswith("word/media/") for name in names)
    assert not any("activeX" in name or "embeddings" in name for name in names)

    document_xml = archive.read("word/document.xml").decode("utf-8")
    core_xml = archive.read("docProps/core.xml").decode("utf-8")
    app_xml = archive.read("docProps/app.xml").decode("utf-8")
    # Personal metadata from the Owner original is scrubbed.
    assert "Tien Thanh" not in core_xml and "Tien Thanh" not in app_xml
    assert "Version 5.1 build 2600" not in app_xml
    # No tracked changes, no external relationship, no field/instruction markup.
    for banned_markup in ("<w:ins ", "<w:del ", "instrText", "TargetMode=\"External\""):
        assert banned_markup not in document_xml
    rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "TargetMode=\"External\"" not in rels and "http" not in rels.replace(
        "http://schemas.openxmlformats.org", ""
    ).replace("http://schemas.microsoft.com", "")
    # Bind slots exist as single-run tokens; Owner dot-leaders are gone from them.
    for token in ("{{ cong_ty }}", "{{ so_bien_ban }}", "{{ ben_a_name }}", "{{ ben_b_name }}",
                  "{{ ngay_dat_hang }}", "{%tr for line in goods_lines %}", "{%tr endfor %}"):
        assert token in document_xml, token
    # Static business wording is preserved (not rewritten). Some phrases are
    # split across runs in the Owner original, so compare the joined text.
    import re as _re

    display_text = "".join(_re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>", document_xml, flags=_re.S))
    joined = display_text.replace("\xa0", " ")
    for static in ("BIÊN BẢN GIAO NHẬN", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT",
                   "Độc lập – Tự do – Hạnh phúc", "BÊN A (Bên nhận hàng):",
                   "BÊN B (Bên giao hàng):", "ĐẠI DIỆN BÊN A", "ĐẠI DIỆN BÊN B",
                   "được lập thành 02 bản"):
        assert static in joined, static


def test_template_availability_fails_closed_on_asset_drift(tmp_path, monkeypatch):
    spec = load_template_spec()
    (tmp_path / "assets/templates/form07").mkdir(parents=True)
    (tmp_path / spec["asset_path"]).write_bytes(b"PK\x03\x04tampered")
    (tmp_path / "assets/templates/form07/FORM07.template.json").write_text(
        (REPO_ROOT / "assets/templates/form07/FORM07.template.json").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    monkeypatch.setattr(form07_mod, "BASE_DIR", tmp_path)
    assert template_asset_available() is False
    with pytest.raises(form07_mod.Form07TemplateError):
        render_form07_docx(build_form07_context(
            delivery=_full_facts()["delivery"], goods_lines=GOODS_LINES,
            agent=_full_facts()["agent"], company=_full_facts()["company"], manual=MANUAL,
        ))


# ---------------------------------------------------------------------------
# 2. Catalog integration (server-owned) + no hard-coded checkbox
# ---------------------------------------------------------------------------

def test_catalog_carries_exactly_one_new_owner_sample():
    rows = {row.id: row for row in SHIPMENT_DOCUMENT_CATALOG}
    assert "FORM_07" in rows
    entry = rows["FORM_07"]
    assert entry.adapter == "FORM07_DOCGEN"
    assert entry.document_type == "FORM_07"
    assert entry.active is True
    # The other five V1 entries are unchanged.
    assert [row.id for row in SHIPMENT_DOCUMENT_CATALOG if row.id != "FORM_07"] == [
        "HDNT", "PL", "HDMB", "PXK", "GNT"
    ]


def test_catalog_availability_is_server_gated():
    unavailable = project_catalog(
        contract_templates_available=False, pxk_template_available=False,
        storage_configured=True, form07_template_available=False,
    )
    by_id = {row["id"]: row for row in unavailable}
    assert by_id["FORM_07"]["available"] is False
    assert "chưa sẵn sàng" in by_id["FORM_07"]["availabilityReason"]

    no_storage = project_catalog(
        contract_templates_available=False, pxk_template_available=False,
        storage_configured=False, form07_template_available=True,
    )
    by_id = {row["id"]: row for row in no_storage}
    assert by_id["FORM_07"]["available"] is False
    assert "Kho chứng từ" in by_id["FORM_07"]["availabilityReason"]

    ready = project_catalog(
        contract_templates_available=False, pxk_template_available=False,
        storage_configured=True, form07_template_available=True,
    )
    by_id = {row["id"]: row for row in ready}
    assert by_id["FORM_07"]["available"] is True
    assert by_id["FORM_07"]["availabilityReason"] == ""


def test_no_hardcoded_form07_checkbox_in_shipment_detail():
    detail_src = (REPO_ROOT / "component/src/pages/ShipmentDetail.tsx").read_text(encoding="utf-8")
    picker_src = (REPO_ROOT / "component/src/components/ShipmentDocumentPicker.tsx").read_text(encoding="utf-8")
    # The ONLY selection surface is the catalog-driven dialog in the picker.
    assert "selection.catalog" in picker_src
    # ShipmentDetail never hard-codes a FORM_07 selection checkbox or option.
    assert 'value="FORM_07"' not in detail_src
    assert "document-template-checkbox" not in detail_src
    assert 'type="checkbox"' not in detail_src


# ---------------------------------------------------------------------------
# 3. Pure canonical readiness — every missing fact fails closed
# ---------------------------------------------------------------------------

def test_readiness_fail_closed_matrix():
    facts = _full_facts()

    ready = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert ready.state == "READY", ready.missing

    undelivered = evaluate_form07_readiness(
        delivery={**facts["delivery"], "is_delivered": False}, goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert undelivered.state == "BLOCKED"
    assert "xác nhận giao" in undelivered.reason

    missing_receiver = evaluate_form07_readiness(
        delivery={**facts["delivery"], "receiver_title": ""}, goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert missing_receiver.state == "MISSING_INFO"
    assert any("receiver_title" in item for item in missing_receiver.missing)

    missing_date = evaluate_form07_readiness(
        delivery={**facts["delivery"], "delivered_at": ""}, goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert missing_date.state == "MISSING_INFO"
    assert any("delivered_at" in item for item in missing_date.missing)

    no_contract = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"], contract=None, manual=MANUAL,
    )
    assert no_contract.state == "BLOCKED"

    ambiguous = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=None, ambiguous_source=True, manual=MANUAL,
    )
    assert ambiguous.state == "BLOCKED"
    assert "đúng một" in ambiguous.reason

    no_goods = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=[],
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert no_goods.state == "MISSING_INFO"
    assert any("Chưa có dòng hàng hóa" in item for item in no_goods.missing)

    unknown_qty = evaluate_form07_readiness(
        delivery=facts["delivery"],
        goods_lines=[{**GOODS_LINES[0], "quantity": ""}],
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert unknown_qty.state == "MISSING_INFO"
    assert any("UNKNOWN" in item for item in unknown_qty.missing)

    missing_party_fact = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent={**facts["agent"], "phone": ""}, company=facts["company"],
        contract=facts["contract"], manual=MANUAL,
    )
    assert missing_party_fact.state == "MISSING_INFO"
    assert any("điện thoại của Đại lý" in item for item in missing_party_fact.missing)

    missing_manual = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual={},
    )
    assert missing_manual.state == "MISSING_INFO"
    assert "Thiếu số biên bản giao nhận (FORM_07)." in missing_manual.missing
    assert "Thiếu ngày đơn đặt hàng (FORM_07)." in missing_manual.missing

    bad_manual_date = evaluate_form07_readiness(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"],
        contract=facts["contract"], manual={"form07_no": "X", "order_date": "không phải ngày"},
    )
    assert bad_manual_date.state == "MISSING_INFO"
    assert any("không hợp lệ" in item for item in bad_manual_date.missing)


# ---------------------------------------------------------------------------
# 4. Deterministic, structurally valid DOCX with repeating goods rows
# ---------------------------------------------------------------------------

def test_render_is_real_deterministic_docx_with_repeating_goods():
    facts = _full_facts()
    context = build_form07_context(
        delivery=facts["delivery"], goods_lines=GOODS_LINES,
        agent=facts["agent"], company=facts["company"], manual=MANUAL,
    )
    first = render_form07_docx(context)
    second = render_form07_docx(context)
    assert first[:2] == b"PK"
    assert hashlib.sha256(first).hexdigest() == hashlib.sha256(second).hexdigest()
    assert scan_package_safety(first)["ok"]

    import docx

    document = docx.Document(io.BytesIO(first))
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    header_text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "BBGN-2026-0295" in header_text
    assert facts["company"]["legal_name"] in header_text
    assert "Kho Thiên Ân" in header_text
    assert "ngày 09 tháng 09 năm 2026" in full_text
    assert "25/09/2026" in full_text
    assert facts["agent"]["legal_name"] in full_text
    assert "A. Kiệt" in full_text and "Thủ kho" in full_text
    assert "Phạm Văn Thành" in full_text and "Giám đốc" in full_text

    goods_table = document.tables[1]
    assert len(goods_table.rows) == 1 + len(GOODS_LINES)
    first_row = [cell.text for cell in goods_table.rows[1].cells]
    assert first_row[0] == "1"
    assert first_row[1] == "Kính nổi 5mm"
    assert first_row[2] == "1200x600"
    assert first_row[3] == "Tấm"
    assert first_row[4] == "100"
    second_row = [cell.text for cell in goods_table.rows[2].cells]
    assert second_row[4] == "20.5"

    document_xml = zipfile.ZipFile(io.BytesIO(first)).read("word/document.xml").decode("utf-8")
    assert "{{" not in document_xml and "{%" not in document_xml
    assert "……" not in document_xml
    # Static wording survives generation (run-joined, as Word displays it).
    import re as _re

    display_text = "".join(_re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>", document_xml, flags=_re.S))
    for static in ("BÊN A (Bên nhận hàng):", "BÊN B (Bên giao hàng):",
                   "Hai bên cùng nhau thống nhất số lượng giao hàng"):
        assert static.replace(" ", "") in display_text.replace(" ", "").replace("\xa0", ""), static


# ---------------------------------------------------------------------------
# 5. Staging dispatcher: selection, RBAC, fail-closed, end-to-end generation
# ---------------------------------------------------------------------------

def _available_policy(*_args, **_kwargs):
    return {
        "catalog": project_catalog(
            contract_templates_available=True,
            pxk_template_available=True,
            storage_configured=True,
            form07_template_available=True,
        ),
        "canSaveSelection": True,
        "canGenerate": True,
        "storageConfigured": True,
    }


def test_selection_persistence_and_deselect_keeps_form07_history(staging, monkeypatch):
    monkeypatch.setattr(vm, "_shipment_document_picker_policy", _available_policy)
    shipment = vm.read_shipment(SHIPMENT_ID)
    assert shipment is not None
    version = int(shipment.get("version") or 0)

    db = database.get_db()
    history = {
        "_id": f"FORM_07::{SHIPMENT_ID}::R1",
        "id": f"FORM_07::{SHIPMENT_ID}::R1",
        "document_id": f"FORM_07::{SHIPMENT_ID}",
        "shipment_id": SHIPMENT_ID,
        "document_type": "FORM_07",
        "generation_revision": 1,
        "status": "GENERATED",
        "output_file_name": f"FORM07_{SHIPMENT_ID}_BBGN-1.docx",
        "output_hash": "fixture-output-hash",
    }
    db["docgen_documents"].insert_one(dict(history))

    saved = _dispatch("shipment_document_selection_save", {
        "shipmentId": SHIPMENT_ID,
        "requestId": "F07-SAVE-1",
        "expectedVersion": version,
        "selectedTemplateIds": ["FORM_07", "PXK"],
    })
    assert saved["ok"], saved
    assert saved["readback"]["selectedDocumentTemplateIds"] == ["PXK", "FORM_07"]

    fresh = vm.read_shipment(SHIPMENT_ID)
    deselect = _dispatch("shipment_document_selection_save", {
        "shipmentId": SHIPMENT_ID,
        "requestId": "F07-SAVE-2",
        "expectedVersion": int(fresh.get("version") or 0),
        "selectedTemplateIds": [],
    }, seq=2)
    assert deselect["ok"], deselect
    latest = vm.read_shipment(SHIPMENT_ID)
    assert latest["document_template_ids"] == []
    # Deselection hides the card but never deletes generation history.
    row = db["docgen_documents"].find_one({"id": history["id"]}, {"_id": 0})
    assert row is not None and row["output_hash"] == "fixture-output-hash"


def test_viewer_cannot_generate_form07(staging, monkeypatch):
    import streamlit as st

    _prepare_source_contract()
    st.session_state["current_user"] = dict(VIEWER_ACTOR)
    before_docs = database.get_db()["docgen_documents"].count_documents({})
    before_attachments = database.get_db()["shipment_attachments"].count_documents({})

    result = _dispatch("form07_generate", {
        "shipmentId": SHIPMENT_ID, "form07No": MANUAL["form07_no"], "orderDate": MANUAL["order_date"],
    })
    assert result["ok"] is False
    assert result["gate"] == "BLOCKED_PERMISSION_DENIED"
    assert database.get_db()["docgen_documents"].count_documents({}) == before_docs
    assert database.get_db()["shipment_attachments"].count_documents({}) == before_attachments


def test_generate_fails_closed_without_delivery_and_manual_identity(staging):
    _prepare_source_contract()

    undelivered = _dispatch("form07_generate", {
        "shipmentId": UNDELIVERED_ID,
        "form07No": MANUAL["form07_no"],
        "orderDate": MANUAL["order_date"],
    })
    assert undelivered["ok"] is False
    assert undelivered["gate"] == "BLOCKED_NOT_DELIVERED"

    missing_manual = _dispatch("form07_generate", {
        "shipmentId": SHIPMENT_ID, "form07No": "", "orderDate": "",
    }, seq=2)
    assert missing_manual["ok"] is False
    assert missing_manual["gate"] == "MISSING_INFO"
    assert any("số biên bản" in item for item in missing_manual["missing"])

    db = database.get_db()
    assert db["docgen_documents"].count_documents({"document_type": "FORM_07"}) == 0
    assert db["shipment_attachments"].count_documents({}) == 0


def test_generate_persist_history_and_remount_download_end_to_end(staging):
    _prepare_source_contract()

    result = _dispatch("form07_generate", {
        "shipmentId": SHIPMENT_ID,
        "form07No": MANUAL["form07_no"],
        "orderDate": MANUAL["order_date"],
    })
    assert result["ok"], result
    persisted = result["persisted"]
    assert persisted["documentId"] == f"FORM_07::{SHIPMENT_ID}"
    assert persisted["recordId"] == f"FORM_07::{SHIPMENT_ID}::R1"
    assert persisted["revision"] == 1
    assert persisted["templateId"] == "FORM_07"
    assert persisted["templateVersion"] == 1
    spec = load_template_spec()
    assert persisted["templateSha256"] == spec["asset_sha256"]
    assert persisted["goodsLineCount"] == 2
    assert persisted["isOfficeFile"] is True
    attachment_id = persisted["attachmentId"]
    assert attachment_id

    # Durable attachment: the EXACT generated bytes are readable back through
    # the existing attachment service (same path the UI download event uses).
    outcome = fetch_attachment_bytes({"id": attachment_id}, shipment_id=SHIPMENT_ID)
    assert outcome["ok"], outcome
    stored_bytes = bytes(outcome["data"])
    assert hashlib.sha256(stored_bytes).hexdigest() == persisted["outputHash"]
    assert stored_bytes[:2] == b"PK"

    # Immutable metadata with bounded provenance.
    db = database.get_db()
    row = db["docgen_documents"].find_one({"id": persisted["recordId"]}, {"_id": 0})
    assert row is not None
    assert row["document_type"] == FORM07_DOCUMENT_TYPE
    assert row["template_id"] == "FORM_07"
    assert row["template_hash"] == spec["asset_sha256"]
    assert row["output_mime"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    snapshot = row["source_snapshot"]
    assert snapshot["form07_no"] == MANUAL["form07_no"]
    assert snapshot["order_date"] == "25/09/2026"
    assert snapshot["contract_id"] == CONTRACT_ID
    assert snapshot["delivered_at"] == "2026-09-09"
    assert snapshot["receiver_name"] == "A. Kiệt"
    assert len(snapshot["goods_lines"]) == 2
    r1_hash = row["output_hash"]

    # Regeneration creates revision R2; R1 stays immutable.
    second = _dispatch("form07_generate", {
        "shipmentId": SHIPMENT_ID,
        "form07No": "BBGN-2026-0295B",
        "orderDate": "2026-09-26",
    }, seq=2)
    assert second["ok"], second
    assert second["persisted"]["revision"] == 2
    r1_after = db["docgen_documents"].find_one({"id": persisted["recordId"]}, {"_id": 0})
    assert r1_after["output_hash"] == r1_hash
    assert db["docgen_documents"].count_documents({"document_type": "FORM_07"}) == 2

    # Read projection: done + history + server readiness READY (identity from
    # the latest durable snapshot), attachment linked by exact output hash.
    fresh = vm.read_shipment(SHIPMENT_ID)
    docs_by_shipment = rd.docgen_documents_by_shipment(
        {
            text_id: dict(row)
            for row in db["docgen_documents"].find({"shipment_id": SHIPMENT_ID}, {"_id": 0})
            for text_id in [rd.text(row.get("id"))]
        }
    )
    projection = rd.shipment_projection(
        fresh,
        contracts={CONTRACT_ID: dict(db["contracts"].find_one({"id": CONTRACT_ID}, {"_id": 0}))},
        docgen_documents=docs_by_shipment,
        attachments={SHIPMENT_ID: [
            dict(r) for r in db["shipment_attachments"].find({"shipment_id": SHIPMENT_ID}, {"_id": 0})
        ]},
    )
    form07_view = projection["form07"]
    assert form07_view["done"] is True
    assert form07_view["number"] == "BBGN-2026-0295B"
    assert form07_view["orderDate"] == "26/09/2026"
    assert form07_view["readiness"]["state"] == "READY"
    assert len(form07_view["documentHistory"]) == 2
    by_revision = {row["revision"]: row for row in form07_view["documentHistory"]}
    assert by_revision[1]["attachmentId"] == attachment_id
    assert by_revision[1]["downloadAvailable"] is True
    assert by_revision[2]["attachmentId"] == second["persisted"]["attachmentId"]
    assert by_revision[2]["mime"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    # Remount/re-entry download: leaving and reopening changes nothing; the
    # EXISTING attachment_download event serves the exact stored bytes and a
    # repeated click does not create a second download row or attachment.
    downloads_before = db["shipment_attachments"].count_documents({})
    download = _dispatch("attachment_download", {
        "shipmentId": SHIPMENT_ID,
        "attachmentId": attachment_id,
        "requestId": "F07-DL-1",
    }, seq=3)
    assert download["ok"], download
    assert download["gate"] == "DOWNLOADED"
    import streamlit as st

    state = st.session_state[vm.ATTACHMENT_DOWNLOAD_KEY]
    assert base64.b64decode(state["base64"]) == stored_bytes
    assert state["seq"] == 1

    repeat = _dispatch("attachment_download", {
        "shipmentId": SHIPMENT_ID,
        "attachmentId": attachment_id,
        "requestId": "F07-DL-2",
    }, seq=4)
    assert repeat["ok"], repeat
    state_after = st.session_state[vm.ATTACHMENT_DOWNLOAD_KEY]
    assert state_after["seq"] == 2
    assert db["shipment_attachments"].count_documents({}) == downloads_before

    # Deselecting FORM_07 after generation hides it from the chosen set but
    # never destroys the durable history.
    deselect = _dispatch("shipment_document_selection_save", {
        "shipmentId": SHIPMENT_ID,
        "requestId": "F07-SAVE-9",
        "expectedVersion": int(vm.read_shipment(SHIPMENT_ID).get("version") or 0),
        "selectedTemplateIds": [],
    }, seq=5)
    assert deselect["ok"], deselect
    assert db["docgen_documents"].count_documents({"document_type": "FORM_07"}) == 2
    still = fetch_attachment_bytes({"id": attachment_id}, shipment_id=SHIPMENT_ID)
    assert still["ok"] and bytes(still["data"]) == stored_bytes
