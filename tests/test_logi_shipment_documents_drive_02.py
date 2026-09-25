"""LOGI_SHIPMENT_DOCUMENTS_DRIVE_REINTEGRATE_02 — shipment attachment slice acceptance.

Ported (and extended) from donor PR #49 onto current main. ONE vertical slice
is proven here against staging/mongomock plus a DETERMINISTIC FAKE DRIVE
ADAPTER — never a real service account, real folder or real Drive:

    shipment attachment upload → safe preprocessing → Drive adapter →
    metadata persistence → shipment document UI readback

Covers:

1. provider contract (folder resolved/created once, upload receives the exact
   post-processing bytes, correct MIME/name, returned backend file identity
   persisted,
   NO binary in Mongo, duplicate SHA refuses the second upload, Drive failure
   leaves no Mongo attachment, metadata failure compensates exactly once,
   unauthorized/oversized requests reach Drive ZERO times);
2. bounded validation (extension + MIME + magic bytes, size ceiling BEFORE any
   Drive call, traversal-safe names);
3. RBAC (upload = existing sales_write boundary, download = existing dashboard
   read boundary; viewer may download but may not upload; an anonymous session
   may do neither);
4. compression honesty (a large image is materially smaller and still opens,
   no upscaling, PDF/DOCX/XLSX stay BYTE-IDENTICAL, optimizer failure never
   blocks an upload);
5. the Drive REST adapter itself (request shaping + error hygiene, via an
   injected stub session);
6. metadata writer guard (binary/base64 payloads are refused by the writer);
7. bounded readback + projection (no Drive id, folder id, SHA or credential in
   the payload the browser receives).

Run:  .venv/bin/python -m pytest tests/test_logi_shipment_documents_drive_02.py -q
"""

from __future__ import annotations

import base64
import io
import json
import logging
import os
import random
import unittest
from typing import Any, Mapping
from unittest.mock import patch

import pytest

from core import database as db
from core import drive_client as dc
from core import shipment_attachments as att
from core.drive_client import DriveError, DriveSettings, GoogleDriveClient
from views import logi_real_data as rd
from views import logi_sandbox
from views import logi_view_model as vm

SHIPMENT_ID = "LO-0306"

# ---------------------------------------------------------------------------
# Deterministic fake Drive adapter (the ONLY Drive surface these tests touch)
# ---------------------------------------------------------------------------


class FakeDrive:
    """Records every Drive interaction; no network, no credential, no Owner data."""

    def __init__(self, *, fail_upload: bool = False, fail_download: bool = False,
                 fail_delete: bool = False) -> None:
        self.settings = DriveSettings(
            root_folder_id="ROOT-FOLDER-FAKE",
            service_account={
                "client_email": "logi-drive@fake.iam.gserviceaccount.com",
                "private_key": "-----BEGIN PRIVATE KEY-----not-a-real-key-----END PRIVATE KEY-----",
                "token_uri": "https://oauth2.googleapis.com/token",
            },
        )
        self.fail_upload = fail_upload
        self.fail_download = fail_download
        self.fail_delete = fail_delete
        #: every ensure_folder call (cache miss or hit) — proves resolution order
        self.resolutions: list[tuple[str, str]] = []
        #: folder creations (a cache hit never lands here)
        self.creations: list[tuple[str, str]] = []
        #: every upload attempt, recorded BEFORE a simulated failure
        self.uploads: list[dict[str, Any]] = []
        self.downloads: list[str] = []
        self.deletes: list[str] = []
        self._folder_cache: dict[tuple[str, str], str] = {}
        self._files: dict[str, bytes] = {}

    # --- folders
    def ensure_folder(self, name: str, parent_id: str) -> str:
        key = (parent_id, name)
        self.resolutions.append(key)
        cached = self._folder_cache.get(key)
        if cached:
            return cached
        folder_id = f"folder-{len(self.creations) + 1}"
        self.creations.append(key)
        self._folder_cache[key] = folder_id
        return folder_id

    # --- files
    def upload_file(self, *, name: str, mime: str, data: bytes, parent_id: str,
                    description: str = "") -> dict[str, Any]:
        self.uploads.append(
            {"name": name, "mime": mime, "data": bytes(data), "parentId": parent_id,
             "description": description}
        )
        if self.fail_upload:
            raise DriveError("Google Drive từ chối yêu cầu (HTTP 403).", status=403)
        file_id = f"drive-file-{len(self.uploads)}"
        self._files[file_id] = bytes(data)
        return {"fileId": file_id, "name": name, "mime": mime, "size": len(data),
                "parents": [parent_id], "webViewLink": ""}

    def download_file(self, file_id: str) -> bytes:
        self.downloads.append(file_id)
        if self.fail_download:
            raise DriveError("Google Drive từ chối yêu cầu (HTTP 404).", status=404)
        return self._files.get(file_id, b"")

    def delete_file(self, file_id: str) -> bool:
        self.deletes.append(file_id)
        if self.fail_delete:
            return False
        self._files.pop(file_id, None)
        return True


class CountingFactory:
    """Counts how often the production client factory was reached."""

    def __init__(self, client: FakeDrive | None) -> None:
        self.client = client
        self.calls = 0

    def __call__(self, *args: Any, **kwargs: Any) -> FakeDrive | None:
        self.calls += 1
        return self.client


# ---------------------------------------------------------------------------
# Deterministic local fixtures (no Owner file, no network)
# ---------------------------------------------------------------------------


def _noise_image_bytes(*, size=(3000, 2000), fmt="JPEG", quality=95) -> bytes:
    """A deterministic photographic-ish image: seeded white noise.

    ``random.Random(42)`` (stable Mersenne Twister) instead of a per-pixel
    Python loop: identical test semantics — incompressible noise that a real
    optimizer must shrink by downscaling + re-encoding — in a fraction of the
    time.
    """
    from PIL import Image

    width, height = size
    raw = random.Random(42).randbytes(width * height * 3)
    image = Image.frombytes("RGB", (width, height), raw)
    buffer = io.BytesIO()
    if fmt.upper() == "PNG":
        image.save(buffer, format="PNG")
    else:
        image.save(buffer, format=fmt, quality=quality)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    """A tiny but structurally valid PDF (kept byte-identical by policy)."""
    body = (
        b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\n"
        b"trailer<</Root 1 0 R>>\n%%EOF\n"
    )
    return body


def _xlsx_bytes() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Tờ khai"
    workbook.active["B2"] = 1234.5
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Hợp đồng nguyên tắc - nội dung thử nghiệm.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload_payload(file_name: str, mime: str, data: bytes, *, shipment_id: str = SHIPMENT_ID,
                    category: str = "customs", request_id: str = "req-1",
                    declared_size: int | None = None) -> dict[str, Any]:
    return {
        "shipmentId": shipment_id,
        "category": category,
        "fileName": file_name,
        "mime": mime,
        "size": len(data) if declared_size is None else declared_size,
        "base64": base64.b64encode(data).decode("ascii"),
        "requestId": request_id,
    }


# ---------------------------------------------------------------------------
# Staging environment (mongomock) — never production, never real Drive
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def pr65_drive_backend(monkeypatch):
    """Keep PR65's focused contract explicit after GridFS became the default."""
    monkeypatch.setenv("ATTACHMENT_STORAGE_BACKEND", "drive_shared")


@pytest.fixture()
def staging():
    import streamlit as st

    logi_sandbox.reset_backend()
    transport = logi_sandbox.ensure_backend()
    assert transport == logi_sandbox.SANDBOX_TRANSPORT, "tests must run on the isolated staging transport"
    st.session_state["current_user"] = {
        "username": "staging.support", "role": "SUPPORT",
        "full_name": "Thu Hằng (Preview/Staging)", "user_id": "STAGING-SUPPORT",
    }
    st.session_state.pop(vm.ATTACHMENT_DOWNLOAD_KEY, None)
    yield transport
    st.session_state.pop(vm.ATTACHMENT_DOWNLOAD_KEY, None)
    logi_sandbox.reset_backend()


@pytest.fixture()
def seq():
    state = {"n": 0}

    def _next() -> int:
        state["n"] += 1
        return state["n"]

    return _next


def _dispatch(name: str, payload: Mapping[str, Any], seq_fn) -> dict[str, Any]:
    trigger = {name: {"seq": seq_fn(), "event": name, "payload": dict(payload)}}
    outcomes = vm.consume_phase_b_triggers({"seq": 0}, trigger)
    assert len(outcomes) == 1
    return outcomes[0]


def _stored(shipment_id: str = SHIPMENT_ID) -> dict[str, dict[str, Any]]:
    return db.load_shipment_attachments(shipment_id)


# ===========================================================================
# 1. Deployment settings — no client, no import, no secret leak
# ===========================================================================


class TestDriveSettings:
    def test_unconfigured_deployment_builds_no_client(self):
        with patch.object(dc, "get_setting", return_value=""), patch.object(dc, "get_raw_setting", return_value=None):
            settings = dc.load_drive_settings()
            assert settings.configured is False
            assert settings.missing == [dc.SERVICE_ACCOUNT_ENV, dc.ROOT_FOLDER_ENV]
            assert dc.build_drive_client() is None

    def test_service_account_accepts_json_string_and_toml_table(self):
        info = {"client_email": "svc@example.iam.gserviceaccount.com",
                "private_key": "-----BEGIN PRIVATE KEY-----\nreal-looking\n-----END PRIVATE KEY-----\n",
                "token_uri": "https://oauth2.googleapis.com/token"}
        from_json = dc._parse_service_account(json.dumps(info))
        from_table = dc._parse_service_account(dict(info))
        assert from_json and from_table
        assert from_json["client_email"] == info["client_email"]
        assert from_json["type"] == "service_account"
        # A malformed payload is reported as unusable instead of half-loaded.
        assert dc._parse_service_account("{not json") is None
        assert dc._parse_service_account({"client_email": "svc@example.com"}) is None

    def test_settings_from_env_are_configured_and_describe_hides_key_material(self):
        info = {"client_email": "svc@example.iam.gserviceaccount.com",
                "private_key": "-----BEGIN PRIVATE KEY-----secret-----END PRIVATE KEY-----",
                "token_uri": "https://oauth2.googleapis.com/token"}
        env = {dc.ROOT_FOLDER_ENV: "ROOT-FOLDER-1", dc.SERVICE_ACCOUNT_ENV: json.dumps(info)}
        with patch.dict(os.environ, env, clear=False):
            settings = dc.load_drive_settings()
        assert settings.configured is True
        assert settings.client_email == info["client_email"]
        described = json.dumps(settings.describe(), ensure_ascii=False)
        assert "secret" not in described and "PRIVATE KEY" not in described
        client = dc.build_drive_client(settings)
        assert isinstance(client, GoogleDriveClient)

    def test_name_sanitizer_refuses_traversal_and_keeps_vietnamese_names(self):
        assert dc.is_unsafe_drive_name("../../etc/passwd") is True
        assert dc.is_unsafe_drive_name("a/b.pdf") is True
        assert dc.is_unsafe_drive_name("a\\b.pdf") is True
        assert dc.is_unsafe_drive_name("..") is True
        assert dc.is_unsafe_drive_name("") is True
        assert dc.is_unsafe_drive_name("Tờ khai hải quan.pdf") is False
        assert dc.sanitize_drive_name("Tờ khai  hải quan.pdf") == "Tờ khai hải quan.pdf"


# ===========================================================================
# 2. The Drive REST adapter itself (stub HTTP session, zero network)
# ===========================================================================


class _StubResponse:
    def __init__(self, payload=None, content: bytes = b"", status: int = 200) -> None:
        self._payload = payload or {}
        self.content = content
        self.status_code = status

    def json(self):
        return self._payload


class _StubSession:
    def __init__(self, responses: list[_StubResponse] | None = None, status: int = 200) -> None:
        self.calls: list[dict[str, Any]] = []
        self._responses = list(responses or [])
        self._status = status

    def request(self, method, url, *, params=None, data=None, headers=None,
                json=None, timeout=None, **kwargs):
        # Production-shaped guard: the real transport is
        # google.auth.transport.requests.AuthorizedSession, a requests.Session
        # subclass, so ONLY real requests keywords are accepted here. An
        # invented kwarg (e.g. the old `json_body`) fails exactly as it would
        # in production instead of being silently recorded.
        if kwargs:
            bad = next(iter(kwargs))
            raise TypeError(
                f"Session.request() got an unexpected keyword argument '{bad}'"
            )
        self.calls.append({"method": method, "url": url, "params": params,
                           "data": data, "headers": headers, "json": json,
                           "timeout": timeout})
        if self._responses:
            return self._responses.pop(0)
        if self._status >= 400:
            return _StubResponse({}, status=self._status)
        return _StubResponse({"id": "FILE-1", "name": "x", "mimeType": "mime", "size": "3"})


def _stub_settings() -> DriveSettings:
    return DriveSettings(root_folder_id="ROOT-FOLDER", service_account={
        "client_email": "svc@example.iam.gserviceaccount.com",
        "private_key": "-----BEGIN PRIVATE KEY-----never-logged-----END PRIVATE KEY-----",
        "token_uri": "https://oauth2.googleapis.com/token",
    })


class TestDriveRestAdapter:
    def test_folder_is_found_then_created_once_and_memoised(self):
        session = _StubSession([
            _StubResponse({"files": []}),          # find → miss
            _StubResponse({"id": "FOLDER-9"}),     # create
        ])
        client = GoogleDriveClient(_stub_settings(), session=session)
        assert client.ensure_folder("shipments", "ROOT-FOLDER") == "FOLDER-9"
        assert client.ensure_folder("shipments", "ROOT-FOLDER") == "FOLDER-9"  # memoised
        assert len(session.calls) == 2, "a resolved folder must not be queried again"
        search = session.calls[0]
        assert search["method"] == "GET" and search["url"].endswith("/drive/v3/files")
        assert "trashed = false" in search["params"]["q"]
        assert "'ROOT-FOLDER' in parents" in search["params"]["q"]
        assert search["params"]["supportsAllDrives"] == "true"
        assert session.calls[1]["json"]["parents"] == ["ROOT-FOLDER"]

    def test_create_folder_executes_through_a_requests_compatible_signature(self):
        # CONTROL_TOWER_REPAIR regression: create_folder once sent
        # `json_body=`, which a real requests.Session.request rejects with
        # TypeError. The hardened stub above enforces the same boundary, so
        # reaching this assertion proves the production-shaped call works.
        session = _StubSession([_StubResponse({"id": "FOLDER-1", "name": "lo-hang"})])
        client = GoogleDriveClient(_stub_settings(), session=session)
        assert client.create_folder("lo-hang", "ROOT-FOLDER") == "FOLDER-1"
        (call,) = session.calls
        assert call["method"] == "POST" and call["url"].endswith("/drive/v3/files")
        assert call["json"] == {
            "name": "lo-hang",
            "mimeType": dc.DRIVE_FOLDER_MIME,
            "parents": ["ROOT-FOLDER"],
        }
        assert call["params"]["supportsAllDrives"] == "true"

    def test_upload_sends_exact_bytes_as_multipart_related(self):
        payload = b"\x89PNG\r\n\x1a\nEXACT-BYTES-\xa9"
        session = _StubSession([_StubResponse({"id": "FILE-7", "name": "ảnh.png",
                                               "mimeType": "image/png", "size": str(len(payload))})])
        client = GoogleDriveClient(_stub_settings(), session=session)
        result = client.upload_file(name="ảnh.png", mime="image/png", data=payload, parent_id="FOLDER-1")
        call = session.calls[0]
        assert call["method"] == "POST"
        assert call["url"].endswith("/upload/drive/v3/files")
        assert call["params"]["uploadType"] == "multipart"
        assert call["params"]["supportsAllDrives"] == "true"
        media_part = call["data"].split(b"Content-Type: image/png\r\n\r\n", 1)[1].split(b"\r\n--", 1)[0]
        assert media_part == payload, "the media part must carry the exact bytes, byte for byte"
        assert call["data"].count(payload) == 1
        assert 'Content-Type: image/png' in call["data"].decode("utf-8", "ignore")
        assert '"name": "ảnh.png"' in call["data"].decode("utf-8", "ignore")
        assert "never-logged" not in json.dumps(call.get("headers") or {})
        assert result["fileId"] == "FILE-7"

    def test_download_uses_alt_media_and_delete_is_bounded(self):
        session = _StubSession([_StubResponse({}, content=b"BYTES"), _StubResponse({}, status=204)])
        client = GoogleDriveClient(_stub_settings(), session=session)
        assert client.download_file("FILE-1") == b"BYTES"
        assert session.calls[0]["params"]["alt"] == "media"
        assert client.delete_file("FILE-1") is True
        assert session.calls[1]["method"] == "DELETE"

    def test_http_failure_becomes_a_credential_free_drive_error(self):
        session = _StubSession(status=403)
        client = GoogleDriveClient(_stub_settings(), session=session)
        with pytest.raises(DriveError) as failure:
            client.upload_file(name="a.pdf", mime="application/pdf", data=b"%PDF-1.4", parent_id="F1")
        message = str(failure.value)
        assert "403" in message
        assert "never-logged" not in message and "PRIVATE KEY" not in message
        assert failure.value.status == 403
        # The compensating delete is fail-closed: an HTTP error returns False.
        assert client.delete_file("FILE-1") is False


# ===========================================================================
# 3. Bounded validation — refusals happen BEFORE any Drive work
# ===========================================================================


class TestUploadValidation:
    def test_allowed_and_rejected_types(self):
        assert att.ALLOWED_EXTENSIONS == ("doc", "docx", "jpeg", "jpg", "pdf", "png", "xls", "xlsx")
        for banned in ("exe", "js", "html", "svg", "bat", "sh", "zip"):
            with pytest.raises(att.UploadRefused) as refusal:
                att.validate_upload(f"tep.{banned}", "", None, b"whatever")
            assert refusal.value.code == att.CODE_TYPE

    def test_content_mismatch_is_refused(self):
        with pytest.raises(att.UploadRefused) as refusal:
            att.validate_upload("bang.pdf", "application/pdf", None, _noise_image_bytes(size=(40, 40)))
        assert refusal.value.code == att.CODE_MISMATCH
        with pytest.raises(att.UploadRefused) as refusal:
            att.validate_upload("anh.jpg", "image/jpeg", None, b"%PDF-1.4")
        assert refusal.value.code == att.CODE_MISMATCH

    def test_mime_mismatch_and_empty_file_are_refused(self):
        with pytest.raises(att.UploadRefused) as refusal:
            att.validate_upload("anh.png", "application/x-msdownload", None, _noise_image_bytes(size=(24, 24)))
        assert refusal.value.code == att.CODE_MISMATCH
        with pytest.raises(att.UploadRefused) as refusal:
            att.validate_upload("trong.pdf", "application/pdf", None, b"")
        assert refusal.value.code == att.CODE_INVALID

    def test_size_ceiling_and_declared_size_mismatch(self):
        oversize = att.MAX_UPLOAD_BYTES + 1
        with pytest.raises(att.UploadRefused) as refusal:
            att.validate_upload("to.pdf", "application/pdf", oversize, b"%PDF-1.4" + b"x" * 10)
        assert refusal.value.code == att.CODE_SIZE
        assert "20.0 MB" in refusal.value.message
        with pytest.raises(att.UploadRefused) as refusal:
            att.validate_upload("le.pdf", "application/pdf", 1, b"%PDF-1.4")
        assert refusal.value.code == att.CODE_INVALID

    def test_traversal_and_control_characters_are_refused(self):
        for name in ("../../etc/passwd.pdf", "a/b.pdf", "a\\b.pdf", "tep\x00.pdf"):
            with pytest.raises(att.UploadRefused) as refusal:
                att.validate_upload(name, "application/pdf", None, _pdf_bytes())
            assert refusal.value.code == att.CODE_INVALID

    def test_base64_body_is_bounded_before_decoding(self):
        too_long = "A" * (att.MAX_BASE64_CHARS + 8)
        with pytest.raises(att.UploadRefused) as refusal:
            att.decode_upload_payload({"base64": too_long})
        assert refusal.value.code == att.CODE_SIZE
        with pytest.raises(att.UploadRefused) as refusal:
            att.decode_upload_payload({"base64": "!!!not-base64!!!"})
        assert refusal.value.code == att.CODE_INVALID
        data, size = att.decode_upload_payload({"base64": base64.b64encode(b"%PDF-1.4").decode()})
        assert data == b"%PDF-1.4" and size == 8


# ===========================================================================
# 4. Provider contract — folder, bytes, persistence, dedupe, failures
# ===========================================================================


class TestProviderContract:
    def test_upload_creates_logistics_customer_period_shipment_hierarchy_and_persists_metadata(self, staging, seq):
        fake = FakeDrive()
        factory = CountingFactory(fake)
        original = _noise_image_bytes(size=(2600, 1800))
        with patch("core.shipment_attachments.build_drive_client", factory):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("Tờ khai hải quan.jpg", "image/jpeg", original,
                                                category="customs"),
                                seq)

        assert outcome["ok"] is True, outcome["error"]
        assert outcome["gate"] == "UPLOADED"
        assert factory.calls == 1

        # Exact Drive hierarchy comes from authoritative shipment context:
        # Logistics / Customer / YYYY / MM / Shipment. No redundant attachments level.
        shipment = vm.read_shipment(SHIPMENT_ID)
        assert shipment is not None
        logistics_name, customer_name, year_name, month_name, shipment_name = att.drive_folder_hierarchy(
            SHIPMENT_ID, shipment
        )
        expected = []
        parent = "ROOT-FOLDER-FAKE"
        for name in (logistics_name, customer_name, year_name, month_name, shipment_name):
            expected.append((parent, name))
            parent = fake._folder_cache[(parent, name)]
        assert fake.creations == expected
        assert len(fake.resolutions) == 5
        assert all(name not in {"shipments", "attachments"} for _parent, name in fake.creations)

        # The upload received the POST-PROCESSING bytes, not the original.
        assert len(fake.uploads) == 1
        uploaded = fake.uploads[0]
        assert uploaded["name"] == "Tờ khai hải quan.jpg"
        assert uploaded["mime"] == "image/jpeg"
        assert len(uploaded["data"]) < len(original)
        assert uploaded["data"] != original
        assert uploaded["parentId"] == fake._folder_cache[
            (fake._folder_cache[
                (fake._folder_cache[
                    (fake._folder_cache[
                        ("ROOT-FOLDER-FAKE", logistics_name)
                    ], customer_name)
                ], year_name)
            ], month_name),
            shipment_name,
        ]

        # Metadata: new rows use the backend-neutral identity.
        rows = _stored()
        assert len(rows) == 1
        record = next(iter(rows.values()))
        assert record["storage_backend"] == "drive_shared"
        assert record["storage_file_id"] == "drive-file-1"
        assert "drive_folder_id" not in record
        assert record["shipment_id"] == SHIPMENT_ID
        assert record["status"] == att.STATUS_AVAILABLE
        assert record["category"] == "customs"
        assert record["original_size"] == len(original)
        assert record["stored_size"] == len(uploaded["data"])
        # SHA-256 is the ORIGINAL upload, even though stored bytes differ.
        assert record["sha256"] == att.sha256_hex(original)
        assert record["uploaded_by"] == "staging.support"

        # READBACK is authoritative: the outcome carries the fresh projection.
        names = [row["name"] for row in outcome["readback"]["attachments"]]
        assert names == ["Tờ khai hải quan.jpg"]
        projected = outcome["readback"]["attachments"][0]
        assert projected["compressionMode"] == att.COMPRESSION_OPTIMIZED
        assert projected["compressionLabel"] == "Ảnh đã tối ưu"
        assert projected["originalSize"] == len(original)
        assert projected["storedSize"] == len(uploaded["data"])
        assert " MB" in projected["sizeLabel"] or " KB" in projected["sizeLabel"]

    def test_upload_legacy_numeric_shipment_id_resolves_and_persists_metadata(self, staging, seq):
        """LOGI_ATTACHMENT_LEGACY_NUMERIC_SHIPMENT_ID_01 (Issue #117).

        Seeds numeric Mongo ID 108313713320, uploads via string UI ID '108313713320',
        proving read_shipment resolves the legacy document without SHIPMENT_NOT_FOUND,
        storage is reached exactly once, and metadata + readback are keyed to '108313713320'.
        """
        legacy_numeric_id = 108313713320
        legacy_string_id = "108313713320"
        database = db.get_db()
        database.shipments.insert_one({
            "id": legacy_numeric_id,
            "customer": "Legacy Customer",
            "is_deleted": False,
            "version": 1,
            "data": {},
        })

        fake = FakeDrive()
        factory = CountingFactory(fake)
        original = _noise_image_bytes(size=(100, 100))
        with patch("core.shipment_attachments.build_drive_client", factory):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh.jpg", "image/jpeg", original,
                                                category="customs",
                                                shipment_id=legacy_string_id),
                                seq)

        assert outcome["gate"] != "SHIPMENT_NOT_FOUND"
        assert outcome["ok"] is True, outcome.get("error")
        assert outcome["gate"] == "UPLOADED"
        assert factory.calls == 1
        assert len(fake.uploads) == 1

        rows = _stored(legacy_string_id)
        assert len(rows) == 1
        record = next(iter(rows.values()))
        assert record["shipment_id"] == legacy_string_id

        # Readback exposes attachment on the legacy shipment
        readback = outcome.get("readback") or {}
        assert readback.get("id") == legacy_string_id
        names = [row["name"] for row in readback.get("attachments", [])]
        assert names == ["anh.jpg"]
        projected = readback.get("attachments", [])[0]
        assert projected["downloadAvailable"] is False

    def test_no_binary_or_base64_is_ever_persisted(self, staging, seq):
        fake = FakeDrive()
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            _dispatch("shipment_attachment_upload",
                      _upload_payload("to-khai.pdf", "application/pdf", _pdf_bytes()), seq)

        raw = db.get_db()["shipment_attachments"].find_one({}, {"_id": 0})
        assert raw is not None
        for key, value in raw.items():
            assert not isinstance(value, (bytes, bytearray, memoryview)), key
            assert str(key).lower() not in {"bytes", "data", "base64", "content", "blob"}
            assert "base64" not in str(value).lower()[:200]
        encoded = json.dumps(raw, ensure_ascii=False, default=str)
        assert len(encoded) < 4096, "attachment metadata must stay bounded"
        assert base64.b64encode(_pdf_bytes()).decode("ascii") not in encoded

    def test_duplicate_sha256_refuses_a_second_drive_upload(self, staging, seq):
        fake = FakeDrive()
        factory = CountingFactory(fake)
        original = _noise_image_bytes(size=(600, 400))
        with patch("core.shipment_attachments.build_drive_client", factory):
            first = _dispatch("shipment_attachment_upload",
                              _upload_payload("anh-1.jpg", "image/jpeg", original, request_id="r1"), seq)
            resolutions_after_first = list(fake.resolutions)
            second = _dispatch("shipment_attachment_upload",
                               _upload_payload("anh-1-copy.jpg", "image/jpeg", original, request_id="r2"), seq)

        assert first["ok"] is True
        assert second["ok"] is True
        assert second["gate"] == "DUPLICATE"
        assert second["duplicate"] is True
        assert len(fake.uploads) == 1, "a duplicate must not create a second Drive copy"
        assert fake.resolutions == resolutions_after_first, "a duplicate must not touch Drive at all"
        assert second["persisted"]["attachmentId"] == first["persisted"]["attachmentId"]
        assert len(_stored()) == 1

    def test_drive_failure_leaves_no_mongo_attachment(self, staging, seq):
        fake = FakeDrive(fail_upload=True)
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh.jpg", "image/jpeg", _noise_image_bytes(size=(300, 300))), seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == att.CODE_DRIVE_FAILED
        assert len(fake.uploads) == 1 and len(fake.deletes) == 0
        assert _stored() == {}
        assert outcome["readback"]["attachments"] == []

    def test_metadata_failure_compensates_with_exactly_one_delete(self, staging, seq):
        fake = FakeDrive()
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)), \
             patch("core.database.save_shipment_attachment", return_value=False), \
             patch(
                 "core.database.probe_shipment_attachment_identity",
                 return_value=(db.ATTACHMENT_PROBE_ABSENT_PROVEN, None),
             ):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh.jpg", "image/jpeg", _noise_image_bytes(size=(300, 300))), seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == att.CODE_METADATA_FAILED
        assert outcome["compensationAttempted"] is True
        assert fake.deletes == ["drive-file-1"], "the just-created file must be deleted exactly once"
        assert _stored() == {}

    def test_metadata_failure_without_cleanup_reports_orphan_risk(self, staging, seq, caplog):
        fake = FakeDrive(fail_delete=True)
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)), \
             patch("core.database.save_shipment_attachment", return_value=False), \
             patch(
                 "core.database.probe_shipment_attachment_identity",
                 return_value=(db.ATTACHMENT_PROBE_ABSENT_PROVEN, None),
             ), \
             caplog.at_level(logging.WARNING, logger="logistics_pro.attachments"):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh.jpg", "image/jpeg", _noise_image_bytes(size=(300, 300))), seq)
        assert outcome["ok"] is False
        assert outcome["orphanRisk"] is True
        assert fake.deletes == ["drive-file-1"]
        assert any("ATTACHMENT_ORPHAN_RISK" in record.message for record in caplog.records)
        assert "PRIVATE KEY" not in caplog.text and "root_folder" not in caplog.text

    def test_unauthorized_role_reaches_drive_zero_times(self, staging, seq):
        import streamlit as st

        fake = FakeDrive()
        factory = CountingFactory(fake)
        st.session_state["current_user"] = {"username": "viewer-1", "role": "viewer",
                                            "full_name": "Chỉ đọc", "user_id": "U-VIEWER"}
        with patch("core.shipment_attachments.build_drive_client", factory):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh.jpg", "image/jpeg", _noise_image_bytes(size=(120, 120))), seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == "BLOCKED_PERMISSION_DENIED"
        assert "sales_write" in outcome["error"]
        assert factory.calls == 0 and fake.uploads == [] and _stored() == {}

    def test_oversized_upload_reaches_drive_zero_times(self, staging, seq):
        fake = FakeDrive()
        factory = CountingFactory(fake)
        big = b"%PDF-1.4" + b"0" * (att.MAX_UPLOAD_BYTES + 64)
        with patch("core.shipment_attachments.build_drive_client", factory):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("to.pdf", "application/pdf", big), seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == att.CODE_SIZE
        assert factory.calls == 0, "the size refusal must happen before any Drive work"
        assert _stored() == {}

    def test_a_lying_declared_size_is_refused_before_drive(self, staging, seq):
        fake = FakeDrive()
        factory = CountingFactory(fake)
        honest = _pdf_bytes()
        # Declares 8 bytes, actually sends ~1 KB of a different document.
        payload = _upload_payload("to-khai.pdf", "application/pdf", honest + b"%" * 1024,
                                  declared_size=8)
        with patch("core.shipment_attachments.build_drive_client", factory):
            outcome = _dispatch("shipment_attachment_upload", payload, seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == att.CODE_INVALID
        assert factory.calls == 0 and _stored() == {}

    def test_unknown_shipment_reaches_drive_zero_times(self, staging, seq):
        fake = FakeDrive()
        factory = CountingFactory(fake)
        with patch("core.shipment_attachments.build_drive_client", factory):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh.jpg", "image/jpeg", _noise_image_bytes(size=(120, 120)),
                                                shipment_id="LO-KHONG-CO"), seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == "SHIPMENT_NOT_FOUND"
        assert factory.calls == 0

    def test_unconfigured_deployment_refuses_honestly_without_creating_bytes(self, staging, seq):
        outcome = _dispatch("shipment_attachment_upload",
                            _upload_payload("anh.jpg", "image/jpeg", _noise_image_bytes(size=(120, 120))), seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == att.CODE_DRIVE_UNCONFIGURED
        assert dc.SERVICE_ACCOUNT_ENV in outcome["error"] and dc.ROOT_FOLDER_ENV in outcome["error"]
        assert _stored() == {}

    def test_readback_survives_a_fresh_view_model_read(self, staging, seq):
        fake = FakeDrive()
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            _dispatch("shipment_attachment_upload",
                      _upload_payload("bang-chung.pdf", "application/pdf", _pdf_bytes()), seq)
        model = vm.build_view_model()
        row = next(row for row in model["shipments"] if row["id"] == SHIPMENT_ID)
        assert [item["name"] for item in row["attachments"]] == ["bang-chung.pdf"]
        # Deleting the metadata (staging only) proves the read is not cached.
        db.get_db()["shipment_attachments"].delete_many({})
        assert rd.attachment_rows(SHIPMENT_ID) == []

    def test_dedupe_scope_is_the_same_shipment_only(self, staging):
        record = {
            "id": "ATT-OTHER-1", "shipment_id": "LO-9999", "original_name": "a.pdf",
            "stored_name": "a.pdf", "mime": "application/pdf", "category": "other",
            "original_size": 8, "stored_size": 8, "sha256": att.sha256_hex(_pdf_bytes()),
            "drive_file_id": "drive-file-other", "drive_folder_id": "folder-x",
            "uploaded_at": att.utc_now_iso(), "uploaded_by": "staging.support",
            "compression_mode": att.COMPRESSION_KEPT, "status": att.STATUS_AVAILABLE,
        }
        assert db.save_shipment_attachment(record) is True
        assert db.find_shipment_attachment_by_sha("LO-9999", record["sha256"]) is not None
        assert db.find_shipment_attachment_by_sha(SHIPMENT_ID, record["sha256"]) is None


# ===========================================================================
# 5. Download — server-authoritative, Drive ids never reach the browser
# ===========================================================================


class TestDownload:
    def _upload_one(self, staging, seq, fake: FakeDrive) -> dict[str, Any]:
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("bang-chung.pdf", "application/pdf", _pdf_bytes()), seq)
        assert outcome["ok"] is True
        return outcome

    def test_download_reads_bytes_from_drive_and_projects_only_metadata(self, staging, seq):
        fake = FakeDrive()
        uploaded = self._upload_one(staging, seq, fake)
        attachment_id = uploaded["persisted"]["attachmentId"]

        fake.downloads.clear()
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            outcome = _dispatch("attachment_download",
                                {"shipmentId": SHIPMENT_ID, "attachmentId": attachment_id,
                                 "requestId": "dl-1"}, seq)

        assert outcome["ok"] is True
        assert fake.downloads == ["drive-file-1"]
        model = vm.build_view_model()
        row = next(row for row in model["shipments"] if row["id"] == SHIPMENT_ID)
        attachment = row["attachments"][0]
        assert attachment["downloadAvailable"] is True
        assert base64.b64decode(attachment["downloadBase64"]) == _pdf_bytes()
        assert attachment["downloadSeq"] == 1
        assert row["attachments"][0]["storedSize"] == len(_pdf_bytes())

        payload = json.dumps(model, ensure_ascii=False, default=str)
        assert "drive-file-1" not in payload
        assert "drive_folder_id" not in payload and "driveFileId" not in payload
        assert "googleapis.com" not in payload
        assert fake.settings.service_account["private_key"] not in payload
        assert "PRIVATE KEY" not in payload
        # The raw SHA-256 stays backend metadata.
        record = next(iter(_stored().values()))
        assert record["sha256"] not in payload

    def test_second_download_bumps_the_seq_so_the_ui_refetches(self, staging, seq):
        fake = FakeDrive()
        uploaded = self._upload_one(staging, seq, fake)
        attachment_id = uploaded["persisted"]["attachmentId"]
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            _dispatch("attachment_download",
                      {"shipmentId": SHIPMENT_ID, "attachmentId": attachment_id, "requestId": "d1"}, seq)
            second = _dispatch("attachment_download",
                               {"shipmentId": SHIPMENT_ID, "attachmentId": attachment_id, "requestId": "d2"}, seq)
        assert second["persisted"]["seq"] == 2
        assert len(fake.downloads) == 2

    def test_read_only_role_may_download_but_may_not_upload(self, staging, seq):
        import streamlit as st

        fake = FakeDrive()
        self._upload_one(staging, seq, fake)
        attachment_id = next(iter(_stored().values()))["id"]

        st.session_state["current_user"] = {"username": "viewer-1", "role": "viewer",
                                            "full_name": "Chỉ đọc", "user_id": "U-VIEWER"}
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            download = _dispatch("attachment_download",
                                 {"shipmentId": SHIPMENT_ID, "attachmentId": attachment_id,
                                  "requestId": "v1"}, seq)
            upload = _dispatch("shipment_attachment_upload",
                               _upload_payload("khac.pdf", "application/pdf", _pdf_bytes() + b"%", request_id="v2"), seq)
        assert download["ok"] is True
        assert upload["ok"] is False and upload["gate"] == "BLOCKED_PERMISSION_DENIED"

    def test_unauthenticated_session_may_neither_download_nor_upload(self, staging, seq):
        import streamlit as st

        fake = FakeDrive()
        self._upload_one(staging, seq, fake)
        attachment_id = next(iter(_stored().values()))["id"]

        st.session_state["current_user"] = {}
        # The setup upload above already used this fake once; the anonymous
        # attempts below must add ZERO further Drive interactions.
        uploads_before = len(fake.uploads)
        downloads_before = len(fake.downloads)
        factory = CountingFactory(fake)
        with patch("core.shipment_attachments.build_drive_client", factory):
            download = _dispatch("attachment_download",
                                 {"shipmentId": SHIPMENT_ID, "attachmentId": attachment_id,
                                  "requestId": "a1"}, seq)
            upload = _dispatch("shipment_attachment_upload",
                               _upload_payload("khac.pdf", "application/pdf", _pdf_bytes(), request_id="a2"), seq)
        assert download["ok"] is False and download["gate"] == "BLOCKED_PERMISSION_DENIED"
        assert upload["ok"] is False and upload["gate"] == "BLOCKED_PERMISSION_DENIED"
        assert factory.calls == 0
        assert len(fake.uploads) == uploads_before and len(fake.downloads) == downloads_before

    def test_missing_and_mismatched_attachments_are_refused(self, staging, seq):
        fake = FakeDrive()
        self._upload_one(staging, seq, fake)
        attachment_id = next(iter(_stored().values()))["id"]
        factory = CountingFactory(fake)
        with patch("core.shipment_attachments.build_drive_client", factory):
            missing = _dispatch("attachment_download",
                                {"shipmentId": SHIPMENT_ID, "attachmentId": "ATT-NOPE", "requestId": "m"}, seq)
            other_shipment = _dispatch("attachment_download",
                                       {"shipmentId": "LO-0307", "attachmentId": attachment_id, "requestId": "o"}, seq)
        assert missing["ok"] is False and missing["gate"] == att.CODE_NOT_FOUND
        assert other_shipment["ok"] is False and other_shipment["gate"] == "SHIPMENT_MISMATCH"
        # Ownership is refused from metadata alone: zero Drive client builds,
        # zero Drive reads for both refusals.
        assert factory.calls == 0 and fake.downloads == []

    def test_drive_download_failure_is_reported_honestly(self, staging, seq):
        fake = FakeDrive(fail_download=True)
        uploaded = self._upload_one(staging, seq, fake)
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            outcome = _dispatch("attachment_download",
                                {"shipmentId": SHIPMENT_ID,
                                 "attachmentId": uploaded["persisted"]["attachmentId"],
                                 "requestId": "df"}, seq)
        assert outcome["ok"] is False
        assert outcome["gate"] == att.CODE_DRIVE_FAILED
        assert "Google Drive lỗi" in outcome["error"]


# ===========================================================================
# 6. Compression honesty
# ===========================================================================


class TestCompression:
    def test_large_image_is_materially_smaller_and_still_opens(self):
        original = _noise_image_bytes(size=(3000, 2000))
        prepared = att.prepare_stored_bytes(original, "jpg", "image/jpeg")
        assert prepared.compression_mode == att.COMPRESSION_OPTIMIZED
        assert len(prepared.data) < len(original) * att.IMAGE_MATERIAL_REDUCTION
        assert len(prepared.data) < len(original) * 0.7, "a document photo must shrink materially"

        from PIL import Image

        with Image.open(io.BytesIO(prepared.data)) as reopened:
            assert reopened.format == "JPEG"
            width, height = reopened.size
            assert max(width, height) <= att.IMAGE_MAX_DIMENSION
            assert width <= 3000 and height <= 2000, "no enlargement is allowed"
            assert width * height > 0
            reopened.load()  # a decodable image, not a corrupt blob
        assert att.sha256_hex(original) != att.sha256_hex(prepared.data)

    def test_png_is_capped_without_upscaling(self):
        from PIL import Image

        original = _noise_image_bytes(size=(2600, 2600), fmt="PNG")
        optimized = att.optimize_image(original, "png")
        assert optimized is not None
        with Image.open(io.BytesIO(optimized)) as reopened:
            assert reopened.format == "PNG"
            assert max(reopened.size) == att.IMAGE_MAX_DIMENSION, "a giant image is capped, never enlarged"

        # The POLICY may still keep the original when the re-encode is not
        # materially smaller — that decision must be honest, never a fake gain.
        prepared = att.prepare_stored_bytes(original, "png", "image/png")
        assert prepared.data in (original, optimized)
        if prepared.compression_mode == att.COMPRESSION_OPTIMIZED:
            assert len(prepared.data) < len(original) * att.IMAGE_MATERIAL_REDUCTION

        for size in ((240, 160), (64, 48)):
            small = _noise_image_bytes(size=size, fmt="PNG")
            result = att.prepare_stored_bytes(small, "png", "image/png")
            if result.compression_mode == att.COMPRESSION_OPTIMIZED:
                assert len(result.data) < len(small) * att.IMAGE_MATERIAL_REDUCTION
                with Image.open(io.BytesIO(result.data)) as reopened:
                    assert reopened.size == size, "a small image is never resized"
            else:
                assert result.data == small, "a kept image keeps its exact original bytes"

    def test_pdf_and_office_inputs_stay_byte_identical(self):
        for name, data, ext, mime in (
            ("bang.pdf", _pdf_bytes(), "pdf", "application/pdf"),
            ("hop-dong.docx", _docx_bytes(), "docx", att.MIME_DOCX),
            ("so-theo-doi.xlsx", _xlsx_bytes(), "xlsx", att.MIME_XLSX),
        ):
            prepared = att.prepare_stored_bytes(data, ext, mime)
            assert prepared.data == data, f"{name} must be kept byte-identical"
            assert prepared.compression_mode == att.COMPRESSION_KEPT
            assert prepared.mime == mime

    def test_exif_orientation_is_applied_without_resizing(self):
        from PIL import Image

        image = Image.new("RGB", (1200, 600), (200, 30, 30))
        exif = Image.Exif()
        exif[274] = 6  # rotate 90° CW
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=92, exif=exif)
        optimized = att.optimize_image(buffer.getvalue(), "jpg")
        assert optimized is not None
        with Image.open(io.BytesIO(optimized)) as reopened:
            assert reopened.size == (600, 1200), "EXIF orientation must be applied"

    def test_optimizer_failure_keeps_the_original_and_never_blocks_upload(self, staging, seq):
        truncated = b"\xff\xd8\xff\xe0" + b"\x00" * 64  # sniffs as JPEG, undecodable
        prepared = att.prepare_stored_bytes(truncated, "jpg", "image/jpeg")
        assert prepared.data == truncated
        assert prepared.compression_mode == att.COMPRESSION_KEPT

        fake = FakeDrive()
        with patch("core.shipment_attachments.build_drive_client", CountingFactory(fake)):
            outcome = _dispatch("shipment_attachment_upload",
                                _upload_payload("anh-hong.jpg", "image/jpeg", truncated), seq)
        assert outcome["ok"] is True
        assert fake.uploads[0]["data"] == truncated
        record = next(iter(_stored().values()))
        assert record["compression_mode"] == att.COMPRESSION_KEPT
        assert record["original_size"] == record["stored_size"] == len(truncated)

    def test_already_optimized_image_reports_kept_original(self):
        small = _noise_image_bytes(size=(320, 240), quality=60)
        prepared = att.prepare_stored_bytes(small, "jpg", "image/jpeg")
        assert prepared.compression_mode == att.COMPRESSION_KEPT
        assert prepared.data == small
        projection = att.project_attachment({
            "id": "ATT-1", "original_name": "a.jpg", "category": "other", "mime": "image/jpeg",
            "original_size": len(small), "stored_size": len(small),
            "compression_mode": att.COMPRESSION_KEPT, "uploaded_at": "2026-09-17T00:00:00+00:00",
            "uploaded_by": "ops", "status": att.STATUS_AVAILABLE,
        })
        assert projection["compressionLabel"] == "Giữ nguyên"
        assert "→" not in projection["sizeLabel"]


# ===========================================================================
# 7. Metadata writer guard — the binary boundary lives in the writer
# ===========================================================================


class TestMetadataWriterGuard:
    def _record(self, **overrides: Any) -> dict[str, Any]:
        record = {
            "id": "ATT-GUARD-1", "shipment_id": SHIPMENT_ID, "original_name": "a.pdf",
            "stored_name": "a.pdf", "mime": "application/pdf", "category": "other",
            "original_size": 8, "stored_size": 8, "sha256": "a" * 64,
            "drive_file_id": "drive-1", "drive_folder_id": "folder-1",
            "uploaded_at": "2026-09-17T00:00:00+00:00", "uploaded_by": "ops",
            "compression_mode": att.COMPRESSION_KEPT, "status": att.STATUS_AVAILABLE,
        }
        record.update(overrides)
        return record

    def test_writer_refuses_binary_and_base64_payloads(self, staging):
        assert db.save_shipment_attachment(self._record(bytes=b"%PDF-1.4")) is False
        assert "metadata" in db.get_last_db_error()
        assert db.save_shipment_attachment(self._record(base64="JVBERi0xLjQ=")) is False
        assert db.save_shipment_attachment(self._record(data=b"x")) is False
        assert db.save_shipment_attachment(self._record(original_name=b"binary-name")) is False
        assert db.load_shipment_attachments() == {}

    def test_writer_requires_identity_and_is_immutable(self, staging):
        assert db.save_shipment_attachment(self._record(drive_file_id="")) is False
        assert db.save_shipment_attachment(self._record()) is True
        assert db.save_shipment_attachment(self._record()) is True  # idempotent same identity
        assert db.save_shipment_attachment(self._record(sha256="b" * 64)) is False
        stored = db.load_shipment_attachment("ATT-GUARD-1")
        assert stored is not None and stored["sha256"] == "a" * 64

    def test_audit_row_is_written_for_an_attachment(self, staging):
        assert db.save_shipment_attachment(self._record()) is True
        rows = db.load_audit_logs("shipment_attachment", "ATT-GUARD-1", limit=5)
        assert rows and rows[0]["action"] == "create"
        assert "sha256" in json.dumps(rows[0], default=str)


# ===========================================================================
# 8. Bounded policy + projection delivered to the UI
# ===========================================================================


class TestUiProjection:
    def test_policy_matches_the_server_constants_and_carries_no_secret(self, staging):
        model = vm.build_view_model()
        policy = model["attachmentPolicy"]
        assert policy["allowUpload"] is True
        assert policy["maxUploadBytes"] == att.MAX_UPLOAD_BYTES
        assert policy["maxUploadLabel"] == "20 MB"
        assert policy["allowedExtensions"] == list(att.ALLOWED_EXTENSIONS)
        assert policy["accept"] == att.ALLOWED_ACCEPT
        assert [item["id"] for item in policy["categories"]] == [item["id"] for item in att.CATEGORIES]
        assert policy["driveConfigured"] is False  # no credential in this environment
        assert "PRIVATE KEY" not in json.dumps(model, default=str)

    def test_viewer_policy_disables_upload_but_keeps_metadata(self, staging):
        import streamlit as st

        st.session_state["current_user"] = {"username": "viewer-1", "role": "viewer",
                                           "full_name": "Chỉ đọc", "user_id": "U-VIEWER"}
        model = vm.build_view_model()
        assert model["attachmentPolicy"]["allowUpload"] is False
        row = next(row for row in model["shipments"] if row["id"] == SHIPMENT_ID)
        assert row["attachments"] == []

    def test_shipment_rows_always_expose_an_attachment_list(self, staging):
        model = vm.build_view_model()
        assert model["shipments"], "staging must provide shipments"
        for row in model["shipments"]:
            assert isinstance(row["attachments"], list)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()


# ===========================================================================
# 11. Owner Drive hierarchy — Logistics / Customer / YYYY / MM / Shipment
# ===========================================================================


class TestOwnerDriveHierarchy:
    @staticmethod
    def _shipment(
        sid: str,
        *,
        customer_id: str = "CUST-001",
        customer: str = "CÔNG TY ABC",
        created_at: str = "2026-09-15",
    ) -> dict[str, Any]:
        return {
            "id": sid,
            "shipment_id": sid,
            "customer_id": customer_id,
            "customer": customer,
            "is_deleted": False,
            "version": 1,
            "data": {"0.0": {"created_at": created_at}},
        }

    def test_folder_names_use_authoritative_customer_and_shipment_period(self):
        shipment = self._shipment("LO-001")
        assert att.drive_folder_hierarchy("LO-001", shipment) == (
            "Logistics",
            "CUST-001 - CÔNG TY ABC",
            "2026",
            "09",
            "LO-001",
        )

        # Missing customer id keeps the real customer name; no fake id is minted.
        shipment = self._shipment("LO-002", customer_id="", customer="Khách hàng Việt")
        assert att.drive_folder_hierarchy("LO-002", shipment)[1] == "Khách hàng Việt"

        # Missing/invalid canonical shipment date never silently becomes today.
        shipment = self._shipment("LO-003", created_at="")
        assert att.drive_folder_hierarchy("LO-003", shipment)[2:4] == (
            att.DRIVE_UNKNOWN_YEAR_FOLDER,
            att.DRIVE_UNKNOWN_MONTH_FOLDER,
        )
        shipment = self._shipment("LO-004", created_at="not-a-date")
        assert att.drive_folder_hierarchy("LO-004", shipment)[2:4] == (
            att.DRIVE_UNKNOWN_YEAR_FOLDER,
            att.DRIVE_UNKNOWN_MONTH_FOLDER,
        )

    def test_customer_year_month_and_shipment_are_distinct_hierarchy_axes(self):
        base = self._shipment("LO-A")
        different_customer = self._shipment("LO-A", customer_id="CUST-002", customer="CÔNG TY XYZ")
        different_year = self._shipment("LO-A", created_at="2025-09-15")
        different_month = self._shipment("LO-A", created_at="2026-10-15")
        different_shipment = self._shipment("LO-B")

        base_path = att.drive_folder_hierarchy("LO-A", base)
        assert att.drive_folder_hierarchy("LO-A", different_customer)[1] != base_path[1]
        assert att.drive_folder_hierarchy("LO-A", different_year)[2] != base_path[2]
        assert att.drive_folder_hierarchy("LO-A", different_month)[3] != base_path[3]
        assert att.drive_folder_hierarchy("LO-B", different_shipment)[4] != base_path[4]

    def test_unsafe_folder_names_use_existing_sanitizer(self):
        shipment = self._shipment(
            "LO/2026\\09",
            customer_id="KH/01",
            customer="CÔNG TY\\ABC",
        )
        hierarchy = att.drive_folder_hierarchy("LO/2026\\09", shipment)
        assert hierarchy[0] == "Logistics"
        assert "/" not in hierarchy[1] and "\\" not in hierarchy[1]
        assert "/" not in hierarchy[4] and "\\" not in hierarchy[4]

    def test_dispatcher_uses_server_shipment_context_and_reuses_same_shipment_folder(
        self, staging, seq
    ):
        sid = "LO-HIER-001"
        authoritative = self._shipment(
            sid,
            customer_id="CUST-777",
            customer="CÔNG TY SERVER",
            created_at="2024-02-20",
        )
        db.get_db().shipments.insert_one(authoritative)

        fake = FakeDrive()
        factory = CountingFactory(fake)
        first_payload = _upload_payload(
            "to-khai.pdf", "application/pdf", _pdf_bytes(),
            shipment_id=sid, request_id="hier-1",
        )
        # Browser-supplied customer/date lies must have zero authority.
        first_payload["customer"] = "CÔNG TY GIẢ"
        first_payload["customerId"] = "FAKE-999"
        first_payload["createdAt"] = "2039-12-31"

        second_payload = _upload_payload(
            "hop-dong.docx", att.MIME_DOCX,
            _docx_bytes(), shipment_id=sid, request_id="hier-2", category="contract",
        )

        with patch("core.shipment_attachments.build_drive_client", factory):
            first = _dispatch("shipment_attachment_upload", first_payload, seq)
            creations_after_first = list(fake.creations)
            second = _dispatch("shipment_attachment_upload", second_payload, seq)

        assert first["ok"] is True, first.get("error")
        assert second["ok"] is True, second.get("error")
        assert len(fake.uploads) == 2
        assert fake.creations == creations_after_first, "same shipment must reuse all parent folders"

        expected_names = (
            "Logistics",
            "CUST-777 - CÔNG TY SERVER",
            "2024",
            "02",
            sid,
        )
        parent = "ROOT-FOLDER-FAKE"
        expected_creations = []
        for name in expected_names:
            expected_creations.append((parent, name))
            parent = fake._folder_cache[(parent, name)]

        assert fake.creations == expected_creations
        assert all(row["parentId"] == parent for row in fake.uploads)
        assert all(name not in {"shipments", "attachments"} for _p, name in fake.creations)
