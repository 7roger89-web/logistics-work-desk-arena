"""Acceptance A/B/D of LOGI_EXCEL_FIDELITY_INLINE_EDIT_V1 on the SYNTHETIC fixture.

This proves the READ-ONLY original-workbook view: real sheet names and order,
cells in their real position, merged cells, column widths, row heights, fonts,
fills, borders, alignment, number formats, source formulas with their cached
results, and an explicit inventory of what is NOT rendered.

It also proves the boundary that keeps it safe: an internal cost sheet is locked
without the business grant, the viewer writes nothing, and the customer XLSX
export never contains the source workbook, a hidden sheet, a defined name or any
internal text — in the visible cells or in the workbook metadata.

The workbook used here is SYNTHETIC (`scripts/make_luc_giac_three_sheet_fixture.py`).
It is NOT the Owner's original Lục Giác file and is not evidence that the real
workbook is reproduced 100%; that gate stays PENDING.
"""
from __future__ import annotations

import base64
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from core.quote_excel_import import parse_luc_giac_workbook
from core.quote_workbook_view import (
    SHEET_CUSTOMER,
    SHEET_INTERNAL,
    SHEET_UNCLASSIFIED,
    format_number,
    overview,
    sheet_page,
)
from scripts.make_luc_giac_three_sheet_fixture import (
    SHEET_CUSTOMER as FIX_CUSTOMER,
    SHEET_INTERNAL_COST,
    SHEET_INTERNAL_DUP,
    build_luc_giac_three_sheet_fixture,
)
from views import logi_c1, logi_sandbox

FIXTURE_NAME = "luc-giac-3sheet-SYNTHETIC.xlsx"


@pytest.fixture
def source() -> bytes:
    return build_luc_giac_three_sheet_fixture()


@pytest.fixture
def proposal(source):
    return parse_luc_giac_workbook(source, FIXTURE_NAME, today=date(2026, 10, 16))


@pytest.fixture
def staging():
    import streamlit as st

    had_user = "current_user" in st.session_state
    previous_user = st.session_state.get("current_user")
    logi_sandbox.reset_backend()
    assert logi_sandbox.ensure_backend() == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    db = logi_sandbox.get_db() if hasattr(logi_sandbox, "get_db") else None
    if db is None:
        from core.database import get_db

        db = get_db()
    db.customers.insert_one({"id": "CUST-LUCGIAC-3SHEET", "name": "Công ty Lục Giác"})
    try:
        yield db
    finally:
        logi_sandbox.reset_backend()
        if had_user:
            st.session_state["current_user"] = previous_user
        else:
            st.session_state.pop("current_user", None)


def dispatch(event, payload):
    outcomes = logi_c1.consume_c1_triggers({"seq": 0}, {event: {"seq": 1, "payload": {"requestId": "test-request", **payload}}})
    assert len(outcomes) == 1
    return outcomes[0]


def payload(details, **extra):
    return {"customerId": "CUST-LUCGIAC-3SHEET", "salesOwner": "Sale test", "goodsDescription": "Bàn họp văn phòng",
            "incoterm": "EXW", "transportMode": "Sea", "excelQuote": details, "fx": {}, **extra}


def imported(source):
    out = dispatch("quote_import_excel", {"customerId": "CUST-LUCGIAC-3SHEET", "fileName": FIXTURE_NAME,
                                          "base64": base64.b64encode(source).decode()})
    assert out["ok"], out
    return out["readback"]


def cell(page, row, column):
    return next((c for c in page["cells"] if c["r"] == row and c["c"] == column), None)


# --------------------------------------------------------------- A: sheet tabs
def test_original_workbook_keeps_real_sheet_names_and_order(source, proposal):
    info = overview(source, proposal)
    assert info["sheet_names"] == [FIX_CUSTOMER, SHEET_INTERNAL_DUP, SHEET_INTERNAL_COST]
    assert [s["index"] for s in info["sheets"]] == [1, 2, 3]
    assert [s["name"] for s in info["sheets"]] == info["sheet_names"]  # tab order == file order
    assert info["customer_sheet"] == FIX_CUSTOMER
    assert [s["classification"] for s in info["sheets"]] == [SHEET_CUSTOMER, SHEET_INTERNAL, SHEET_INTERNAL]
    # The parser read the CUSTOMER sheet, not "the third sheet" and not the copy.
    assert proposal["source"]["goods_sheet"] == FIX_CUSTOMER
    assert proposal["source"]["table_sheets"] == [FIX_CUSTOMER, SHEET_INTERNAL_DUP]


def test_customer_sheet_view_keeps_position_merges_sizes_and_styles(source, proposal):
    # Rendering under the grant: the raw view path always passes it explicitly.
    page = sheet_page(source, FIX_CUSTOMER, proposal=proposal, can_view_internal=True)
    assert page["read_only"] is True and page["locked"] is False
    assert page["classification"] == SHEET_CUSTOMER
    assert {"min_row": 2, "min_col": 1, "max_row": 2, "max_col": 8, "ref": "A2:H2"} in page["merges"]
    assert cell(page, 2, 1)["value"] == "BÁO GIÁ DỊCH VỤ LOGISTICS"
    assert cell(page, 2, 1)["font"]["bold"] is True and cell(page, 2, 1)["fill"] == "#D9E2F3"
    assert cell(page, 2, 1)["align"]["h"] == "center"
    assert cell(page, 7, 6)["value"] == "Số lượng" and cell(page, 7, 6)["fill"] == "#E4EBF7"
    assert "b:thin" in cell(page, 8, 1)["border"] and "r:thin" in cell(page, 8, 1)["border"]
    assert cell(page, 8, 7)["nf"] == '#,##0'
    assert page["column_widths"]["A"] == 34.0 and page["column_widths"]["B"] == 22.0
    assert page["frozen_panes"] == "A8"
    # percentage and currency formats travel with the cell
    assert cell(page, 21, 2)["nf"] == "0%" and cell(page, 21, 2)["display"] == "80%"
    assert cell(page, 21, 2)["value"] == 0.8


def test_source_formulas_keep_their_cached_results(source, proposal):
    page = sheet_page(source, FIX_CUSTOMER, proposal=proposal, can_view_internal=True)
    formula = cell(page, 13, 6)
    assert formula["kind"] == "formula" and formula["f"] == "=SUM(F8:F12)"
    assert formula["value"] == 27 and "unknown" not in formula
    total = cell(page, 13, 7)
    assert total["f"] == "=SUMPRODUCT(F8:F12,G8:G12)"
    assert total["value"] == 607180 and "RMB" in total["display"]


def test_formula_without_cache_is_unknown_never_zero(source, proposal):
    page = sheet_page(source, FIX_CUSTOMER, proposal=proposal, can_view_internal=True)
    unknown = cell(page, 22, 2)
    assert unknown["f"] == "=SUM(F8:F12)*500" and unknown["unknown"] is True
    assert unknown["value"] == "" and unknown["value"] != 0
    assert "display" not in unknown


def test_internal_cost_sheet_locked_without_grant_and_readable_with_it(source, proposal):
    locked = sheet_page(source, SHEET_INTERNAL_COST, proposal=proposal, can_view_internal=False)
    assert locked["locked"] is True and locked["cells"] == []
    assert locked["classification"] == SHEET_INTERNAL and "commercial_write" in locked["reason"]
    assert locked["row_count"] > 0  # the user still learns the sheet exists
    granted = sheet_page(source, SHEET_INTERNAL_COST, proposal=proposal, can_view_internal=True)
    assert granted["locked"] is False
    labels = [c["value"] for c in granted["cells"] if c.get("kind") == "value"]
    assert "LCC Hải Phòng" in labels and 12500000 in labels


def test_source_cells_are_default_deny_on_the_customer_sheet_too(source, proposal):
    """The module default must be deny: content costs a grant, for every sheet."""
    page = sheet_page(source, FIX_CUSTOMER, proposal=proposal)          # default can_view_internal=False
    assert page["locked"] is True and page["cells"] == []
    assert page["classification"] == SHEET_CUSTOMER
    assert "commercial_write" in page["reason"]
    assert sheet_page(source, FIX_CUSTOMER, proposal=proposal, can_view_internal=True)["locked"] is False


def test_viewer_reports_what_it_does_not_render(tmp_path, proposal):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = FIX_CUSTOMER
    ws["A1"] = "Khách gửi"
    ws["B1"] = 10
    ws["B2"] = "=A1"
    ws.conditional_formatting.add("B1:B2", CellIsRule(operator="greaterThan", formula=["5"]))
    validation = DataValidation(type="list", formula1='"A,B"')
    ws.add_data_validation(validation)
    validation.add("B1")
    ws.conditional_formatting  # keep the rule alive on the sheet
    wb.defined_names.add(__import__("openpyxl").workbook.defined_name.DefinedName("VUNG_GIA", attr_text=f"{FIX_CUSTOMER}!$B$1"))
    hidden = wb.create_sheet("Tinh noi bo")
    hidden["A1"] = "PRIVATE-COST-SHEET"
    hidden.sheet_state = "hidden"
    target = Path(tmp_path) / "mixed.xlsx"
    wb.save(target)
    data = target.read_bytes()

    info = overview(data, None)
    names = info["sheet_names"]
    assert names == [FIX_CUSTOMER, "Tinh noi bo"]
    hidden_info = next(s for s in info["sheets"] if s["name"] == "Tinh noi bo")
    assert hidden_info["state"] == "hidden" and hidden_info["not_rendered"]["sheet_hidden"] == 1
    customer_info = info["sheets"][0]
    assert customer_info["not_rendered"]["conditional_formatting"] == 1
    assert customer_info["not_rendered"]["data_validation"] == 1
    assert customer_info["not_rendered"]["defined_names"] == 1
    # The limitation is declared, never hidden behind a "100% preserved" claim.
    fidelity = info["fidelity"]
    assert fidelity["claim"].count("KHÔNG") == 1
    assert any("Hình ảnh" in item for item in fidelity["not_rendered"])
    assert any("Định dạng có điều kiện" in item for item in fidelity["not_rendered"])
    assert any("Ô gộp" in item for item in fidelity["supported"])


def test_bounded_number_format_renderer_is_honest_about_unknown_codes():
    assert format_number(607180, '#,##0" RMB"').endswith("RMB")
    assert format_number(0.8, "0%") == "80%"
    assert format_number(1234.5, "#,##0.00").replace("\u00a0", ",") == "1,234.50"
    assert format_number(date(2026, 11, 15), "dd/mm/yyyy") == "15/11/2026"
    assert format_number(5, "[$$-409]#,##0.00") in (None, "5.00")  # no fake precision claimed
    assert format_number("text", "#,##0") is None


# --------------------------------------------------------- RBAC / handler path
def test_view_event_is_read_only_and_never_returns_the_file(staging, source):
    imported_source = imported(source)
    created = dispatch("quote_create", payload(imported_source))
    assert created["ok"], created
    qid = created["id"]
    quotes_before = staging.quotes.count_documents({})

    out = dispatch("quote_source_workbook_view", {"id": qid})
    assert out["ok"], out
    readback = out["readback"]
    assert readback["read_only"] is True
    assert readback["overview"]["sheet_names"] == [FIX_CUSTOMER, SHEET_INTERNAL_DUP, SHEET_INTERNAL_COST]
    assert readback["page"] is None  # no sheet asked for yet
    assert "document" not in out and "downloadBase64" not in str(out)
    assert staging.quotes.count_documents({}) == quotes_before  # nothing written

    page = dispatch("quote_source_workbook_view", {"id": qid, "sheet": FIX_CUSTOMER})
    assert page["ok"], page
    assert page["readback"]["page"]["sheet"] == FIX_CUSTOMER
    assert any(c.get("value") == "BÁO GIÁ DỊCH VỤ LOGISTICS" for c in page["readback"]["page"]["cells"])


def test_view_event_locks_internal_sheet_for_read_only_role(staging, source):
    import streamlit as st

    created = dispatch("quote_create", payload(imported(source)))
    assert created["ok"], created
    qid = created["id"]
    st.session_state["current_user"] = {"username": "viewer-test", "role": "viewer"}
    # `commercial_view` is enough to ASK, never to receive source cells — not even
    # from the sheet the parser called "customer". This assertion used to read
    # "commercial_view is enough for the customer sheet", which is exactly the
    # hole: a one-tab supplier file keeps its cost figures inside that sheet.
    customer_sheet = dispatch("quote_source_workbook_view", {"id": qid, "sheet": FIX_CUSTOMER})
    assert customer_sheet["ok"] is True                       # metadata, not content
    assert customer_sheet["readback"]["overview"]["sheet_names"] == [FIX_CUSTOMER, SHEET_INTERNAL_DUP, SHEET_INTERNAL_COST]
    customer_page = customer_sheet["readback"]["page"]
    assert customer_page["locked"] is True and customer_page["cells"] == [] and customer_page["merges"] == []
    assert customer_page["row_count"] > 0                     # the sheet is still disclosed as existing
    assert "commercial_write" in customer_page["reason"]
    internal = dispatch("quote_source_workbook_view", {"id": qid, "sheet": SHEET_INTERNAL_COST})
    assert internal["ok"] is True and internal["readback"]["page"]["locked"] is True
    assert internal["readback"]["page"]["cells"] == []
    st.session_state["current_user"] = {}  # no authenticated user at all
    refused = dispatch("quote_source_workbook_view", {"id": qid})
    assert refused["ok"] is False and "RBAC" in refused["error"]


def test_view_event_refuses_bytes_that_do_not_match_the_stored_sha256(staging, source, monkeypatch):
    from core.gridfs_client import GridFSAttachmentStore

    created = dispatch("quote_create", payload(imported(source)))
    assert created["ok"], created
    monkeypatch.setattr(GridFSAttachmentStore, "download_file", lambda self, file_id: b"tampered-bytes")
    out = dispatch("quote_source_workbook_view", {"id": created["id"], "sheet": FIX_CUSTOMER})
    assert out["ok"] is False and "SHA256" in out["error"]


# --------------------------------------------- export must not carry the source
def test_customer_export_excludes_source_sheets_hidden_sheets_and_metadata(tmp_path):
    """Acceptance I: no cost sheet, no hidden sheet, no defined name, no metadata."""
    from openpyxl.workbook.defined_name import DefinedName

    wb = load_workbook(BytesIO(build_luc_giac_three_sheet_fixture()))
    hidden = wb.create_sheet("Nội bộ ẩn")
    hidden["A1"] = "PRIVATE-COST-SHEET"
    hidden["A2"] = 987654321
    hidden.sheet_state = "hidden"
    wb.defined_names.add(DefinedName("NOI_BO_GIA", attr_text="'Nội bộ ẩn'!$A$2"))
    wb.properties.title = "PRIVATE-WORKBOOK-METADATA"
    wb.properties.creator = "PRIVATE-AUTHOR"
    target = Path(tmp_path) / "with_hidden.xlsx"
    wb.save(target)
    data = target.read_bytes()

    from scripts.make_luc_giac_three_sheet_fixture import build_luc_giac_three_sheet_fixture as _build  # noqa: F401

    proposal = parse_luc_giac_workbook(data, "with_hidden.xlsx")
    from core.quotations import calculate_quote, customer_quote_snapshot
    from core.excel_utils import generate_quote_excel

    quote = calculate_quote({"excel_quote": proposal, "status": "DRAFT", "customer_name": "Công ty Lục Giác"})
    buffer = generate_quote_excel(quote)
    exported = buffer.getvalue()
    assert b"PRIVATE-COST-SHEET" not in exported and b"987654321" not in exported
    assert b"PRIVATE-WORKBOOK-METADATA" not in exported and b"PRIVATE-AUTHOR" not in exported

    reopened = load_workbook(BytesIO(exported))
    assert reopened.sheetnames == ["Bao gia"]
    assert all(sheet.sheet_state == "visible" for sheet in reopened.worksheets)
    assert list(reopened.defined_names) == []
    cells = [c for sheet in reopened for row in sheet for c in row]
    assert not any(c.data_type == "f" for c in cells)
    text = "\n".join(str(c.value or "") for c in cells)
    assert "31800000" not in text and "12500000" not in text and "PRIVATE" not in text
    assert "Tính nội bộ" not in text
    # The customer sheet's own facts are still there.
    assert "BÁO GIÁ DỊCH VỤ LOGISTICS" in text and "Thanh toán 80% trước khi giao hàng." in text
    with ZipFile(BytesIO(exported)) as archive:
        names = archive.namelist()
        assert sum(1 for name in names if name.startswith("xl/worksheets/sheet")) == 1
        package_text = "".join(archive.read(name).decode("utf-8", "ignore") for name in names if name.endswith(".xml"))
    assert "PRIVATE" not in package_text and "Nội bộ ẩn" not in package_text
    assert customer_quote_snapshot(quote)["text"]["values"]["doc.title"] == "BÁO GIÁ DỊCH VỤ LOGISTICS"


def test_unclassified_sheet_is_not_silently_called_internal(tmp_path, proposal):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Phu luc"
    wb.active["A1"] = "Phụ lục kỹ thuật"
    target = Path(tmp_path) / "appendix.xlsx"
    wb.save(target)
    info = overview(target.read_bytes(), None)
    assert info["sheets"][0]["classification"] == SHEET_UNCLASSIFIED
