"""Build the bounded real-input fixture for LOGI_GETFLY_DSDH_IMPORT_PARITY_01.

The workbook deliberately uses merged order/date/customer cells. openpyxl writes
blank cells below the first row of each order, so the real parser's forward-fill
path is exercised without a pasted-table shortcut.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


HEADERS = [
    "Mã ĐH",
    "Ngày đặt hàng",
    "Tên KH",
    "Mã KH",
    "Mã SP",
    "Tên SP",
    "Mô tả",
    "Số lượng",
    "Giá vốn",
    "Giá bán",
    "Thành tiền",
    "Lợi nhuận SP",
    "Trạng thái",
]

ROWS = [
    ["LO-0306", "05/09/2026", "An Phát JSC", "CUST-ANPHAT", "VC-001", "Cước xe", "Chuyến chính", 1, "100000", "150000", "150000", "50000", "Đang xử lý"],
    [None, None, None, None, "DG-002", "Đóng gói", "Kiện gỗ", 2, "50000", "75000", "150000", "50000", None],
    ["GF-DSDH-001", "06/09/2026", "Đức Long Co.", None, "VC-NEW", "Vận chuyển mới", "Hàng mẫu", 1, "200000", "300000", "300000", "100000", "Đang xử lý"],
    [None, None, None, None, "DV-NEW", "Bốc xếp", "Tại kho", 3, "10000", "20000", "60000", "30000", None],
]


def build_fixture(output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DSDH"
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
    for row_number, values in enumerate(ROWS, start=2):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
    # The source shape uses merged cells for order-level facts. The parser sees
    # the lower cells as blank and must attach them to the right order.
    for column in range(1, 5):
        sheet.merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
        sheet.merge_cells(start_row=4, start_column=column, end_row=5, end_column=column)
    workbook.save(output)
    return output


if __name__ == "__main__":
    import sys

    destination = sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/DSDH.xlsx"
    print(build_fixture(destination))
