"""Synthetic LOG-DOCGEN-01 fixtures, built in memory; NOT legal templates.

Only the mapped cell/run topology and OOXML features under regression are
modelled. These fixtures prove programmatic behaviour, not Native Office
rendering or fidelity of the owner's private legal templates. Never commit
emitted XLSX/DOCX bytes or install them in DOCGEN_TEMPLATE_DIR.
"""

from __future__ import annotations

import io
import zipfile

import docx
import openpyxl
from docx.shared import Cm, Pt
from lxml import etree
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"x": SHEET_NS}
SHEET_PARTS = {
    "HDMB": "xl/worksheets/sheet1.xml",
    "pl": "xl/worksheets/sheet2.xml",
    "pxk": "xl/worksheets/sheet3.xml",
}
FORMULAS = {
    "HDMB": {"F25": "SUM(F24:F24)", "F26": "F25*8%", "F27": "SUM(F25:F26)"},
    "pl": {"F28": "SUM(F10:F27)", "F29": "F28*8%", "F30": "SUM(F28:F29)"},
}


def xml_bytes(element) -> bytes:
    return etree.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)


def zip_parts(raw: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert archive.testzip() is None
        return {name: archive.read(name) for name in archive.namelist()}


def pack_parts(parts: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def build_synthetic_xlsx(
    *, shared_strings: bool = True, existing_defined_names: bool = True,
    pxk_formula_inputs: bool = True,
) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_PARTS:
        sheet = workbook.create_sheet(name)
        # Explicit empty cells: the narrow engine does not invent unmapped cells.
        for row in sheet.iter_rows(min_row=1, max_row=44, min_col=1, max_col=6):
            for cell in row:
                cell.value = ""
        sheet["A1"] = f"SYNTHETIC {name} — NOT A LEGAL DOCUMENT"
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        for coord, formula in FORMULAS.get(name, {}).items():
            sheet[coord] = "=" + formula

    for coord in ("E24", "F24"):
        cell = workbook["HDMB"][coord]
        cell.value = 0
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        cell.number_format = "#,##0"
        cell.font = Font(name="Times New Roman", size=11)
        cell.fill = PatternFill("solid", fgColor="EEEEEE")
        cell.border = Border(bottom=Side(style="thin"))

    pl = workbook["pl"]
    pl.row_dimensions[6].height = 36
    pl["A6"].alignment = Alignment(wrap_text=True, vertical="center")
    for row in (1, 2, 3, 4, 5, 6, 7, 8, 31, 33):
        pl.merge_cells(f"A{row}:F{row}")
    if existing_defined_names:
        workbook.defined_names.add(DefinedName("synthetic_keep", attr_text="'HDMB'!$A$1"))
        pl.print_title_rows = "$8:$8"  # The repair must replace this, not duplicate it.

    pxk = workbook["pxk"]
    for row in (6, 7, 10, 11):
        pxk[f"A{row}"].alignment = Alignment(vertical="center", wrap_text=False)
        pxk.row_dimensions[row].height = 22
    for row in (6, 7):
        pxk.merge_cells(f"A{row}:D{row}")
    for target, source in (("B15", "B24"), ("C15", "C24"), ("D15", "D24")):
        pxk[target] = f"=HDMB!{source}" if pxk_formula_inputs else 0

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    parts = zip_parts(buffer.getvalue())
    rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    # Match the supported template's workbook-relative target convention.
    for rel in rels:
        if rel.get("Target", "").startswith("/xl/"):
            rel.set("Target", rel.get("Target")[4:])
    types = etree.fromstring(parts["[Content_Types].xml"])
    chain = etree.Element(f"{{{SHEET_NS}}}calcChain", nsmap={None: SHEET_NS})
    strings = etree.Element(f"{{{SHEET_NS}}}sst", nsmap={None: SHEET_NS})
    string_indices: dict[str, int] = {}
    string_count = 0

    for sheet_id, path in enumerate(SHEET_PARTS.values(), start=1):
        sheet = etree.fromstring(parts[path])
        for cell in sheet.findall(".//x:c", NS):
            if cell.find("x:f", NS) is not None:
                # Stale cache deliberately differs from every expected total.
                value = cell.find("x:v", NS)
                if value is None:
                    value = etree.SubElement(cell, f"{{{SHEET_NS}}}v")
                value.text = "777"
                etree.SubElement(chain, f"{{{SHEET_NS}}}c", r=cell.get("r"), i=str(sheet_id))
            if shared_strings and cell.get("t") == "inlineStr":
                text = "".join(t.text or "" for t in cell.findall(".//x:t", NS))
                if text not in string_indices:
                    string_indices[text] = len(string_indices)
                    item = etree.SubElement(strings, f"{{{SHEET_NS}}}si")
                    etree.SubElement(item, f"{{{SHEET_NS}}}t").text = text
                for child in list(cell):
                    cell.remove(child)
                cell.set("t", "s")
                etree.SubElement(cell, f"{{{SHEET_NS}}}v").text = str(string_indices[text])
                string_count += 1
        parts[path] = xml_bytes(sheet)

    for stem in ["calcChain"] + (["sharedStrings"] if shared_strings else []):
        etree.SubElement(rels, f"{{{PKG_REL_NS}}}Relationship", {
            "Id": f"rSynthetic{stem}", "Type": f"{OFFICE_REL_NS}/{stem}", "Target": f"{stem}.xml",
        })
        etree.SubElement(types, f"{{{CONTENT_NS}}}Override", {
            "PartName": f"/xl/{stem}.xml",
            "ContentType": f"application/vnd.openxmlformats-officedocument.spreadsheetml.{stem}+xml",
        })
    if shared_strings:
        strings.set("count", str(string_count))
        strings.set("uniqueCount", str(len(string_indices)))
        parts["xl/sharedStrings.xml"] = xml_bytes(strings)
    parts["xl/calcChain.xml"] = xml_bytes(chain)
    parts["xl/_rels/workbook.xml.rels"] = xml_bytes(rels)
    parts["[Content_Types].xml"] = xml_bytes(types)
    return pack_parts(parts)


def build_synthetic_docx() -> bytes:
    document = docx.Document()
    runs_by_paragraph = {
        2: ["Số:", " ", "", "SYNTHETIC-NUMBER"],
        4: ["Ngày ", "…", " tháng ", "…", " năm ", "…", " tại ", "", "SYNTHETIC"],
        5: ["Bên A:", " ", "", "SYNTHETIC AGENT"],
        6: ["Địa chỉ: ", "SYNTHETIC"],
        7: ["Điện thoại: ", "SYNTHETIC"],
        8: ["Mã số thuế: ", "SYNTHETIC"],
        9: ["Đại diện: Ông", "/Bà ….." + " " * 90, " " * 90 + "Chức vụ: ", "….." + " " * 90],
        11: ["Bên B:", " ", "", "SYNTHETIC COMPANY"],
        12: ["Địa chỉ: ", "SYNTHETIC"],
        13: ["Điện thoại: SYNTHETIC"],
        14: ["Mã số thuế: ", "SYNTHETIC"],
        15: ["Đại diện", " : ", "SYNTHETIC REP", " " * 90, " " * 90, "Chức vụ: ….."],
        22: ["SYNTHETIC", "", "", "", " từ ", "…", " đến ", "…"],
        30: ["Tài khoản", ": ", "SYNTHETIC", " tại ", "", "SYNTHETIC BANK"],
    }
    for index in range(31):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        for run_index, text in enumerate(runs_by_paragraph.get(index, [f"SYNTHETIC paragraph {index}"])):
            run = paragraph.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.bold = run_index == 0
            run.italic = run_index == 1
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
