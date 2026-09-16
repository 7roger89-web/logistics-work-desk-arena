"""LOGI_GETFLY_DSDH_IMPORT_PARITY_01 — real XLSX GetFly import.

The acceptance boundary is deliberately small:

    DSDH.xlsx bytes → existing core.getfly_reader → server preview in session
    → explicit operator choices → save_shipment_atomic → fresh readback.

No preview collection, direct Mongo mutation, customer-master side effect,
client-normalized commit data, or inferred VAT is allowed. All tests use the
isolated mongomock staging transport.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import openpyxl
import pytest
import streamlit as st

from core.getfly_reader import process_getfly_excel
from views import logi_sandbox
from views import logi_view_model as vm


FIXTURE = Path(__file__).parent / "fixtures" / "DSDH.xlsx"


def _event(seq: int, name: str, payload: dict) -> dict:
    return {name: {"seq": seq, "event": name, "payload": payload}}


def _xlsx_payload(path: Path = FIXTURE, *, request_id: str = "GF-PREVIEW-TEST") -> dict:
    raw = path.read_bytes()
    return {
        "filename": path.name,
        "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "size": len(raw),
        "contentBase64": base64.b64encode(raw).decode("ascii"),
        "requestId": request_id,
    }


@pytest.fixture()
def staging():
    logi_sandbox.reset_backend()
    logi_sandbox.ensure_backend()
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    st.session_state.pop(vm.GETFLY_PREVIEW_KEY, None)
    st.session_state.pop(vm.LAST_WRITES_KEY, None)
    yield
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    st.session_state.pop(vm.GETFLY_PREVIEW_KEY, None)
    logi_sandbox.reset_backend()


def _preview(seq: int = 1) -> dict:
    outcomes = vm.consume_phase_b_triggers(
        {"seq": 0}, _event(seq, "getfly_preview", _xlsx_payload())
    )
    assert len(outcomes) == 1, outcomes
    assert outcomes[0]["ok"], outcomes[0]
    return vm._getfly_preview_projection()


def _commit(preview: dict, *, mode: str, seq: int = 2, **payload) -> dict:
    body = {
        "previewId": preview["previewId"],
        "mode": mode,
        "requestId": payload.pop("requestId", f"GF-COMMIT-{seq}"),
        **payload,
    }
    outcomes = vm.consume_phase_b_triggers(
        {"seq": 0}, _event(seq, "getfly_import", body)
    )
    assert len(outcomes) == 1, outcomes
    return outcomes[0]


def test_real_xlsx_parser_and_preview_have_zero_writes(staging):
    raw = FIXTURE.read_bytes()
    parsed, error = process_getfly_excel(io.BytesIO(raw))
    assert error is None
    assert list(parsed["order_id"].dropna().unique()) == ["LO-0306", "GF-DSDH-001"]
    assert parsed.loc[1, "order_id"] == "LO-0306", "merged order cell must be forward-filled"
    assert parsed.loc[3, "order_id"] == "GF-DSDH-001"
    assert parsed.loc[1, "product_note"] == "Kiện gỗ"

    before_ids = set(vm.load_shipments())
    preview = _preview()
    after_ids = set(vm.load_shipments())

    assert before_ids == after_ids, "preview must perform no shipment mutation"
    assert "contentBase64" not in st.session_state[vm.GETFLY_PREVIEW_KEY]
    assert preview["filename"] == "DSDH.xlsx"
    assert len(preview["normalizedRows"]) == 4
    assert preview["normalizedRows"][2]["line_id"] == "GF-DSDH-001::L1"
    assert preview["summary"] == {"orders": 2, "revenue": 660000.0, "profit": 230000.0}
    assert [order["orderId"] for order in preview["orders"]] == ["LO-0306", "GF-DSDH-001"]
    assert preview["orders"][0]["existing"] is True
    assert preview["orders"][1]["existing"] is False
    assert preview["orders"][0]["serviceLines"][1]["description"] == "Kiện gỗ"
    assert preview["orders"][0]["customerMatch"] == "MATCHED_BY_CODE"
    assert preview["orders"][1]["customerMatch"] == "MATCHED_BY_NAME"
    # The same server projection is present on the final root view-model that
    # React receives; this is not a side-channel or client-local preview.
    model = vm.build_view_model()
    assert model["getfly"]["previewId"] == preview["previewId"]
    assert model["getfly"]["fingerprint"] == preview["fingerprint"]
    assert len(model["getfly"]["normalizedRows"]) == 4


def test_mode_a_creates_only_missing_shipment_and_replay_skips(staging):
    preview = _preview()
    out = _commit(preview, mode=vm.GETFLY_MODE_SHIPMENT_ONLY)
    assert out["ok"], out
    assert [row["status"] for row in out["results"]] == [
        "SKIPPED_ALREADY_EXISTS", "CREATED"
    ]

    fresh = vm.read_shipment("GF-DSDH-001")
    assert fresh is not None
    assert fresh["customer"] == "Đức Long Co."
    assert fresh["customer_id"] == "CUST-DUCLONG"
    assert fresh["data"]["0.0"]["order_date"] == "2026-09-06"
    assert fresh["data"]["0.0"]["source_status"] == "Đang xử lý"
    # Source did not state any operational/commercial fact; none is invented.
    assert fresh["data"]["0.0"].get("mode", "") == ""
    assert fresh["data"]["0.0"].get("incoterm", "") == ""
    assert fresh["data"]["1.2"] == {}
    assert fresh["data"]["2.4"] == {}
    assert fresh["data"]["2.7"] == {}
    assert fresh["data"]["entrustment"] == {}
    # Mode A does not book source service prices/costs as Finance Actual.
    assert fresh["data"]["3.4"]["details"] == []

    replay = _commit(preview, mode=vm.GETFLY_MODE_SHIPMENT_ONLY, seq=3)
    assert replay["ok"], replay
    assert [row["status"] for row in replay["results"]] == [
        "SKIPPED_ALREADY_EXISTS", "SKIPPED_ALREADY_EXISTS"
    ]
    assert len(vm.read_shipment("GF-DSDH-001")["data"]["3.4"]["details"]) == 0


def test_mode_b_refuses_blank_vat_before_mutation_then_accepts_zero_and_eight(staging):
    preview = _preview()
    new_order = next(order for order in preview["orders"] if not order["existing"])
    first, second = new_order["serviceLines"]

    blank_vat = _commit(
        preview,
        mode=vm.GETFLY_MODE_SHIPMENT_AND_FINANCE,
        vatByLineId={first["lineId"]: "0"},
    )
    assert not blank_vat["ok"]
    refused = next(row for row in blank_vat["results"] if row["orderId"] == "GF-DSDH-001")
    assert refused["status"] == "REFUSED_VALIDATION"
    assert "VAT" in " ".join(refused["errors"])
    assert vm.read_shipment("GF-DSDH-001") is None, "blank VAT must refuse before shipment mutation"

    accepted = _commit(
        preview,
        mode=vm.GETFLY_MODE_SHIPMENT_AND_FINANCE,
        seq=3,
        vatByLineId={first["lineId"]: "0", second["lineId"]: "8"},
    )
    assert accepted["ok"], accepted
    assert next(row for row in accepted["results"] if row["orderId"] == "GF-DSDH-001")["status"] == "CREATED"

    fresh = vm.read_shipment("GF-DSDH-001")
    details = fresh["data"]["3.4"]["details"]
    assert len(details) == 2
    assert details[0]["code"] == "VC-NEW"
    assert details[0]["qty"] == 1.0
    assert details[0]["cost"] == 200000.0
    assert details[0]["price"] == 300000.0
    assert details[0]["vat_line"] == 0.0
    assert details[0]["note"] == "Hàng mẫu"
    assert details[1]["qty"] == 3.0
    assert details[1]["cost"] == 10000.0
    assert details[1]["price"] == 20000.0
    assert details[1]["vat_line"] == 8.0
    assert fresh["data"]["3.4"]["out"] == 364800.0
    assert fresh["data"]["3.4"]["in"] == 230000.0
    assert fresh["data"]["3.4"]["output_vat"] == 4800.0
    assert vm.read_shipment("GF-DSDH-001")["data"]["3.4"]["details"] == details


@pytest.mark.parametrize(
    ("column", "value", "error_fragment"),
    [
        (10, None, "Giá bán"),       # explicit absence is UNKNOWN, not zero
        (9, "not-a-number", "Giá vốn"),  # invalid cost is not coerced to zero
    ],
)
def test_mode_b_refuses_missing_or_invalid_source_money_before_mutation(
    staging, column, value, error_fragment
):
    workbook = openpyxl.load_workbook(FIXTURE)
    sheet = workbook.active
    # The first data row is the existing order; the second order is the
    # missing-order target and has two complete lines in the fixture.
    sheet.cell(row=4, column=column).value = value
    stream = io.BytesIO()
    workbook.save(stream)
    raw = stream.getvalue()
    outcome = vm._handle_getfly_preview({
        "filename": "DSDH.xlsx", "size": len(raw),
        "contentBase64": base64.b64encode(raw).decode(), "requestId": "MONEY-1",
    })
    assert outcome["ok"], outcome
    preview = vm._getfly_preview_projection()
    new_order = next(order for order in preview["orders"] if not order["existing"])
    vat = {line["lineId"]: "8" for line in new_order["serviceLines"]}
    refused = _commit(
        preview,
        mode=vm.GETFLY_MODE_SHIPMENT_AND_FINANCE,
        vatByLineId=vat,
        requestId="MONEY-COMMIT",
    )
    assert not refused["ok"]
    result = next(row for row in refused["results"] if row["orderId"] == "GF-DSDH-001")
    assert result["status"] == "REFUSED_VALIDATION"
    assert error_fragment in " ".join(result["errors"])
    assert vm.read_shipment("GF-DSDH-001") is None


def test_failed_canonical_write_is_reported_per_order(staging, monkeypatch):
    preview = _preview()
    import core.database as database

    monkeypatch.setattr(database, "save_shipment_atomic", lambda *args, **kwargs: False)
    out = _commit(preview, mode=vm.GETFLY_MODE_SHIPMENT_ONLY)
    assert [row["status"] for row in out["results"]] == [
        "SKIPPED_ALREADY_EXISTS", "FAILED_WRITE"
    ]
    assert vm.read_shipment("GF-DSDH-001") is None


def test_server_ignores_client_normalized_order_data(staging):
    preview = _preview()
    malicious = _commit(
        preview,
        mode=vm.GETFLY_MODE_SHIPMENT_ONLY,
        orders=[{"orderId": "GF-NOT-IN-PREVIEW", "customerName": "Invented"}],
    )
    assert malicious["ok"]
    assert {row["orderId"] for row in malicious["results"]} == {"LO-0306", "GF-DSDH-001"}
    assert vm.read_shipment("GF-NOT-IN-PREVIEW") is None
    assert vm.read_shipment("GF-DSDH-001")["customer"] == "Đức Long Co."


def test_viewer_commit_is_refused_with_zero_mutation(staging):
    preview = _preview()
    before = set(vm.load_shipments())
    st.session_state["current_user"] = {
        "username": "viewer.getfly", "full_name": "Viewer", "role": "viewer", "user_id": "VIEWER-1"
    }
    out = _commit(preview, mode=vm.GETFLY_MODE_SHIPMENT_ONLY)
    assert not out["ok"]
    assert "RBAC" in out["error"]
    assert all(row["status"] == "REFUSED_VALIDATION" for row in out["results"])
    assert set(vm.load_shipments()) == before


def test_malformed_file_and_no_valid_order_id_fail_honestly(staging):
    bad = vm._handle_getfly_preview({
        "filename": "DSDH.xlsx", "size": 8,
        "contentBase64": base64.b64encode(b"not xlsx").decode(), "requestId": "BAD-1",
    })
    assert not bad["ok"]
    assert "Không thể đọc DSDH.xlsx" in bad["error"]
    assert set(vm.load_shipments()) == {"LO-0306", "LO-0307", "LO-0308", "LO-0299", "LO-0295"}

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Mã ĐH", "Ngày đặt hàng", "Tên KH", "Tên SP", "Giá bán"])
    sheet.append([None, "06/09/2026", "Đức Long Co.", "Dịch vụ", 10])
    stream = io.BytesIO()
    workbook.save(stream)
    raw = stream.getvalue()
    no_id = vm._handle_getfly_preview({
        "filename": "DSDH.xlsx", "size": len(raw),
        "contentBase64": base64.b64encode(raw).decode(), "requestId": "BAD-2",
    })
    assert not no_id["ok"]
    assert "Mã ĐH" in no_id["error"]
    assert set(vm.load_shipments()) == {"LO-0306", "LO-0307", "LO-0308", "LO-0299", "LO-0295"}


def test_unknown_customer_is_not_fuzzy_linked(staging):
    from core.database import load_customers

    before_customer_ids = set(load_customers())
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Mã ĐH", "Ngày đặt hàng", "Tên KH", "Mã KH", "Tên SP"])
    # A bad source code must not be rescued by a familiar-looking name;
    # name-only exact matching is for rows that have no source code.
    sheet.append(["GF-UNKNOWN", "06/09/2026", "An Phát JSC", "MISSING-CODE", "Dịch vụ"])
    stream = io.BytesIO()
    workbook.save(stream)
    raw = stream.getvalue()
    outcome = vm._handle_getfly_preview({
        "filename": "DSDH.xlsx", "size": len(raw),
        "contentBase64": base64.b64encode(raw).decode(), "requestId": "UNKNOWN-1",
    })
    assert outcome["ok"], outcome
    preview = vm._getfly_preview_projection()
    order = preview["orders"][0]
    assert order["customerMatch"] == "UNKNOWN"
    assert order["customerId"] == ""
    assert set(load_customers()) == before_customer_ids, "preview must not auto-create customer masters"
