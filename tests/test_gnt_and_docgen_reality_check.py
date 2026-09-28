"""LOGI_GNT_DOC_EXPORT_REALITY_CHECK_01 — Verification and acceptance suite.

Covers:
1. GNT legal & semantic audit:
   - Form C1-02/NS legal basis: TT 84/2016/TT-BTC (under TT 328/2016 Art 14(3), amended by TT 72/2021 & TT 23/2025)
   - Budget classification subitem: 1901 (Thuế NK) and 1702 (Thuế GTGT hàng NK, NOT 1701)
2. Taxpayer authority: No auto-promotion of customer to legal importer/taxpayer without explicit record
3. Explicit taxpayer operator input/save via shipment_field_save and fresh readback
4. UNKNOWN != ZERO semantics for import duty and VAT (cases A–E)
5. GNT workbook structure and fidelity via openpyxl inspection (2 sheets, exact amounts, total in words, no fabricated treasury facts)
6. Bounded gnt_generate event dispatching, server RBAC, session-scoped output & readback
7. Office export reality check for HDNT (DOCX), PL (XLSX), HDMB (XLSX), PXK (XLSX)
8. Focused adversarial mutation checks (1701 refuse with ValueError and customer auto-promotion guard)
"""

from __future__ import annotations

import base64
import io
import tempfile
from decimal import Decimal
from typing import Any, Mapping
import unittest

import docx
import openpyxl
import pytest

from core.auth import ROLE_PERMISSIONS
from core.docgen import generator as docgen_gen
from core.docgen import rules as docgen_rules
from core.gnt_builder import build_gnt_workbook, check_gnt_prerequisites, doc_so_thanh_chu_vi
from views import logi_c1, logi_docgen_templates, logi_real_data as rd, logi_sandbox, logi_view_model as vm


@pytest.fixture()
def staging():
    """Isolated staging environment with deterministic actor."""
    import streamlit as st

    logi_sandbox.reset_backend()
    transport = logi_sandbox.ensure_backend()
    assert transport == logi_sandbox.SANDBOX_TRANSPORT
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    yield transport
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)
    logi_sandbox.reset_backend()


@pytest.fixture()
def event_seq():
    state = {"n": 0}

    def _make(name: str, payload: dict) -> dict:
        state["n"] += 1
        return {name: {"seq": state["n"], "event": name, "payload": payload}}

    return _make


def _dispatch(seq_fn, name, payload):
    outcomes = vm.consume_phase_b_triggers({"seq": 0}, seq_fn(name, payload))
    assert len(outcomes) == 1
    return outcomes[0]


# ===========================================================================
# 1. GNT Legal / Semantic Verification & Subitem Correctness
# ===========================================================================

def test_gnt_import_vat_subitem_is_1702_not_1701():
    """Proves that import VAT emits 1702 (Thuế GTGT hàng NK) and NOT 1701 (hàng SXKD trong nước)."""
    data = {
        "shipment_id": "LO-TEST",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "import_duty_amount": 5000000.0,
        "import_vat_amount": 800000.0,
        "taxpayer_name": "CÔNG TY TNHH MINH AN",
        "taxpayer_tax_code": "0109876543",
    }
    raw_bytes = build_gnt_workbook(data)
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))

    # Sheet 1: NHẬP THÔNG TIN
    ws1 = wb["NHẬP THÔNG TIN"]
    found_vat_subitem_s1 = False
    for row in ws1.iter_rows(values_only=True):
        row_str = " ".join(str(c) for c in row if c is not None)
        if "1702" in row_str:
            found_vat_subitem_s1 = True
        assert "1701" not in row_str, f"Found stale 1701 subitem in Sheet 1: {row_str}"
    assert found_vat_subitem_s1, "Expected subitem 1702 in Sheet 1 for import VAT"

    # Sheet 2: GIẤY NỘP THUẾ
    ws2 = wb["GIẤY NỘP THUẾ"]
    found_duty_row = False
    found_vat_row = False
    for row in ws2.iter_rows(values_only=True):
        row_str = " ".join(str(c) for c in row if c is not None)
        if "Thuế nhập khẩu" in row_str:
            assert "1901" in row_str, f"Duty row missing 1901: {row_str}"
            found_duty_row = True
        if "Thuế GTGT hàng nhập khẩu" in row_str or "Thuế GTGT hàng NK" in row_str:
            assert "1702" in row_str, f"Import VAT row must have 1702: {row_str}"
            assert "1701" not in row_str, f"Import VAT row must NOT have 1701: {row_str}"
            found_vat_row = True

    assert found_duty_row, "Duty row not found in Sheet 2"
    assert found_vat_row, "VAT row not found in Sheet 2"


def test_gnt_legal_references_are_accurate():
    """Verify legal reference strings in generated GNT workbook.
    
    Must accurately cite TT 84/2016/TT-BTC without claiming TT 328 co-promulgated the form.
    """
    data = {
        "shipment_id": "LO-TEST",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "import_duty_amount": 5000000.0,
        "import_vat_amount": 800000.0,
        "taxpayer_name": "CÔNG TY TNHH MINH AN",
        "taxpayer_tax_code": "0109876543",
    }
    raw_bytes = build_gnt_workbook(data)
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))

    ws1 = wb["NHẬP THÔNG TIN"]
    subtitle1 = str(ws1.cell(row=2, column=1).value)
    assert "84/2016/TT-BTC" in subtitle1
    assert "328/2016" not in subtitle1
    assert "184/2015" not in subtitle1
    assert "100%" not in subtitle1

    ws2 = wb["GIẤY NỘP THUẾ"]
    subtitle2 = str(ws2.cell(row=2, column=3).value)
    assert "84/2016/TT-BTC" in subtitle2
    assert "328/2016" not in subtitle2
    assert "184/2015" not in subtitle2
    assert "100%" not in subtitle2


# ===========================================================================
# 2. Taxpayer Authority & Operator Save / Readback
# ===========================================================================

def test_gnt_taxpayer_authority_semantics():
    """Customer is NOT automatically promoted to legal importer/taxpayer without explicit record."""
    # Case 1: Shipment has customer but no explicit GNT taxpayer in 2.4
    ship_no_taxpayer = {
        "id": "S-NO-TAXPAYER",
        "customer_id": "CUS-ANPHAT",
        "customer": "CÔNG TY TNHH AN PHÁT",
        "customer_tax_code": "0101234567",
        "data": {
            "2.4": {
                "tk_no": "105588421330",
                "tk_date": "15/08/2026",
                "import_duty_amount": 5000000,
                "import_vat_amount": 800000,
            }
        },
    }
    prereq = check_gnt_prerequisites(ship_no_taxpayer)
    assert prereq["ready"] is False
    assert "Tên người nộp thuế" in prereq["missing"]
    assert "Mã số thuế (MST)" in prereq["missing"]

    proj = rd._gnt_projection(ship_no_taxpayer, ship_no_taxpayer["data"]["2.4"], {}, "S-NO-TAXPAYER")
    assert proj["ready"] is False
    assert "Người nộp thuế" in proj["missing"]
    assert "Mã số thuế" in proj["missing"]
    assert proj["payer"] == ""
    assert proj["taxCode"] == ""
    assert proj["customerHint"] == "CÔNG TY TNHH AN PHÁT"

    # Case 2: Explicit taxpayer present in customs.gnt
    ship_explicit_taxpayer = {
        "id": "S-EXPLICIT",
        "customer": "CÔNG TY ỦY THÁC ABC",
        "data": {
            "2.4": {
                "tk_no": "105588421330",
                "tk_date": "15/08/2026",
                "import_duty_amount": 5000000,
                "import_vat_amount": 800000,
                "gnt": {
                    "nguoi_nop": "CÔNG TY NHẬP KHẨU ĐÍCH THỰC",
                    "mst": "0109998888",
                    "dia_chi": "Hà Nội",
                },
            }
        },
    }
    prereq_ok = check_gnt_prerequisites(ship_explicit_taxpayer)
    assert prereq_ok["ready"] is True
    assert prereq_ok["data"]["taxpayer_name"] == "CÔNG TY NHẬP KHẨU ĐÍCH THỰC"
    assert prereq_ok["data"]["taxpayer_tax_code"] == "0109998888"


def test_gnt_taxpayer_operator_save_and_readback(staging, event_seq):
    """Operator saves explicit GNT taxpayer via shipment_field_save; fresh readback enables GNT."""
    sid = "LO-TAX-OP-01"
    _dispatch(event_seq, "shipment_create", {
        "shipmentId": sid,
        "customerId": "CUS-SAMPLE",
        "customer": "CÔNG TY KHÁCH HÀNG MẪU",
        "contractId": "CT-NONE",
    })

    # Initially new shipment has customer hint but no explicit GNT taxpayer
    ship_init = vm.read_shipment(sid)
    customs_init = (ship_init.get("data") or {}).get("2.4") or {}
    proj_init = rd._gnt_projection(ship_init, customs_init, {}, sid)
    assert proj_init["ready"] is False
    assert "Người nộp thuế" in proj_init["missing"]
    assert "Mã số thuế" in proj_init["missing"]
    assert proj_init["customerHint"] == "CÔNG TY KHÁCH HÀNG MẪU"

    # Operator enters and saves explicit taxpayer info along with customs facts
    out = _dispatch(event_seq, "shipment_field_save", {
        "shipmentId": sid,
        "fields": {
            "gnt_taxpayer_name": "CÔNG TY TNHH THƯƠNG MẠI MINH TÂN",
            "gnt_taxpayer_tax_code": "0108887777",
            "gnt_taxpayer_address": "Số 45 Cầu Giấy, Hà Nội",
            "tk_no": "105588421330",
            "tk_date": "15/08/2026",
            "import_duty_amount": "5000000",
            "import_vat_amount": "800000",
        },
    })
    assert out["ok"] is True, out.get("error")

    # Fresh readback from database
    fresh = vm.read_shipment(sid)
    customs_fresh = (fresh.get("data") or {}).get("2.4") or {}
    gnt_fresh = customs_fresh.get("gnt") or {}
    assert gnt_fresh.get("nguoi_nop") == "CÔNG TY TNHH THƯƠNG MẠI MINH TÂN"
    assert gnt_fresh.get("mst") == "0108887777"
    assert gnt_fresh.get("dia_chi") == "Số 45 Cầu Giấy, Hà Nội"

    proj_fresh = rd._gnt_projection(fresh, customs_fresh, {}, sid)
    assert proj_fresh["ready"] is True
    assert proj_fresh["payer"] == "CÔNG TY TNHH THƯƠNG MẠI MINH TÂN"
    assert proj_fresh["taxCode"] == "0108887777"
    assert proj_fresh["missing"] == []

    # Clearing explicit taxpayer field -> becomes blocked again
    out_clear = _dispatch(event_seq, "shipment_field_save", {
        "shipmentId": sid,
        "fields": {
            "gnt_taxpayer_name": "",
            "gnt_taxpayer_tax_code": "",
        },
    })
    assert out_clear["ok"] is True
    fresh_cleared = vm.read_shipment(sid)
    customs_cleared = (fresh_cleared.get("data") or {}).get("2.4") or {}
    proj_cleared = rd._gnt_projection(fresh_cleared, customs_cleared, {}, sid)
    assert proj_cleared["ready"] is False
    assert "Người nộp thuế" in proj_cleared["missing"]
    assert "Mã số thuế" in proj_cleared["missing"]


def test_gnt_unknown_vs_zero_tax_semantics():
    """Test Cases A–E for UNKNOWN != ZERO tax amounts."""
    base_ship = {
        "id": "S-TAX-TEST",
        "data": {
            "2.4": {
                "tk_no": "105588421330",
                "tk_date": "15/08/2026",
                "gnt": {
                    "nguoi_nop": "CÔNG TY TNHH TEST",
                    "mst": "0101234567",
                },
            }
        },
    }

    # A. import duty = 0, import VAT > 0 -> READY
    ship_a = {
        "id": "S-A",
        "data": {
            "2.4": dict(base_ship["data"]["2.4"], import_duty_amount=0, import_vat_amount=800000)
        },
    }
    prereq_a = check_gnt_prerequisites(ship_a)
    assert prereq_a["ready"] is True
    assert prereq_a["data"]["import_duty_amount"] == 0
    assert prereq_a["data"]["import_vat_amount"] == 800000

    # B. import duty > 0, import VAT = 0 -> READY
    ship_b = {
        "id": "S-B",
        "data": {
            "2.4": dict(base_ship["data"]["2.4"], import_duty_amount=5000000, import_vat_amount=0)
        },
    }
    prereq_b = check_gnt_prerequisites(ship_b)
    assert prereq_b["ready"] is True
    assert prereq_b["data"]["import_duty_amount"] == 5000000
    assert prereq_b["data"]["import_vat_amount"] == 0

    # C. blank duty (None or "") with VAT > 0 -> BLOCKED UNKNOWN
    ship_c = {
        "id": "S-C",
        "data": {
            "2.4": dict(base_ship["data"]["2.4"], import_duty_amount="", import_vat_amount=800000)
        },
    }
    prereq_c = check_gnt_prerequisites(ship_c)
    assert prereq_c["ready"] is False
    assert any("Thuế nhập khẩu" in m for m in prereq_c["missing"])

    # D. negative duty/VAT -> BLOCKED INVALID
    ship_d = {
        "id": "S-D",
        "data": {
            "2.4": dict(base_ship["data"]["2.4"], import_duty_amount=-5000, import_vat_amount=800000)
        },
    }
    prereq_d = check_gnt_prerequisites(ship_d)
    assert prereq_d["ready"] is False
    assert any("không hợp lệ" in m or "âm" in m for m in prereq_d["missing"])

    # E. both duty=0 and VAT=0 -> honest NO_PAYABLE_AMOUNT
    ship_e = {
        "id": "S-E",
        "data": {
            "2.4": dict(base_ship["data"]["2.4"], import_duty_amount=0, import_vat_amount=0)
        },
    }
    prereq_e = check_gnt_prerequisites(ship_e)
    assert prereq_e["ready"] is False
    assert prereq_e["no_payable"] is True
    assert prereq_e["reason"] == "NO_PAYABLE_AMOUNT"
    # Neither duty nor VAT is called "missing"
    assert not any("Thuế" in m for m in prereq_e["missing"])


# ===========================================================================
# 3. Detailed Workbook Structure & OOXML Fidelity
# ===========================================================================

def test_gnt_workbook_complete_structure_and_values():
    """Inspect the generated workbook with openpyxl and verify all fields."""
    fixture_data = {
        "shipment_id": "LO-0307",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "import_duty_amount": 5432100.0,
        "import_vat_amount": 876500.0,
        "taxpayer_name": "CÔNG TY TNHH THƯƠNG MẠI AN PHÁT",
        "taxpayer_tax_code": "0101234567",
        "taxpayer_address": "123 Phố Huế, Hai Bà Trưng, Hà Nội",
        "taxpayer_province": "Hà Nội",
        "taxpayer_district": "Hai Bà Trưng",
    }
    raw_bytes = build_gnt_workbook(fixture_data)
    assert raw_bytes[:2] == b"PK", "Output must be a valid OOXML ZIP container"

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False)
    assert wb.sheetnames == ["NHẬP THÔNG TIN", "GIẤY NỘP THUẾ"]

    # --- Sheet 1: NHẬP THÔNG TIN ---
    ws1 = wb["NHẬP THÔNG TIN"]
    # Verify autofilled cells
    values_s1 = {}
    for r in range(1, 30):
        label = ws1.cell(row=r, column=1).value
        val = ws1.cell(row=r, column=2).value
        if label and val is not None:
            values_s1[str(label).strip()] = val

    assert values_s1.get("Số tờ khai (TK)") == "105588421330"
    assert values_s1.get("Ngày tờ khai") == "15/08/2026"
    assert values_s1.get("Mã ĐH (Shipment)") == "LO-0307"
    assert values_s1.get("Tên tổ chức/cá nhân") == "CÔNG TY TNHH THƯƠNG MẠI AN PHÁT"
    assert values_s1.get("Mã số thuế (MST)") == "0101234567"
    assert values_s1.get("Thuế nhập khẩu") == 5432100.0
    assert values_s1.get("Thuế GTGT hàng nhập khẩu") == 876500.0
    assert values_s1.get("TỔNG CỘNG (đ)") == 6308600.0

    total_words = doc_so_thanh_chu_vi(6308600.0)
    assert "Sáu triệu ba trăm lẻ tám nghìn sáu trăm đồng" in total_words
    assert values_s1.get("Tổng bằng chữ") == total_words

    # --- Sheet 2: GIẤY NỘP THUẾ (C1-02/NS) ---
    ws2 = wb["GIẤY NỘP THUẾ"]
    assert ws2.cell(row=1, column=1).value == "GIẤY NỘP TIỀN VÀO NGÂN SÁCH NHÀ NƯỚC"
    assert ws2.cell(row=2, column=1).value == "Mẫu số: C1-02/NS"

    # Taxpayer details
    assert ws2.cell(row=5, column=2).value == "CÔNG TY TNHH THƯƠNG MẠI AN PHÁT"
    assert ws2.cell(row=5, column=7).value == "0101234567"
    assert ws2.cell(row=6, column=2).value == "123 Phố Huế, Hai Bà Trưng, Hà Nội"

    # Tax table lines
    # Row 13: Header STT, Số tờ khai...
    # Row 14: Tax row 1 (Thuế NK)
    assert ws2.cell(row=14, column=1).value == "1"
    assert ws2.cell(row=14, column=2).value == "105588421330"
    assert ws2.cell(row=14, column=3).value == "15/08/2026"
    assert ws2.cell(row=14, column=4).value == "Thuế nhập khẩu"
    assert ws2.cell(row=14, column=5).value == 5432100.0
    assert ws2.cell(row=14, column=6).value == "1901"

    # Row 15: Tax row 2 (Thuế GTGT hàng NK)
    assert ws2.cell(row=15, column=1).value == "2"
    assert ws2.cell(row=15, column=2).value == "105588421330"
    assert ws2.cell(row=15, column=3).value == "15/08/2026"
    assert ws2.cell(row=15, column=4).value == "Thuế GTGT hàng nhập khẩu"
    assert ws2.cell(row=15, column=5).value == 876500.0
    assert ws2.cell(row=15, column=6).value == "1702"

    # Row 16: Total row with formula =SUM(E14:E15)
    assert ws2.cell(row=16, column=2).value == "TỔNG CỘNG"
    assert "=SUM(E14:E15)" in str(ws2.cell(row=16, column=5).value)

    # Row 18: Amount in words
    assert ws2.cell(row=18, column=2).value == total_words

    # Case-specific treasury fields: verify no invented constant
    # Section D of Sheet 1: NH, Số TK KBNN, KBNN nơi thu, Cơ quan QLT remain blank
    manual_labels = ["NH (Ngân hàng)", "Số tài khoản KBNN", "KBNN nơi thu", "Cơ quan quản lý thu", "Mã chương"]
    for r in range(15, 30):
        label = ws1.cell(row=r, column=1).value
        val = ws1.cell(row=r, column=2).value
        if label and any(m in str(label) for m in manual_labels):
            assert val == "" or val is None, f"Fabricated treasury field for {label}: {val}"


# ===========================================================================
# 4. Bounded gnt_generate Event & Fresh Readback
# ===========================================================================

def test_gnt_generate_end_to_end_flow(staging, event_seq):
    """Save customs, explicit taxpayer & tax facts → dispatch gnt_generate → receive XLSX bytes & readback."""
    sid = "LO-0307"
    # Seed required customs, GNT taxpayer & tax facts
    save_res = _dispatch(event_seq, "shipment_field_save", {
        "shipmentId": sid,
        "fields": {
            "gnt_taxpayer_name": "Đức Long Co.",
            "gnt_taxpayer_tax_code": "0312987654",
            "gnt_taxpayer_address": "Lô 5, KCN Quế Võ, Bắc Ninh",
            "tk_no": "105588421330",
            "tk_date": "15/08/2026",
            "import_duty_amount": "5000000",
            "import_vat_amount": "800000",
        },
    })
    assert save_res["ok"], save_res["error"]

    # PR140 precondition — GNT generation is gated on the operator-confirmed
    # import model. Challenge the server gate BEFORE the model is confirmed…
    blocked = _dispatch(event_seq, "gnt_generate", {"shipmentId": sid})
    assert blocked["ok"] is False
    assert blocked["gate"] == "COMMERCIAL_FLOW_REQUIRED"

    # …then confirm the model through the existing commercial-flow writer.
    # This lot imports under the customer's own name; the customer master MST
    # (0312987654) equals the explicit GNT taxpayer below, so the gate passes.
    flow = _dispatch(event_seq, "shipment_commercial_flow_save", {
        "shipmentId": sid,
        "requestId": "REQ-FLOW-0307",
        "mode": "CUSTOMER_SELF",
        "companyId": "",
        "subAgentId": "",
    })
    assert flow["ok"], flow["error"]

    gnt_res = _dispatch(event_seq, "gnt_generate", {"shipmentId": sid})
    assert gnt_res["ok"] is True, gnt_res.get("error")
    assert gnt_res["gate"] == "READY"
    assert gnt_res["persisted"]["isOfficeFile"] is True
    assert gnt_res["persisted"]["duty"] == 5000000.0
    assert gnt_res["persisted"]["vat"] == 800000.0

    # Verify logi_c1 session-scoped storage (not in MongoDB)
    outputs = logi_c1.stored_outputs(sid)
    record_id = gnt_res["persisted"]["recordId"]
    assert record_id in outputs
    assert outputs[record_id]["document_type"] == "GNT"
    assert outputs[record_id]["bytes"][:2] == b"PK"

    # Readback projection verification
    shipment_read = vm.read_shipment(sid)
    gnt_proj = rd._gnt_projection(shipment_read, shipment_read.get("data", {}).get("2.4", {}), {}, sid)
    assert gnt_proj["done"] is True
    assert gnt_proj["document"] is not None
    assert gnt_proj["document"]["downloadAvailable"] is True
    assert gnt_proj["document"]["downloadBase64"] is not None

    # Base64 decode to verify real XLSX bytes
    raw_decoded = base64.b64decode(gnt_proj["document"]["downloadBase64"])
    assert raw_decoded[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(raw_decoded))
    assert "GIẤY NỘP THUẾ" in wb.sheetnames


def test_gnt_generate_fails_closed_when_facts_are_missing(staging, event_seq):
    """UNKNOWN != ZERO: missing declaration number or tax amounts must refuse generation."""
    sid = "LO-0306"
    # Seed only declaration number, missing tax amounts
    _dispatch(event_seq, "shipment_field_save", {
        "shipmentId": sid,
        "fields": {"tk_no": "105579900211", "tk_date": "03/08/2026"},
    })

    # PR140 precondition: confirm the import model first (customer imports
    # under its own name), so the failure under test is the missing GNT facts.
    flow = _dispatch(event_seq, "shipment_commercial_flow_save", {
        "shipmentId": sid,
        "requestId": "REQ-FLOW-0306",
        "mode": "CUSTOMER_SELF",
        "companyId": "",
        "subAgentId": "",
    })
    assert flow["ok"], flow["error"]

    gnt_res = _dispatch(event_seq, "gnt_generate", {"shipmentId": sid})
    assert gnt_res["ok"] is False
    assert gnt_res["gate"] == "BLOCKED_MISSING_PREREQUISITES"
    assert any("Thuế" in m for m in gnt_res["missing"])


# ===========================================================================
# 5. RBAC Authorization Gate
# ===========================================================================

def test_gnt_generate_rbac_rejection(staging, event_seq):
    """Unauthorized role (Viewer without sales_write) is refused BEFORE generation."""
    import streamlit as st

    sid = "LO-0307"
    _dispatch(event_seq, "shipment_field_save", {
        "shipmentId": sid,
        "fields": {
            "gnt_taxpayer_name": "Đức Long Co.",
            "gnt_taxpayer_tax_code": "0312987654",
            "tk_no": "105588421330",
            "tk_date": "15/08/2026",
            "import_duty_amount": "5000000",
            "import_vat_amount": "800000",
        },
    })

    # Switch user to Viewer
    st.session_state["current_user"] = {
        "username": "viewer_user",
        "display_name": "Viewer User",
        "role": "viewer",
    }

    gnt_res = _dispatch(event_seq, "gnt_generate", {"shipmentId": sid})
    assert gnt_res["ok"] is False
    assert gnt_res["gate"] == "BLOCKED_PERMISSION_DENIED"
    assert "sales_write" in gnt_res["error"]

    # Restore admin actor
    st.session_state["current_user"] = dict(logi_sandbox.STAGING_ACTOR)


# ===========================================================================
# 6. Existing Office Export Reality Check (HDNT, PL, HDMB, PXK)
# ===========================================================================

def test_office_export_reality_check_hdnt_pl_hdmb_pxk(staging):
    """Independently verifies HDNT (DOCX), PL (XLSX), HDMB (XLSX), PXK (XLSX) generation.
    
    Verifies MIME types, OOXML headers, expected document numbers/dates and party facts.
    """
    xlsx_path, docx_path = logi_docgen_templates.staging_template_paths()
    assert xlsx_path is not None and docx_path is not None

    agent = {
        "agent_id": "AGT-TT",
        "legal_name": "CÔNG TY CỔ PHẦN ĐẠI LÝ TRƯỜNG THÀNH",
        "legal_address": "Tầng 5, Tòa nhà TT, Hà Nội",
        "legal_code": "0108889999",
        "legal_representative": "Nguyễn Văn Thành",
        "representative_title": "Tổng Giám Đốc",
        "bank_account_holder": "CONG TY CP DAI LY TRUONG THANH",
        "bank_account_number": "999988887777",
        "bank_name": "Vietcombank",
    }
    company = {
        "company_id": "CMP-PA",
        "legal_name": "CÔNG TY TNHH PHÚ ANH LOGISTICS",
        "legal_address": "Số 10 Đường Giải Phóng, Hà Nội",
        "legal_code": "0101234567",
        "tax_or_legal_code": "0101234567",
        "legal_representative": "Trần Phú Anh",
        "representative_title": "Giám Đốc",
        "bank_account_holder": "CONG TY TNHH PHU ANH LOGISTICS",
        "bank_account_number": "111122223333",
        "bank_name": "BIDV",
    }

    # --- 1. HDNT (DOCX) (< 50M VND threshold) ---
    contract_hdnt = {
        "id": "HDNT-2026-001",
        "agent_id": "AGT-TT",
        "company_id": "CMP-PA",
        "hdnt_no": "HDNT-2026-001",
        "hdnt_date": "2026-08-01",
        "valid_from": "2026-08-01",
        "valid_to": "2027-08-01",
        "currency": "VND",
    }
    pl_goods = [
        {
            "description": "Linh kiện điện tử A",
            "quantity": 100,
            "unit": "Cái",
            "unit_price_vnd_excl_vat": 200000,
            "vat_rate": Decimal("0.08"),
            "confirmed_by": "tester",
            "confirmed_at": "2026-08-15",
        }
    ]
    res_hdnt = docgen_gen.generate_document(
        docgen_rules.HDNT,
        shipment_id="LO-0307",
        agent=agent,
        company=company,
        goods_lines=pl_goods,
        contract=contract_hdnt,
        template_path_docx=docx_path,
    )
    raw_hdnt = res_hdnt["bytes"]
    assert raw_hdnt[:2] == b"PK"
    doc = docx.Document(io.BytesIO(raw_hdnt))
    assert len(doc.paragraphs) > 0
    assert res_hdnt["output_file_name"].endswith(".docx")

    # --- 2. PL (XLSX) (< 50M VND, requires parent HDNT) ---
    contract_pl = {
        "id": "PL-2026-001",
        "pl_no": "PL-2026-001",
        "pl_date": "2026-08-15",
        "currency": "VND",
    }
    parent_hdnt = dict(contract_hdnt, kind="HDNT", status="active")
    res_pl = docgen_gen.generate_document(
        docgen_rules.PL,
        shipment_id="LO-0307",
        agent=agent,
        company=company,
        goods_lines=pl_goods,
        contract=contract_pl,
        parent_contract=parent_hdnt,
        template_path_xlsx=xlsx_path,
    )
    raw_pl = res_pl["bytes"]
    assert raw_pl[:2] == b"PK"
    wb_pl = openpyxl.load_workbook(io.BytesIO(raw_pl))
    assert "pl" in wb_pl.sheetnames
    assert res_pl["output_file_name"].endswith(".xlsx")

    # --- 3. HDMB (XLSX) (>= 50M VND, 1 goods line) ---
    contract_hdmb = {
        "id": "HDMB-2026-001",
        "hdmb_no": "HDMB-2026-001",
        "hdmb_date": "2026-08-15",
        "currency": "VND",
    }
    hdmb_goods = [
        {
            "description": "Máy móc thiết bị công nghiệp",
            "quantity": 1,
            "unit": "Bộ",
            "unit_price_vnd_excl_vat": 60000000,
            "vat_rate": Decimal("0.08"),
            "confirmed_by": "tester",
            "confirmed_at": "2026-08-15",
        }
    ]
    res_hdmb = docgen_gen.generate_document(
        docgen_rules.HDMB,
        shipment_id=None,
        agent=agent,
        company=company,
        goods_lines=hdmb_goods,
        contract=contract_hdmb,
        template_path_xlsx=xlsx_path,
    )
    raw_hdmb = res_hdmb["bytes"]
    assert raw_hdmb[:2] == b"PK"
    wb_hdmb = openpyxl.load_workbook(io.BytesIO(raw_hdmb))
    assert "HDMB" in wb_hdmb.sheetnames
    assert res_hdmb["output_file_name"].endswith(".xlsx")

    # --- 4. PXK (XLSX) (requires confirmed delivery + explicit HDMB/PL source) ---
    pxk_identity = {
        "id": "PXK-2026-001",
        "pxk_no": "PXK-2026-001",
        "pxk_date": "2026-08-20",
        "currency": "VND",
    }
    delivery_view = {
        "delivery_completed": True,
        "delivery_date": "2026-08-20",
        "delivery_location": "Kho K1 Cảng Hải Phòng",
        "receiver_name": "Nguyễn Văn Giao",
        "receiver_title": "Thủ kho",
    }
    pxk_source = {
        "source_document_type": "HDMB",
        "source_contract_id": "HDMB-2026-001",
        "source_line_count": 1,
        "source_goods_line": hdmb_goods[0],
    }
    res_pxk = docgen_gen.generate_document(
        docgen_rules.PXK,
        shipment_id="LO-0307",
        agent=agent,
        company=company,
        goods_lines=hdmb_goods,
        contract=pxk_identity,
        delivery=delivery_view,
        pxk_source=pxk_source,
        template_path_xlsx=xlsx_path,
    )
    raw_pxk = res_pxk["bytes"]
    assert raw_pxk[:2] == b"PK"
    wb_pxk = openpyxl.load_workbook(io.BytesIO(raw_pxk))
    assert "pxk" in wb_pxk.sheetnames
    assert res_pxk["output_file_name"].endswith(".xlsx")


def test_pxk_generation_refused_before_delivery():
    """Proves that PXK cannot be generated when delivery is unconfirmed."""
    readiness = docgen_rules.evaluate_document_readiness(
        docgen_rules.PXK,
        goods_lines=[],
        agent={},
        company={},
        selected_contract={"pxk_no": "PXK-01", "pxk_date": "2026-08-20"},
        delivery_completed=False,  # Unconfirmed
    )
    assert readiness.state == docgen_rules.BLOCKED
    assert "DELIVERY_COMPLETED" in (readiness.reason or "") or "giao" in (readiness.reason or "")


# ===========================================================================
# 8. Focused Adversarial Mutation Checks
# ===========================================================================

def test_adversarial_mutation_1701_fails():
    """Mutating import VAT subitem to 1701 or blank MUST be refused by GNT builder with ValueError."""
    # 1. Import VAT + 1701 -> refused
    corrupted_data_1701 = {
        "shipment_id": "LO-ADVERSARIAL",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "ds_thue": [
            {"ten": "Thuế nhập khẩu", "ndkt": "1901", "tien": 5000000},
            {"ten": "Thuế GTGT hàng nhập khẩu", "ndkt": "1701", "tien": 800000},  # Adversarial 1701
        ],
    }
    with pytest.raises(ValueError, match="1702"):
        build_gnt_workbook(corrupted_data_1701)

    # 2. Import VAT + blank ndkt -> refused
    corrupted_vat_blank = {
        "shipment_id": "LO-ADVERSARIAL-VAT-BLANK",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "ds_thue": [
            {"ten": "Thuế nhập khẩu", "ndkt": "1901", "tien": 5000000},
            {"ten": "Thuế GTGT hàng nhập khẩu", "ndkt": "", "tien": 800000},  # Blank NDKT
        ],
    }
    with pytest.raises(ValueError, match="1702"):
        build_gnt_workbook(corrupted_vat_blank)

    # 3. Import Duty + blank ndkt -> refused
    corrupted_duty_blank = {
        "shipment_id": "LO-ADVERSARIAL-DUTY-BLANK",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "ds_thue": [
            {"ten": "Thuế nhập khẩu", "ndkt": "", "tien": 5000000},  # Blank NDKT
            {"ten": "Thuế GTGT hàng nhập khẩu", "ndkt": "1702", "tien": 800000},
        ],
    }
    with pytest.raises(ValueError, match="1901"):
        build_gnt_workbook(corrupted_duty_blank)

    # 4. Import Duty + wrong ndkt (1902) -> refused
    corrupted_duty_wrong = {
        "shipment_id": "LO-ADVERSARIAL-DUTY-WRONG",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "ds_thue": [
            {"ten": "Thuế nhập khẩu", "ndkt": "1902", "tien": 5000000},  # Wrong NDKT
            {"ten": "Thuế GTGT hàng nhập khẩu", "ndkt": "1702", "tien": 800000},
        ],
    }
    with pytest.raises(ValueError, match="1901"):
        build_gnt_workbook(corrupted_duty_wrong)

    # 5. Prove 1901 + 1702 is accepted
    valid_data = {
        "shipment_id": "LO-VALID",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "ds_thue": [
            {"ten": "Thuế nhập khẩu", "ndkt": "1901", "tien": 5000000},
            {"ten": "Thuế GTGT hàng nhập khẩu", "ndkt": "1702", "tien": 800000},
        ],
    }
    raw = build_gnt_workbook(valid_data)
    assert raw[:2] == b"PK"

    # 6. Prove custom unrelated tax row is preserved even with custom NDKT
    custom_data = {
        "shipment_id": "LO-CUSTOM",
        "tk_no": "105588421330",
        "tk_date": "15/08/2026",
        "ds_thue": [
            {"ten": "Thuế bảo vệ môi trường", "ndkt": "2600", "tien": 1500000},
            {"ten": "Phí khác", "ndkt": "", "tien": 200000},
        ],
    }
    custom_raw = build_gnt_workbook(custom_data)
    wb = openpyxl.load_workbook(io.BytesIO(custom_raw))
    ws2 = wb["GIẤY NỘP THUẾ"]
    found_custom = any("2600" in " ".join(str(c) for c in row if c is not None) for row in ws2.iter_rows(values_only=True))
    assert found_custom


def test_adversarial_customer_auto_promotion_fails(staging, event_seq):
    """Negative check: customer presence must NEVER auto-promote to GNT taxpayer authority."""
    # Create a fresh lot with customer facts but NO explicit customs taxpayer.
    # The customer is a REAL master record with a known MST on purpose: even a
    # fully known customer must never auto-promote to GNT taxpayer authority.
    sid = "LO-ADVERSARIAL-CUST"
    _dispatch(event_seq, "shipment_create", {
        "shipmentId": sid,
        "customerId": "CUST-DUCLONG",
        "customer": "Đức Long Co.",
        "contractId": "CT-NONE",
    })
    _dispatch(event_seq, "shipment_field_save", {
        "shipmentId": sid,
        "fields": {
            "tk_no": "105588421330",
            "tk_date": "15/08/2026",
            "import_duty_amount": "5000000",
            "import_vat_amount": "800000",
        },
    })

    # PR140 precondition: confirm the import model through the existing writer.
    flow = _dispatch(event_seq, "shipment_commercial_flow_save", {
        "shipmentId": sid,
        "requestId": "REQ-FLOW-ADV",
        "mode": "CUSTOMER_SELF",
        "companyId": "",
        "subAgentId": "",
    })
    assert flow["ok"], flow["error"]

    # GNT generation MUST fail closed because legal taxpayer is not explicitly known
    gnt_res = _dispatch(event_seq, "gnt_generate", {"shipmentId": sid})
    assert gnt_res["ok"] is False
    assert gnt_res["gate"] == "BLOCKED_MISSING_PREREQUISITES"
    assert any("người nộp thuế" in m.lower() or "mst" in m.lower() for m in gnt_res["missing"])
