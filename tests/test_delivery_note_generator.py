"""Focused evidence for the bounded Phiếu giao hàng XLSX generator."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from core.delivery_note import (
    DeliveryNoteError,
    generate_delivery_note_xlsx,
    inspect_delivery_note_source,
)


def _shipment(*, delivered: bool = True, goods_description: str = "Máy móc công nghiệp"):
    return {
        "id": "LO-DN-001",
        "customer_id": "CUS-DN",
        "customer": "CÔNG TY KHÁCH HÀNG",
        "data": {
            "1.2": {
                "cargo_summary": {
                    "goods_description": goods_description,
                    "package_count": 12,
                    "package_type": "Kiện",
                    "gross_weight_kg": 1250.5,
                    "cbm": 8.25,
                    "dimensions": "120x80x100 cm",
                    "note": "Giữ khô",
                },
                "transport_legs": [
                    {
                        "id": "LEG-1",
                        "sequence": 1,
                        "mode": "Truck",
                        "origin": "Hải Phòng",
                        "destination": "Hà Nội",
                        "units": [
                            {
                                "id": "TRUCK-1",
                                "kind": "truck",
                                "number": "29H-123.45",
                                "type_size": "5T",
                                "driver_name": "Nguyễn Văn A",
                                "driver_phone": "0901000000",
                            }
                        ],
                    }
                ],
            },
            "2.7": {
                "is_delivered": delivered,
                "delivered_at": "2026-09-30" if delivered else "",
                "delivery_date": "2026-09-30",
                "delivery_location": "Kho khách hàng Hà Nội",
                "receiver_name": "Trần Văn Nhận",
                "receiver_title": "Thủ kho",
                "phone_receiver": "0902000000",
            },
        },
    }


def _template(path: Path, *, with_sheet: bool = True) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "HDMB"
    if with_sheet:
        ws = wb.create_sheet("delivery_note")
        for cell in (
            "B4", "B5", "B6", "B7", "B8", "B10", "B11", "B12", "B13", "B14",
            "B16", "B17", "B18", "B19", "B20", "B21", "B22", "B24", "B25", "B26",
        ):
            ws[cell] = "..."
    wb.save(path)


def test_delivery_note_source_keeps_unknowns_and_requires_delivery():
    source = inspect_delivery_note_source(_shipment(delivered=False), customer={"tax_code": "0101234567"})

    assert source["ready"] is False
    assert any("chưa xác nhận" in item.lower() for item in source["missing"])
    assert source["package_count"] == 12
    assert source["gross_weight_kg"] == 1250.5
    assert source["customer_tax_code"] == "0101234567"


def test_delivery_note_blocks_missing_goods_description():
    source = inspect_delivery_note_source(_shipment(goods_description=""))

    assert source["ready"] is False
    assert any("mô tả hàng hóa" in item.lower() for item in source["missing"])


def test_delivery_note_requires_owner_template_sheet(tmp_path):
    path = tmp_path / "owner.xlsx"
    _template(path, with_sheet=False)

    with pytest.raises(DeliveryNoteError) as exc:
        generate_delivery_note_xlsx(
            str(path),
            shipment=_shipment(),
            customer={"tax_code": "0101234567"},
            delivery_note_no="PGH-001",
            delivery_note_date="2026-09-30",
        )

    assert exc.value.code == "TEMPLATE_SHEET_MISSING"


def test_delivery_note_generates_real_xlsx_from_canonical_facts(tmp_path):
    path = tmp_path / "owner.xlsx"
    _template(path)

    result = generate_delivery_note_xlsx(
        str(path),
        shipment=_shipment(),
        customer={"tax_code": "0101234567"},
        delivery_note_no="PGH-001/2026",
        delivery_note_date="2026-09-30",
    )

    assert result["bytes"][:2] == b"PK"
    assert result["document_type"] == "DELIVERY_NOTE"
    assert result["output_file_name"] == "PHIEU_GIAO_HANG_PGH-001-2026.xlsx"
    assert result["source_snapshot"]["delivery_note_no"] == "PGH-001/2026"
    assert result["source_snapshot"]["delivery_note_date"] == "2026-09-30"

    wb = openpyxl.load_workbook(io.BytesIO(result["bytes"]), data_only=False)
    ws = wb["delivery_note"]
    assert ws["B4"].value == "PGH-001/2026"
    assert ws["B5"].value == "30/09/2026"
    assert ws["B6"].value == "LO-DN-001"
    assert ws["B7"].value == "CÔNG TY KHÁCH HÀNG"
    assert ws["B8"].value == "0101234567"
    assert ws["B10"].value == "30/09/2026"
    assert ws["B11"].value == "Kho khách hàng Hà Nội"
    assert ws["B12"].value == "Trần Văn Nhận"
    assert ws["B16"].value == "Máy móc công nghiệp"
    assert ws["B17"].value == 12
    assert ws["B19"].value == 1250.5
    assert ws["B20"].value == 8.25
    assert ws["B24"].value == "29H-123.45"
    assert ws["B25"].value == "Nguyễn Văn A"
