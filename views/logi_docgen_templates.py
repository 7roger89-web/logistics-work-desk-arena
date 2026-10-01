"""Sandbox/staging QA template provider for LOG-DOCGEN-01 (Phase C1).

The Owner's legal Office templates are NEVER committed and ``DOCGEN_TEMPLATE_DIR``
defaults to ``None`` (the owner gate in ``core.config``).  For the isolated
staging transport ONLY, this module synthesizes minimal-but-REAL OOXML
templates — an ``.xlsx`` workbook containing the ``HDMB`` / ``pl`` / ``pxk`` /
``delivery_note`` sheets with the mapped cells, and a ``.docx`` with the mapped HDNT paragraphs —
so the real generation service (``core.docgen.generator``) can be exercised
end-to-end in Preview.

These files are inputs for QA, not business outputs.  The generated documents
are still produced by the canonical service from the contract's real facts and
are real Office files.  Production always resolves templates from
``DOCGEN_TEMPLATE_DIR``; when that is unset the service fails closed.
"""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any

_XLSX_CELLS: dict[str, dict[str, str]] = {
    "HDMB": {
        "A4": "...", "A5": "...", "A8": "...", "A9": "...", "A10": "...",
        "A11": "...", "A12": "...", "A13": "...", "A15": "...", "A16": "...",
        "A17": "...", "A18": "...", "A19": "...", "A20": "...", "B24": "...",
        "C24": "...", "D24": "...", "E24": "...", "F24": "...", "F25": "...",
        "F26": "...", "F27": "...", "A28": "...", "A35": "...", "A36": "...",
        "A37": "...", "A38": "...", "A40": "...", "A41": "...",
    },
    "pl": {
        "A4": "...", "A5": "...", "A6": "...", "A31": "...", "A33": "...",
        "F28": "...", "F29": "...", "F30": "...",
    },
    "pxk": {
        "A4": "...", "A5": "...", "A6": "...", "A7": "...", "A9": "...",
        "A10": "...", "A11": "...", "B15": "...", "C15": "...", "D15": "...",
    },
    # QA-only cell contract for core.delivery_note. Production wording/layout
    # remains in the Owner-controlled workbook under DOCGEN_TEMPLATE_DIR.
    "delivery_note": {
        "B4": "...", "B5": "...", "B6": "...", "B7": "...", "B8": "...",
        "B10": "...", "B11": "...", "B12": "...", "B13": "...", "B14": "...",
        "B16": "...", "B17": "...", "B18": "...", "B19": "...", "B20": "...",
        "B21": "...", "B22": "...", "B24": "...", "B25": "...", "B26": "...",
    },
}
# pl goods rows A10..F27 (18 lines) — the engine writes rows 10..27.
for _row in range(10, 28):
    for _col in ("A", "B", "C", "D", "E", "F"):
        _XLSX_CELLS["pl"][f"{_col}{_row}"] = "..."

# HDNT paragraph run-index requirements from build_docx_run_edits:
# para2:run3, para4:runs1/3/5/8, para5:run3, para6:run1, para7:run1, para8:run1,
# para9:runs1/2/3, para11:run3, para12:run1, para13:run0, para14:run1,
# para15:runs2/3/4/5, para22:runs5/7, para30:runs2/5.  Max index is 8 (para 4).
_DOCX_PARAGRAPHS = 31
_DOCX_RUNS_PER_PARAGRAPH = 9


def _write_xlsx_template(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    sheet_names = ["HDMB", "pl", "pxk", "delivery_note"]
    wb.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        wb.create_sheet(title=name)
    for name in sheet_names:
        ws = wb[name]
        for coord, value in _XLSX_CELLS.get(name, {}).items():
            ws[coord] = value
    wb.save(str(path))


def _write_docx_template(path: Path) -> None:
    import docx

    doc = docx.Document()
    for _ in range(_DOCX_PARAGRAPHS):
        paragraph = doc.add_paragraph()
        for _ in range(_DOCX_RUNS_PER_PARAGRAPH):
            paragraph.add_run("…")
    doc.save(str(path))


def staging_template_dir() -> Path:
    return Path(tempfile.gettempdir()) / "logi-docgen-staging-templates"


def staging_template_paths() -> tuple[str | None, str | None]:
    """Return ``(xlsx_path, docx_path)`` for the staging transport.

    Creates the synthetic templates on first use.  Never used in production:
    the caller only calls this when the transport is the sandbox AND
    ``DOCGEN_TEMPLATE_DIR`` is unset.
    """
    base = staging_template_dir()
    try:
        base.mkdir(parents=True, exist_ok=True)
    except OSError:  # pragma: no cover - temp dir always available
        return None, None
    xlsx = base / "truong-thanh-agent.xlsx"
    docx = base / "HDNT-TT-agent.docx"
    try:
        # A prior sandbox process may already have created the pre-delivery-note
        # workbook. Refresh the QA asset when its sheet contract is stale.
        refresh_xlsx = not xlsx.exists()
        if not refresh_xlsx:
            from openpyxl import load_workbook
            probe = load_workbook(str(xlsx), read_only=True, data_only=False)
            try:
                refresh_xlsx = "delivery_note" not in probe.sheetnames
            finally:
                probe.close()
        if refresh_xlsx:
            _write_xlsx_template(xlsx)
        if not docx.exists():
            _write_docx_template(docx)
    except Exception:  # noqa: BLE001 - template synth must never crash the page
        return None, None
    if not xlsx.exists() or not docx.exists():
        return None, None
    return str(xlsx), str(docx)
