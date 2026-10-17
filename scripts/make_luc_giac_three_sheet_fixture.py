"""SYNTHETIC three-sheet Lục Giác-SHAPED workbook for LOGI_EXCEL_FIDELITY_INLINE_EDIT_V1.

`SYNTHETIC` is not a formality: this file is built here, it is NOT the Owner's
original Lục Giác workbook, and it is never evidence that the original is
preserved 100%. It exists so the acceptance checks can run without the real
customer file and without uploading anything into the Agent.

Sheet names, order and shape mirror the three sheets described in the task:
  1. `BÁO GIÁ`                      — the sheet sent to the customer
  2. `LG nhập - Đại lý TT L Xuất`   — duplicated internal calculation
  3. `Dự kiến Chi phí`              — internal cost summary (never published)

It carries the formatting the viewer must reproduce (merged title, fonts,
fills, borders, alignment, number formats, column widths, row heights, freeze
panes) and BOTH formula states: one formula WITH a cached result and one WITHOUT
a cached result (which must stay UNKNOWN, never 0).

Usage:
  .venv/bin/python scripts/make_luc_giac_three_sheet_fixture.py .tmp/luc-giac-3sheet.xlsx
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET_CUSTOMER = "BÁO GIÁ"
SHEET_INTERNAL_DUP = "LG nhập - Đại lý TT L Xuất"
SHEET_INTERNAL_COST = "Dự kiến Chi phí"

# 5 rows / 27 pcs / 607,180 RMB — the figures the task fixes for acceptance C.
ROWS = (
    (5200, 1500, 760, 16, 13055),
    (7500, 1600, 760, 1, 18850),
    (12000, 1600, 760, 3, 30150),
    (16000, 1600, 760, 6, 40200),
    (19000, 1600, 760, 1, 47800),
)
HEADERS = ("STT", "Tên hàng", "Model", "Vật liệu", "Kích thước (mm)", "Số lượng", "Đơn giá (RMB)", "Ghi chú")
FIXTURE_BANNER = "FIXTURE MÔ PHỎNG — KHÔNG PHẢI WORKBOOK GỐC LỤC GIÁC"

THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(name="Arial", size=16, bold=True, color="1F3864")
HEAD_FILL = PatternFill("solid", fgColor="E4EBF7")
WARN_FONT = Font(name="Arial", size=10, bold=True, color="9C0006")


def _goods_rows(sheet, *, price_delta: int = 0, start_row: int = 8) -> None:
    for index, (length, width, height, qty, unit_price) in enumerate(ROWS, 1):
        row = start_row + index - 1
        sheet.append([
            index,
            "Bàn họp văn phòng",
            "MS-C0242S" if index == 1 else None,
            "MDF + veneer" if index == 1 else None,
            f"{length} × {width} × {height}",
            qty,
            unit_price + price_delta,
            "Chưa xác nhận model/vật liệu" if index > 1 else None,
        ])
        for column in range(1, 9):
            sheet.cell(row=row, column=column).border = BORDER
        sheet.cell(row=row, column=6).alignment = Alignment(horizontal="center")
        sheet.cell(row=row, column=7).number_format = '#,##0'


def build_luc_giac_three_sheet_fixture() -> bytes:
    wb = Workbook()

    # ---------------------------------------------------------------- sheet 1
    ws = wb.active
    ws.title = SHEET_CUSTOMER
    ws.merge_cells("A1:H1")
    ws["A1"] = FIXTURE_BANNER
    ws["A1"].font = WARN_FONT
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2")
    ws["A2"] = "BÁO GIÁ DỊCH VỤ LOGISTICS"
    ws["A2"].font = TITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].fill = PatternFill("solid", fgColor="D9E2F3")
    ws.row_dimensions[2].height = 28
    ws["A3"], ws["B3"] = "Kính gửi", "Công ty Lục Giác"
    ws.merge_cells("B3:E3")
    ws["A4"], ws["B4"] = "Nhà cung cấp", "Guangdong Hongye Furniture Group Co., Ltd."
    ws["A5"], ws["B5"] = "Ngày báo giá", "2026-10-16"
    ws["B5"].number_format = "dd/mm/yyyy"
    ws["A6"], ws["B6"] = "Hiệu lực đến", "15/11/2026"
    for row in (3, 4, 5, 6):
        ws.cell(row=row, column=1).font = Font(name="Arial", size=10, bold=True)
    header_row = 7
    for column, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row, column=column, value=title)
        cell.font = Font(name="Arial", size=10, bold=True, color="1F3864")
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30
    _goods_rows(ws, start_row=header_row + 1)
    total_row = header_row + len(ROWS) + 1
    ws.cell(row=total_row, column=5, value="TỔNG").font = Font(bold=True)
    ws.cell(row=total_row, column=6, value="=SUM(F8:F12)").font = Font(bold=True)
    ws.cell(row=total_row, column=7, value="=SUMPRODUCT(F8:F12,G8:G12)").font = Font(bold=True)
    ws.cell(row=total_row, column=7).number_format = '#,##0" RMB"'
    ws.cell(row=total_row, column=8, value="Tổng theo công thức nguồn").alignment = Alignment(wrap_text=True)
    terms_row = total_row + 2
    ws.cell(row=terms_row, column=1, value="ĐIỀU KIỆN & THỜI GIAN").font = Font(bold=True, color="1F3864")
    for offset, text in enumerate((
        "Thanh toán 80% trước khi giao hàng.",
        "Thời gian sản xuất dự kiến 30 ngày.",
        "Thời gian nhập khẩu dự kiến 14–20 ngày.",
        "Có bản vẽ thiết kế sau một tuần kể từ khi đặt cọc.",
    ), 1):
        ws.cell(row=terms_row + offset, column=1, value=text)
    percent_row = terms_row + 6
    ws.cell(row=percent_row, column=1, value="Tỷ lệ tạm ứng")
    ws.cell(row=percent_row, column=2, value=0.8).number_format = "0%"
    # A formula the workbook never cached: the viewer must report UNKNOWN, not 0.
    ws.cell(row=percent_row + 1, column=1, value="Dự phòng phát sinh (công thức chưa có giá trị cache)")
    ws.cell(row=percent_row + 1, column=2, value="=SUM(F8:F12)*500")
    ws.column_dimensions["A"].width = 34
    for column in "BCDEFGH":
        ws.column_dimensions[column].width = 22
    ws.freeze_panes = "A8"

    # ---------------------------------------------------------------- sheet 2
    internal = wb.create_sheet(SHEET_INTERNAL_DUP)
    internal["A1"] = "TÍNH NỘI BỘ — GIÁ MUA, KHÔNG GỬI KHÁCH"
    internal["A1"].font = WARN_FONT
    internal.append([])
    internal.append(list(HEADERS))
    for column in range(1, 9):
        cell = internal.cell(row=3, column=column)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
    _goods_rows(internal, price_delta=-1055, start_row=4)
    internal["A10"], internal["B10"] = "Vận chuyển EXW Nansha–Hải Phòng", 31800000
    internal["A11"], internal["B11"] = "Phí đại lý thanh toán/thủ tục", "2%"
    internal["A12"], internal["B12"] = "Thông quan Việt Nam", 3500000
    internal["B10"].number_format = '#,##0" VND"'
    internal.column_dimensions["A"].width = 38
    internal.column_dimensions["B"].width = 22

    # ---------------------------------------------------------------- sheet 3
    cost = wb.create_sheet(SHEET_INTERNAL_COST)
    cost["A1"] = "DỰ KIẾN CHI PHÍ — BẢNG TỔNG HỢP THAM KHẢO NỘI BỘ"
    cost["A1"].font = WARN_FONT
    cost.append([])
    cost.append(["Khoản mục", "Giá trị", "Ghi chú"])
    for column in range(1, 4):
        cost.cell(row=3, column=column).font = Font(bold=True)
        cost.cell(row=3, column=column).fill = HEAD_FILL
    summary = (
        ("Container", "20 feet", "Chưa xác nhận tháo rời"),
        ("CBM", 25, "Khác với 3 CBM ở sheet khách"),
        ("Trọng lượng", 3900, "kg"),
        ("RMB/USD", 6.78, "Tỷ giá cũ trong file"),
        ("USD/VND", 26500, "Tỷ giá cũ trong file"),
        ("HS tham khảo", "940330", "Chỉ tham khảo"),
        ("Thuế NK tham khảo", "5%", ""),
        ("VAT NK tham khảo", "8%", ""),
        ("LCC Hải Phòng", 12500000, ""),
        ("Giao nhận/nâng hạ kho lẻ", 3000000, ""),
        ("Vận chuyển Hải Phòng–Hà Nội", 5500000, ""),
        ("Nâng hạ/giao Thanh Trì nếu phát sinh", 12000000, "Có điều kiện"),
        ("Phí đại lý 2% phân bổ làm tròn", 47500000, "Chênh lệch làm tròn"),
    )
    for label, value, note in summary:
        cost.append([label, value, note])
    cost.append(["Chưa bao gồm VAT 8%", None, "Chưa rõ VAT dịch vụ hay VAT nhập khẩu"])
    for row in range(4, 4 + len(summary)):
        if isinstance(cost.cell(row=row, column=2).value, int):
            cost.cell(row=row, column=2).number_format = '#,##0'
    cost["B5"].number_format = "0.00"
    cost.column_dimensions["A"].width = 40
    cost.column_dimensions["B"].width = 18
    cost.column_dimensions["C"].width = 34

    output = io.BytesIO()
    wb.save(output)
    return _inject_cached_values(output.getvalue())


def _inject_cached_values(data: bytes) -> bytes:
    """Give the SOURCE formulas their real cached results, like Excel does.

    openpyxl cannot write a formula together with a cached value, so the sheet
    XML is patched after saving: `=SUM(F8:F12)` → cached 27 and
    `=SUMPRODUCT(F8:F12,G8:G12)` → cached 607180. The third formula in the
    fixture is deliberately LEFT UNCACHED so the UNKNOWN path is testable.
    """
    cached = {"F13": "27", "G13": "607180"}
    source = ZipFile(io.BytesIO(data))
    target = io.BytesIO()
    with ZipFile(target, "w", ZIP_DEFLATED) as out:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = content.decode("utf-8")
                for reference, value in cached.items():
                    # openpyxl writes `<c r="F13" s="9"><f>SUM(F8:F12)</f><v /></c>`;
                    # the cached result replaces the empty <v/>.
                    pattern = re.compile(r'(<c r="%s"[^>]*>)(<f>.*?</f>)(<v\s*/>|<v>.*?</v>)?(</c>)' % reference, re.S)
                    xml = pattern.sub(lambda m: f"{m.group(1)}{m.group(2)}<v>{value}</v>{m.group(4)}", xml, count=1)
                content = xml.encode("utf-8")
            out.writestr(item, content)
    return target.getvalue()


def write(target: Path) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build_luc_giac_three_sheet_fixture())
    return target


if __name__ == "__main__":
    destination = Path(sys.argv[1] if len(sys.argv) > 1 else ".tmp/luc-giac-3sheet-SYNTHETIC.xlsx")
    write(destination)
    print(f"SYNTHETIC_FIXTURE: {destination} ({destination.stat().st_size} bytes); "
          f"sheets={[SHEET_CUSTOMER, SHEET_INTERNAL_DUP, SHEET_INTERNAL_COST]}; "
          "expected 5 goods rows / 27 pcs / 607180 RMB")
