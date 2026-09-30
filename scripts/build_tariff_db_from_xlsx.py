#!/usr/bin/env python3
"""Build the fast 2026 tariff SQLite database from the governed XLSX source.

Only rows in sheet BT2026 that carry a 4-8 digit HS code are indexed. Runtime
never opens the XLSX: exact/prefix HS lookup uses a B-tree index and Vietnamese
text lookup uses FTS5. The source workbook hash is embedded in tariff_meta.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import lzma
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

COLUMNS: list[tuple[str, str]] = [
    ("_source_row", "INTEGER"),
    ("V", "TEXT"),
    ("Mã HS", "TEXT"),
    ("Mô tả tiếng Việt", "TEXT"),
    ("Đơn vị tính", "TEXT"),
    ("NK TT", "TEXT"),
    ("NK ưu đãi", "TEXT"),
    ("VAT", "TEXT"),
    ("ACFTA", "TEXT"),
    ("ATIGA", "TEXT"),
    ("AJCEP", "TEXT"),
    ("VJEPA", "TEXT"),
    ("AKFTA", "TEXT"),
    ("AANZFTA", "TEXT"),
    ("AIFTA", "TEXT"),
    ("VKFTA", "TEXT"),
    ("VCFTA", "TEXT"),
    ("VN-EAEU", "TEXT"),
    ("CPTPP", "TEXT"),
    ("AHKFTA", "TEXT"),
    ("VNCU", "TEXT"),
    ("EVFTA", "TEXT"),
    ("UKVFTA", "TEXT"),
    ("VN-LAO", "TEXT"),
    ("VN-CAM", "TEXT"),
    ("VIFTA", "TEXT"),
    ("RCEP A", "TEXT"),
    ("RCEP B", "TEXT"),
    ("RCEP C", "TEXT"),
    ("RCEP D", "TEXT"),
    ("RCEP E", "TEXT"),
    ("RCEP F", "TEXT"),
    ("TTĐB", "TEXT"),
    ("XK", "TEXT"),
    ("XK CPTPP", "TEXT"),
    ("XK EVFTA", "TEXT"),
    ("XK UKVFTA", "TEXT"),
    ("Thuế BVMT", "TEXT"),
    ("Chính sách mặt hàng", "TEXT"),
    ("Giảm VAT", "TEXT"),
]


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _hs_code(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    text = text.replace(".", "").replace(" ", "")
    if text.isdigit() and 4 <= len(text) <= 8:
        return text
    return None


def build(source: Path, destination: Path, *, sheet: str = "BT2026") -> int:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        destination.unlink()

    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    ws = workbook[sheet]

    rows: list[list[Any]] = []
    for source_row, row in enumerate(
        ws.iter_rows(min_row=4, max_col=39, values_only=True),
        start=4,
    ):
        code = _hs_code(row[1])
        if not code:
            continue
        values: list[Any] = [
            source_row,
            _clean(row[0]),
            code,
            _clean(row[2]),
            _clean(row[3]),
        ]
        values.extend(_clean(row[i]) for i in range(4, 25))
        # RCEP has six A-F subcolumns in XLSX columns 26-31.
        values.extend(_clean(row[i]) for i in range(25, 31))
        values.extend(_clean(row[i]) for i in range(31, 39))
        if len(values) != len(COLUMNS):
            raise RuntimeError(
                f"Unexpected BT2026 schema at row {source_row}: "
                f"{len(values)} values, expected {len(COLUMNS)}."
            )
        rows.append(values)

    with sqlite3.connect(destination) as conn:
        conn.executescript(
            """
            PRAGMA journal_mode=OFF;
            PRAGMA synchronous=OFF;
            PRAGMA temp_store=MEMORY;
            PRAGMA page_size=4096;
            """
        )
        definition = ",\n".join(f'"{name}" {kind}' for name, kind in COLUMNS)
        conn.execute(
            f"CREATE TABLE tariff_table (id INTEGER PRIMARY KEY, {definition})"
        )
        column_sql = ",".join(f'"{name}"' for name, _ in COLUMNS)
        placeholders = ",".join("?" for _ in COLUMNS)
        conn.executemany(
            f"INSERT INTO tariff_table ({column_sql}) VALUES ({placeholders})",
            rows,
        )
        conn.execute('CREATE INDEX idx_tariff_hs ON tariff_table("Mã HS")')
        conn.execute(
            'CREATE INDEX idx_tariff_hs_len '
            'ON tariff_table(length("Mã HS"), "Mã HS")'
        )

        conn.execute(
            """
            CREATE VIRTUAL TABLE tariff_fts USING fts5(
                description_vi,
                policy,
                vat_reduction,
                content='tariff_table',
                content_rowid='id',
                tokenize='unicode61 remove_diacritics 2'
            )
            """
        )
        conn.execute(
            """
            INSERT INTO tariff_fts(rowid, description_vi, policy, vat_reduction)
            SELECT
                id,
                coalesce("Mô tả tiếng Việt", ''),
                coalesce("Chính sách mặt hàng", ''),
                coalesce("Giảm VAT", '')
            FROM tariff_table
            """
        )
        conn.execute(
            "CREATE TABLE tariff_meta "
            "(key TEXT PRIMARY KEY, value TEXT NOT NULL)"
        )
        meta = {
            "source_filename": source.name,
            "source_sha256": source_sha,
            "source_sheet": sheet,
            "source_rows": str(ws.max_row),
            "indexed_rows": str(len(rows)),
            "year": "2026",
            "schema_version": "2",
            "search": "B-tree exact/prefix HS + FTS5 Vietnamese",
        }
        conn.executemany(
            "INSERT INTO tariff_meta(key, value) VALUES (?, ?)",
            list(meta.items()),
        )
        conn.execute("ANALYZE")
        conn.execute("PRAGMA optimize")
        conn.commit()

    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path, help="BIEU THUE XNK 2026.xlsx")
    parser.add_argument(
        "--database",
        type=Path,
        default=Path("tariff.db"),
        help="Output SQLite database.",
    )
    parser.add_argument(
        "--bundle",
        type=Path,
        default=None,
        help="Optional XZ-compressed normalized CSV for repository bundling.",
    )
    parser.add_argument(
        "--meta",
        type=Path,
        default=None,
        help="Optional JSON metadata sidecar for the bundled data.",
    )
    args = parser.parse_args()

    count = build(args.source, args.database)
    print(f"Built {count:,} indexed tariff rows: {args.database}")

    if args.bundle is not None:
        args.bundle.parent.mkdir(parents=True, exist_ok=True)
        with sqlite3.connect(args.database) as conn:
            columns = [
                row[1]
                for row in conn.execute("PRAGMA table_info(tariff_table)")
                if row[1] != "id"
            ]
            text = io.StringIO(newline="")
            writer = csv.writer(text, lineterminator="\n")
            writer.writerow(columns)
            select = ",".join(f'"{column}"' for column in columns)
            for row in conn.execute(f"SELECT {select} FROM tariff_table"):
                writer.writerow(["" if value is None else value for value in row])
        args.bundle.write_bytes(
            lzma.compress(text.getvalue().encode("utf-8"), preset=9)
        )
        print(f"Compressed normalized bundle: {args.bundle}")

    if args.meta is not None:
        args.meta.parent.mkdir(parents=True, exist_ok=True)
        meta = {
            "source_filename": args.source.name,
            "source_sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
            "source_sheet": "BT2026",
            "year": "2026",
            "indexed_rows": count,
            "schema_version": 2,
            "search": "B-tree exact/prefix HS + FTS5 Vietnamese",
            "bundle_sha256": (
                hashlib.sha256(args.bundle.read_bytes()).hexdigest()
                if args.bundle is not None and args.bundle.exists()
                else None
            ),
        }
        args.meta.write_text(
            json.dumps(meta, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"Metadata: {args.meta}")


if __name__ == "__main__":
    main()
