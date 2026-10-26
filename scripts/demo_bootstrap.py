#!/usr/bin/env python3
"""Bootstrap dữ liệu demo cục bộ cho một clone mới — không secret, không production.

Mục tiêu: clone sạch, chạy script này, ``streamlit run app.py`` mở được app với
dữ liệu mẫu trên transport staging in-memory. Script tái dùng các thành phần sẵn
có, không dựng hệ thống lưu trữ/thứ hai nào:

1. ``scripts/initialize_tariff_db.py`` tạo ``tariff.db`` rỗng đúng schema ứng dụng
   truy vấn (không nhúng dữ liệu biểu thuế — dữ liệu thật là tài sản ngoài git).
2. Sinh fixture PLACEHOLDER cho GNT / MA_NGAN_SACH / template DOCX, gắn nhãn rõ
   là bản mẫu không phải template pháp lý. Các placeholder này KHÔNG được gắn vào
   ``DOCGEN_TEMPLATE_DIR``: docgen giữ nguyên hành vi fail-closed
   (``TEMPLATE_ASSET_STORAGE_OWNER_GATE``) khi thiếu template thật.
3. Seed transport ``mongomock`` cô lập qua ``views.logi_sandbox`` +
   ``views.logi_staging_seed`` (cùng seeder app dùng cho Preview/Staging).
4. Từ chối chạy khi phát hiện ``MONGO_URI`` thật — không đụng dữ liệu production.

Chạy lại an toàn (idempotent): seeder bỏ qua document đã tồn tại, các file
placeholder được ghi đè với cùng nội dung.
"""
from __future__ import annotations

import argparse
import json
import sqlite3
import subprocess
import sys
from pathlib import Path
from typing import Any

BASE_DIR = Path(__file__).resolve().parents[1]
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

PLACEHOLDER_MARKER = "DEMO_PLACEHOLDER_V1 — KHONG PHAI TEMPLATE PHAP LY"
SANDBOX_PLACEHOLDER_URI = "mongodb://mongomock.local/logistics_staging"


class BootstrapError(RuntimeError):
    """Dùng cho mọi chặn an toàn của bootstrap."""


def _refuse_if_real_mongo() -> None:
    """Chốt chặn production: không seed/ghi gì khi đang trỏ MongoDB thật."""
    from core import config as core_config

    uri = str(getattr(core_config, "MONGO_URI", "") or "").strip()
    if uri and uri != SANDBOX_PLACEHOLDER_URI:
        raise BootstrapError(
            "Phát hiện MONGO_URI thật. demo_bootstrap chỉ chạy ở chế độ không có "
            "cấu hình production; hãy bỏ biến MONGO_URI/MONGO_DB_NAME trong shell "
            "và chạy lại. Không có thao tác nào được thực hiện."
        )


def bootstrap_tariff_db(destination: Path) -> Path:
    """Tạo tariff.db rỗng bằng chính script initializer sẵn có (một nguồn schema)."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    initializer = BASE_DIR / "scripts" / "initialize_tariff_db.py"
    subprocess.run(
        [sys.executable, str(initializer), "--database", str(destination)],
        cwd=BASE_DIR,
        check=True,
        capture_output=True,
        text=True,
    )
    with sqlite3.connect(destination) as conn:
        tables = {row[0] for row in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        )}
    if "tariff_table" not in tables:
        raise BootstrapError(f"initializer không tạo được tariff_table tại {destination}")
    return destination


def _write_placeholder_xlsx(path: Path, sheet_name: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DEMO_PLACEHOLDER"
    sheet["A1"] = PLACEHOLDER_MARKER
    sheet["A2"] = "File mẫu cho demo clone; template thật do Owner cấp ngoài git."
    sheet["A3"] = f"Sheet nghiệp vụ tham chiếu: {sheet_name}"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_placeholder_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("DEMO PLACEHOLDER — KHÔNG PHẢI TEMPLATE PHÁP LÝ", level=1)
    document.add_paragraph(PLACEHOLDER_MARKER)
    document.add_paragraph(
        "docgen vẫn fail-closed (TEMPLATE_ASSET_STORAGE_OWNER_GATE) cho đến khi "
        "Owner đặt template HDNT/PXK thật vào DOCGEN_TEMPLATE_DIR."
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def bootstrap_placeholders(templates_dir: Path) -> dict[str, str]:
    """Sinh các placeholder có gắn nhãn; tuyệt đối không tự nhận là template đủ điều kiện."""
    files = {
        "gnt_placeholder": templates_dir / "GNT_MAU.placeholder.xlsx",
        "treasury_placeholder": templates_dir / "MA_NGAN_SACH.placeholder.xlsx",
        "docx_placeholder": templates_dir / "docgen" / "HDNT-TT.placeholder.docx",
    }
    _write_placeholder_xlsx(files["gnt_placeholder"], "GNT")
    _write_placeholder_xlsx(files["treasury_placeholder"], "MA_NGAN_SACH")
    _write_placeholder_docx(files["docx_placeholder"])

    readme = templates_dir / "README.md"
    readme.write_text(
        "# Demo placeholders — KHÔNG phải template pháp lý\n\n"
        "Các file trong thư mục này do `scripts/demo_bootstrap.py` sinh ra để một\n"
        "bản clone mới có fixture mẫu. Chúng KHÔNG được nối vào\n"
        "`DOCGEN_TEMPLATE_DIR`; sinh chứng từ thật vẫn bị chặn fail-closed bởi\n"
        "`TEMPLATE_ASSET_STORAGE_OWNER_GATE` cho đến khi Owner đặt template thật\n"
        "(HDNT-TT-agent.docx, truong-thanh-agent.xlsx) và trỏ `DOCGEN_TEMPLATE_DIR`.\n\n"
        "Muốn thử đường dẫn GNT/MA_NGAN_SACH cục bộ, đặt biến môi trường\n"
        "`GNT_TEMPLATE_PATH` / `TREASURY_FILE_PATH` trỏ tới các file placeholder này.\n",
        encoding="utf-8",
    )
    def _display_path(path: Path) -> str:
        try:
            return str(path.relative_to(BASE_DIR))
        except ValueError:
            return str(path)

    manifest = {
        "marker": PLACEHOLDER_MARKER,
        "legal_templates": False,
        "docgen_owner_gate": "TEMPLATE_ASSET_STORAGE_OWNER_GATE",
        "wired_to_docgen_template_dir": False,
        "files": {key: _display_path(path) for key, path in files.items()},
    }
    (templates_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return {key: str(path) for key, path in files.items()}


def bootstrap_mongomock_seed() -> dict[str, Any]:
    """Cài transport mongomock cô lập và seed dữ liệu mẫu qua seeder chính chủ."""
    from views import logi_sandbox

    logi_sandbox.reset_backend()
    transport = logi_sandbox.ensure_backend(force_reseed=True)
    if transport != logi_sandbox.SANDBOX_TRANSPORT:
        raise BootstrapError(f"mongomock transport không hoạt động, nhận: {transport}")

    from core.database import get_db

    db = get_db()
    if db is None:
        raise BootstrapError("không lấy được staging database sau khi ensure_backend")
    counts = {
        name: db[name].count_documents({})
        for name in ("customers", "contracts", "quotes", "users", "tasks", "shipments")
    }
    return {"transport": transport, "counts": counts}


def bootstrap(workdir: Path) -> dict[str, Any]:
    """Thực hiện toàn bộ các bước bootstrap cục bộ, trả về báo cáo cấu trúc."""
    _refuse_if_real_mongo()
    workdir.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {"workdir": str(workdir)}
    report["tariff_db"] = str(
        bootstrap_tariff_db(workdir / "data" / "tariff.db")
    )
    report["placeholders"] = bootstrap_placeholders(workdir / "templates")
    report["mongomock"] = bootstrap_mongomock_seed()

    from core import config as core_config

    report["docgen_template_dir"] = str(core_config.DOCGEN_TEMPLATE_DIR) if core_config.DOCGEN_TEMPLATE_DIR else None
    report["docgen_fail_closed_preserved"] = core_config.DOCGEN_TEMPLATE_DIR is None
    return report


def _print_instructions(report: dict[str, Any]) -> None:
    print("=" * 72)
    print("DEMO BOOTSTRAP HOÀN TẤT — không secret, không đụng production")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print("-" * 72)
    print("Bước tiếp theo:")
    print("  1. streamlit run app.py   # mở app với dữ liệu staging mongomock")
    print("  2. Phiên dev đơn lẻ có thể tắt auth: AUTH_REQUIRED=false streamlit run app.py")
    print("  3. Template chứng từ pháp lý vẫn cần Owner gate")
    print("     (TEMPLATE_ASSET_STORAGE_OWNER_GATE); xem demo_bootstrap/templates/README.md")
    print("=" * 72)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Bootstrap dữ liệu demo cục bộ cho clone mới.")
    parser.add_argument(
        "--workdir",
        type=Path,
        default=BASE_DIR / "demo_bootstrap",
        help="Thư mục chứa artifact demo (mặc định: ./demo_bootstrap, đã gitignore).",
    )
    args = parser.parse_args(argv)
    report = bootstrap(args.workdir.resolve())
    _print_instructions(report)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except BootstrapError as exc:
        print(f"BOOTSTRAP BLOCKED: {exc}", file=sys.stderr)
        raise SystemExit(2)
