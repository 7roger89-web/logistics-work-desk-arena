"""Generate SYNTHETIC acceptance XLSX; never claims to be the original file.

Usage: .venv/bin/python scripts/make_luc_giac_quote_fixture.py .tmp/luc-giac-fixture.xlsx
No customer uploads/binary fixture needs to be committed.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROWS = [(5200, 1500, 760, 16, 13055), (7500, 1600, 760, 1, 18850),
        (12000, 1600, 760, 3, 30150), (16000, 1600, 760, 6, 40200), (19000, 1600, 760, 1, 47800)]


def build_luc_giac_fixture() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Khach gui"
    ws.merge_cells("A1:H1")
    ws["A1"] = "FIXTURE MÔ PHỎNG — KHÔNG PHẢI WORKBOOK GỐC LỤC GIÁC"
    ws["A2"], ws["B2"] = "Khách hàng", "Công ty Lục Giác"
    ws["A3"], ws["B3"] = "Nhà cung cấp", "Guangdong Hongye Furniture Group Co., Ltd."
    headers = ["STT", "Tên hàng", "Model", "Vật liệu", "Kích thước (mm)", "Số lượng", "Đơn giá (RMB)", "Ghi chú"]
    ws.append([])
    ws.append([])
    ws.append(headers)
    for i, (length, width, height, qty, unit) in enumerate(ROWS, 1):
        ws.append([i, "Bàn họp văn phòng", "MS-C0242S" if i == 1 else None, "MDF + veneer" if i == 1 else None,
                   f"{length} × {width} × {height}", qty, unit, "Chưa xác nhận model/vật liệu" if i > 1 else None])
    ws.append(["TỔNG", None, None, None, None, 27, 607180])
    for text in ("Thanh toán 80% trước khi giao hàng.", "Thời gian sản xuất dự kiến 30 ngày.",
                 "Thời gian nhập khẩu dự kiến 14–20 ngày.", "Có bản vẽ thiết kế sau một tuần kể từ khi đặt cọc."):
        ws.append([text])
    internal = wb.create_sheet("Tinh noi bo")
    internal.append(["FIXTURE MÔ PHỎNG — DỮ LIỆU CŨ, KHÔNG PHẢI GIÁ HIỆN HÀNH"])
    internal.append(headers)
    for i, (length, width, height, qty, unit) in enumerate(ROWS, 1):
        internal.append([i, "Bàn họp văn phòng", None, None,
                         f"{1200 if length == 12000 else length} × {7600 if i == 2 else width} × {height}", qty, unit])
    for label, value in [("Container", "20 feet"), ("CBM", 25), ("Trọng lượng", 3900),
                         ("RMB/USD", 6.78), ("USD/VND", 26500), ("HS tham khảo", "940330"),
                         ("Thuế NK tham khảo", "5%"), ("VAT NK tham khảo", "8%"), ("Hiệu lực đến", "31/07/2026"),
                         ("Vận chuyển EXW Nansha–Hải Phòng", 31800000), ("Phí đại lý thanh toán/thủ tục", "2%"),
                         ("Thông quan Việt Nam", 3500000), ("LCC Hải Phòng", 12500000),
                         ("Giao nhận/nâng hạ kho lẻ", 3000000), ("Vận chuyển Hải Phòng–Hà Nội", 5500000),
                         ("Nâng hạ/giao Thanh Trì nếu phát sinh", 12000000)]:
        internal.append([label, value])
    summary = wb.create_sheet("Tong hop")
    for row in [("FIXTURE MÔ PHỎNG",), ("CBM", 3), ("Trọng lượng", 400),
                ("Phí đại lý 2% phân bổ làm tròn", 47500000), ("Chưa bao gồm VAT 8%",),
                ("Dòng thiếu công thức", "=#REF!")]:
        summary.append(row)
    for sheet in wb:
        sheet.freeze_panes = "F7" if sheet == ws else "A2"
        for col in "ABCDEFGH":
            sheet.column_dimensions[col].width = 24
        sheet.column_dimensions["A"].width = 38
        sheet["A1"].font = Font(bold=True, color="9C0006")
    for cell in ws[6]:
        cell.fill = PatternFill("solid", fgColor="E4EBF7")
        cell.font = Font(bold=True)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else ".tmp/luc-giac-fixture.xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build_luc_giac_fixture())
    print(f"SYNTHETIC_FIXTURE: {target} ({target.stat().st_size} bytes); expected 5 rows / 27 pcs / 607180 RMB")
