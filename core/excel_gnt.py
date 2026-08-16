# FILE: core/excel_gnt.py
import io
import os
import openpyxl
from openpyxl.styles import Border, Side, Alignment, Font
from copy import copy
import streamlit as st
from core.config import GNT_TEMPLATE_PATH
from core.utils import safe_float

TEMPLATE_FILE = str(GNT_TEMPLATE_PATH)


def parse_tax_amount_vnd(value):
    """Chuẩn hóa số tiền GNT về số nguyên VNĐ, không làm sai float có đuôi .0."""
    return int(round(safe_float(value)))


# --- 1. HÀM ĐỌC SỐ TIỀN THUẦN VIỆT (KHÔNG CẦN THƯ VIỆN NGOÀI) ---
def doc_so_thanh_chu_vi(n):
    if n == 0: return "Không đồng"
    
    chu_so = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
    tien_to = ["", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ"]
    
    def doc_ba_so(so):
        tram = so // 100
        chuc = (so % 100) // 10
        don_vi = so % 10
        kq = ""
        
        if tram != 0:
            kq += chu_so[tram] + " trăm "
            if chuc == 0 and don_vi != 0: kq += "linh "
        
        if chuc != 0:
            if chuc == 1: kq += "mười "
            else: kq += chu_so[chuc] + " mươi "
            
        if don_vi != 0:
            if chuc != 0 and chuc != 1:
                if don_vi == 1: kq += "mốt"
                elif don_vi == 5: kq += "lăm"
                else: kq += chu_so[don_vi]
            else:
                kq += chu_so[don_vi]
        return kq

    # Xử lý số lớn
    n = int(n)
    parts = []
    while n > 0:
        parts.append(n % 1000)
        n //= 1000
        
    kq_parts = []
    for i, part in enumerate(parts):
        if part == 0: continue
        doan = doc_ba_so(part).strip()
        tien = tien_to[i]
        kq_parts.append(f"{doan} {tien}".strip())
        
    final_text = " ".join(reversed(kq_parts)).strip()
    
    # Capitalize chữ cái đầu
    final_text = final_text[0].upper() + final_text[1:]
    return final_text + " đồng./."

def find_row_by_text(sheet, text_to_find, col_idx=None):
    """Tìm ô chứa text"""
    for row in sheet.iter_rows():
        for cell in row:
            if col_idx and cell.column != col_idx: continue
            if cell.value and isinstance(cell.value, str) and text_to_find.lower() in cell.value.lower():
                return cell
    return None

def copy_style(source_cell, target_cell):
    """Copy định dạng từ ô này sang ô kia"""
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = copy(source_cell.number_format)
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)

def create_gnt_excel(data):
    """Tạo file GNT từ template"""
    if not os.path.exists(TEMPLATE_FILE):
        st.error(f"❌ Không tìm thấy file mẫu '{TEMPLATE_FILE}'. Vui lòng upload file mẫu vào thư mục gốc.")
        return None

    try:
        wb = openpyxl.load_workbook(TEMPLATE_FILE)
        ws = wb.active 
        
        # --- 1. ĐIỀN HEADER ---
        fill_map = {
            "Người nộp thuế": data.get('nguoi_nop', ''),
            "Mã số thuế": data.get('mst', ''),
            "Địa chỉ": data.get('dia_chi', ''),
            "Vào tài khoản của KBNN": data.get('ten_kbnn', ''),
            "Tên cơ quan quản lý thu": data.get('ten_cq_thu', ''),
            "Mã cơ quan quản lý thu": data.get('ma_cq_thu', '') # Thêm nếu cần
        }

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for key, val in fill_map.items():
                        # Tìm chính xác cụm từ khóa (VD: "Mã số thuế:")
                        if key in cell.value:
                            # Tách phần tiêu đề tĩnh và điền giá trị mới
                            parts = cell.value.split(":")
                            prefix = parts[0] + ": "
                            cell.value = prefix + str(val)

        # --- 2. CHECKBOX TÀI KHOẢN ---
        tk_cell = find_row_by_text(ws, "TK thu NSNN")
        if tk_cell:
            txt = tk_cell.value.replace("[x]", "[ ]").replace("[X]", "[ ]")
            tk_thu = str(data.get('tk_thu', '7111'))
            
            if tk_thu == '7111':
                txt = txt.replace("TK thu NSNN [ ]", "TK thu NSNN [x]")
            elif tk_thu == '3511':
                txt = txt.replace("TK tạm thu [ ]", "TK tạm thu [x]")
            tk_cell.value = txt

        # --- 3. BẢNG THUẾ (DYNAMIC) ---
        header_cell = find_row_by_text(ws, "Nội dung các khoản nộp NSNN")
        if not header_cell:
            st.error("❌ Template sai định dạng: Không tìm thấy cột 'Nội dung các khoản nộp NSNN'")
            return None
            
        start_row = header_cell.row + 1
        taxes = data.get('ds_thue', [])
        if not taxes: taxes = [{'ten': '', 'tien': 0, 'ndkt': ''}]

        # Chèn dòng nếu nhiều hơn 1 khoản thuế
        if len(taxes) > 1:
            ws.insert_rows(start_row + 1, amount=len(taxes) - 1)

        real_total_money = 0
        
        for i, tax in enumerate(taxes):
            curr_r = start_row + i
            
            # Giá trị từ Streamlit thường là float (VD: 1000000.0); không được
            # xóa dấu chấm theo kiểu chuỗi vì sẽ làm số tiền tăng gấp 10 lần.
            val = parse_tax_amount_vnd(tax.get('tien', 0))
            real_total_money += val

            # Ghi dữ liệu & Định dạng
            # Cột A (1): STT
            c1 = ws.cell(row=curr_r, column=1, value=i + 1)
            c1.alignment = Alignment(horizontal='center', vertical='center')
            
            # Cột B (2): Số tờ khai
            c2 = ws.cell(row=curr_r, column=2, value=data.get('so_tk', ''))
            c2.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            # Cột C (3): Kỳ thuế (Ngày TK)
            c3 = ws.cell(row=curr_r, column=3, value=data.get('ngay_tk', ''))
            c3.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            # Cột D (4): Nội dung
            c4 = ws.cell(row=curr_r, column=4, value=tax.get('ten', ''))
            c4.alignment = Alignment(vertical='center', wrap_text=True)
            
            # Cột F (6): Số tiền
            c6 = ws.cell(row=curr_r, column=6, value=val)
            c6.number_format = '#,##0'
            c6.alignment = Alignment(horizontal='right', vertical='center')
            
            # Cột H (8): Mã NDKT
            c8 = ws.cell(row=curr_r, column=8, value=tax.get('ndkt', ''))
            c8.alignment = Alignment(horizontal='center', vertical='center')

            # Kẻ khung cho dòng hiện tại (Từ cột 1 -> 9)
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            for col in range(1, 10):
                ws.cell(row=curr_r, column=col).border = thin_border

        # --- 4. TỔNG TIỀN & CHỮ ---
        # Dòng tổng cộng nằm ngay sau dòng thuế cuối cùng
        total_row_idx = start_row + len(taxes)
        
        # Ghi công thức SUM (Excel tự tính khi mở file)
        # =SUM(F{start_row}:F{start_row + len - 1})
        sum_formula = f"=SUM(F{start_row}:F{total_row_idx - 1})"
        
        c_total = ws.cell(row=total_row_idx, column=6)
        c_total.value = real_total_money # Ghi giá trị tĩnh trước phòng khi lỗi công thức
        # Nếu muốn dùng công thức: c_total.value = sum_formula 
        # (Lưu ý: Openpyxl ghi công thức nhưng cần phần mềm Excel mở ra để tính lại)
        
        c_total.number_format = '#,##0'
        c_total.font = Font(bold=True)

        # Ghi bằng chữ
        text_cell = find_row_by_text(ws, "Tổng số tiền ghi bằng chữ")
        if text_cell:
            text_vn = doc_so_thanh_chu_vi(real_total_money)
            text_cell.value = f"Tổng số tiền ghi bằng chữ: {text_vn}"

        # --- 5. XUẤT ---
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    except Exception as e:
        st.error(f"Lỗi tạo file GNT: {e}")
        return None