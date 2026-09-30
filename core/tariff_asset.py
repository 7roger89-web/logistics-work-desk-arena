"""Build the tariff search SQLite asset from an Admin-supplied XLSX.

The request path never reads Excel. This module is called only when ADMIN
replaces the governed tariff workbook: it extracts BT20xx rows, builds SQLite
B-tree + FTS5 indexes, compresses the resulting DB, and returns the bytes for
durable Mongo storage.
"""
from __future__ import annotations

import hashlib
import io
import lzma
import os
import re
import sqlite3
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MAX_SOURCE_BYTES = 12 * 1024 * 1024
MAX_INDEXED_ROWS = 100_000

COLUMNS: list[tuple[str, str]] = [
    ("_source_row", "INTEGER"),
    ("V", "TEXT"),
    ("Mã HS", "TEXT"),
    ("Mô tả tiếng Việt", "TEXT"),
    ("Đơn vị tính", "TEXT"),
    ("NK TT", "TEXT"),
    ("NK ưu đãi", "TEXT"),
    ("VAT", "TEXT"),
    ("ACFTA", "TEXT"),
    ("ATIGA", "TEXT"),
    ("AJCEP", "TEXT"),
    ("VJEPA", "TEXT"),
    ("AKFTA", "TEXT"),
    ("AANZFTA", "TEXT"),
    ("AIFTA", "TEXT"),
    ("VKFTA", "TEXT"),
    ("VCFTA", "TEXT"),
    ("VN-EAEU", "TEXT"),
    ("CPTPP", "TEXT"),
    ("AHKFTA", "TEXT"),
    ("VNCU", "TEXT"),
    ("EVFTA", "TEXT"),
    ("UKVFTA", "TEXT"),
    ("VN-LAO", "TEXT"),
    ("VN-CAM", "TEXT"),
    ("VIFTA", "TEXT"),
    ("RCEP A", "TEXT"),
    ("RCEP B", "TEXT"),
    ("RCEP C", "TEXT"),
    ("RCEP D", "TEXT"),
    ("RCEP E", "TEXT"),
    ("RCEP F", "TEXT"),
    ("TTĐB", "TEXT"),
    ("XK", "TEXT"),
    ("XK CPTPP", "TEXT"),
    ("XK EVFTA", "TEXT"),
    ("XK UKVFTA", "TEXT"),
    ("Thuế BVMT", "TEXT"),
    ("Chính sách mặt hàng", "TEXT"),
    ("Giảm VAT", "TEXT"),
]


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _hs_code(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    text = text.replace(".", "").replace(" ", "")
    if text.isdigit() and 4 <= len(text) <= 8:
        return text
    return None


def _pick_sheet(workbook: Any) -> tuple[str, str | None]:
    candidates = [
        name for name in workbook.sheetnames
        if re.fullmatch(r"BT(?:20)?\d{2}", str(name).strip(), flags=re.IGNORECASE)
    ]
    if "BT2026" in workbook.sheetnames:
        name = "BT2026"
    elif candidates:
        name = candidates[0]
    else:
        raise ValueError(
            "Không tìm thấy sheet biểu thuế dạng BT2026/BT20xx trong file."
        )
    match = re.search(r"(20\d{2})", name)
    return name, match.group(1) if match else None


def _source_values(row: tuple[Any, ...], source_row: int) -> list[Any] | None:
    code = _hs_code(row[1] if len(row) > 1 else None)
    if not code:
        return None

    padded = list(row[:39]) + [None] * max(0, 39 - len(row))
    values: list[Any] = [
        source_row,
        _clean(padded[0]),
        code,
        _clean(padded[2]),
        _clean(padded[3]),
    ]
    values.extend(_clean(padded[i]) for i in range(4, 25))
    values.extend(_clean(padded[i]) for i in range(25, 31))
    values.extend(_clean(padded[i]) for i in range(31, 39))
    if len(values) != len(COLUMNS):
        raise RuntimeError(
            f"Schema biểu thuế không khớp tại dòng {source_row}: "
            f"{len(values)} cột, cần {len(COLUMNS)}."
        )
    return values


def build_tariff_asset(
    source_bytes: bytes,
    *,
    source_filename: str,
) -> dict[str, Any]:
    raw = bytes(source_bytes or b"")
    if not raw:
        raise ValueError("File Excel rỗng.")
    if len(raw) > MAX_SOURCE_BYTES:
        raise ValueError("File Excel vượt giới hạn 12 MB.")
    if not str(source_filename or "").lower().endswith(".xlsx"):
        raise ValueError("BIỂU THUẾ chỉ nhận file .xlsx.")

    source_sha = hashlib.sha256(raw).hexdigest()
    workbook = load_workbook(
        io.BytesIO(raw),
        read_only=True,
        data_only=True,
    )
    sheet_name, year = _pick_sheet(workbook)
    ws = workbook[sheet_name]

    fd, temp_name = tempfile.mkstemp(prefix="tariff-build-", suffix=".db")
    os.close(fd)
    db_path = Path(temp_name)
    indexed_rows = 0

    try:
        with sqlite3.connect(db_path) as conn:
            conn.executescript(
                """
                PRAGMA journal_mode=OFF;
                PRAGMA synchronous=OFF;
                PRAGMA temp_store=MEMORY;
                PRAGMA page_size=4096;
                """
            )
            definition = ",\n".join(
                f'"{name}" {kind}' for name, kind in COLUMNS
            )
            conn.execute(
                f"CREATE TABLE tariff_table "
                f"(id INTEGER PRIMARY KEY, {definition})"
            )
            column_sql = ",".join(f'"{name}"' for name, _ in COLUMNS)
            placeholders = ",".join("?" for _ in COLUMNS)
            insert_sql = (
                f"INSERT INTO tariff_table ({column_sql}) "
                f"VALUES ({placeholders})"
            )

            batch: list[list[Any]] = []
            for source_row, row in enumerate(
                ws.iter_rows(min_row=4, max_col=39, values_only=True),
                start=4,
            ):
                values = _source_values(tuple(row), source_row)
                if values is None:
                    continue
                indexed_rows += 1
                if indexed_rows > MAX_INDEXED_ROWS:
                    raise ValueError(
                        f"File có hơn {MAX_INDEXED_ROWS:,} dòng mã HS; "
                        "vượt giới hạn an toàn."
                    )
                batch.append(values)
                if len(batch) >= 1000:
                    conn.executemany(insert_sql, batch)
                    batch.clear()
            if batch:
                conn.executemany(insert_sql, batch)

            if indexed_rows == 0:
                raise ValueError(
                    "Không tìm thấy dòng mã HS hợp lệ trong sheet biểu thuế."
                )

            conn.execute(
                'CREATE INDEX idx_tariff_hs ON tariff_table("Mã HS")'
            )
            conn.execute(
                'CREATE INDEX idx_tariff_hs_len '
                'ON tariff_table(length("Mã HS"), "Mã HS")'
            )
            conn.execute(
                """
                CREATE VIRTUAL TABLE tariff_fts USING fts5(
                    description_vi,
                    policy,
                    vat_reduction,
                    content='tariff_table',
                    content_rowid='id',
                    tokenize='unicode61 remove_diacritics 2'
                )
                """
            )
            conn.execute(
                """
                INSERT INTO tariff_fts(
                    rowid, description_vi, policy, vat_reduction
                )
                SELECT
                    id,
                    coalesce("Mô tả tiếng Việt", ''),
                    coalesce("Chính sách mặt hàng", ''),
                    coalesce("Giảm VAT", '')
                FROM tariff_table
                """
            )

            metadata: dict[str, Any] = {
                "source_filename": str(source_filename),
                "source_sha256": source_sha,
                "source_sheet": sheet_name,
                "source_rows": int(ws.max_row or 0),
                "indexed_rows": indexed_rows,
                "year": year,
                "schema_version": 2,
                "search": "B-tree exact/prefix HS + FTS5 Vietnamese",
            }
            conn.execute(
                "CREATE TABLE tariff_meta "
                "(key TEXT PRIMARY KEY, value TEXT NOT NULL)"
            )
            conn.executemany(
                "INSERT INTO tariff_meta(key, value) VALUES (?, ?)",
                [
                    (str(key), str(value))
                    for key, value in metadata.items()
                    if value is not None
                ],
            )
            conn.execute("ANALYZE")
            conn.execute("PRAGMA optimize")
            conn.commit()

        db_bytes = db_path.read_bytes()
        db_sha = hashlib.sha256(db_bytes).hexdigest()
        compressed = lzma.compress(db_bytes, preset=9)
        metadata["db_bytes"] = len(db_bytes)
        metadata["compressed_bytes"] = len(compressed)
        metadata["db_sha256"] = db_sha
        return {
            "payload_lzma": compressed,
            "db_sha256": db_sha,
            "metadata": metadata,
        }
    finally:
        try:
            workbook.close()
        except Exception:
            pass
        try:
            db_path.unlink(missing_ok=True)
        except OSError:
            pass
