"""Bounded real-file text layer for the trade-document canary.

Run 2 proves only the smallest container-reading seam:
real synthetic XLSX/PDF bytes -> deterministic located text -> existing
``core.document_extraction.extract_document(..., adapter=...)`` and
``core.trade_document_crosscheck``.

This module is deliberately *not* a document platform.  It reads visible XLSX
cell text with the existing spreadsheet stack and born-digital PDF text with a
single PDF parser.  It never OCRs, never calls a provider, never follows links
or actions, and never evaluates spreadsheet formulas.
"""

from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass, field
from enum import Enum
from importlib import metadata
from io import BytesIO
from typing import Any, Mapping

from core.document_extraction import (
    DocumentArtifact,
    DocumentSecurityError,
    MAX_PAGE_COUNT,
    MAX_PAYLOAD_BYTES,
    MAX_TEXT_CHARS,
    sanitize_filename,
)
from core.trade_compliance import DocumentType

TEXT_NOT_AVAILABLE_NO_OCR = "TEXT_NOT_AVAILABLE_NO_OCR"
XLSX_FORMULA_UNSUPPORTED = "XLSX_FORMULA_UNSUPPORTED"

MAX_XLSX_CELLS = 50_000
MAX_PDF_CONTENT_STREAM_BYTES = 2 * 1024 * 1024


class LocatedTextFileType(str, Enum):
    XLSX = "xlsx"
    PDF = "pdf"


@dataclass(frozen=True)
class TextReaderLimits:
    """Small explicit bounds for untrusted real-file containers."""

    max_file_bytes: int = MAX_PAYLOAD_BYTES
    max_text_chars: int = MAX_TEXT_CHARS
    max_pdf_pages: int = MAX_PAGE_COUNT
    max_pdf_content_stream_bytes: int = MAX_PDF_CONTENT_STREAM_BYTES
    max_xlsx_cells: int = MAX_XLSX_CELLS


@dataclass(frozen=True)
class LocatedTextFragment:
    """One extracted text fragment with its real source location."""

    text: str
    page_number: int | None = None
    sheet_name: str | None = None
    cell_coordinate: str | None = None
    state: str = "OK"

    def location(self) -> dict[str, Any]:
        payload: dict[str, Any] = {"state": self.state}
        if self.page_number is not None:
            payload["page_number"] = self.page_number
        if self.sheet_name is not None:
            payload["sheet_name"] = self.sheet_name
        if self.cell_coordinate is not None:
            payload["cell_coordinate"] = self.cell_coordinate
        return payload


@dataclass(frozen=True)
class LocatedTextDocument:
    """Located text output consumed by the existing extraction/cross-check core."""

    source_filename: str
    source_sha256: str
    file_type: LocatedTextFileType
    reader_method: str
    reader_version: str
    fragments: tuple[LocatedTextFragment, ...]
    pages: tuple[int, ...] = ()
    sheets: tuple[str, ...] = ()
    warnings: tuple[str, ...] = ()
    unsupported_state: str | None = None
    text: str = ""
    line_locations: Mapping[int, Mapping[str, Any]] = field(default_factory=dict)

    def to_document_artifact(self, document_type: DocumentType | str | None = None) -> DocumentArtifact:
        """Adapt located text into the existing LOG-DOC01 artifact seam.

        ``content`` is the deterministic text layer, while metadata carries the
        raw container SHA-256 so Run 2 provenance can use the actual XLSX/PDF
        bytes rather than the derived text hash.  Run 1 callers without this
        metadata keep their original ``content_sha256`` behavior.
        """

        return DocumentArtifact(
            filename=self.source_filename,
            content=self.text,
            document_type=document_type,
            page_count=len(self.pages) or None,
            metadata={
                "source_sha256": self.source_sha256,
                "source_file_type": self.file_type.value,
                "reader_method": self.reader_method,
                "reader_version": self.reader_version,
                "warnings": list(self.warnings),
                "unsupported_state": self.unsupported_state,
                "line_locations": {int(key): dict(value) for key, value in self.line_locations.items()},
                "pages": list(self.pages),
                "sheets": list(self.sheets),
            },
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_filename": self.source_filename,
            "source_sha256": self.source_sha256,
            "file_type": self.file_type.value,
            "reader_method": self.reader_method,
            "reader_version": self.reader_version,
            "pages": list(self.pages),
            "sheets": list(self.sheets),
            "fragments": [
                {"text": fragment.text, "location": fragment.location()}
                for fragment in self.fragments
            ],
            "warnings": list(self.warnings),
            "unsupported_state": self.unsupported_state,
            "text": self.text,
        }


def _extension(filename: str) -> str:
    safe_name = sanitize_filename(filename)
    return "." + safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""


def _reader_version(package: str) -> str:
    try:
        return metadata.version(package)
    except metadata.PackageNotFoundError:  # pragma: no cover - environment guard
        return "unknown"


def _check_common_bounds(data: bytes, limits: TextReaderLimits) -> None:
    if not isinstance(data, bytes):
        raise DocumentSecurityError("document bytes are required")
    if len(data) > limits.max_file_bytes:
        raise DocumentSecurityError("document exceeds the bounded size limit")


def _cell_value_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _contains_vba_project(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            return any(name.lower().endswith("vbaproject.bin") for name in archive.namelist())
    except zipfile.BadZipFile as exc:
        raise DocumentSecurityError("xlsx container is corrupt or unreadable") from exc


def _read_xlsx_text(data: bytes, filename: str, limits: TextReaderLimits) -> LocatedTextDocument:
    extension = _extension(filename)
    if extension == ".xlsm":
        raise DocumentSecurityError("macro-enabled XLSM is rejected for this canary")
    if extension != ".xlsx":
        raise DocumentSecurityError("xlsx reader only accepts .xlsx files")
    if _contains_vba_project(data):
        raise DocumentSecurityError("workbook macro project content is rejected")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentSecurityError("openpyxl is required for xlsx text reading") from exc

    reader_version = _reader_version("openpyxl")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # fail closed on untrusted containers
        raise DocumentSecurityError("xlsx container is corrupt or unreadable") from exc

    fragments: list[LocatedTextFragment] = []
    warnings: list[str] = []
    lines: list[str] = []
    line_locations: dict[int, dict[str, Any]] = {}
    visible_sheets: list[str] = []
    cell_count = 0
    formula_seen = False
    marker_seen_by_sheet: dict[str, bool] = {}

    for worksheet in workbook.worksheets:
        if worksheet.sheet_state != "visible":
            continue
        visible_sheets.append(worksheet.title)
        marker_seen_by_sheet[worksheet.title] = False
        lines.append(f"[SHEET {worksheet.title}]")
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cell_count += 1
                if cell_count > limits.max_xlsx_cells:
                    raise DocumentSecurityError("xlsx cell count exceeds the bounded limit")
                text = _cell_value_to_text(cell.value)
                if not text:
                    continue
                state = "OK"
                if text.startswith("="):
                    formula_seen = True
                    state = XLSX_FORMULA_UNSUPPORTED
                    warnings.append(f"{XLSX_FORMULA_UNSUPPORTED}:{worksheet.title}!{cell.coordinate}")
                if "SYNTHETIC — NOT A REAL SHIPMENT" in text:
                    marker_seen_by_sheet[worksheet.title] = True
                line_number = len(lines) + 1
                lines.append(text)
                line_locations[line_number] = {
                    "file_type": LocatedTextFileType.XLSX.value,
                    "sheet_name": worksheet.title,
                    "cell_coordinate": cell.coordinate,
                    "reader_method": "openpyxl.load_workbook(data_only=False)",
                    "reader_version": reader_version,
                    "state": state,
                }
                fragments.append(
                    LocatedTextFragment(
                        text=text,
                        sheet_name=worksheet.title,
                        cell_coordinate=cell.coordinate,
                        state=state,
                    )
                )

    for sheet_name, marker_seen in marker_seen_by_sheet.items():
        if not marker_seen:
            warnings.append(f"SYNTHETIC_MARKER_NOT_FOUND:{sheet_name}")

    text = "\n".join(lines)
    if len(text) > limits.max_text_chars:
        raise DocumentSecurityError("xlsx text exceeds the bounded size limit")

    return LocatedTextDocument(
        source_filename=sanitize_filename(filename),
        source_sha256=hashlib.sha256(data).hexdigest(),
        file_type=LocatedTextFileType.XLSX,
        reader_method="openpyxl.load_workbook(data_only=False)",
        reader_version=reader_version,
        fragments=tuple(fragments),
        pages=(),
        sheets=tuple(visible_sheets),
        warnings=tuple(dict.fromkeys(warnings)),
        unsupported_state="UNSUPPORTED_FORMULA_VALUES_PRESENT" if formula_seen else None,
        text=text,
        line_locations=line_locations,
    )


def _pdf_content_stream_size(page: Any) -> int:
    contents = page.get_contents()
    if contents is None:
        return 0
    if isinstance(contents, list):
        return sum(len(item.get_data()) for item in contents)
    return len(contents.get_data())


def _read_pdf_text(data: bytes, filename: str, limits: TextReaderLimits) -> LocatedTextDocument:
    if _extension(filename) != ".pdf":
        raise DocumentSecurityError("pdf reader only accepts .pdf files")

    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentSecurityError("pypdf is required for pdf text reading") from exc

    reader_version = _reader_version("pypdf")
    try:
        reader = PdfReader(BytesIO(data), strict=True)
    except Exception as exc:
        raise DocumentSecurityError("pdf container is corrupt or unreadable") from exc

    if reader.is_encrypted:
        raise DocumentSecurityError("encrypted pdf is rejected for this canary")
    page_count = len(reader.pages)
    if page_count > limits.max_pdf_pages:
        raise DocumentSecurityError("pdf page count exceeds the bounded limit")

    fragments: list[LocatedTextFragment] = []
    warnings: list[str] = []
    lines: list[str] = []
    line_locations: dict[int, dict[str, Any]] = {}
    pages: list[int] = []

    for index, page in enumerate(reader.pages, start=1):
        pages.append(index)
        try:
            stream_size = _pdf_content_stream_size(page)
        except Exception as exc:
            raise DocumentSecurityError("pdf content stream is unreadable") from exc
        if stream_size > limits.max_pdf_content_stream_bytes:
            raise DocumentSecurityError("pdf content stream exceeds the bounded limit")
        lines.append(f"[PAGE {index}]")
        try:
            extracted = page.extract_text() or ""
        except Exception as exc:
            raise DocumentSecurityError("pdf text extraction failed") from exc
        page_lines = [line.strip() for line in extracted.splitlines() if line.strip()]
        if not page_lines:
            warnings.append(f"PAGE_TEXT_NOT_AVAILABLE_NO_OCR:{index}")
            continue
        for page_line in page_lines:
            line_number = len(lines) + 1
            lines.append(page_line)
            line_locations[line_number] = {
                "file_type": LocatedTextFileType.PDF.value,
                "page_number": index,
                "reader_method": "pypdf.PdfReader.page.extract_text",
                "reader_version": reader_version,
                "state": "OK",
            }
            fragments.append(LocatedTextFragment(text=page_line, page_number=index))

    text = "\n".join(lines)
    if len(text) > limits.max_text_chars:
        raise DocumentSecurityError("pdf text exceeds the bounded size limit")
    unsupported_state = None if fragments else TEXT_NOT_AVAILABLE_NO_OCR
    if unsupported_state:
        warnings.append(TEXT_NOT_AVAILABLE_NO_OCR)

    return LocatedTextDocument(
        source_filename=sanitize_filename(filename),
        source_sha256=hashlib.sha256(data).hexdigest(),
        file_type=LocatedTextFileType.PDF,
        reader_method="pypdf.PdfReader.page.extract_text",
        reader_version=reader_version,
        fragments=tuple(fragments),
        pages=tuple(pages),
        sheets=(),
        warnings=tuple(dict.fromkeys(warnings)),
        unsupported_state=unsupported_state,
        text=text,
        line_locations=line_locations,
    )


def read_document_text(
    artifact_bytes: bytes,
    filename: str,
    *,
    limits: TextReaderLimits | None = None,
) -> LocatedTextDocument:
    """Read deterministic located text from a real XLSX or born-digital PDF.

    The function consumes actual bytes, records their SHA-256, and returns an
    explicit no-OCR state for PDF files whose selectable text layer is blank.
    Unsupported/corrupt/over-bound inputs fail closed with
    :class:`DocumentSecurityError`.
    """

    active_limits = limits or TextReaderLimits()
    _check_common_bounds(artifact_bytes, active_limits)
    extension = _extension(filename)
    if extension in {".xlsx", ".xlsm"}:
        return _read_xlsx_text(artifact_bytes, filename, active_limits)
    if extension == ".pdf":
        return _read_pdf_text(artifact_bytes, filename, active_limits)
    raise DocumentSecurityError("unsupported real-file type for located text")


__all__ = [
    "LocatedTextDocument",
    "LocatedTextFileType",
    "LocatedTextFragment",
    "MAX_PDF_CONTENT_STREAM_BYTES",
    "MAX_XLSX_CELLS",
    "TEXT_NOT_AVAILABLE_NO_OCR",
    "TextReaderLimits",
    "XLSX_FORMULA_UNSUPPORTED",
    "read_document_text",
]
