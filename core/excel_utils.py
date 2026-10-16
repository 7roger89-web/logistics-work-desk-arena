# FILE: core/excel_utils.py
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date
from core.config import GETFLY_SERVICES_LIST
from core.utils import safe_float, parse_date
from core.logic import calculate_from_shipment

# --- STYLES ---
BORDER_THIN = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
FONT_BOLD = Font(name='Arial', bold=True, size=10)
FONT_HEADER = Font(name='Arial', bold=True, color='FFFFFF')
FILL_HEADER = PatternFill(start_color="0F2C59", end_color="0F2C59", fill_type="solid")
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')

def apply_header_style(cell):
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.alignment = ALIGN_CENTER
    cell.border = BORDER_THIN

def auto_fit_columns(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
            except (TypeError, ValueError):
                continue
        adjusted_width = (max_length + 2) * 1.1
        ws.column_dimensions[column_letter].width = min(adjusted_width, 50)

def find_best_match_code(input_name):
    if not input_name: return "DVK"
    clean_input = input_name.lower().strip()
    for s in GETFLY_SERVICES_LIST:
        if s['name'].lower().strip() == clean_input: return s['code']
    for s in GETFLY_SERVICES_LIST:
        s_name = s['name'].lower().strip()
        if clean_input in s_name or s_name in clean_input: return s['code']
    return "DVK"

def generate_mini_export_excel(shipment):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Debit Note"

    ws['A1'] = "LOGISTICS PRO SYSTEM"; ws['A1'].font = Font(size=16, bold=True, color="0F2C59")
    ws['A2'] = "BẢNG KÊ CHI TIẾT DỊCH VỤ (DEBIT NOTE)"; ws['A2'].font = Font(size=14, bold=True)
    
    d = shipment.get('data', {})
    sid = shipment.get('id')
    cust = shipment.get('customer')
    
    ws['A4'] = f"Mã Lô: {sid}"; ws['A4'].font = FONT_BOLD
    ws['D4'] = f"Khách hàng: {cust}"; ws['D4'].font = FONT_BOLD
    ws['A5'] = f"Ngày tàu: {d.get('1.2',{}).get('etd')}"; 
    ws['D5'] = f"Số Bill: {d.get('1.2',{}).get('mbl')}"

    headers = ["STT", "Tên Dịch vụ", "Mã", "ĐVT", "SL", "Đơn giá (Vốn)", "Đơn giá (Bán)", "Thành tiền (Trước VAT)", "VAT", "Tổng Tiền", "Ghi chú"]
    start_row = 7
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        apply_header_style(cell)

    details = d.get('3.4', {}).get('details', [])
    curr_row = start_row + 1
    total_final = 0

    for i, item in enumerate(details, 1):
        qty = safe_float(item.get('qty', 1))
        p_in = safe_float(item.get('cost', item.get('price_in', 0)))
        p_out = safe_float(item.get('price', item.get('price_out', 0)))
        vat_pct = safe_float(item.get('vat_line', item.get('vat', 0)))
        
        pre_tax = qty * p_out
        vat_val = pre_tax * vat_pct / 100
        line_total = pre_tax + vat_val
        total_final += line_total
        
        code = item.get('code', '')
        name = item.get('name_full', item.get('name',''))
        if not code or code == "DV": code = find_best_match_code(name)

        row_data = [
            i, name, code, "Lần", 
            qty, p_in, p_out, pre_tax, f"{vat_pct}%", line_total, item.get('note', '')
        ]
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.border = BORDER_THIN
            if col_idx in [6, 7, 8, 10]: cell.number_format = '#,##0'
            if col_idx == 1: cell.alignment = ALIGN_CENTER
        curr_row += 1

    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=9)
    total_label = ws.cell(row=curr_row, column=1, value="TỔNG CỘNG PHẢI THU:")
    total_label.font = FONT_BOLD; total_label.alignment = ALIGN_RIGHT
    
    total_val = ws.cell(row=curr_row, column=10, value=total_final)
    total_val.font = Font(bold=True, color="C62828"); total_val.number_format = '#,##0'
    total_val.border = BORDER_THIN

    auto_fit_columns(ws)
    wb.save(output)
    output.seek(0)
    return output

def generate_commission_excel(target_list):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Commission Report"

    headers = ["Mã Lô", "Khách Hàng", "Ngày Tạo", "Doanh Thu", "Chi Phí", "Lợi Nhuận Gộp", "Hoa Hồng (Chi)", "Thực Lãi Net", "Ghi Chú"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        apply_header_style(cell)

    curr_row = 2
    for sid, s in target_list:
        d = s.get('data', {})
        summary = calculate_from_shipment(s)
        revenue_net = safe_float(summary.get('revenue_net', 0))
        cost = safe_float(summary.get('cost', 0))
        comm_info = d.get('3.4', {}).get('commission', {})
        comm_val = safe_float(summary.get('commission', 0))
        profit_before_commission = revenue_net - cost
        net_profit = safe_float(summary.get('profit', 0))

        row = [
            sid, s.get('client_short') or s.get('customer') or s.get('client', ''), d.get('0.0', {}).get('created_at', ''),
            revenue_net, cost, profit_before_commission, comm_val, net_profit, comm_info.get('note', '')
        ]
        
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.border = BORDER_THIN
            if col_idx in [4, 5, 6, 7, 8]: cell.number_format = '#,##0'
        curr_row += 1

    auto_fit_columns(ws)
    wb.save(output)
    output.seek(0)
    return output

# --- 3. XUẤT FILE IMPORT GETFLY (Fix Lỗi Tên Cột & Thêm Tên NV) ---
def generate_getfly_excel(shipments_dict, assignee_name="", customers_dict=None):
    if customers_dict is None:
        from core.database import load_customers
        customers_dict = load_customers()
    db_cust = customers_dict
    
    # ĐÃ SỬA: "Địa chỉ giao hàng" -> "Địa chỉ", "CẤP ĐỘ RỦI RO " -> "CẤP ĐỘ RỦI RO"
    ALL_COLUMNS = [
        "Ngày đặt hàng", "Ngày tạo", "Mã ĐH", "Người thực hiện", 
        "Mã KH", "Tên KH", "Số ĐT", "Email", "Địa chỉ", 
        "Người liên hệ", "Điện thoại người liên hệ", "Tỉnh/ TP", "Quận/ Huyện", 
        "Nguồn khách hàng", "Mã vận đơn", "Trạng thái vận đơn", "Mã kho", 
        "Mã SP", "Tên sản phẩm", "Mô tả", "Đơn vị", "Nhóm sản phẩm", 
        "Số lượng", "Giá vốn", "Giá bán", "CK(%)", "CK(Đ)", "VAT(%)", 
        "Thành tiền", "Lợi nhuận SP", "Doanh số", "Chiết khấu (đ)", 
        "Doanh thu sau CK", "Phí vận chuyển (đ)", "Phí lắp đặt (đ)", 
        "Doanh thu trước thuế", "VAT (đ)", "Doanh thu", "Đã thanh toán", 
        "Còn lại", "Lợi nhuận", "Điều khoản đơn hàng", "Hình thức thanh toán", 
        "Hình thức thanh toán từng lần", "Ngày thanh toán", "Nguồn đơn hàng", 
        "Người tạo", "Trạng thái", "Nhóm khách hàng", "REFUN", 
        "HOA HỒNG KH GỬI", "CẤP ĐỘ RỦI RO"
    ]

    CODE_TO_OFFICIAL = {s['code']: s for s in GETFLY_SERVICES_LIST}
    data_rows = []
    
    c_map = {v.get('name', '').strip(): v for v in db_cust.values()} if db_cust else {}

    for sid, s in shipments_dict.items():
        d = s.get('data', {})
        cust_name = (s.get('customer') or s.get('client') or '').strip()
        c_info = db_cust.get(str(s.get('customer_id', '')), {}) or c_map.get(cust_name, {})
        
        ma_kh_chuan = c_info.get('id', '')
        
        created_at = parse_date(d.get('0.0', {}).get('created_at'))
        date_str = created_at.strftime('%d/%m/%Y') if created_at else ""
        
        details = d.get('3.4', {}).get('details', [])
        summary = calculate_from_shipment(s)
        paid_amt = safe_float(summary.get('paid', 0))
        total_rev = safe_float(summary.get('receivable', 0))
        total_profit = safe_float(summary.get('profit', 0))
        total_debt = safe_float(summary.get('debt', 0))
        
        base_row = {k: "" for k in ALL_COLUMNS}
        
        if not details: 
            h_row = base_row.copy()
            h_row.update({
                "Mã ĐH": sid, "Ngày đặt hàng": date_str, "Ngày tạo": date_str,
                "Người thực hiện": assignee_name, # Gắn tên Người thực hiện
                "Mã KH": ma_kh_chuan, "Tên KH": cust_name, 
                "Số ĐT": c_info.get('phone', ''), "Email": c_info.get('email', ''),
                "Địa chỉ": c_info.get('address', ''),
                "Doanh thu": total_rev, "Lợi nhuận": total_profit, "Đã thanh toán": paid_amt,
                "Trạng thái": "Hoàn thành" if s.get('status') in ['done', 'Done'] else "Mới",
                "Còn lại": total_debt,
                "Tên sản phẩm": "Dịch vụ Logistics", "Mã SP": "DVK", 
                "Số lượng": 1, "Thành tiền": total_rev
            })
            data_rows.append(h_row)
        else:
            for i, item in enumerate(details):
                row = base_row.copy()
                
                qty = safe_float(item.get('qty', 1))
                p_out = safe_float(item.get('price', item.get('price_out', 0)))
                p_in = safe_float(item.get('cost', item.get('price_in', 0)))
                vat_pct = safe_float(item.get('vat_line', item.get('vat', 0)))
                line_rev = qty * p_out
                line_profit = (p_out - p_in) * qty
                
                current_code = str(item.get('code', '')).strip()
                current_name = str(item.get('name_full', item.get('name', ''))).strip()
                
                if "(" in current_name and ")" in current_name:
                    extracted_code = current_name.split("(")[-1].replace(")", "").strip()
                    current_name = current_name.split("(")[0].strip()
                    if not current_code or current_code in ["None", ""]:
                        current_code = extracted_code
                
                if not current_code or current_code in ["DV", "DV-KHAC", "None"]:
                    current_code = find_best_match_code(current_name)
                
                unit_val, group_val = "Lần", "DỊCH VỤ XNK"
                if current_code in CODE_TO_OFFICIAL:
                    current_name = CODE_TO_OFFICIAL[current_code]['name']
                    unit_val = CODE_TO_OFFICIAL[current_code]['unit']
                    group_val = CODE_TO_OFFICIAL[current_code]['group']

                row.update({
                    "Mã SP": current_code,
                    "Tên sản phẩm": current_name,
                    "Mô tả": item.get('note', ''),
                    "Đơn vị": unit_val, 
                    "Nhóm sản phẩm": group_val,
                    "Số lượng": qty,
                    "Giá vốn": p_in, "Giá bán": p_out, "VAT(%)": vat_pct,
                    "Thành tiền": line_rev, "Doanh số": line_rev, "Lợi nhuận SP": line_profit,
                })

                if i == 0:
                    row.update({
                        "Mã ĐH": sid, "Ngày đặt hàng": date_str, "Ngày tạo": date_str,
                        "Người thực hiện": assignee_name, # Gắn tên Người thực hiện
                        "Mã KH": ma_kh_chuan, "Tên KH": cust_name, 
                        "Số ĐT": c_info.get('phone', ''), "Email": c_info.get('email', ''),
                        "Địa chỉ": c_info.get('address', ''),
                        "Doanh thu": total_rev, "Lợi nhuận": total_profit, "Đã thanh toán": paid_amt,
                        "Trạng thái": "Hoàn thành" if s.get('status') in ['done', 'Done'] else "Mới",
                        "Còn lại": total_debt
                    })
                
                data_rows.append(row)

    if not data_rows: return None
    df = pd.DataFrame(data_rows)
    df = df.reindex(columns=ALL_COLUMNS)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
        ws = writer.sheets['Sheet1']
        auto_fit_columns(ws)
        
    output.seek(0)
    return output

def _write_key_value_sheet(title, sections):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="0F2C59")
    row = 3
    for heading, values in sections:
        ws.cell(row=row, column=1, value=heading).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=row, column=1).fill = FILL_HEADER
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
        for key, value in values:
            ws.cell(row=row, column=1, value=str(key)).font = FONT_BOLD
            ws.cell(row=row, column=2, value=value)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            row += 1
        row += 1
    auto_fit_columns(ws)
    wb.save(output)
    output.seek(0)
    return output


def _quote_lines(quote):
    """Canonical service lines of a LOG-C12 quote, with the legacy shapes."""
    lines = quote.get("service_lines")
    if isinstance(lines, list) and lines:
        return lines
    return quote.get("details", quote.get("items", [])) or []


class _CustomerText:
    """The Sale's editable customer wording for ONE quotation.

    LOGI_EXCEL_FIDELITY_INLINE_EDIT_V1 — the Step-3 preview and this exporter
    read the same ``customer_quote_snapshot()["text"]`` map. ``get()`` returns the
    edited value when the quote carries the map, and today's template literal
    otherwise, so a legacy quote exports byte-for-byte as before. Amounts are
    never part of this map: they stay backend-computed numbers.
    """

    def __init__(self, quote):
        from core.quote_customer_text import text_values

        self.values = text_values(quote)

    @property
    def mapped(self) -> bool:
        return bool(self.values)

    def get(self, region_id: str, default: str) -> str:
        value = self.values.get(region_id)
        return value if isinstance(value, str) and value.strip() else default


def _line_field(line, *names, default=None):
    for name in names:
        if name in line and line.get(name) not in (None, ""):
            return line.get(name)
    return default


def _quote_display_date(quote):
    """The quotation's OWN date, never silently "today".

    Precedence: an explicit quote date → the canonical ``created_at`` the quote
    was written with → today only when the document genuinely carries neither.
    Re-exporting a June quote in September must not restamp it as September.
    """
    for key in ("date", "quote_date", "issued_date"):
        value = str(quote.get(key) or "").strip()
        if value:
            return value
    created = str(quote.get("created_at") or "").strip()
    if created:
        # Canonical timestamps are ISO ("2026-06-12T09:31:07+00:00"); the date
        # part is what a quotation shows. Anything else travels verbatim.
        return created.split("T")[0].strip() or created
    return str(date.today())


def generate_quote_excel(quote, *, company=None, customer=None, logo_path=None):
    """Customer quotation XLSX in the Owner's real template hierarchy.

    Rows 1-5   company header (Trường Thành authority) + BÁO GIÁ
    Rows 7-10  From/Attn/Tel/Mobile  |  To/Email/Mobile/Date
    Row 11     Subject
    Rows 12-15 intro / goods / route / service statement
    Row 17     table headers (STT · Chi tiết · Đơn giá · Số lượng · Thành tiền ·
               VAT · thành tiền · Đơn vị tiền · Phát sinh (nếu có))
    then       one row per service line, per-currency subtotals, FX disclosure,
               Ghi chú / closing / signature.

    Truth rules enforced here:
      * every line keeps its OWN currency; different currencies are NEVER summed
        into one number — there is one subtotal per currency code;
      * a VND conversion is only shown when the quote itself stores a quote-time
        FX rate, and that exact rate is disclosed;
      * conditional / estimated wording travels in "Phát sinh (nếu có)" so an
        optional charge is visible, not buried;
      * an approximate amount stays visibly approximate;
      * NO formulas are written at all, so the workbook cannot carry #REF!,
        #DIV/0!, #VALUE! or #NAME? — every cell is a literal value;
      * no Owner-local PC and no Office automation: pure openpyxl.
    """
    quote = dict(quote or {})
    excel_customer = bool(quote.get("excel_quote") or "goods" in quote)
    if quote.get("excel_quote"):
        from core.quotations import customer_quote_snapshot
        quote = customer_quote_snapshot(quote)
    company = dict(company or {})
    customer = dict(customer or {})
    details = _quote_lines(quote)
    ctext = _CustomerText(quote)

    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Bao gia"

    title_font = Font(name="Arial", size=14, bold=True, color="0F2C59")
    small = Font(name="Arial", size=9)
    label = Font(name="Arial", size=10, bold=True)

    # ---- rows 1-5: company authority header -------------------------------
    # LOGI_QUOTE_COST_WORKFLOW_UX_01_REPAIR_01 — read the REAL field names the
    # `authority_companies` master stores (legal_name / legal_address /
    # legal_representative / representative_title / phone / email). The older
    # `address` / `representative` guesses never matched, so a fully populated
    # authority record still rendered blank. Legacy aliases stay as fallbacks;
    # no duplicate company model is introduced.
    company_name = (company.get("legal_name") or company.get("name")
                    or "CÔNG TY TNHH GIAO NHẬN QUỐC TẾ TRƯỜNG THÀNH")
    company_address = company.get("legal_address") or company.get("address") or ""
    company_rep = company.get("legal_representative") or company.get("representative") or ""
    company_rep_title = company.get("representative_title") or ""
    company_phone = company.get("phone") or ""
    company_email = company.get("email") or ""
    ws["A1"] = company_name
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="0F2C59")
    ws["A2"] = company_address
    ws["A2"].font = small
    ws["A3"] = company_email
    ws["A3"].font = small
    ws["A4"] = company.get("website", "")
    ws["A4"].font = small
    # The document title is the FIRST editable customer region. With the text
    # map present it is authoritative for both the title row and the Subject
    # line; without the map the historical literals are preserved.
    document_title = ctext.get("doc.title", "BÁO GIÁ" if not ctext.mapped else "BÁO GIÁ DỊCH VỤ LOGISTICS")
    ws["A5"] = document_title
    ws["A5"].font = title_font
    ws["A5"].alignment = ALIGN_CENTER

    if logo_path:
        try:  # a missing/unsupported logo must never break the export
            from openpyxl.drawing.image import Image as XLImage

            image = XLImage(str(logo_path))
            image.width, image.height = 120, 48
            ws.add_image(image, "J1")
        except Exception:  # noqa: BLE001
            pass

    # ---- rows 7-10: From / To block ---------------------------------------
    # The "To" block is enriched from the EXISTING customer master (resolved by
    # the quote's canonical customer_id by the caller). A field the master does
    # not hold stays blank — nothing is fabricated and no customer is mutated.
    customer_name = (quote.get("customer_name") or quote.get("customer")
                     or customer.get("name") or quote.get("client", ""))
    customer_contact = (quote.get("attn") or quote.get("customer_contact")
                        or customer.get("contact_person") or "")
    customer_email = quote.get("customer_email") or customer.get("email") or ""
    customer_phone = quote.get("customer_phone") or customer.get("phone") or ""

    sender = company.get("contact_name") or company_rep
    if sender and company_rep_title:
        sender = f"{sender} ({company_rep_title})"
    left = [
        ("From", sender),
        ("Attn", customer_contact),
        ("Tel/Fax", company_phone),
        ("Mobile", company.get("mobile") or company_phone),
    ]
    right = [
        ("To", customer_name),
        ("Email", customer_email),
        ("Mobile", customer_phone),
        ("Date", _quote_display_date(quote)),
    ]
    for idx, ((lk, lv), (rk, rv)) in enumerate(zip(left, right)):
        row = 7 + idx
        ws.cell(row=row, column=1, value=lk).font = label
        ws.cell(row=row, column=2, value=lv)
        ws.cell(row=row, column=6, value=rk).font = label
        ws.cell(row=row, column=7, value=rv)

    # ---- row 11: subject; rows 12-15 the narrative ------------------------
    origin = str(quote.get("origin") or "")
    destination = str(quote.get("destination") or "")
    route = str(quote.get("route") or "")
    route_text = " - ".join([p for p in (origin, destination) if p]) or route
    incoterm = str(quote.get("incoterm") or "")
    subject_title = document_title if ctext.mapped else "BÁO GIÁ DỊCH VỤ LOGISTICS"
    ws["A11"] = "Subject: " + subject_title + (f" — {route_text}" if route_text else "")
    ws["A11"].font = label
    ws["A12"] = ctext.get("doc.greeting", "Trân trọng cảm ơn Quý khách đã quan tâm dịch vụ của chúng tôi. Chúng tôi xin gửi báo giá như sau:")
    ws["A13"] = "Mặt hàng: " + (ctext.get("doc.goods_description", str(quote.get("goods_description") or "")) or "—")
    ws["A14"] = "Tuyến: " + (ctext.get("doc.route", route_text) or "—") + (f" · Điều kiện giao hàng: {incoterm}" if incoterm else "")
    ws["A15"] = ctext.get("sec.service_summary", "Dịch vụ: giao nhận / vận chuyển / thủ tục theo các khoản mục dưới đây.")

    # ---- row 17: table headers -------------------------------------------
    from core.quote_customer_text import SERVICE_COLUMNS

    headers = [ctext.get(f"col.service.{key}", default) for key, default in SERVICE_COLUMNS]
    header_row = 17
    if excel_customer:
        ws["A15"] = ctext.get("doc.option_line",
                            f"Phương án {quote.get('option', '')}: " + ("Khách tự đứng tên nhập khẩu" if quote.get("option") == "A" else "Dịch vụ nhập khẩu ủy thác"))
        ws["A16"] = "Phạm vi dịch vụ đã xác nhận: " + (ctext.get("scope", str(quote.get("scope") or "")) or "Chưa xác nhận")
        row = 18
        shipment = quote.get("shipment") or {}
        public_route = " → ".join(str(shipment.get(k) or "Chưa xác nhận") for k in ("pol", "pod"))
        transport = " / ".join(str(v) for v in (quote.get("transport_mode"), shipment.get("load_model"), shipment.get("equipment")) if v)
        for text in (f"Cảng / điểm: {public_route} · {transport}",
                     "Nơi lấy hàng: " + str(shipment["pickup"]) if shipment.get("pickup") else "",
                     "Giao tại: " + str(shipment["delivery"]) if shipment.get("delivery") else "",
                     "Ngày giao dự kiến: " + str(quote["expected_delivery_date"]) if quote.get("expected_delivery_date") else ""):
            if text:
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
                ws.cell(row, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[row].height = 30
                row += 1
        ws.cell(row, 1, ctext.get("sec.goods", "DANH MỤC HÀNG HÓA — giá theo bảng khách gửi")).font = FONT_BOLD
        row += 1
        goods_column_defaults = ("STT", "Hàng hóa / Model / Vật liệu", "Dài mm", "Rộng mm", "Cao mm", "Số lượng", "Đơn giá", "Tiền tệ", "Thành tiền")
        goods_column_keys = ("stt", "name", "length", "width", "height", "qty", "unit_price", "currency", "line_total")
        for col, (key, default) in enumerate(zip(goods_column_keys, goods_column_defaults), 1):
            apply_header_style(ws.cell(row, col, ctext.get(f"col.goods.{key}", default)))
        from core.quote_customer_text import line_key

        for i, goods in enumerate(quote.get("goods") or [], 1):
            row += 1
            key = line_key("goods", goods, i - 1, fields=("description", "model", "material", "length_mm", "width_mm", "height_mm"))
            goods_name = " · ".join(value for value in (
                ctext.get(f"goods.{key}.name", str(goods.get("description") or "")),
                ctext.get(f"goods.{key}.model", str(goods.get("model") or "")),
                ctext.get(f"goods.{key}.material", str(goods.get("material") or "")),
            ) if value)
            values = [i, goods_name,
                      goods.get("length_mm"), goods.get("width_mm"), goods.get("height_mm"), goods.get("qty"),
                      goods.get("unit_price"), goods.get("currency"), goods.get("line_total")]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, "Chưa biết" if value is None else value).border = BORDER_THIN
            # Goods notes are the reviewed customer goods description; source
            # provenance and internal warning notes never enter this allowlist.
            goods_note = ctext.get(f"goods.{key}.note", str(goods.get("note") or ""))
            if goods_note:
                row += 1
                ws.cell(row, 2, "Ghi chú hàng: " + goods_note)
        for total in quote.get("goods_calculation", {}).get("goods_totals") or []:
            row += 1
            ws.cell(row, 2, "Tổng tiền hàng " + total["currency"]).font = FONT_BOLD
            ws.cell(row, 9, total["amount"] if total["amount"] is not None else "Chưa xác định")
        row += 2
        ws.cell(row, 1, ctext.get("sec.services", "PHÍ DỊCH VỤ — chỉ giá bán đã xác nhận")).font = FONT_BOLD
        header_row = row + 1
    for idx, header in enumerate(headers, 1):
        apply_header_style(ws.cell(row=header_row, column=idx, value=header))

    totals_by_currency = {}
    row = header_row + 1
    for index, item in enumerate(details, 1):
        item = dict(item)
        name = _line_field(item, "service_name", "name_full", "name", "label", default="")
        if excel_customer:
            # The published service name / public note are editable regions; the
            # amount columns below stay backend-computed numbers.
            line_id = line_key("service", item, index - 1, fields=("service_name", "unit", "currency"))
            name = ctext.get(f"service.{line_id}.name", str(name))
        unit_price = safe_float(_line_field(item, "sell_unit", "price", "price_out", "amount", default=0))
        qty = safe_float(_line_field(item, "qty", default=1))
        # UNKNOWN VAT is not 0% VAT. A quote line whose tax has not been
        # confirmed must print "—", because a customer reading "0%" would read
        # it as a commitment that no VAT applies. An explicit 0 still prints
        # "0%", so the two states never look alike.
        raw_vat = _line_field(item, "tax_pct", "vat", "vat_line", default=None)
        vat_known = raw_vat is not None
        vat_pct = safe_float(raw_vat) if vat_known else None
        currency = str(_line_field(item, "currency", "ccy", default="VND")).upper()
        pre_tax = item.get("sell_total") if excel_customer else unit_price * qty
        with_tax = ((pre_tax + item["tax_total"]) if item.get("tax_total") is not None and pre_tax is not None else None) if excel_customer else (pre_tax * (1 + vat_pct / 100) if vat_known else None)
        vat_known = with_tax is not None if excel_customer else vat_known
        bucket = totals_by_currency.setdefault(currency, {"total": 0.0, "known": True})
        if vat_known:
            bucket["total"] += with_tax
        else:
            bucket["known"] = False
            # the pre-tax amount is still a real fact of this currency
            bucket["total"] += pre_tax

        note = str(item.get("public_note") or "") if excel_customer else str(item.get("note") or "")
        if excel_customer:
            note = ctext.get(f"service.{line_key('service', item, index - 1, fields=('service_name', 'unit', 'currency'))}.public_note", note)
        condition = str(item.get("condition") or "")
        status = str(item.get("status") or "")
        contingent = []
        if condition == "IF_APPLICABLE" or "nếu có" in note.lower() or "phát sinh" in note.lower():
            contingent.append("Có thể phát sinh (nếu có)")
        if status == "ESTIMATE_REQUIRES_CONFIRMATION" or "khoảng" in note.lower() or "ước tính" in note.lower():
            contingent.append("Ước tính — cần xác nhận, chưa phải chi phí chắc chắn")
        if note:
            contingent.append(note)

        values = [
            index, name, unit_price, qty, pre_tax,
            f"{vat_pct:g}%" if vat_known else "—",
            with_tax if vat_known else "—",
            currency, " · ".join(contingent),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.border = BORDER_THIN
            if col_idx in {3, 5, 7}:
                cell.number_format = "#,##0"
            if col_idx in {1, 6, 8}:
                cell.alignment = ALIGN_CENTER
            if col_idx == 9:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    # ---- per-currency subtotals (NEVER one mixed-currency number) ---------
    row += 1
    ws.cell(row=row, column=2, value=ctext.get("sec.totals", "TỔNG CỘNG THEO TỪNG LOẠI TIỀN (không quy đổi gộp)")).font = FONT_BOLD
    row += 1
    for currency in sorted(totals_by_currency):
        bucket = totals_by_currency[currency]
        label_text = f"Tổng {currency}" if bucket["known"] else f"Tổng {currency} (chưa gồm VAT — còn dòng chưa xác định VAT)"
        ws.cell(row=row, column=2, value=label_text).font = FONT_BOLD
        total_cell = ws.cell(row=row, column=7, value=bucket["total"])
        total_cell.number_format = "#,##0"
        total_cell.font = FONT_BOLD
        ws.cell(row=row, column=8, value=currency).alignment = ALIGN_CENTER
        row += 1

    # ---- FX disclosure: only a rate the quote really stores ---------------
    fx = quote.get("fx_snapshot") if isinstance(quote.get("fx_snapshot"), dict) else {}
    rates = {str(code).upper(): safe_float(rate)
             for code, rate in dict(fx.get("rates_to_vnd") or {}).items()}
    used = [code for code in sorted(totals_by_currency)
            if code != "VND" and rates.get(code)]
    if used:
        row += 1
        ws.cell(row=row, column=2, value="Tỷ giá báo giá đã dùng để quy đổi VND:").font = FONT_BOLD
        row += 1
        for code in used:
            ws.cell(row=row, column=2, value=f"{code} → VND = {rates[code]:,.0f}")
            converted = ws.cell(row=row, column=7, value=totals_by_currency[code]["total"] * rates[code])
            converted.number_format = "#,##0"
            ws.cell(row=row, column=8, value="VND").alignment = ALIGN_CENTER
            row += 1
        ws.cell(row=row, column=2,
                value=str(fx.get("source") or "Tỷ giá nhập khi báo giá; không phải tỷ giá giao dịch thực tế."))
        row += 1
    elif any(code != "VND" for code in totals_by_currency):
        row += 1
        ws.cell(row=row, column=2,
                value="Báo giá có dòng ngoại tệ nhưng chưa có tỷ giá báo giá — không quy đổi VND.").font = FONT_BOLD
        row += 1

    if excel_customer:
        row += 1
        customs = quote.get("customs") or {}
        ws.cell(row, 2, ctext.get("sec.customs", "THUẾ NHẬP KHẨU DỰ KIẾN (không phải kết luận pháp lý)")).font = FONT_BOLD
        row += 1
        if customs.get("calculable"):
            ws.cell(row, 2, f"Cơ sở hải quan xác nhận riêng (VND); Thuế NK {customs.get('duty_pct')}% / VAT NK {customs.get('vat_pct')}%")
            ws.cell(row, 7, customs.get("basis_vnd"))
            row += 1
            for text, field in (("Thuế NK dự kiến", "duty_vnd"), ("VAT NK dự kiến", "vat_vnd")):
                ws.cell(row, 2, text)
                ws.cell(row, 7, customs.get(field))
                ws.cell(row, 8, "VND")
                row += 1
        else:
            ws.cell(row, 2, "Chưa xác định — trị giá hải quan / thuế suất chưa được xác nhận.")
            row += 1
        calculation = quote.get("goods_calculation") or {}
        for text, field in (("Tạm tính hàng + dịch vụ trước mọi thuế (VND)", "before_tax_vnd"),
                            ("TỔNG DỰ TOÁN gồm thuế đã xác định (VND)", "estimate_vnd")):
            ws.cell(row, 2, text).font = FONT_BOLD
            amount = calculation.get(field)
            ws.cell(row, 7, amount if amount is not None else "Chưa xác định")
            row += 1
        ws.cell(row, 2, "Đây là dự toán; không phải hóa đơn, khoản công nợ hay chứng từ xác nhận thuế.")
        row += 1

    # ---- Ghi chú / closing / signature ------------------------------------
    row += 1
    ws.cell(row=row, column=1, value=ctext.get("sec.notes", "Ghi chú:")).font = FONT_BOLD
    row += 1
    # CUSTOMER-FACING notes only. `quote["assumptions"]` is an INTERNAL control
    # list (agent price reconciliation, customs policy, …) and is deliberately
    # NOT printed onto a document the customer receives.
    notes = [str(n) for n in (quote.get("customer_notes") or []) if str(n).strip()]
    if not notes:
        notes = ["Báo giá là dự toán, chưa phải hóa đơn hoặc nghĩa vụ thanh toán."]
    for note in notes:
        ws.cell(row=row, column=1, value="- " + note)
        row += 1
    ws.cell(row=row, column=1,
            value="- Các khoản ghi “Phát sinh (nếu có)” chỉ tính khi thực tế phát sinh.")
    row += 1
    ws.cell(row=row, column=1,
            value="- Các khoản ghi “Ước tính” chưa phải chi phí chắc chắn, sẽ xác nhận lại trước khi thực hiện.")
    # An optional extra paragraph the Sale added in the Step-3 preview. Empty by
    # default, so an untouched quotation prints exactly as before.
    addendum = ctext.get("doc.addendum", "")
    if addendum:
        row += 1
        for line_text in str(addendum).split("\n"):
            ws.cell(row=row, column=1, value=line_text)
            row += 1
    row += 2
    ws.cell(row=row, column=1, value=ctext.get("doc.closing", "Rất mong được hợp tác cùng Quý khách."))
    row += 2
    ws.cell(row=row, column=6, value=company_name).font = FONT_BOLD

    for cells in ws:
        for cell in cells:
            if cell.data_type == "f":
                cell.data_type = "s"  # customer text beginning '=' is literal, never executable
    auto_fit_columns(ws)
    wb.save(output)
    output.seek(0)
    return output


def generate_sales_contract_excel(contract):
    contract = contract or {}
    return _write_key_value_sheet("HỢP ĐỒNG / SALES CONTRACT", [
        ("THÔNG TIN HỢP ĐỒNG", [
            ("Số hợp đồng", contract.get("contract_no", contract.get("id", ""))),
            ("Ngày ký", contract.get("date", "")),
            ("Đối tác", contract.get("client", "")),
            ("Giá trị", contract.get("value", 0)),
            ("Tiền tệ", contract.get("currency", "USD")),
            ("Điều khoản", contract.get("note", "")),
        ])
    ])


def generate_debit_note_excel(shipment):
    return generate_mini_export_excel(shipment)